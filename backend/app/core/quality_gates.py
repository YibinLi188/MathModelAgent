"""阶段质量闸门。

质量闸门不判断模型优劣，而是阻止明显不完整或不可审计的结果继续流向
论文阶段。所有失败都以可读错误返回，便于前端和验收报告处理。
"""

from __future__ import annotations

import json
import math
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.schemas.A2A import CoderToWriter, ModelerToCoder, WriterResponse


class QualityGateError(ValueError):
    """结果未满足阶段契约。"""


def _contains_finite_number(value: Any) -> bool:
    if isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return math.isfinite(float(value))
    if isinstance(value, dict):
        return any(_contains_finite_number(item) for item in value.values())
    if isinstance(value, list):
        return any(_contains_finite_number(item) for item in value)
    return False


def _finite_number_values(value: Any) -> list[float]:
    if isinstance(value, bool):
        return []
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return [float(value)]
    if isinstance(value, dict):
        return [
            number for item in value.values() for number in _finite_number_values(item)
        ]
    if isinstance(value, list):
        return [number for item in value for number in _finite_number_values(item)]
    return []


def _to_metres(value: str, unit: str) -> float:
    number = float(value)
    return number / 100 if unit.lower() == "cm" else number


def _planned_result_files(plan_text: str) -> set[str]:
    """Return result workbooks that the plan actually promises to create."""
    planned: set[str] = set()
    for match in re.finditer(r"result\d+\.xlsx", plan_text, re.IGNORECASE):
        clause_start = max(
            plan_text.rfind(separator, 0, match.start())
            for separator in ("。", "；", ";", "\n")
        )
        prefix = plan_text[clause_start + 1 : match.start()][-48:]
        negates_creation = re.search(
            r"(?:不|无需|无须|不得|禁止)(?:再|另行|额外|单独)?\s*"
            r"(?:生成|输出|写入|保存|创建|提交|提供|包含|需要|要求)?"
            r"[^。；;\n]{0,16}$|题面[^。；;\n]{0,8}(?:未|不)要求\s*$",
            prefix,
            re.IGNORECASE,
        )
        if not negates_creation:
            planned.add(match.group(0).lower())
    return planned


def _has_affirmative_internal_tangency(text: str) -> bool:
    for match in re.finditer(r"内切(?:分支)?", text, re.IGNORECASE):
        clause_start = max(
            text.rfind(separator, 0, match.start())
            for separator in ("。", "；", ";", "\n")
        )
        prefix = text[clause_start + 1 : match.start()][-24:]
        if re.search(
            r"(?:不|不得|不能|不可|禁止|避免|严禁|排除)"
            r"[^。；;\n]{0,12}$",
            prefix,
        ):
            continue
        return True
    return False


def derive_source_guardrails(source_text: str) -> dict[str, Any] | None:
    """Derive a narrow serial-member invariant directly from statement facts.

    This is not an answer key. It only activates when the statement itself gives a
    serial member count, two end markers per member, shared-handle connections,
    member lengths and the marker offset from each nearest end.
    """
    entity_match = re.search(
        r"由\s*(\d+)\s*节(?:板凳|构件|杆件|链节).*?组成", source_text
    )
    two_markers = re.search(r"每\s*节.{0,20}(?:两个|2\s*个)孔", source_text, re.DOTALL)
    shared_connection = re.search(
        r"相邻.{0,30}(?:通过.{0,12})?(?:把手|铰链|连接件).{0,12}连接",
        source_text,
    )
    offset_match = re.search(
        r"孔的中心距离最近的板头\s*(\d+(?:\.\d+)?)\s*(cm|m)",
        source_text,
        re.IGNORECASE,
    )
    lengths = re.findall(
        r"([\u4e00-\u9fff、和]{1,12})的板长(?:均)?为\s*"
        r"(\d+(?:\.\d+)?)\s*(cm|m)",
        source_text,
        re.IGNORECASE,
    )
    if not (
        entity_match and two_markers and shared_connection and offset_match and lengths
    ):
        return None

    entity_count = int(entity_match.group(1))
    offset_m = _to_metres(offset_match.group(1), offset_match.group(2))
    member_constraints = []
    for label, value, unit in lengths:
        length_m = _to_metres(value, unit)
        distance_m = length_m - 2 * offset_m
        if distance_m <= 0 or distance_m > length_m:
            raise QualityGateError("题面构件尺寸无法形成有效的两端内部节点距离")
        member_constraints.append(
            {
                "label": label,
                "source_length_m": length_m,
                "end_offset_m": offset_m,
                "node_distance_m": distance_m,
            }
        )
    guardrails: dict[str, Any] = {
        "kind": "serial_two_end_marker_members",
        "entity_count": entity_count,
        "node_count": entity_count + 1,
        "shared_connection_nodes": True,
        "member_constraints": member_constraints,
    }
    role_counts = re.search(
        r"第\s*1\s*节为龙头，?\s*后面\s*(\d+)\s*节为龙身，?\s*"
        r"最后\s*1\s*节为龙尾",
        source_text,
    )
    if role_counts:
        body_count = int(role_counts.group(1))
        if body_count + 2 == entity_count:
            guardrails["handle_identity_contract"] = {
                "total_output_handles": entity_count + 1,
                "head_front_node": 1,
                "body_front_nodes": [2, body_count + 1],
                "body_indices": [1, body_count],
                "tail_front_node": body_count + 2,
                "tail_rear_node": body_count + 3,
                "tail_rear_included_in_total": True,
            }
    width_match = re.search(
        r"板宽(?:均)?为\s*(\d+(?:\.\d+)?)\s*(cm|m)",
        source_text,
        re.IGNORECASE,
    )
    if width_match:
        guardrails["member_width_m"] = _to_metres(
            width_match.group(1), width_match.group(2)
        )
    initial_spiral = re.search(
        r"问题\s*1[^。]{0,100}螺距为\s*(\d+(?:\.\d+)?)\s*(cm|m)"
        r".*?初始(?:时)?[^。]{0,100}第\s*(\d+)\s*圈",
        source_text,
        re.DOTALL | re.IGNORECASE,
    )
    if initial_spiral:
        pitch_m = _to_metres(initial_spiral.group(1), initial_spiral.group(2))
        turn_count = int(initial_spiral.group(3))
        guardrails["initial_spiral"] = {
            "pitch_m": pitch_m,
            "turn_count": turn_count,
            "radius_m": pitch_m * turn_count,
        }
    if re.search(
        r"问题\s*2.{0,300}(?:板凳之间不发生碰撞|不能再继续盘入)",
        source_text,
        re.DOTALL,
    ) and re.search(
        r"问题\s*3.{0,300}龙头前把手.{0,100}盘入到调头空间的边界",
        source_text,
        re.DOTALL,
    ):
        guardrails["cross_question_feasibility"] = {
            "target_subject": "龙头前把手",
            "target_event": "到达调头空间边界",
            "required": ["到达前存在正行程", "全链有限板体无碰撞直至目标事件"],
            "forbidden": ["初始位置等于目标边界的零行程解", "只验证轨迹曲线相交"],
        }
    if re.search(
        r"两段圆弧.{0,50}相切.{0,100}(?:半径.{0,20}2\s*倍|2\s*[:：]\s*1)|"
        r"(?:半径.{0,20}2\s*倍|2\s*[:：]\s*1).{0,100}两段圆弧.{0,50}相切",
        source_text,
        re.DOTALL,
    ):
        guardrails["piecewise_curve_continuity"] = {
            "required": ["位置与切向 C1 连续", "连接点曲率按 1/R 分别计算并明确跳变"],
            "forbidden": ["把相切称为曲率连续或 C2 连续"],
        }
        if re.search(r"S\s*形", source_text, re.IGNORECASE):
            guardrails["s_curve_tangency_contract"] = {
                "required": ["反向转向的外切分支", "圆心距 R1+R2"],
                "forbidden": ["内切分支", "R1-R2", "R1±R2"],
            }
    source_ques4 = re.search(
        r"问题\s*4(?P<body>.*?)(?=问题\s*5|$)", source_text, re.DOTALL
    )
    if source_ques4 and re.search(
        r"(?:能否|是否)[^。；]{0,50}(?:缩短|变短)", source_ques4.group("body")
    ):
        guardrails["shortening_comparison_contract"] = {
            "required": [
                "由题面调整前双圆弧构型及全部相切约束复算基线长度 L_base",
                "按实际相切点与连接点决定的圆心角计算 L=R1·Δφ1+R2·Δφ2",
                "将优化候选 L_opt 与 L_base 作数值比较后再回答能否缩短",
            ],
            "forbidden": ["任取一个可行半径或任意相切构型充当调整前基线"],
        }
    source_ques5 = re.search(r"问题\s*5(?P<body>.*)$", source_text, re.DOTALL)
    if (
        source_ques5
        and "盘入螺线" in source_text
        and "盘出螺线" in source_text
        and re.search(r"最大[^。；]{0,30}速度", source_ques5.group("body"))
    ):
        guardrails["unbounded_path_extremum_contract"] = {
            "required": [
                "覆盖全部节点和高曲率有限区间",
                "给出截断外解析定量上界，或逐次扩大外侧截断并验证极值稳定且新增外侧壳层不再增大",
            ],
            "forbidden": ["仅凭曲率趋零、速度倍率趋一宣称已获得全路径极值"],
        }
    return guardrails


def format_source_guardrails(source_text: str) -> str:
    guardrails = derive_source_guardrails(source_text)
    if guardrails is None:
        return "未触发可确定的串联构件程序约束。"
    return "程序从题面原值直接推导的前置约束（不是参考答案）：" + json.dumps(
        guardrails, ensure_ascii=False, separators=(",", ":")
    )


def derive_source_figure_requirements(source_text: str) -> list[dict[str, Any]]:
    """Find source figures that carry an initial/boundary geometry fact.

    The check is deliberately narrow: decorative or explanatory figures do not
    become mandatory merely because the statement mentions them. A figure is
    mandatory when the same sentence uses it to define an initial position,
    named point, coordinate or direction.
    """
    requirements: dict[int, dict[str, Any]] = {}
    for match in re.finditer(r"(?:见图|图见)\s*(\d+)", source_text):
        sentence_start = max(
            source_text.rfind(mark, 0, match.start()) for mark in "。！？；"
        )
        sentence_end_candidates = [
            index
            for mark in "。！？；"
            if (index := source_text.find(mark, match.end())) >= 0
        ]
        sentence_end = min(sentence_end_candidates, default=len(source_text))
        context = source_text[sentence_start + 1 : sentence_end + 1].strip()
        if not re.search(
            r"(?:初始|起点|坐标|方向|方位|位于|位置|"
            r"边界|圆形区域|半径|直径|[A-Za-zＡ-Ｚａ-ｚ]\s*点)",
            context,
        ):
            continue
        figure_number = int(match.group(1))
        requirements[figure_number] = {
            "figure": figure_number,
            "context": re.sub(r"\s+", " ", context),
        }
    return list(requirements.values())


def format_source_figure_requirements(source_text: str) -> str:
    requirements = derive_source_figure_requirements(source_text)
    if not requirements:
        return "题面未触发必须核验的关键图示。"
    return "题面关键图示核验要求：" + json.dumps(
        requirements, ensure_ascii=False, separators=(",", ":")
    )


def _require_guardrail_numbers(
    value: Any, source_text: str, *, allow_centimetres: bool = False
) -> None:
    guardrails = derive_source_guardrails(source_text)
    if guardrails is None:
        return
    numbers = _finite_number_values(value)
    node_count = float(guardrails["node_count"])
    expected_distances = [
        float(item["node_distance_m"]) for item in guardrails["member_constraints"]
    ]
    missing = []
    if not any(math.isclose(actual, node_count, abs_tol=1e-6) for actual in numbers):
        missing.append(node_count)
    for distance in expected_distances:
        candidates = [distance]
        if allow_centimetres:
            candidates.append(distance * 100)
        if not any(
            math.isclose(actual, candidate, abs_tol=1e-6)
            for actual in numbers
            for candidate in candidates
        ):
            missing.append(distance)
    if missing:
        raise QualityGateError(
            f"串联构件证据未包含题面可直接推导的节点数/内部节点距离: {missing}"
        )


def _reject_guardrail_contradictions(
    value: Any, source_text: str, *, subtask_title: str | None = None
) -> None:
    guardrails = derive_source_guardrails(source_text)
    if guardrails is None:
        return
    text = (
        value
        if isinstance(value, str)
        else json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    )
    text_fields: list[str] = []

    def collect_text_fields(item: Any) -> None:
        if isinstance(item, str):
            text_fields.append(item)
        elif isinstance(item, dict):
            for child in item.values():
                collect_text_fields(child)
        elif isinstance(item, (list, tuple)):
            for child in item:
                collect_text_fields(child)

    collect_text_fields(value)
    if not text_fields:
        text_fields.append(text)

    body_count_match = re.search(r"后面\s*(\d+)\s*节为龙身", source_text, re.IGNORECASE)
    for field in text_fields:
        for arithmetic in re.finditer(
            r"(?P<lhs>\d+(?:\.\d+)?(?:\s*[×xX*]\s*\d+(?:\.\d+)?)*"
            r"(?:\s*[+＋\-−]\s*\d+(?:\.\d+)?"
            r"(?:\s*[×xX*]\s*\d+(?:\.\d+)?)*?){1,5})"
            r"\s*=\s*(?P<rhs>\d+(?:\.\d+)?)\s*(?:m|s|元|个)?",
            field,
        ):
            tokens = re.split(
                r"\s*([+\-])\s*",
                arithmetic.group("lhs").replace("＋", "+").replace("−", "-"),
            )

            def product_value(term: str) -> float:
                product = 1.0
                for factor in re.split(r"\s*[×xX*]\s*", term):
                    product *= float(factor)
                return product

            computed = product_value(tokens[0])
            for operator, term in zip(tokens[1::2], tokens[2::2], strict=True):
                term_value = product_value(term)
                computed = (
                    computed + term_value
                    if operator == "+"
                    else computed - term_value
                )
            claimed = float(arithmetic.group("rhs"))
            if not math.isclose(computed, claimed, rel_tol=1e-9, abs_tol=0.005):
                raise QualityGateError(
                    "建模方案中的算术等式不成立："
                    f"{arithmetic.group('lhs')}={computed:g}，不是 {claimed:g}"
                )

        if body_count_match:
            body_count = int(body_count_match.group(1))
            expected_tail_front = body_count + 2
            expected_tail_rear = body_count + 3
            for named_body in re.finditer(
                r"第\s*(?P<body>\d+)\s*(?:节|条)?龙身(?:前|后)?(?:把手)?",
                field,
                re.IGNORECASE,
            ):
                body_number = int(named_body.group("body"))
                if body_number > body_count:
                    raise QualityGateError(
                        "龙身节号超出题面范围："
                        f"题面只有 {body_count} 节龙身，不能引用第 {body_number} 节龙身；"
                        "龙尾前把手必须按独立节点列出"
                    )
            explicit_body_mappings = [
                (int(match.group("node")), int(match.group("body")))
                for pattern in (
                    r"节点\s*(?P<node>\d+)\s*(?:=|为)\s*"
                    r"第\s*(?P<body>\d+)\s*节龙身前把手",
                    r"第\s*(?P<body>\d+)\s*节龙身前把手\s*"
                    r"(?:=|为)\s*节点\s*(?P<node>\d+)",
                )
                for match in re.finditer(pattern, field, re.IGNORECASE)
            ]
            for node_index, body_index in explicit_body_mappings:
                expected_body_node = body_index + 1
                if node_index != expected_body_node:
                    raise QualityGateError(
                        "显式龙身节点映射与串联拓扑矛盾："
                        f"第 {body_index} 节龙身前把手应为节点 "
                        f"{expected_body_node}，而非节点 {node_index}"
                    )
            for mapping in re.finditer(
                r"[ki]\s*=\s*(?P<lower>\d+)\s*(?:\.\.|…|至|[-−])\s*"
                r"(?P<upper>\d+)\s*(?:为|=)\s*第\s*[ki]\s*节龙身(?:的)?前把手",
                field,
                re.IGNORECASE,
            ):
                upper = int(mapping.group("upper"))
                if upper > body_count:
                    raise QualityGateError(
                        "龙身把手索引超出题面实体数："
                        f"题面只有 {body_count} 节龙身，不能把 k 扩展到 {upper}；"
                        "龙尾前、后把手必须作为独立节点列出"
                    )
            for pattern, expected_node, label in (
                (
                    r"节点\s*(?P<node>\d+)\s*(?:=|为)"
                    r"[^，,。；\n]{0,40}龙尾前把手",
                    expected_tail_front,
                    "龙尾前把手",
                ),
                (
                    r"龙尾前把手[^，,。；\n]{0,24}(?:=|为)?\s*节点\s*"
                    r"(?P<node>\d+)",
                    expected_tail_front,
                    "龙尾前把手",
                ),
                (
                    r"节点\s*(?P<node>\d+)\s*(?:=|为)"
                    r"[^，,。；\n]{0,40}龙尾后把手",
                    expected_tail_rear,
                    "龙尾后把手",
                ),
                (
                    r"龙尾后把手[^，,。；\n]{0,24}(?:=|为|是)\s*节点\s*"
                    r"(?P<node>\d+)",
                    expected_tail_rear,
                    "龙尾后把手",
                ),
            ):
                for mapping in re.finditer(pattern, field, re.IGNORECASE):
                    actual_node = int(mapping.group("node"))
                    if actual_node != expected_node:
                        raise QualityGateError(
                            f"显式{label}节点映射与串联拓扑矛盾："
                            f"应为节点 {expected_node}，而非节点 {actual_node}"
                        )
            equates_tail_front_with_last_body_front = re.search(
                rf"龙尾前把手[^。；\n]{{0,20}}(?:即|就是|等同于?|与)"
                rf"[^。；\n]{{0,20}}第\s*{body_count}\s*节龙身(?:的)?前把手|"
                rf"第\s*{body_count}\s*节龙身(?:的)?前把手[^。；\n]{{0,20}}"
                r"(?:即|就是|等同于?|与)[^。；\n]{0,20}龙尾前把手",
                field,
                re.IGNORECASE,
            )
            if equates_tail_front_with_last_body_front:
                raise QualityGateError(
                    "龙尾前把手不能等同于最后一节龙身前把手："
                    f"第 {body_count} 节龙身前把手是节点 {body_count + 1}，"
                    f"龙尾前把手是下一节点 {body_count + 2}"
                )
            for mapping in re.finditer(
                r"节点\s*(?P<node_lower>\d+)\s*(?:~|～|至|[-−])\s*"
                r"(?P<node_upper>\d+)[^。；\n]{0,24}(?:依次)?(?:为|=)?\s*"
                r"第\s*(?P<body_lower>\d+)\s*(?:~|～|至|[-−])\s*"
                r"(?P<body_upper>\d+)\s*节龙身(?:的)?前把手",
                field,
                re.IGNORECASE,
            ):
                node_lower = int(mapping.group("node_lower"))
                node_upper = int(mapping.group("node_upper"))
                body_lower = int(mapping.group("body_lower"))
                body_upper = int(mapping.group("body_upper"))
                if (
                    node_lower != body_lower + 1
                    or node_upper != body_upper + 1
                    or node_upper - node_lower != body_upper - body_lower
                    or body_lower != 1
                    or body_upper > body_count
                ):
                    raise QualityGateError(
                        "显式龙身节点区间与串联拓扑矛盾："
                        f"第 1 至 {body_count} 节龙身前把手只能对应节点 "
                        f"2 至 {body_count + 1}"
                    )
            wrong_tail_front_node = body_count + 1
            if re.search(
                rf"节点\s*{wrong_tail_front_node}\s*(?:=|为)"
                r"[^，,。；\n]{0,55}龙尾前把手",
                field,
                re.IGNORECASE,
            ):
                raise QualityGateError(
                    "节点身份映射把最后一节龙身的前把手误作龙尾前把手："
                    f"节点 {wrong_tail_front_node} 是第 {body_count} 节龙身前把手，"
                    f"龙尾前把手应为节点 {body_count + 2}"
                )
            if re.search(
                r"龙尾前把手[^，,。；\n]{0,24}(?:=|为)?\s*节点\s*"
                rf"{wrong_tail_front_node}(?!\d)",
                field,
                re.IGNORECASE,
            ):
                raise QualityGateError(
                    "节点身份映射把最后一节龙身的前把手误作龙尾前把手："
                    f"节点 {wrong_tail_front_node} 是第 {body_count} 节龙身前把手，"
                    f"龙尾前把手应为节点 {body_count + 2}"
                )
            for mapping in re.finditer(
                r"节点\s*k(?:\s*[（(][^）)\n]{0,30}[）)])?\s*(?:为|=)\s*"
                r"第\s*k\s*[-−]\s*1\s*节(?:龙身)?前把手",
                field,
                re.IGNORECASE,
            ):
                window = field[
                    max(0, mapping.start() - 40) : min(len(field), mapping.end() + 45)
                ]
                bounded_range = re.search(
                    r"(?P<lower>\d+)\s*[≤<]=?\s*k\s*[≤<]=?\s*(?P<upper>\d+)"
                    r"|k\s*(?:=|∈)\s*(?P<lower_alt>\d+)\s*"
                    r"(?:\.\.|…|至|[-−])\s*(?P<upper_alt>\d+)"
                    r"|k\s*(?:=|∈)\s*[\[(]\s*(?P<lower_bracket>\d+)\s*"
                    r"[,，]\s*(?P<upper_bracket>\d+)\s*[\])]",
                    window,
                    re.IGNORECASE,
                )
                if bounded_range is None:
                    raise QualityGateError(
                        "节点身份映射缺少龙身区间：`节点 k=第 k-1 节前把手` "
                        f"只能限定在 2≤k≤{body_count + 1}；龙头前把手、龙尾前把手和"
                        "龙尾后把手必须单列，不能把无界 k 公式套到链首尾"
                    )
                lower = int(
                    bounded_range.group("lower")
                    or bounded_range.group("lower_alt")
                    or bounded_range.group("lower_bracket")
                )
                upper = int(
                    bounded_range.group("upper")
                    or bounded_range.group("upper_alt")
                    or bounded_range.group("upper_bracket")
                )
                if lower != 2 or upper > body_count + 1:
                    raise QualityGateError(
                        "节点身份区间把龙尾前把手误标为龙身："
                        f"题面只有 {body_count} 节龙身，龙身前把手节点范围最多到 "
                        f"{body_count + 1}，其后须单列龙尾前把手"
                    )

        for decomposition in re.finditer(
            r"(?P<total>\d+)\s*个?节点\s*[（(](?P<terms>[^）)\n]{0,100})[）)]",
            field,
        ):
            counts = [
                int(match.group(1))
                for match in re.finditer(
                    r"(\d+)\s*个?(?:龙头|龙身|龙尾)",
                    decomposition.group("terms"),
                )
            ]
            if len(counts) >= 2 and sum(counts) != int(decomposition.group("total")):
                raise QualityGateError(
                    "节点拓扑分解与声明总数不一致："
                    f"分项合计 {sum(counts)}，声明 {decomposition.group('total')}；"
                    "须逐项列全龙头、龙身、龙尾的前后把手身份"
                )

        for segment in re.split(r"[。；;\n]+", field):
            treats_adjacent_as_collision = re.search(
                r"(?<!不)(?<!非)相邻(?:两)?(?:节)?板凳[^。；;\n]{0,35}"
                r"(?:碰撞|干涉|相交)|"
                r"(?:碰撞|干涉|相交)[^。；;\n]{0,35}"
                r"(?<!不)(?<!非)相邻(?:两)?(?:节)?板凳",
                segment,
            )
            allows_shared_connection = re.search(
                r"(?:允许|不计|不算(?:作|为)?|排除|不视为)[^。；;\n]{0,24}"
                r"(?:相邻|共享把手|装配接触)|"
                r"(?:相邻|共享把手|装配接触)[^。；;\n]{0,24}"
                r"(?:允许|不计|不算(?:作|为)?|排除|不视为)",
                segment,
            )
            if treats_adjacent_as_collision and not allows_shared_connection:
                raise QualityGateError(
                    "碰撞定义把相邻板凳的共享把手装配关系直接计为干涉；"
                    "须排除允许的相邻连接接触，并对非相邻实体轮廓判定首次碰撞"
                )

    def collect_question_text(item: Any, question_number: int) -> list[str]:
        scoped: list[str] = []

        def flatten(child: Any) -> None:
            if isinstance(child, str):
                scoped.append(child)
            elif isinstance(child, dict):
                for nested in child.values():
                    flatten(nested)
            elif isinstance(child, (list, tuple)):
                for nested in child:
                    flatten(nested)

        def visit(child: Any) -> None:
            if not isinstance(child, dict):
                return
            subtask = str(child.get("subtask", "")).strip().lower()
            if subtask in {
                f"ques{question_number}",
                f"q{question_number}",
                f"problem{question_number}",
                f"问题{question_number}",
            }:
                flatten(child)
                return
            key_pattern = re.compile(
                rf"^(?:ques|q|problem|问题)\s*{question_number}$",
                re.IGNORECASE,
            )
            for key, nested in child.items():
                if key_pattern.match(str(key).strip()):
                    flatten(nested)
                else:
                    visit(nested)

        visit(item)
        return scoped

    patterns = (
        r"(?:conn(?:ection)?(?:[_\s-]*(?:head|body|distance|length))?|"
        r"连接(?:点|距离|长度|间距)?)"
        r"[^。；;\n]{0,45}(?:一半之和|半长之和|"
        r"\([^()]{0,30}\+[^()]{0,30}\)\s*/\s*2)",
        r"(?:沿[^。；,，、]{0,20})?弧长[^。；,，、]{0,30}"
        r"(?:恒等于|等于|=|近似(?:等于|用)?)\s*"
        r"(?:对应|两点|相邻)?\s*(?:刚性(?:构件)?弦长|弦长|直线距离|"
        r"节点距离|节点距|孔距|把手间距)",
        r"弧长(?:坐标)?(?:差|间隔)?[^。；,，、]{0,18}(?:=|取为)\s*"
        r"(?:2\.86|1\.65|孔距|弦长)",
        r"(?:测地|路径)距离(?:约束)?[^。；]{0,80}"
        r"s\s*\([^)]*\)\s*-\s*s\s*\([^)]*\)\s*=\s*(?:d|[12]\.(?:65|86))",
        r"(?:√|sqrt)\s*\([^。；\n]{0,80}\)\s*(?:Δ\s*θ|d\s*θ)\s*=\s*"
        r"(?:节点(?:间)?距(?:离)?|孔距|[12]\.(?:65|86))",
        r"\|\s*r\s*\([^)]*\)\s*-\s*r\s*\([^)]*\)\s*\|\s*=\s*"
        r"(?:L|d|[12]\.(?:65|86))",
        r"(?:2\.86[^。；]{0,30}1\.65|1\.65[^。；]{0,30}2\.86)"
        r"[^。；]{0,20}(?:交替|轮换)",
        r"(?:所有|全部|任意|各)(?:[^。；]{0,20})?把手"
        r"(?:[^。；]{0,25})?(?:速度|速率)(?:大小|模长|模)?"
        r"(?:[^。；]{0,12})?(?:相同|相等|均为|恒为)(?:为|是)?"
        r"\s*1(?:\.0+)?\s*m/s",
        r"(?:所有|全部|任意|各)(?:[^。；]{0,20})?把手"
        r"(?:[^。；]{0,25})?(?:等速|同速)",
        r"(?:全链|整条链|所有节点|全部节点|各节点)"
        r"[^。；]{0,25}(?:速度|速率)(?:大小|模长|模)?\s*"
        r"(?:均|都|全部)?(?<!不)(?<!非)(?:相同|相等|均为|恒为|"
        r"等于(?:龙头|先导节点))"
        r"[^。；]{0,16}(?:1(?:\.0+)?\s*m/s|(?:龙头|先导节点)(?:速度|速率))?",
    )
    def is_negated_uniform_speed(match: re.Match[str]) -> bool:
        if not re.search(r"速度|速率|等速|同速", match.group(0)):
            return False
        sentence_start = max(
            text.rfind(delimiter, 0, match.start()) for delimiter in "。；;\n"
        )
        prefix = text[sentence_start + 1 : match.start()]
        return bool(
            re.search(
                r"(?:不可|不能|不得|禁止|并非|不是|不应|不宜)"
                r"(?:直接)?(?:令|设|假定|认为|断言)?[^。；;\n]{0,12}$",
                prefix,
            )
        )

    matched_guardrail = next(
        (
            match
            for pattern in patterns
            if (match := re.search(pattern, text, re.IGNORECASE))
            and not is_negated_uniform_speed(match)
        ),
        None,
    )
    if matched_guardrail:
        raise QualityGateError(
            "串联刚性构件证据使用跨构件半长平均、违反固定弦长/弧长关系，"
            "或无依据断言所有节点等速；命中片段：" + matched_guardrail.group(0)[:160]
        )
    expected_node_count = int(guardrails["node_count"])
    structured_counts: list[int] = []

    def collect_structured_counts(item: Any) -> None:
        if isinstance(item, dict):
            for key, child in item.items():
                normalized_key = str(key).strip().lower()
                if (
                    normalized_key
                    in {
                        "node_count",
                        "handle_count",
                        "hub_count",
                        "connection_node_count",
                    }
                    and isinstance(child, (int, float))
                    and not isinstance(child, bool)
                ):
                    if math.isfinite(float(child)) and float(child).is_integer():
                        structured_counts.append(int(child))
                collect_structured_counts(child)
        elif isinstance(item, (list, tuple)):
            for child in item:
                collect_structured_counts(child)

    collect_structured_counts(value)
    declared_counts: set[int] = set()
    for segment in re.split(r"[。；;\n]+", text):
        if re.search(r"论文|表\s*[12]|关键节点|指定节点", segment):
            continue
        declared_counts.update(
            int(match.group(1))
            for match in re.finditer(
                r"(?:共|全部|总计|得到|得)(?:有)?\s*(\d+)\s*个?"
                r"(?:不同)?(?:把手(?:中心|位置)?|连接点|节点)",
                segment,
            )
        )
        declared_counts.update(
            int(match.group(1))
            for match in re.finditer(
                r"(?<![\d.])(\d+)\s*个?节点(?:拓扑)?链",
                segment,
            )
        )
        declared_counts.update(
            int(match.group(1))
            for match in re.finditer(
                r"(?:每(?:个|一)(?:时刻|时间点)|每帧|每行)"
                r"[^。；;\n]{0,20}?(\d+)\s*个?对象",
                segment,
            )
        )
    declared_counts.update(structured_counts)
    declared_counts.update(
        int(match.group("count"))
        for match in re.finditer(
            r"\d+\s*[×xX*]\s*(?P<count>\d+)\s*"
            r"(?:个)?(?:把手|节点|位置|速度|对象)",
            text,
            re.IGNORECASE,
        )
    )
    tail_rear_ordinals = {
        int(match.group("count"))
        for match in re.finditer(
            r"龙尾后把手[^。；\n]{0,35}第\s*(?P<count>\d+)\s*个?"
            r"(?:待)?输出对象",
            text,
            re.IGNORECASE,
        )
    }
    declared_counts.update(tail_rear_ordinals)
    if re.search(
        rf"{expected_node_count}\s*个?节点\s*[+＋]\s*(?:1\s*个?)?龙尾后把手",
        text,
        re.IGNORECASE,
    ):
        declared_counts.add(expected_node_count + 1)
    wrong_counts = sorted(declared_counts - {expected_node_count})
    if wrong_counts:
        raise QualityGateError(
            "串联刚性构件证据同时声明了与题面拓扑矛盾的把手/节点总数: "
            f"{wrong_counts}；题面可直接推导为 {expected_node_count}"
        )
    expected_entity_count = int(guardrails["entity_count"])
    declared_entity_counts = {
        int(match.group(1))
        for match in re.finditer(
            r"(?:共|全部|总计|全链)?\s*(\d+)\s*个?"
            r"(?:有向)?(?:矩形板体|矩形实体|矩形|板体|板凳实体)",
            text,
        )
    }
    wrong_entity_counts = sorted(declared_entity_counts - {expected_entity_count})
    if wrong_entity_counts:
        raise QualityGateError(
            "碰撞实体总数与题面构件拓扑矛盾: "
            f"{wrong_entity_counts}；{expected_entity_count} 节板凳应形成 "
            f"{expected_entity_count} 个有限尺寸实体"
        )
    for pair_count in re.finditer(r"(\d+)\s*[×xX*·]\s*(\d+)\s*/\s*2\s*对", text):
        left, right = map(int, pair_count.groups())
        if (left * right) % 2:
            raise QualityGateError(
                "碰撞候选对数量公式得到半整数，说明实体数或排除相邻对的计数错误；"
                f"{left}×{right}/2 不是整数"
            )
    for field in text_fields:
        first_index = re.search(r"龙头前把手(?:为|是|[=＝])节点\s*(\d+)", field)
        last_index = re.search(r"龙尾后把手(?:为|是|[=＝])节点\s*(\d+)", field)
        if first_index and last_index:
            expected_last_index = int(first_index.group(1)) + int(
                guardrails["entity_count"]
            )
            if int(last_index.group(1)) != expected_last_index:
                raise QualityGateError(
                    "串联节点编号端点不闭合："
                    f"从 {first_index.group(1)} 开始时，龙尾后把手应编号 "
                    f"{expected_last_index}，而非 {last_index.group(1)}"
                )
    expected_remaining_segments = int(guardrails["entity_count"]) - 1
    declared_remaining_segments = {
        int(match.group(1))
        for match in re.finditer(r"(?:后续|其余)\s*(\d+)\s*段", text)
    }
    wrong_remaining_segments = sorted(
        declared_remaining_segments - {expected_remaining_segments}
    )
    if wrong_remaining_segments:
        raise QualityGateError(
            "串联构件的异质首段/统一后续段数量与题面拓扑矛盾: "
            f"{wrong_remaining_segments}；首段之外应有 {expected_remaining_segments} 段"
        )
    repeated_distance = float(guardrails["member_constraints"][1]["node_distance_m"])
    expected_last_segment_index = int(guardrails["entity_count"])
    for segment in re.split(r"[。；;\n]+", text):
        normalized_segment = segment.translate(
            str.maketrans("₀₁₂₃₄₅₆₇₈₉", "0123456789")
        )
        if not re.search(
            rf"(?<![\d.]){re.escape(f'{repeated_distance:.2f}')}(?:0*)\s*m?",
            normalized_segment,
        ):
            continue
        for match in re.finditer(
            r"(?:i|I)\s*=\s*(\d+)\s*(?:\.\.|…|~|～|至|到|-)\s*(\d+)",
            normalized_segment,
        ):
            start_index, end_index = map(int, match.groups())
            uses_forward_pair = re.search(
                r"(?:P|节点)\s*_?\{?i\}?[^。；\n]{0,20}"
                r"(?:P|节点)\s*_?\{?i\s*\+\s*1\}?|"
                r"(?:P|节点)\s*_?\{?i\s*\+\s*1\}?[^。；\n]{0,20}"
                r"(?:P|节点)\s*_?\{?i\}?|"
                r"第\s*i\s*节点[^。；\n]{0,20}第\s*i\s*\+\s*1\s*节点|"
                r"第\s*i\s*\+\s*1\s*节点[^。；\n]{0,20}第\s*i\s*节点",
                segment,
                re.IGNORECASE,
            )
            if (
                uses_forward_pair
                and start_index == 2
                and end_index != expected_last_segment_index
            ):
                raise QualityGateError(
                    "串联构件统一后续段索引未覆盖完整拓扑："
                    f"i=2..{end_index}，应为 i=2..{expected_last_segment_index}"
                )
        uses_backward_pair = re.search(
            r"(?:P|节点)\s*_?\{?i\}?[^。；\n]{0,24}"
            r"(?:P|节点)\s*_?\{?i\s*[-−]\s*1\}?|"
            r"(?:P|节点)\s*_?\{?i\s*[-−]\s*1\}?[^。；\n]{0,24}"
            r"(?:P|节点)\s*_?\{?i\}?",
            normalized_segment,
            re.IGNORECASE,
        )
        shifted_backward_lengths = re.search(
            r"(?:d|L)\s*_?\{?1\}?\s*=\s*2\.86\s*m?[^。；\n]{0,45}"
            r"(?:d|L)\s*_?\{?i\}?\s*=\s*1\.65\s*m?[^。；\n]{0,25}"
            r"i\s*(?:≥|>=)\s*2|"
            r"(?:d|L)\s*_?\{?1\}?\s*=\s*2\.86\s*m?[^。；\n]{0,45}"
            r"i\s*(?:≥|>=)\s*2[^。；\n]{0,25}"
            r"(?:d|L)\s*_?\{?i\}?\s*=\s*1\.65\s*m?",
            normalized_segment,
            re.IGNORECASE,
        )
        if uses_backward_pair and shifted_backward_lengths:
            raise QualityGateError(
                "节点距离下标与端点公式错位：使用 P_i-P_{i-1}=d_i 且 "
                "i=2..224 时，首段应为 d_2=2.86m、统一后续段应从 i=3 开始；"
                "若保留 d_1=2.86m，则应改用 P_{i+1}-P_i=d_i、i=1..223"
            )
        shifted_explicit_backward_ranges = re.search(
            r"(?:d|L)\s*_?\{?1\}?\s*=\s*2\.86\s*m?[^。；\n]{0,55}"
            r"(?:d|L)\s*_?\{?2\}?\s*(?:\.\.|…|~|～|至|到|-)[^。；\n]{0,8}"
            rf"(?:d|L)?\s*_?\{{?(?:{expected_last_segment_index}|{expected_node_count})\}}?"
            r"\s*=\s*1\.65\s*m?",
            normalized_segment,
            re.IGNORECASE,
        )
        if uses_backward_pair and shifted_explicit_backward_ranges:
            raise QualityGateError(
                "节点距离显式下标与后向端点公式整体错移：若 d_i=|P_i-P_{i-1}|、"
                f"i=2..{expected_node_count}，则 d_2=2.86m，"
                f"d_3..d_{expected_node_count}=1.65m；不能写成 d_1 与"
                f" d_2..d_{expected_last_segment_index}"
            )
    normalized_text = text.translate(str.maketrans("₀₁₂₃₄₅₆₇₈₉", "0123456789"))
    uses_backward_distance_definition = re.search(
        r"(?:P|节点)\s*_?\{?i\}?[^。\n]{0,30}"
        r"(?:P|节点)\s*_?\{?i\s*[-−]\s*1\}?|"
        r"(?:P|节点)\s*_?\{?i\s*[-−]\s*1\}?[^。\n]{0,30}"
        r"(?:P|节点)\s*_?\{?i\}?",
        normalized_text,
        re.IGNORECASE,
    )
    uses_forward_distance_definition = re.search(
        r"(?:P|节点)\s*_?\{?i\}?[^。\n]{0,30}"
        r"(?:P|节点)\s*_?\{?i\s*\+\s*1\}?|"
        r"(?:P|节点)\s*_?\{?i\s*\+\s*1\}?[^。\n]{0,30}"
        r"(?:P|节点)\s*_?\{?i\}?",
        normalized_text,
        re.IGNORECASE,
    )
    uses_shifted_explicit_forward_range = re.search(
        r"(?:d|L)\s*_?\{?2\}?\s*=\s*2\.86\s*m?[^。\n]{0,70}"
        r"(?:d|L)\s*_?\{?3\}?\s*(?:\.\.|…|~|～|至|到|-)[^。\n]{0,8}"
        rf"(?:d|L)?\s*_?\{{?{expected_node_count}\}}?\s*=\s*1\.65\s*m?",
        normalized_text,
        re.IGNORECASE,
    )
    if uses_forward_distance_definition and uses_shifted_explicit_forward_range:
        raise QualityGateError(
            "节点距离显式下标与前向端点公式整体错移：若 d_i=|P_{i+1}-P_i|、"
            f"i=1..{expected_last_segment_index}，则 d_1=2.86m，"
            f"d_2..d_{expected_last_segment_index}=1.65m；若使用 d_2 与"
            f" d_3..d_{expected_node_count}，端点公式必须改为 P_i-P_{{i-1}}=d_i"
        )
    uses_shifted_explicit_distance_range = re.search(
        r"(?:d|L)\s*_?\{?1\}?\s*=\s*2\.86\s*m?[^。\n]{0,70}"
        r"(?:d|L)\s*_?\{?2\}?\s*(?:\.\.|…|~|～|至|到|-)[^。\n]{0,8}"
        rf"(?:d|L)?\s*_?\{{?{expected_last_segment_index}\}}?\s*=\s*1\.65\s*m?",
        normalized_text,
        re.IGNORECASE,
    )
    if uses_backward_distance_definition and uses_shifted_explicit_distance_range:
        raise QualityGateError(
            "节点距离显式下标与后向端点公式整体错移：若 d_i=|P_i-P_{i-1}|、"
            f"i=2..{expected_node_count}，则 d_2=2.86m，"
            f"d_3..d_{expected_node_count}=1.65m；不能写成 d_1 与"
            f" d_2..d_{expected_last_segment_index}"
        )
    total_rigid_segment_patterns = (
        r"(?:构成|总计)\s*(\d+)\s*个?(?:刚性|节点间|连接)?段",
        r"共有\s*(\d+)\s*个?(?:刚性|节点间|连接)段",
        r"(?:板内)?刚性段(?:的)?总数\s*(?:为|是|=|：|:)?\s*(\d+)\s*段",
        r"(?:验证|核对|复核)?\s*总段数\s*(?:为|是|=|：|:)?\s*"
        r"(?:\d+\s*[+＋]\s*\d+\s*=\s*)?(\d+)\s*段",
        rf"{expected_node_count}\s*个?节点(?:间|之间).{{0,16}}共\s*(\d+)\s*段",
        r"共\s*(\d+)\s*段\s*[:：]\s*1\s*段[^。；\n]{0,40}"
        r"(?:2\.86|龙头)",
    )
    declared_total_rigid_segments = {
        int(match.group(1))
        for pattern in total_rigid_segment_patterns
        for match in re.finditer(pattern, text)
    }
    describes_member_segment_lengths = all(
        re.search(rf"(?<![\d.]){re.escape(f'{distance:.2f}')}\s*m?", text)
        for distance in (
            float(item["node_distance_m"])
            for item in guardrails["member_constraints"][:2]
        )
    )
    wrong_total_rigid_segments = sorted(
        declared_total_rigid_segments - {int(guardrails["entity_count"])}
    )
    if describes_member_segment_lengths and wrong_total_rigid_segments:
        raise QualityGateError(
            "串联刚性段总数与节点拓扑矛盾: "
            f"{wrong_total_rigid_segments}；{expected_node_count} 个节点之间应有 "
            f"{int(guardrails['entity_count'])} 个板内刚性段"
        )

    member_distances = [
        float(item["node_distance_m"]) for item in guardrails["member_constraints"]
    ]
    if len(member_distances) >= 2:
        first_distance = re.escape(f"{member_distances[0]:.2f}")
        repeated_distance = re.escape(f"{member_distances[1]:.2f}")
        repeated_counts = {
            int(match.group(1) or match.group(2))
            for match in re.finditer(
                rf"{first_distance}\s*\+\s*(?:(\d+)\s*[×xX*]\s*"
                rf"{repeated_distance}|{repeated_distance}\s*[×xX*]\s*(\d+))",
                text,
            )
        }
        wrong_formula_counts = sorted(repeated_counts - {expected_remaining_segments})
        if wrong_formula_counts:
            raise QualityGateError(
                "串联构件总长度公式中的统一后续段乘数与题面拓扑矛盾: "
                f"{wrong_formula_counts}；首段之外应乘 {expected_remaining_segments}，"
                "因为龙身构件数之外还须计入龙尾自身的一段"
            )

    initial_spiral = guardrails.get("initial_spiral")
    if initial_spiral:
        turn_count = int(initial_spiral["turn_count"])
        pitch_m = float(initial_spiral["pitch_m"])
        expected_radius_m = float(initial_spiral["radius_m"])
        wrong_turn_formula = re.compile(
            rf"{turn_count}\s*[×xX*]\s*(?:p|{pitch_m:g})\s*/\s*"
            r"\(?\s*2\s*(?:π|pi)\s*\)?",
            re.IGNORECASE,
        )
        for field in text_fields:
            for segment in re.split(r"[。；;\n]+", field):
                if wrong_turn_formula.search(segment):
                    raise QualityGateError(
                        "等距螺线圈数换算重复除以 2π：题面的第 N 圈已经表示 "
                        "θ/(2π)=N，因此该圈半径应为 N×螺距"
                    )
                radius_context = re.search(
                    rf"第\s*{turn_count}\s*圈[^。；;\n]{{0,50}}?"
                    r"(?:半径|极径)(?P<value>[^。；;\n]*)",
                    segment,
                )
                if not radius_context:
                    continue
                radius_prefix = segment[
                    radius_context.start() : radius_context.start("value")
                ]
                if re.search(r"调头|转向|边界|区域", radius_prefix):
                    continue
                radius_text = radius_context.group("value")
                radius_text = re.split(r"(?:半径|极径)", radius_text, maxsplit=1)[0]
                radius_text = re.split(r"[,，、)）]", radius_text, maxsplit=1)[0]
                declared_metres = [
                    float(number)
                    for number in re.findall(
                        r"(?<![\d.])(\d+(?:\.\d+)?)\s*m(?![a-zA-Z])",
                        radius_text,
                        re.IGNORECASE,
                    )
                ]
                if declared_metres and not math.isclose(
                    declared_metres[-1], expected_radius_m, abs_tol=0.01
                ):
                    raise QualityGateError(
                        "题面初始圈数与声明半径不一致："
                        f"第 {turn_count} 圈、螺距 {pitch_m:g} m 对应半径 "
                        f"{expected_radius_m:g} m，而非 {declared_metres[-1]:g} m"
                    )

        source_ques1 = re.search(
            r"问题\s*1(?P<body>.*?)(?=问题\s*2|$)", source_text, re.DOTALL
        )
        if source_ques1:
            source_ques1_text = source_ques1.group("body")
            duration_match = re.search(
                r"(?:到|至)\s*(\d+(?:\.\d+)?)\s*s\s*为止",
                source_ques1_text,
                re.IGNORECASE,
            )
            speed_match = re.search(
                r"龙头前把手[^。；\n]{0,60}速度[^。；\n]{0,25}"
                r"(\d+(?:\.\d+)?)\s*m\s*/\s*s",
                source_ques1_text,
                re.IGNORECASE,
            )
            if duration_match and speed_match:
                duration_s = float(duration_match.group(1))
                speed_mps = float(speed_match.group(1))
                travel_m = duration_s * speed_mps
                theta_initial = 2 * math.pi * turn_count
                spiral_scale = pitch_m / (2 * math.pi)

                def spiral_arc_primitive(theta: float) -> float:
                    return (
                        0.5
                        * spiral_scale
                        * (theta * math.sqrt(theta * theta + 1) + math.asinh(theta))
                    )

                available_arc = spiral_arc_primitive(theta_initial)
                if 0 < travel_m < available_arc:
                    target_arc = available_arc - travel_m
                    lower, upper = 0.0, theta_initial
                    for _ in range(80):
                        midpoint = (lower + upper) / 2
                        if spiral_arc_primitive(midpoint) < target_arc:
                            lower = midpoint
                        else:
                            upper = midpoint
                    expected_turns = (theta_initial - (lower + upper) / 2) / (
                        2 * math.pi
                    )
                    ques1_text = " ".join(collect_question_text(value, 1))
                    for segment in re.split(r"[。；;\n]+", ques1_text):
                        if not (
                            re.search(rf"{duration_s:g}\s*s", segment, re.IGNORECASE)
                            or re.search(rf"{travel_m:g}\s*m", segment, re.IGNORECASE)
                        ):
                            continue
                        claimed_turns = re.search(
                            r"(?:对应|累计|行进|盘入)[^。；\n]{0,30}?"
                            r"(?:约|大约)?\s*(\d+(?:\.\d+)?)\s*圈",
                            segment,
                        )
                        if claimed_turns and not math.isclose(
                            float(claimed_turns.group(1)),
                            expected_turns,
                            rel_tol=0.1,
                            abs_tol=0.5,
                        ):
                            raise QualityGateError(
                                "等距螺线定时运动的行进圈数与弧长积分矛盾："
                                f"{duration_s:g} s、{speed_mps:g} m/s 对应约 "
                                f"{expected_turns:.2f} 圈，而非 "
                                f"{float(claimed_turns.group(1)):g} 圈"
                            )

    uses_outward_increasing_theta = re.search(
        r"r\s*=\s*(?!-)[^,，。；;]{0,60}(?:θ|theta)", text, re.IGNORECASE
    )
    if (
        "等距螺线" in source_text
        and "盘入" in source_text
        and uses_outward_increasing_theta
    ):
        inward_increases_theta = (
            r"(?:(?<!\{)龙头(?:前把手)?(?!的?(?:外侧|内侧))|前把手|盘入)"
            r"(?:(?!(?:后续|后方|向后|向外|沿链|沿\s*(?:θ|theta)|龙身|龙尾|盘出|递减|减小)).){0,100}"
            r"(?:极角|角度|θ|theta)"
            r"(?:(?!(?:后续|后方|向后|向外|沿链|沿\s*(?:θ|theta)|龙身|龙尾|盘出|递减|减小)).){0,25}"
            r"(?:增大|递增|增加|单调上升)"
            r"(?![^。；;\n]{0,12}盘出)",
            r"(?:极角|角度|θ|theta)"
            r"(?:(?!(?:后续|后方|向后|向外|沿链|沿\s*(?:θ|theta)|龙身|龙尾|盘出|递减|减小)).){0,25}"
            r"(?:随时间)?(?:单调)?(?:增大|递增|增加)"
            r"(?:(?!(?:后续|后方|向后|向外|沿链|沿\s*(?:θ|theta)|龙身|龙尾|盘出|递减|减小)).){0,100}"
            r"(?:龙头|前把手|盘入)",
        )
        trailing_decreases_theta = (
            r"(?:θ|theta)_?\{?\s*[iIkK]\s*\+\s*1\s*\}?\s*<\s*"
            r"(?:θ|theta)_?\{?\s*[iIkK]\s*\}?",
            r"(?:后续|后方|其余|从节点\s*2|第\s*[ki]\s*个节点)"
            r".{0,120}(?:较小\s*(?:极角|θ)|"
            r"(?:θ|theta)[_a-z0-9]*\s*<\s*(?:θ|theta))",
            r"(?:沿螺线)?向(?:减小|较小)\s*(?:r|极径|极角|θ|theta)"
            r".{0,80}(?:搜索|求根|后续|后方|节点)",
            r"(?:龙头后把手|后续.{0,25}(?:节点|把手)|后方.{0,25}(?:节点|把手))"
            r".{0,100}(?:位于|沿|向).{0,20}(?:半径方向)?内侧",
            r"(?:链尾|龙尾).{0,50}(?:位于|处于|延伸至|伸入)"
            r".{0,20}(?:半径方向)?内侧",
            r"(?:θ|theta)_?\{?[iIkK]\s*\+\s*1\}?\s*=\s*"
            r"(?:θ|theta)_?\{?[iIkK]\}?\s*-\s*(?:Δ\s*(?:θ|theta)|d(?:θ|theta))",
        )
        direction_segments = [
            segment.strip()
            for field in text_fields
            for segment in re.split(r"[。；;\n]+", field)
            if segment.strip()
        ]
        positive_inward_rate = re.search(
            r"d\s*(?:θ|theta)\s*/\s*d\s*t\s*=\s*\+?\s*(?:v|1(?:\.0+)?)"
            r"\s*/",
            text,
            re.IGNORECASE,
        )
        if positive_inward_rate:
            raise QualityGateError(
                "盘入角速度公式符号矛盾：采用 r=cθ（c>0）且龙头向中心盘入时，"
                "正的速率大小 v 对应 dθ/dt<0，不能写成 dθ/dt=+v/...；"
                "命中片段：" + positive_inward_rate.group(0)[:120]
            )
        negative_start_decreases = re.search(
            r"(?:θ|theta)(?:\s*_?\s*0)?\s*=\s*-\s*\d+(?:\.\d+)?\s*(?:π|pi)?"
            r"[^。；;\n]{0,80}(?:继续|随后|随时间)?\s*(?:递减|减小|下降)",
            text,
            re.IGNORECASE,
        )
        if negative_start_decreases:
            raise QualityGateError(
                "盘入参数方向矛盾：采用 r=cθ（c>0）时，从负角初值继续减小会使"
                "半径绝对值向外增大；应使用与正半径一致的参数支路并令 θ 随盘入递减"
            )
        def is_nonasserted_theta_direction(
            segment: str, pattern: str, match: re.Match[str]
        ) -> bool:
            prefix = segment[: match.start()]
            if re.search(
                r"(?:不得|不能|禁止|不应|不可)(?:直接)?"
                r"(?:写成|写为|采用|令|取)?[^。；;\n]{0,12}$",
                prefix,
            ):
                return True
            if re.search(
                r"(?:沿\s*)?(?:θ|theta)[^。；;\n]{0,10}"
                r"(?:增大|递增)[^。；;\n]{0,10}(?:取根|求根|搜索)",
                match.group(0),
                re.IGNORECASE,
            ) and re.search(
                r"(?:后方|后续|其余|链尾)[^。；;\n]{0,45}(?:外侧|向外)|"
                r"(?:外侧|向外)[^。；;\n]{0,45}(?:后方|后续|其余|链尾)",
                segment,
            ):
                return True
            if pattern != trailing_decreases_theta[0]:
                return False
            return bool(
                re.search(r"S\s*\(\s*(?:θ|theta)\s*\(\s*t\s*\)\s*\)", segment, re.I)
                and re.search(r"(?:v\s*[·*×]?\s*t|弧长方程)", segment, re.I)
                and re.fullmatch(pattern, match.group(0), re.I)
            )

        direction_match = next(
            (
                match
                for segment in direction_segments
                for pattern in (*inward_increases_theta, *trailing_decreases_theta)
                if (match := re.search(pattern, segment, re.IGNORECASE))
                and not is_nonasserted_theta_direction(segment, pattern, match)
            ),
            None,
        )
        if direction_match:
            raise QualityGateError(
                "盘入方向或串联跟随次序矛盾：采用 r=cθ（c>0）时，"
                "龙头向中心盘入使 θ 递减，后续把手位于外侧且 θ 应逐点增大；"
                "命中片段：" + direction_match.group(0)[:160]
            )

    if "碰撞" in source_text and re.search(r"板宽|宽.{0,8}(?:cm|m)", source_text):
        point_threshold = re.compile(
            r"(?:节点|把手|中心)[^。；;\n]{0,45}(?:距离|间距)"
            r"[^。；;\n]{0,35}(?:小于|低于|不大于|<=|≤)"
            r"[^。；;\n]{0,18}(?:0\.30|0\.3|板宽)",
            re.IGNORECASE,
        )
        exact_body_collision = re.compile(
            r"(?:(?:有向)?矩形|板凳(?:实体|轮廓|多边形)|板体(?:轮廓|多边形)|"
            r"分离轴|SAT)[^。；;\n]{0,100}(?:相交|重叠|碰撞|精确判定)|"
            r"(?:相交|重叠|碰撞|精确判定)[^。；;\n]{0,100}"
            r"(?:(?:有向)?矩形|板凳(?:实体|轮廓|多边形)|板体(?:轮廓|多边形)|"
            r"分离轴|SAT)",
            re.IGNORECASE,
        )
        coarse_only = re.compile(r"粗筛|预筛|候选对|空间索引", re.IGNORECASE)
        node_distance_symbol = re.search(
            r"L(?:_?i|_?1)[^。；;\n]{0,45}(?:2\.86|1\.65)|"
            r"(?:2\.86|1\.65)[^。；;\n]{0,45}L(?:_?i|_?1)",
            text,
            re.IGNORECASE,
        )
        for field in text_fields:
            if point_threshold.search(field) and not (
                coarse_only.search(field) and exact_body_collision.search(field)
            ):
                raise QualityGateError(
                    "碰撞判据把把手/节点中心距阈值当成板凳实体碰撞："
                    "节点距离只能用于粗筛，最终必须检查具有长度、宽度和朝向的板体相交"
                )
            if node_distance_symbol and re.search(
                r"(?:有向)?矩形[^。；;\n]{0,35}(?:长|长度)\s*(?:为|=)?\s*L_?i|"
                r"(?:板凳|板体)[^。；;\n]{0,25}(?:长|长度)\s*(?:为|=)?\s*L_?i",
                field,
                re.IGNORECASE,
            ):
                raise QualityGateError(
                    "碰撞矩形复用了刚性节点距符号 L_i 作为板体全长；"
                    "节点距离与板体外形长度必须使用不同变量并分别取题面原值"
                )
            if re.search(
                r"(?:碰撞对)?中心距[^。；;\n]{0,30}(?:最小值|接触值)?"
                r"[^。；;\n]{0,20}(?:接近|等于|达到)[^。；;\n]{0,20}"
                r"两(?:板)?半长之和|"
                r"两(?:板)?半长之和[^。；;\n]{0,30}(?:中心距|碰撞)",
                field,
            ):
                raise QualityGateError(
                    "有向矩形接触没有统一的‘中心距等于两板半长之和’阈值；"
                    "中心距只能作保守粗筛，首次接触和最小间隙必须由带朝向的实体轮廓/SAT复核"
                )
            wrong_body_length = re.search(
                r"(?:有向矩形|矩形实体|板凳实体|板体轮廓)[^。；;\n]{0,80}"
                r"(?:长|长度)\s*(?:为|=|取)?\s*(?:2\.86|1\.65)|"
                r"(?:半长轴|长半轴)\s*(?:为|=|取)?\s*"
                r"(?:孔距|d\s*/\s*2|(?:2\.86|1\.65)\s*/\s*2)|"
                r"(?:板凳|板体)(?:长度|长)[^。；;\n]{0,20}"
                r"(?:2\.86|1\.65)",
                field,
                re.IGNORECASE,
            )
            if wrong_body_length:
                source_lengths = [
                    float(item["source_length_m"])
                    for item in guardrails["member_constraints"]
                ]
                raise QualityGateError(
                    "碰撞/干涉实体把孔中心距误作板凳外形长度："
                    f"板体长度应使用题面原长 {source_lengths} m，"
                    "2.86/1.65 m 只用于两孔间刚性距离"
                )
            wrong_bounding_radius = re.search(
                r"(?:包围圆|半对角线)[^。；;\n]{0,100}"
                r"\(?\s*(?:2\.86|1\.65)\s*\)?\s*(?:\^\s*2|²)|"
                r"(?:sqrt|√)\s*\([^。；;\n]{0,20}\(?\s*(?:2\.86|1\.65)\s*\)?"
                r"\s*(?:\^\s*2|²)[^。；;\n]{0,30}\)\s*/\s*2",
                field,
                re.IGNORECASE,
            )
            if wrong_bounding_radius:
                raise QualityGateError(
                    "碰撞粗筛的包围圆也必须使用题面板体原长及板宽计算半对角线；"
                    "2.86/1.65 m 是孔中心距，遗漏两端外伸会漏掉真实碰撞候选"
                )
            wrong_body_length_semantics = re.search(
                r"(?<!不)(?:孔距|两孔(?:中心)?距)[^。；;\n]{0,12}"
                r"(?:决定|作为|用作|定义)[^。；;\n]{0,12}"
                r"(?:矩形|板凳|板体)(?:长度|长)|"
                r"(?:矩形|板凳|板体)(?:长度|长)[^。；;\n]{0,12}"
                r"(?:取|等于|采用)[^。；;\n]{0,8}(?:孔距|两孔(?:中心)?距)",
                field,
            )
            if wrong_body_length_semantics:
                raise QualityGateError(
                    "碰撞实体长度口径错误：孔中心距只定义运动学刚性约束，"
                    "矩形/板体长度必须使用含两端外伸的题面板长"
                )
            if re.search(
                r"(?:矩形|板凳|板体)?中心(?:为|取)[^。；;\n]{0,8}两把手终点", field
            ):
                raise QualityGateError(
                    "碰撞矩形中心定义错误：应取该板两孔/把手坐标的中点，"
                    "不能写成“两把手终点”"
                )
            source_lengths = [
                float(item["source_length_m"])
                for item in guardrails["member_constraints"]
            ]
            for length_claim in re.finditer(
                r"(?:板凳|矩形)[^。；;\n]{0,24}?"
                r"(?<![步弧链])(?:长|长度)\s*(?:为|=)?\s*"
                r"(?P<first>\d+(?:\.\d+)?)"
                r"(?:\s*/\s*(?P<second>\d+(?:\.\d+)?))?\s*(?P<unit>cm|m)",
                field,
                re.IGNORECASE,
            ):
                divisor = 100.0 if length_claim.group("unit").lower() == "cm" else 1.0
                declared_lengths = [
                    float(length_claim.group("first")) / divisor,
                    *(
                        [float(length_claim.group("second")) / divisor]
                        if length_claim.group("second")
                        else []
                    ),
                ]
                wrong_lengths = [
                    length
                    for length in declared_lengths
                    if not any(
                        math.isclose(length, expected, abs_tol=0.005)
                        for expected in source_lengths
                    )
                ]
                if wrong_lengths:
                    raise QualityGateError(
                        "碰撞矩形外形长度与题面原长不一致："
                        f"声明 {wrong_lengths} m，题面板长为 {source_lengths} m"
                    )
        ques2_text = value.get("ques2", "") if isinstance(value, dict) else ""
        member_width_m = guardrails.get("member_width_m")
        if ques2_text and member_width_m:
            maximum_length = max(source_lengths)
            conservative_safe_centre_cutoff = 2 * math.hypot(
                maximum_length / 2, float(member_width_m) / 2
            )
            for segment in re.split(r"[。；;\n]+", ques2_text):
                if not re.search(r"粗筛|预筛|候选对|空间索引", segment):
                    continue
                half_width_cutoff = re.search(
                    r"(?:两板|板体|矩形)[^。；]{0,12}半宽(?:之和)?|"
                    r"半宽(?:之和)?[^。；]{0,12}(?:阈值|粗筛|候选)",
                    segment,
                )
                if half_width_cutoff and re.search(
                    r"(?:半对角线|包围圆半径)", segment, re.IGNORECASE
                ):
                    half_width_cutoff = None
                threshold_match = re.search(
                    r"(?:距离|间距|阈值)[^。；]{0,24}?"
                    r"(?:<|小于|低于|不超过|取|为)?\s*"
                    r"(?P<value>\d+(?:\.\d+)?)\s*(?P<unit>cm|m)",
                    segment,
                    re.IGNORECASE,
                )
                unsafe_numeric_cutoff = False
                declared_cutoff = None
                if threshold_match:
                    declared_cutoff = _to_metres(
                        threshold_match.group("value"), threshold_match.group("unit")
                    )
                    unsafe_numeric_cutoff = (
                        declared_cutoff + 1e-9 < conservative_safe_centre_cutoff
                    )
                if half_width_cutoff or unsafe_numeric_cutoff:
                    detail = (
                        f"声明中心距阈值 {declared_cutoff:g} m，"
                        if declared_cutoff is not None
                        else "声明按两板半宽设置中心距阈值，"
                    )
                    raise QualityGateError(
                        "碰撞候选粗筛不保证无漏检："
                        + detail
                        + "有向矩形中心距的安全包围界至少应使用两板半对角线/"
                        "包围圆半径之和（统一常数粗筛的保守界约为 "
                        f"{conservative_safe_centre_cutoff:.3f} m），或使用不裁剪真实重叠对的空间索引"
                    )
        vague_collision_broad_phase = re.search(
            r"(?:粗筛|预筛|候选对)[^。；\n]{0,45}"
            r"(?:节点包围圆|空间网格)|"
            r"(?:节点包围圆|空间网格)[^。；\n]{0,45}(?:粗筛|预筛|候选对)",
            ques2_text,
            re.IGNORECASE,
        )
        certifies_broad_phase_coverage = re.search(
            r"(?:半对角线|AABB|轴对齐包围盒|板体包围盒|完整覆盖板体|"
            r"不漏检|无漏检|保守包围)",
            ques2_text,
            re.IGNORECASE,
        )
        if vague_collision_broad_phase and not certifies_broad_phase_coverage:
            raise QualityGateError(
                "碰撞粗筛只写节点包围圆/空间网格，未证明候选生成无漏检；"
                "须用题面板体原长和宽度构造半对角线包围圆，或将每个板体的"
                "完整 AABB/保守包围范围插入空间索引后再做 SAT 精判"
            )
        handles_connected_neighbours = re.search(
            r"(?:仅|只)?(?:检查|判定|检测)?[^。；;\n]{0,8}非相邻"
            r"(?:板凳|板体|构件|矩形|对)?|"
            r"\|?\s*i\s*-\s*j\s*\|?\s*>\s*1|"
            r"(?:非相邻|排除|跳过|不计|允许)[^。；;\n]{0,30}"
            r"(?:相邻|共享把手|连接处)|"
            r"(?:相邻|共享把手|连接处|装配接触)[^。；;\n]{0,30}"
            r"(?:排除|跳过|不计|允许(?:接触|重叠|装配)|"
            r"为允许(?:装配)?|不视为碰撞|正常连接)",
            ques2_text,
        )
        if (
            ques2_text
            and guardrails.get("shared_connection_nodes")
            and exact_body_collision.search(ques2_text)
            and not handles_connected_neighbours
        ):
            raise QualityGateError(
                "串联板体碰撞范围未处理连接邻居：相邻构件在共享把手处的"
                "装配接触/重叠必须明确允许或排除，首次碰撞应对非相邻板体"
                "（及题面另有规定的邻接干涉）按真实轮廓判定"
            )
        presupposes_sat_heuristic_order = re.search(
            r"(?:确认|保证|应有|必然)[^。；\n]{0,12}SAT(?:结果|时刻)?"
            r"[^。；\n]{0,16}(?:不早于|不晚于|早于|晚于)"
            r"[^。；\n]{0,16}(?:中心点|中心距)|"
            r"SAT(?:结果|时刻)?[^。；\n]{0,16}(?:不早于|不晚于|早于|晚于)"
            r"[^。；\n]{0,16}(?:中心点|中心距)",
            ques2_text,
            re.IGNORECASE,
        )
        if presupposes_sat_heuristic_order:
            raise QualityGateError(
                "不能预设实体 SAT 与中心点/中心距启发式的碰撞时刻先后；"
                "启发式阈值未证明为充分或必要条件时，两者不存在固定排序，"
                "只能分别计算并报告实际差异"
            )
        presupposes_collision_location = re.search(
            r"碰撞[^。；\n]{0,10}(?:必然|一定|必定)"
            r"[^。；\n]{0,16}(?:发生|出现|位于)|"
            r"(?:必然|一定|必定)[^。；\n]{0,10}(?:发生|出现)"
            r"[^。；\n]{0,18}碰撞",
            ques2_text,
        )
        if presupposes_collision_location and not re.search(
            r"(?:不预设|待验证|检验是否)[^。；\n]{0,20}"
            r"(?:碰撞|必然|发生)",
            ques2_text,
        ):
            raise QualityGateError(
                "首次碰撞的位置/阶段不能在实体计算前预设为必然结论；"
                "应扫描完整事件域后由 SAT 活动约束识别碰撞对和时刻，"
                "高曲率只能作为重点检查区域"
            )

    ques1_text = "\n".join(collect_question_text(value, 1))
    malformed_archimedean_arc_primitive = re.search(
        r"S\s*\(\s*(?:θ|theta)\s*\)\s*=\s*0\.5\s*[×*·]?\s*"
        r"(?:sqrt|√)\s*\(",
        ques1_text,
        re.IGNORECASE,
    )
    if malformed_archimedean_arc_primitive:
        raise QualityGateError(
            "等距螺线弧长原函数漏乘参数：r=aθ时应使用"
            "S(θ)=a/2·[θ√(1+θ²)+asinh(θ)]（允许等价对数式）；"
            "不能从0.5·√(...)直接起式"
        )
    reversed_speed_projection = None
    for segment in re.split(r"[。；;\n]+", ques1_text):
        candidate = re.search(
            r"v_?\{?\s*(?:i|k)\s*\+\s*1\s*\}?\s*=\s*"
            r"v_?\{?\s*(?:i|k)\s*\}?[^。；\n]{0,80}"
            r"(?:弦向量|Δ\s*P)[^/。；\n]{0,35}"
            r"(?:切向|τ)_?\{?\s*(?:i|k)\s*\+\s*1\s*\}?"
            r"[^。；\n]{0,24}/[^。；\n]{0,60}"
            r"(?:切向|τ)_?\{?\s*(?:i|k)\s*\}?",
            segment,
            re.IGNORECASE,
        )
        if candidate and not re.search(
            r"(?:不得|不能|不可|不应|错误|错写|并非)[^。；\n]{0,35}"
            r"v_?\{?\s*(?:i|k)\s*\+\s*1",
            segment[: candidate.start() + 1],
            re.IGNORECASE,
        ):
            reversed_speed_projection = candidate
            break
    if reversed_speed_projection:
        raise QualityGateError(
            "刚性弦长速度投影比的分子分母写反：由弦长约束求导，"
            "v_{i+1}/v_i 应等于弦向量在节点 i 切向上的投影除以"
            "在节点 i+1 切向上的投影"
        )
    cosine_only_speed = re.search(
        r"v_?\{?\s*(?:i|k)\s*\+\s*1\s*\}?\s*=\s*"
        r"v_?\{?\s*(?:i|k)\s*\}?\s*(?:[·×*])?\s*"
        r"(?:cos|余弦)\s*\([^)]*(?:Δ|差)[^)]*\)",
        ques1_text,
        re.IGNORECASE,
    )
    if cosine_only_speed:
        raise QualityGateError(
            "刚性链速度递推不能简化为两端切线夹角的单个余弦；"
            "必须使用弦向量与两端切向的投影比，并检查分母接近零的状态"
        )
    unsupported_monotone_speed = re.search(
        r"(?:速度|速率)[^。；\n]{0,18}(?:沿链|从龙头[^。；\n]{0,12}龙尾)"
        r"[^。；\n]{0,18}单调(?:衰减|递减|下降|不增)|"
        r"(?:沿链|从龙头[^。；\n]{0,12}龙尾)[^。；\n]{0,18}"
        r"(?:速度|速率)[^。；\n]{0,12}单调(?:衰减|递减|下降|不增)",
        ques1_text,
        re.IGNORECASE,
    )
    if unsupported_monotone_speed and not re.search(
        r"(?:待验证|检验是否|不预设|由结果判定|可能不)", ques1_text
    ):
        raise QualityGateError(
            "刚性链沿曲线路径的速度传播不能在计算前预设速度沿链单调衰减；"
            "须由弦向量与两端切向投影的数值结果判定，并允许局部放大或非单调"
        )
    for segment in re.split(r"[。；;\n]+", ques1_text):
        all_nodes_unit_arc_increment = re.search(
            r"S\s*\(\s*(?:θ|theta)\s*[_ ]?[ik]\s*\(\s*t\s*\)\s*\)"
            r"[^。；;\n]{0,50}S\s*\(\s*(?:θ|theta)\s*[_ ]?[ik]\s*"
            r"\(\s*t\s*[-−]\s*1\s*\)\s*\)[^。；;\n]{0,30}"
            r"(?:[-−]\s*1|1\s*m(?:/s)?|v\s*Δ?t)",
            segment,
            re.IGNORECASE,
        )
        limits_arc_increment_to_head = re.search(
            r"(?:龙头|节点\s*1|[ik]\s*=\s*1)", segment, re.IGNORECASE
        )
        if all_nodes_unit_arc_increment and not limits_arc_increment_to_head:
            raise QualityGateError(
                "问题一把任意节点 θ_i 的逐秒弧长增量都按龙头1m/s校验；"
                "题面只固定龙头速度，单位弧长闭合残差只能用于节点1，"
                "其余节点须与各自速度积分核对"
            )

    unequal_tangent_arcs = re.search(
        r"两段圆弧.{0,50}相切.{0,100}(?:半径.{0,20}2\s*倍|2\s*[:：]\s*1)|"
        r"(?:半径.{0,20}2\s*倍|2\s*[:：]\s*1).{0,100}两段圆弧.{0,50}相切",
        source_text,
        re.DOTALL,
    )
    if unequal_tangent_arcs:
        ques4_text = value.get("ques4", "") if isinstance(value, dict) else ""
        ques4_contract_text = (
            (value.get("eda", "") + "\n" + ques4_text)
            if isinstance(value, dict)
            else ques4_text
        )
        ques4_contract_text = ques4_contract_text.translate(str.maketrans("₁₂", "12"))
        states_c1 = re.search(
            r"(?:C\s*1|一阶|切向|切线).{0,15}(?:连续|相同|一致)|"
            r"(?:位置|坐标)\s*[+＋/、与和]\s*(?:切向|切线)\s*C\s*1|"
            r"C\s*1\s*(?:但|而)?非\s*C\s*2",
            ques4_contract_text,
            re.IGNORECASE,
        )
        states_discontinuity = re.search(
            r"(?:不|非|不能|无法|并不).{0,6}曲率连续|"
            r"曲率连续.{0,12}(?:不成立|不满足|仅为假设)|"
            r"曲率(?:跳变|不连续)|曲率[^。；\n]{0,80}(?:存在|发生|并|且)?跳变|"
            r"(?:不得|禁止)[^。；\n]{0,18}曲率连续|"
            r"1\s*/\s*R.{0,30}(?:不等|不同|跳变)|"
            r"(?:非|不满足)\s*C\s*2",
            ques4_contract_text,
            re.IGNORECASE,
        )
        for segment in (
            segment.strip()
            for field in text_fields
            for segment in re.split(r"[。；;\n]+", field)
            if segment.strip()
        ):
            if "曲率连续" not in segment:
                continue
            assertion_text = re.sub(
                r"曲率连续性?(?:审计|检查|检验|核验|验证|分析)",
                "",
                segment,
            )
            if "曲率连续" not in assertion_text:
                continue
            segment_states_discontinuity = re.search(
                r"(?:不|非|不能|无法|并不).{0,6}曲率连续|"
                r"曲率连续.{0,12}(?:不成立|不满足|仅为假设)|曲率(?:跳变|不连续)",
                assertion_text,
            )
            segment_states_discontinuity = segment_states_discontinuity or re.search(
                r"(?:不得|禁止)[^。；\n]{0,18}曲率连续|"
                r"曲率[^。；\n]{0,35}(?:存在|发生)?跳变",
                assertion_text,
            )
            if not segment_states_discontinuity:
                raise QualityGateError(
                    "固定不等半径的相切圆弧只能保证位置/切向 C1 连续，"
                    "两侧曲率分别为半径倒数，不能声明曲率连续"
                )
        if ques4_text and not (states_c1 and states_discontinuity):
            raise QualityGateError(
                "固定不等半径的相切圆弧方案必须同时明确：位置/切向 C1 连续，"
                "以及连接点两侧曲率 1/R 不同并发生跳变；不能只写相切残差"
            )
        source_ques4 = re.search(
            r"问题\s*4(?P<body>.*?)(?=问题\s*5|$)", source_text, re.DOTALL
        )
        source_asks_if_shorter = source_ques4 and re.search(
            r"(?:能否|是否)[^。；]{0,50}(?:缩短|变短)",
            source_ques4.group("body"),
        )
        preanswers_shortening = re.search(
            r"(?:题目输出契约|输出要求|预期结论)[^。；\n]{0,80}"
            r"(?:回答[^。；\n]{0,18})?(?:为肯定|肯定回答|可以缩短|能够缩短)",
            ques4_text,
        )
        if source_asks_if_shorter and preanswers_shortening:
            raise QualityGateError(
                "问题四询问能否缩短时，不能在优化计算前把答案预设为肯定；"
                "须先复算题面基线和约束内候选解，再由长度差及数值误差决定回答"
            )
        if source_asks_if_shorter:
            ambiguous_turning_baseline = re.search(
                r"(?:基线|基准)[^。；\n]{0,80}(?:唯一\s*或\s*优选|"
                r"(?:任意|某个)?可行解[^。；\n]{0,18}(?:优选|择优))|"
                r"(?:唯一\s*或\s*优选|(?:任意|某个)?可行解[^。；\n]{0,18}"
                r"(?:优选|择优))[^。；\n]{0,80}(?:基线|基准)",
                ques4_text,
                re.IGNORECASE,
            )
            if ambiguous_turning_baseline:
                raise QualityGateError(
                    "问题四的调整前基线不能定义成‘唯一或优选’的任意可行解；"
                    "须由题面调整前几何条件唯一复算并报告相切残差，"
                    "若存在多解则必须给出题面授权的选择准则，不能在候选族中择优充当基线"
                )
            fixes_turn_endpoints = re.search(
                r"(?:固定|给定|锁定)[^。；\n]{0,18}"
                r"(?:入口(?:点|切点)\s*M[^。；\n]{0,45}出口(?:点|切点)\s*N|"
                r"出口(?:点|切点)\s*N[^。；\n]{0,45}入口(?:点|切点)\s*M|"
                r"两端点|端点\s*M\s*[,，、和与]\s*N)|"
                r"(?:盘入|入口)边界点\s*M[^。；\n]{0,100}"
                r"(?:盘出|出口)边界点\s*N|"
                r"(?:盘出|出口)边界点\s*N[^。；\n]{0,100}"
                r"(?:盘入|入口)边界点\s*M",
                ques4_text,
                re.IGNORECASE,
            )
            fixes_endpoint_tangents = re.search(
                r"M\s*处[^。；\n]{0,24}(?:切向|切线)[^。；\n]{0,45}"
                r"N\s*处[^。；\n]{0,24}(?:切向|切线)|"
                r"N\s*处[^。；\n]{0,24}(?:切向|切线)[^。；\n]{0,45}"
                r"M\s*处[^。；\n]{0,24}(?:切向|切线)",
                ques4_text,
                re.IGNORECASE,
            )
            searches_radius_after_fixing_endpoints = re.search(
                r"(?:黄金分割|网格(?:搜索|扫描)?|扫描|二分|连续优化)"
                r"[^。；\n]{0,45}R\s*[_ ]?2|"
                r"R\s*[_ ]?2[^。；\n]{0,45}"
                r"(?:黄金分割|网格(?:搜索|扫描)?|扫描|二分|连续优化)",
                ques4_text,
                re.IGNORECASE,
            )
            denies_fixed_endpoints = re.search(
                r"(?:入口|出口|两端|M\s*[,，、和与]\s*N)"
                r"[^。；\n]{0,25}(?:并非|不是|不作|不设为|不固定|允许|可以)"
                r"[^。；\n]{0,16}(?:固定|锁定)?(?:端点|切点)|"
                r"(?:切点|端点|M\s*[,，、和与]\s*N)"
                r"[^。；\n]{0,25}(?:沿|可沿)[^。；\n]{0,20}螺线(?:移动|滑动)",
                ques4_text,
                re.IGNORECASE,
            )
            denies_radius_search = re.search(
                r"(?:不再|无需|不需|不作|禁止)[^。；\n]{0,8}"
                r"(?:扫描|搜索|二分|连续优化)[^。；\n]{0,12}R\s*[_ ]?2|"
                r"R\s*[_ ]?2[^。；\n]{0,12}"
                r"(?:无需|不需|不作|禁止)[^。；\n]{0,8}(?:扫描|搜索|优化)",
                ques4_text,
                re.IGNORECASE,
            )
            if (
                fixes_turn_endpoints
                and fixes_endpoint_tangents
                and searches_radius_after_fixing_endpoints
                and not denies_fixed_endpoints
                and not denies_radius_search
            ):
                raise QualityGateError(
                    "分段相切曲线优化的自由度定义矛盾：既固定入口/出口点及其切向，"
                    "又固定半径比和两弧相切关系后，不能仍把 R2 当作独立连续变量扫描；"
                    "应先列未知量与独立约束。若端点固定则直接解约束并检验可行解是否唯一；"
                    "若切点可沿源曲线移动，则须明确其参数和移动域，不能同时称为固定边界点"
                )
            assumes_semicircular_arcs = re.search(
                r"(?:目标|弧长|L)[^。；;\n]{0,45}"
                r"(?:π\s*R\s*1\s*[+＋]\s*π\s*R\s*2|3\s*π\s*R\s*2)|"
                r"(?:若|按|视为|假定)[^。；;\n]{0,8}半圆"
                r"[^。；;\n]{0,40}(?:目标|弧长|L)",
                ques4_text,
                re.IGNORECASE,
            )
            if assumes_semicircular_arcs:
                raise QualityGateError(
                    "问题四的两段调头圆弧不能未求解就假定为半圆；"
                    "弧长必须由实际相切点与连接点决定的圆心角计算，"
                    "即 L=R1·Δφ1+R2·Δφ2，不能直接写成 3πR2"
                )
            has_baseline_length = re.search(
                r"(?:基线|基准|原(?:调头)?曲线|调整前)[^。；\n]{0,35}"
                r"(?:长度|弧长|L\s*[_ ]?(?:base|0))|"
                r"(?:长度|弧长|L\s*[_ ]?(?:base|0))[^。；\n]{0,35}"
                r"(?:基线|基准|原(?:调头)?曲线|调整前)",
                ques4_text,
                re.IGNORECASE,
            )
            compares_shortening = re.search(
                r"(?:长度差|弧长差|缩短(?:量|比例)|比较|对比)|"
                r"L\s*[_ ]?(?:opt|min)[^。；\n]{0,20}"
                r"(?:<|小于|减去|与)[^。；\n]{0,20}L\s*[_ ]?(?:base|0)",
                ques4_text,
                re.IGNORECASE,
            )
            if subtask_title != "eda" and not (
                has_baseline_length and compares_shortening
            ):
                raise QualityGateError(
                    "问题四询问能否缩短，必须复算一条题面约束内的调整前基线长度，"
                    "并与优化候选的长度作数值比较；可行 R2 集不能代替基线"
                )
            arbitrary_feasible_baseline = re.search(
                r"(?:取|任选|选择)(?:一|一个|任意)?[^。；\n]{0,18}"
                r"(?:满足|符合)[^。；\n]{0,12}(?:相切|约束)[^。；\n]{0,18}"
                r"(?:R\s*2|半径|构型)[^。；\n]{0,12}(?:基线|基准)|"
                r"(?:基线|基准)[^。；\n]{0,12}(?:取|任选|选择)(?:一|一个|任意)?"
                r"[^。；\n]{0,18}(?:可行|相切)",
                ques4_text,
                re.IGNORECASE,
            )
            if arbitrary_feasible_baseline and not re.search(
                r"(?:题面|调整前|原调头曲线)[^。；\n]{0,35}"
                r"(?:复算|反解|还原)[^。；\n]{0,20}(?:基线|基准)|"
                r"(?:基线|基准)[^。；\n]{0,20}(?:由|按|根据)"
                r"[^。；\n]{0,20}(?:题面|调整前|原调头曲线)",
                ques4_text,
                re.IGNORECASE,
            ):
                raise QualityGateError(
                    "问题四不能任取一个可行半径或相切构型充当调整前基线；"
                    "必须由题面调整前双圆弧构型及全部相切约束复算 L_base"
                )
        invents_source_r2_baseline = re.search(
            r"题面(?:给出|给定|提供)[^。；\n]{0,18}"
            r"(?:R\s*2|后一段(?:圆弧)?半径)[^。；\n]{0,18}基准|"
            r"题面[^。；\n]{0,18}(?:R\s*2|后一段(?:圆弧)?半径)基准",
            ques4_text,
            re.IGNORECASE,
        )
        source_has_r2_baseline = source_ques4 and re.search(
            r"R\s*2\s*=|后一段(?:圆弧)?半径\s*(?:为|=)\s*\d",
            source_ques4.group("body"),
            re.IGNORECASE,
        )
        if invents_source_r2_baseline and not source_has_r2_baseline:
            raise QualityGateError(
                "问题四题面没有给出数值 R2 基准；"
                "只能由题面几何约束复算规范基线，不能把未声明的半径说成来源事实"
            )
        naked_numeric_r2_baseline = re.search(
            r"(?:基准(?:配置|方案)?|对照)[^。；\n]{0,24}"
            r"R\s*2\s*=\s*\d+(?:\.\d+)?\s*m?(?!\s*[:：])|"
            r"R\s*2\s*=\s*\d+(?:\.\d+)?\s*m?(?!\s*[:：])[^。；\n]{0,24}"
            r"(?:基准(?:配置|方案)?|对照)",
            ques4_text,
            re.IGNORECASE,
        )
        derives_numeric_r2_baseline = re.search(
            r"(?:由|根据)[^。；\n]{0,30}(?:相切|题面几何|非线性方程)"
            r"[^。；\n]{0,30}(?:复算|求解|反解|得到)[^。；\n]{0,20}R\s*2|"
            r"R\s*2[^。；\n]{0,30}(?:由|根据)[^。；\n]{0,20}"
            r"(?:相切|题面几何)[^。；\n]{0,20}(?:复算|求解|反解)",
            ques4_text,
            re.IGNORECASE,
        )
        if (
            naked_numeric_r2_baseline
            and not source_has_r2_baseline
            and not derives_numeric_r2_baseline
        ):
            raise QualityGateError(
                "问题四题面只固定双圆弧半径比，没有给数值 R2 基准；"
                "裸写一个 R2 数值作基准/对照属于来源幻觉，必须由相切几何复算并记录残差"
            )
        if re.search(
            r"等半径[^。；\n]{0,12}基线|基线[^。；\n]{0,12}等半径", ques4_text
        ):
            raise QualityGateError(
                "问题四题面基线的两段圆弧半径比为 2:1，不能改成等半径基线；"
                "若做约束外对照，必须明确其不是题面可行基线"
            )
        opens_fixed_ratio = re.search(
            r"优化[^。；\n]{0,10}可调(?:整)?(?:半径)?比(?:值|例)|"
            r"决策变量[^。；\n]{0,20}(?:R1\s*/\s*R2|半径比|比值)|"
            r"(?:R1\s*/\s*R2|半径比)[^。；\n]{0,20}(?:作为|为)?决策变量|"
            r"(?:再|然后)?(?:放开|开放|解除)\s*"
            r"(?:圆弧)?(?:半径)?比(?:值|例)?",
            ques4_text,
            re.IGNORECASE,
        )
        denies_opening_fixed_ratio = re.search(
            r"(?:不再|不|不得|禁止)[^。；\n]{0,8}(?:放开|开放|解除)"
            r"[^。；\n]{0,12}(?:半径)?比(?:值|例)?",
            ques4_text,
            re.IGNORECASE,
        )
        labels_open_ratio_infeasible = re.search(
            r"(?:开放|放开|解除)[^。；\n]{0,12}(?:半径)?比(?:值|例)?"
            r"[^。；\n]{0,18}(?:不可行|约束外|仅作对照|对照)|"
            r"(?:不可行|约束外|仅作对照)[^。；\n]{0,18}"
            r"(?:开放|放开|解除)[^。；\n]{0,12}(?:半径)?比(?:值|例)?",
            ques4_text,
            re.IGNORECASE,
        )
        if (
            opens_fixed_ratio
            and not denies_opening_fixed_ratio
            and not labels_open_ratio_infeasible
        ):
            raise QualityGateError(
                "问题四题面将前后两段圆弧半径比固定为 2:1；"
                "只能在该约束内调整圆弧几何，不能把 R1/R2 比值作为决策变量"
            )
        forces_entire_chain_inside_turning_circle = re.search(
            r"(?:整条链|全链|整条龙|全部(?:节点|把手|板体)|"
            r"所有(?:节点|把手|板体)|链板体)"
            r"[^。；\n]{0,70}(?:运动)?(?:全程)?[^。；\n]{0,20}"
            r"(?:不越出|不得越出|限制在|位于|全部进入)"
            r"[^。；\n]{0,25}(?:调头空间|调头圆|圆内|边界)",
            ques4_text,
        )
        if forces_entire_chain_inside_turning_circle and not re.search(
            r"(?:不要求|不需|无需|并非|不是)[^。；\n]{0,20}"
            r"(?:整条链|全链|整条龙|全部|所有|链板体)",
            ques4_text,
        ):
            raise QualityGateError(
                "问题四把调头空间的轨迹约束擅自扩大到整条链："
                "两侧龙身本来就沿盘入/盘出螺线延伸到圆外；应约束调头轨迹及"
                "题面明确指定的对象，不能要求整条链板体全程位于4.5 m圆内"
            )
        overconstrains_supporting_circle = re.search(
            r"(?:圆心到(?:调头)?圆心距离|两圆心距离|圆心距)"
            r"\s*[+＋]\s*(?:各自)?半径\s*(?:<=|≤)\s*4\.5\s*m?",
            ques4_text,
        )
        if overconstrains_supporting_circle:
            raise QualityGateError(
                "问题四只要求实际调头圆弧位于调头空间内，不能用“支撑圆圆心距+整圆半径"
                "≤4.5m”强迫整个支撑圆都在区域内；应对实际弧段求解析极值或逐点加密检查"
            )
        sample_only_arc_containment = re.search(
            r"(?:全(?:路径|曲线)|所有点|整条调头曲线)[^。；\n]{0,45}"
            r"(?:位于|在)[^。；\n]{0,18}(?:4\.5\s*m|调头(?:空间|圆))"
            r"[^。；\n]{0,55}(?:采样|离散)\s*\d+\s*点|"
            r"(?:采样|离散)\s*\d+\s*点[^。；\n]{0,55}"
            r"(?:最大极径|全(?:路径|曲线)|所有点)[^。；\n]{0,35}"
            r"(?:<=|≤)\s*4\.5",
            ques4_text,
            re.IGNORECASE,
        )
        certifies_continuous_arc_containment = re.search(
            r"(?:极径|距离)[^。；\n]{0,30}(?:解析求导|解析极值|驻点)|"
            r"(?:区间算术|区间界|严格上界|采样误差界|弦高误差界)|"
            r"(?:端点|圆心连线方向)[^。；\n]{0,45}(?:极值点|解析最大)",
            ques4_text,
            re.IGNORECASE,
        )
        if sample_only_arc_containment and not certifies_continuous_arc_containment:
            raise QualityGateError(
                "问题四仅用有限采样点声称整段圆弧都位于调头圆内；"
                "连续圆弧的区域约束须解析检查端点与极径驻点，或给出严格的区间/"
                "采样误差界，固定点数采样不能排除采样间越界"
            )
        source_requires_turning_path_containment = source_ques4 and re.search(
            r"调头空间内[^。；\n]{0,20}(?:完成)?调头",
            source_ques4.group("body"),
        )
        states_turning_path_containment = re.search(
            r"(?:全(?:路径|曲线)|调头(?:路径|曲线)|实际弧段|两段圆弧)"
            r"[^。；\n]{0,55}(?:位于|限制在|不越出|不超出|满足)"
            r"[^。；\n]{0,30}(?:4\.5\s*m|调头空间|调头圆|圆内)",
            ques4_text,
        )
        if source_requires_turning_path_containment and not (
            states_turning_path_containment and certifies_continuous_arc_containment
        ):
            raise QualityGateError(
                "问题四缺少调头曲线位于题面调头空间内的连续区域证书；"
                "须对实际两段圆弧的极径解析检查端点与驻点，或给出严格区间/"
                "采样误差界，不能只画出9m圆或省略该可行性约束"
            )
        source_defines_turn_start_as_zero = source_ques4 and re.search(
            r"以[^。；\n]{0,30}调头开始[^。；\n]{0,20}(?:零时刻|0\s*时刻)|"
            r"以调头开始时间为零时刻",
            source_ques4.group("body"),
        )
        maps_zero_to_turn_start = re.search(
            r"(?:t\s*=\s*0|零时刻)[^。；\n]{0,25}"
            r"(?:调头开始|开始调头|进入调头)|"
            r"(?:调头开始|开始调头)[^。；\n]{0,25}"
            r"(?:t\s*=\s*0|零时刻)",
            ques4_text,
            re.IGNORECASE,
        )
        if source_defines_turn_start_as_zero and not maps_zero_to_turn_start:
            raise QualityGateError(
                "问题四遗漏题面时间原点：必须明确 t=0 对应调头开始，"
                "再据此向前生成盘入段、向后生成调头及盘出段；"
                "只写 -100~100s 不能保证时间与路径事件对齐"
            )
        wrong_same_arc_collinearity = re.search(
            r"O\s*[1₁]\s*[,，、]?\s*P\s*[1₁]\s*[,，、]?\s*P"
            r"[^。；\n]{0,12}(?:三点)?共线|"
            r"O\s*[1₁]\s*[,，、]?\s*P\s*[,，、]?\s*P\s*[1₁]"
            r"[^。；\n]{0,12}(?:三点)?共线|"
            r"O\s*[2₂]\s*[,，、]?\s*P\s*[2₂]\s*[,，、]?\s*P"
            r"[^。；\n]{0,12}(?:三点)?共线|"
            r"O\s*[2₂]\s*[,，、]?\s*P\s*[,，、]?\s*P\s*[2₂]"
            r"[^。；\n]{0,12}(?:三点)?共线|"
            r"(?:圆心)?O\s*[1₁][^。；\n]{0,18}(?:入口)?切点\s*P\s*[1₁]"
            r"[^。；\n]{0,18}(?:公共)?切点\s*P[^。；\n]{0,8}(?:三点)?共线|"
            r"(?:圆心)?O\s*[2₂][^。；\n]{0,18}(?:出口)?切点\s*P\s*[2₂]"
            r"[^。；\n]{0,18}(?:公共)?切点\s*P[^。；\n]{0,8}(?:三点)?共线",
            ques4_text,
            re.IGNORECASE,
        )
        if wrong_same_arc_collinearity:
            raise QualityGateError(
                "双圆弧几何错误：同一圆弧的圆心、入口/出口切点和公共切点"
                "通常不是三点共线；半径应分别连接圆心与两个不同弧上点，"
                "共线的是两圆心与公共切点"
            )
        if re.search(
            r"圆心距\s*=\s*R\s*1\s*(?:±|\+\s*/\s*-|或)\s*R\s*2|"
            r"圆心距\s*=\s*R\s*1\s*[-+]\s*R\s*2[^。；\n]{0,20}"
            r"(?:视|根据)[^。；\n]{0,8}(?:内|外)切",
            ques4_text,
            re.IGNORECASE,
        ):
            raise QualityGateError(
                "S 形双圆弧在连接点曲率转向相反，必须确定为外切分支并使用"
                "圆心距 R1+R2；不能把 R1±R2 留给代码阶段任意选择"
            )
        source_requires_s_curve = source_ques4 and re.search(
            r"S\s*形[^。；\n]{0,40}两段圆弧|"
            r"两段圆弧[^。；\n]{0,40}S\s*形",
            source_ques4.group("body"),
            re.IGNORECASE,
        )
        if source_requires_s_curve:
            if _has_affirmative_internal_tangency(ques4_contract_text):
                raise QualityGateError(
                    "S 形反向转向的双圆弧应选择外切分支并使用圆心距 R1+R2；"
                    "同一建模合同不得把该构型称为内切"
                )
            selects_external_tangency = re.search(
                r"外切(?:分支)?", ques4_contract_text, re.IGNORECASE
            )
            uses_external_center_distance = re.search(
                r"圆心距[^。；\n]{0,12}R\s*1\s*\+\s*R\s*2|"
                r"R\s*1\s*\+\s*R\s*2[^。；\n]{0,12}圆心距",
                ques4_contract_text,
                re.IGNORECASE,
            )
            defines_double_and_single_radius = re.search(
                r"(?:前段|前一段)[^。；\n]{0,10}半径\s*2\s*R"
                r"[^。；\n]{0,25}(?:后段|后一段)[^。；\n]{0,10}半径\s*R\b|"
                r"R\s*1\s*=\s*2\s*R[^。；\n]{0,25}R\s*2\s*=\s*R\b",
                ques4_contract_text,
                re.IGNORECASE,
            )
            uses_external_center_distance = uses_external_center_distance or (
                defines_double_and_single_radius
                and re.search(
                    r"圆心距\s*(?:=|为)\s*3\s*R\b|3\s*R\s*(?:为|是)圆心距",
                    ques4_contract_text,
                    re.IGNORECASE,
                )
            )
            if subtask_title != "eda" and not (
                selects_external_tangency and uses_external_center_distance
            ):
                raise QualityGateError(
                    "S 形反向转向的双圆弧必须在建模合同中明确选择"
                    "外切分支并使用圆心距 R1+R2；仅写“两弧相切”不足以"
                    "排除错误的内切构型"
                )
        labels_arc_endpoints = (
            re.search(r"(?:起点\s*P\s*1|P\s*1[^。；\n]{0,12}起点)", ques4_text, re.I)
            and re.search(
                r"(?:终点\s*P\s*3|P\s*3[^。；\n]{0,12}终点)", ques4_text, re.I
            )
            and re.search(
                r"(?:两弧[^。；\n]{0,12}P\s*2[^。；\n]{0,12}相切|"
                r"P\s*2[^。；\n]{0,12}(?:两弧)?切点)",
                ques4_text,
                re.I,
            )
        )
        if labels_arc_endpoints and (
            re.search(
                r"P\s*1\s*P\s*2\s*=\s*2\s*[ρrho].{0,30}"
                r"P\s*2\s*P\s*3\s*=\s*[ρrho]",
                ques4_text,
                re.I | re.DOTALL,
            )
            or re.search(
                r"P\s*1[^。；\n]{0,8}P\s*2[^。；\n]{0,8}P\s*3"
                r"[^。；\n]{0,8}共线",
                ques4_text,
                re.I,
            )
        ):
            raise QualityGateError(
                "双圆弧几何把弧端点、连接点和圆心混淆：半径约束应连接各圆心与其弧上点，"
                "相切时共线的是两圆心与公共切点；不能令起点 P1、连接点 P2、终点 P3 共线，"
                "也不能把弧端点弦长 P1P2/P2P3 直接等于圆弧半径"
            )
        sensitivity_text = (
            value.get("sensitivity_analysis", "") if isinstance(value, dict) else ""
        )
        perturbs_fixed_ratio = re.search(
            r"(?:圆弧)?半径比[^。；\n]{0,30}"
            r"(?:[±＋+-]\s*\d+(?:\.\d+)?\s*%|扰动|扫描|变化)",
            sensitivity_text,
            re.IGNORECASE,
        )
        if perturbs_fixed_ratio and not re.search(
            r"(?:圆弧)?半径比(?:值|例)?[^。；\n]{0,28}"
            r"(?:不扰动|不作扰动|不得扰动|禁止扰动|保持固定|固定为|"
            r"不改变|不得改变)|"
            r"(?:圆弧)?半径比[^。；\n]{0,60}(?:来源常量|固定物理量)"
            r"[^。；\n]{0,24}(?:不做|不作|不得|禁止)[^。；\n]{0,8}扰动|"
            r"(?:保持|固定)[^。；\n]{0,18}(?:圆弧)?半径比"
            r"[^。；\n]{0,24}(?:不变|固定|不扰动)",
            sensitivity_text,
            re.IGNORECASE,
        ):
            raise QualityGateError(
                "问题四的 2:1 圆弧半径比是题面固定约束，不能在敏感性分析中对该比值作扰动；"
                "应扰动约束内的尺度、切点、数值容差或来源不确定参数"
            )
        if re.search(
            r"(?:单圆弧|直线连接)[^。；\n]{0,18}基线|"
            r"基线[^。；\n]{0,18}(?:单圆弧|直线连接)",
            ques4_text,
        ) and not re.search(
            r"(?:单圆弧|直线连接)[^。；\n]{0,25}(?:约束外|不可行|仅作对照)|"
            r"(?:约束外|不可行|仅作对照)[^。；\n]{0,25}(?:单圆弧|直线连接)",
            ques4_text,
        ):
            raise QualityGateError(
                "问题四基线必须是题面给定的 2:1 双圆弧相切路径；"
                "单圆弧或直线连接不满足该候选族，不能冒充可行基线"
            )
        q4_claim_text = ques4_text + "\n" + sensitivity_text
        if re.search(
            r"多(?:起点|初值|次重启)[^。；\n]{0,35}"
            r"(?:确认|证明|保证)[^。；\n]{0,12}全局(?:最小|最优)|"
            r"(?:确认|证明|保证)[^。；\n]{0,12}全局(?:最小|最优)"
            r"[^。；\n]{0,35}多(?:起点|初值|次重启)",
            q4_claim_text,
        ):
            raise QualityGateError(
                "多起点/多次重启只能提供局部解稳定性证据，不能单独证明全局最优；"
                "须给可验证下界、穷举覆盖或把结论降为当前候选族内最优解"
            )
        grid_claims_global_optimum = re.search(
            r"(?:小步长|离散|网格)[^。；\n]{0,24}(?:穷举|扫描)"
            r"[^。；\n]{0,24}(?:确认|证明|保证)[^。；\n]{0,12}"
            r"(?:全局(?:最小|最优)|无更优)|"
            r"穷举[^。；\n]{0,18}(?:小步长|离散|网格)"
            r"[^。；\n]{0,24}(?:确认|证明|保证)[^。；\n]{0,12}"
            r"(?:全局(?:最小|最优)|无更优)",
            ques4_text,
        )
        has_global_certificate = re.search(
            r"(?:区间算术|解析(?:下界|证明)|可验证下界|凸性证明|完备枚举|"
            r"离散误差界|网格间误差界)",
            ques4_text,
        )
        if grid_claims_global_optimum and not has_global_certificate:
            raise QualityGateError(
                "有限网格或小步长扫描不能单独证明连续域全局最优或‘无更优’；"
                "须给网格间误差界、可验证下界/凸性证书，或把结论降为网格与局部精修所得候选解"
            )
        normalized_q4_formula = ques4_text.translate(str.maketrans("₁₂", "12"))
        for length_formula in re.finditer(
            r"L(?:\s*\([^)]*\))?\s*=\s*"
            r"(?P<first>\d+(?:\.\d+)?)\s*(?:ρ|rho)\s*[·*×]?\s*"
            r"(?:Δ\s*[φθ]|d(?:elta)?\s*(?:phi|theta))\s*1"
            r"\s*[+＋]\s*(?P<second>\d+(?:\.\d+)?)\s*(?:ρ|rho)\s*[·*×]?\s*"
            r"(?:Δ\s*[φθ]|d(?:elta)?\s*(?:phi|theta))\s*2",
            normalized_q4_formula,
            re.IGNORECASE,
        ):
            coefficients = (
                float(length_formula.group("first")),
                float(length_formula.group("second")),
            )
            if not (
                math.isclose(coefficients[0], 2.0)
                and math.isclose(coefficients[1], 1.0)
            ):
                raise QualityGateError(
                    "双圆弧长度系数与R1=2ρ、R2=ρ矛盾："
                    f"读到 {coefficients[0]:g}ρΔφ1+{coefficients[1]:g}ρΔφ2，"
                    "正确式应为2ρΔφ1+ρΔφ2"
                )
        if re.search(
            r"(?:L(?:_s)?|弧长)[^。；\n=]{0,20}=.{0,45}?"
            r"(?:0\.5\s*[ρrho]|[ρrho]\s*/\s*2).{0,30}?"
            r"(?:theta|θ)[_ ]?(?:span)?2",
            ques4_text,
            re.IGNORECASE,
        ):
            raise QualityGateError(
                "双圆弧长度公式与半径定义不一致：若 R1=2ρ、R2=ρ，"
                "则 L=2ρ·Δθ1+ρ·Δθ2，不能把第二段写成 0.5ρ·Δθ2"
            )
        if source_asks_if_shorter and subtask_title != "eda":
            normalized_ques4_text = ques4_text.translate(str.maketrans("₁₂", "12"))
            has_actual_central_angle_length = re.search(
                r"L\s*(?:[_ ]?(?:base|opt|0))?\s*=\s*"
                r"R\s*1\s*[·*×]?\s*(?:Δ\s*[φθ]|d(?:elta)?\s*(?:phi|theta))\s*1"
                r"\s*[+＋]\s*R\s*2\s*[·*×]?\s*"
                r"(?:Δ\s*[φθ]|d(?:elta)?\s*(?:phi|theta))\s*2|"
                r"(?:实际|由)[^。；\n]{0,30}(?:相切点|切点)"
                r"[^。；\n]{0,30}(?:圆心角|中心角)"
                r"[^。；\n]{0,35}(?:半径[^。；\n]{0,12}(?:乘|乘以)|"
                r"弧长|长度)[^。；\n]{0,20}(?:相加|求和)",
                normalized_ques4_text,
                re.IGNORECASE,
            )
            has_reduced_central_angle_length = re.search(
                r"R\s*1\s*=\s*2\s*[·*×]?\s*R\s*2",
                normalized_ques4_text,
                re.IGNORECASE,
            ) and re.search(
                r"(?:L\s*(?:[_ ]?(?:base|opt|0))?|弧长)\s*=\s*"
                r"2\s*[·*×]?\s*R\s*2\s*[·*×]?\s*"
                r"(?:Δ\s*[φθ]|d(?:elta)?\s*(?:phi|theta))\s*1"
                r"\s*[+＋]\s*R\s*2\s*[·*×]?\s*"
                r"(?:Δ\s*[φθ]|d(?:elta)?\s*(?:phi|theta))\s*2",
                normalized_ques4_text,
                re.IGNORECASE,
            )
            reduced_radius = r"(?:ρ|rho|r(?!\s*[12]))"
            has_scalar_reduced_central_angle_length = re.search(
                rf"(?:后(?:一)?段(?:圆弧)?半径|R\s*2)\s*(?:为|是|=|：|:)?\s*"
                rf"{reduced_radius}|"
                rf"(?:前(?:一)?段(?:圆弧)?半径|R\s*1)\s*=\s*2\s*[·*×]?\s*"
                rf"{reduced_radius}",
                normalized_ques4_text,
                re.IGNORECASE,
            ) and re.search(
                rf"(?:L\s*(?:(?:[_ ]?(?:base|opt|0))|\(\s*R\s*\))?|"
                rf"(?:目标|总)?(?:弧长|长度))\s*=\s*"
                rf"2\s*[·*×]?\s*{reduced_radius}\s*[·*×]?\s*"
                r"(?:Δ\s*[φθ]|d(?:elta)?\s*(?:phi|theta))\s*1"
                r"(?:\s*\(\s*R\s*\))?"
                rf"\s*[+＋]\s*{reduced_radius}\s*[·*×]?\s*"
                r"(?:Δ\s*[φθ]|d(?:elta)?\s*(?:phi|theta))\s*2"
                r"(?:\s*\(\s*R\s*\))?",
                normalized_ques4_text,
                re.IGNORECASE,
            )
            if not (
                has_actual_central_angle_length
                or has_reduced_central_angle_length
                or has_scalar_reduced_central_angle_length
            ):
                raise QualityGateError(
                    "问题四的基线与候选双圆弧长度必须由实际切点决定的圆心角计算；"
                    "须明确 L=R1·Δφ1+R2·Δφ2，不能只写求解一个未定义口径的 L_base/L_opt"
                )
        arc_curvature_text = "\n".join(
            text
            for text in (
                ques4_text,
                value.get("ques5", "") if isinstance(value, dict) else "",
                value.get("sensitivity_analysis", "")
                if isinstance(value, dict)
                else "",
            )
            if text
        )
        if re.search(
            r"曲率(?:最小|较小)[^。；\n]{0,24}(?:1\s*/\s*[ρrho])",
            arc_curvature_text,
            re.IGNORECASE,
        ):
            raise QualityGateError(
                "圆弧曲率大小写反：R1=2ρ、R2=ρ 时，1/ρ 是较大曲率，1/(2ρ) 才是较小曲率"
            )
        if re.search(
            r"1\s*/\s*(?:R[_ ]?2|[ρrho])\s*"
            r"(?:→|->|到|跳变(?:至|为)?)\s*1\s*/\s*"
            r"(?:R[_ ]?1|\(?\s*2\s*[*·]?\s*(?:R[_ ]?2|[ρrho])\s*\)?)",
            arc_curvature_text,
            re.IGNORECASE,
        ):
            raise QualityGateError(
                "双圆弧沿行进方向的曲率跳变顺序写反：前弧 R1=2R2、后弧 R2 时，"
                "应由 1/R1=1/(2R2) 跳变到 1/R2，不能反向书写"
            )
        ques5_text = value.get("ques5", "") if isinstance(value, dict) else ""
        unsupported_bottleneck_claim = re.search(
            r"(?:尾部节点|龙尾|节点\s*224)[^。；\n]{0,28}"
            r"(?:通常|必然|实际|确定)[^。；\n]{0,12}(?:最高|最大|更高|峰值|瓶颈)|"
            r"(?:实际|必然|按规定)[^。；\n]{0,12}k[_ ]?max\s*>\s*1|"
            r"标注[^。；\n]{0,12}(?:节点\s*224|尾部)[^。；\n]{0,12}峰值",
            ques5_text,
            re.IGNORECASE,
        )
        if unsupported_bottleneck_claim and not re.search(
            r"(?:待计算|候选|不得预设|仅作示意)", ques5_text
        ):
            raise QualityGateError(
                "最大速度问在求解前预设了瓶颈节点或 k_max 大小；"
                "必须遍历题面要求的全部把手与完整时间域后再由数值证据识别活动约束"
            )

    sensitivity_text = (
        value.get("sensitivity_analysis", "") if isinstance(value, dict) else ""
    )
    sensitivity_sentences = re.split(r"[。；;\n]", sensitivity_text)
    modeling_sentences = re.split(
        r"[。；;\n]",
        "\n".join(
            text
            for key, text in value.items()
            if key != "eda" and isinstance(text, str)
        )
        if isinstance(value, dict)
        else sensitivity_text,
    )
    fixed_sensitivity_conflicts: list[str] = []
    perturbation_marker = (
        r"(?:(?<![eE])[±＋+\-]\s*\d+(?:\.\d+)?\s*"
        r"(?:%|cm|mm|m(?!\s*/\s*s))|扰动|变化)"
    )
    fixed_geometry_pattern = (
        r"(?:板长|板体长度|孔位|孔心(?:偏移|距离)?|孔距|板宽|板体宽度)"
    )

    def denies_all_fixed_source_perturbations(sentence: str) -> bool:
        global_denial = re.search(
            r"(?:所有|全部)?(?:题面)?固定(?:的)?"
            r"(?:参数|物理量|约束|事实)[^。；;\n]{0,180}"
            r"(?:不(?:做|作)(?:任何)?(?:敏感性)?扰动|不纳入敏感性(?:分析)?|"
            r"(?:均)?作为确定性常量)",
            sentence,
        )
        explicit_positive = re.search(
            r"(?:板长|板体长度|孔位|孔距|板宽|龙头(?:前把手)?速度|"
            r"调头空间直径|速度限制|限速|螺距)"
            r"[^。；;\n]{0,35}(?:[±＋]\s*\d|"
            r"(?:允许|进行|纳入)[^。；;\n]{0,12}(?:扰动|变化)|"
            r"作为[^。；;\n]{0,8}(?:扰动|敏感性)(?:变量|参数)?)",
            sentence,
        )
        return bool(global_denial and not explicit_positive)

    if guardrails.get("member_constraints") and any(
        re.search(
            fixed_geometry_pattern + rf"[^。；;\n]{{0,28}}{perturbation_marker}",
            sentence,
            re.IGNORECASE,
        )
        and not denies_all_fixed_source_perturbations(sentence)
        and not re.search(
            r"(?:不动|不扰动|不可扰动|不得扰动|不作扰动|禁止扰动|保持)"
            rf"[^。；;\n]{{0,45}}{fixed_geometry_pattern}[^。；;\n]{{0,12}}(?:固定)?|"
            rf"{fixed_geometry_pattern}[^。；;\n]{{0,35}}"
            r"(?:保持固定|不扰动|不可扰动|不得扰动|不作扰动)",
            sentence,
            re.IGNORECASE,
        )
        for sentence in modeling_sentences
    ):
        fixed_sensitivity_conflicts.append("题面固定构件尺寸/孔位")
    source_has_fixed_leader_speed = re.search(
        r"龙头前把手[^。；\n]{0,60}(?:速度|行进速度)"
        r"[^。；\n]{0,30}(?:始终)?保持\s*\d+(?:\.\d+)?\s*m\s*/\s*s",
        source_text,
        re.IGNORECASE,
    )
    if source_has_fixed_leader_speed and any(
        re.search(
            r"(?:龙头(?:前把手)?速度|问题\s*1(?:\s*/\s*4)?"
            r"[^。；;\n]{0,20}龙头[^。；;\n]{0,10}速度)"
            rf"[^。；;\n]{{0,45}}{perturbation_marker}",
            sentence,
            re.IGNORECASE,
        )
        and not denies_all_fixed_source_perturbations(sentence)
        and not re.search(
            r"(?:不扰动|不作扰动|保持固定|禁止扰动)[^。；;\n]{0,30}"
            r"(?:龙头(?:前把手)?速度|问题\s*1(?:\s*/\s*4)?"
            r"[^。；;\n]{0,12}龙头[^。；;\n]{0,10}速度)|"
            r"(?:龙头(?:前把手)?速度|问题\s*1(?:\s*[-~/、和]\s*2)?"
            r"[^。；;\n]{0,20}龙头[^。；;\n]{0,10}速度)[^。；;\n]{0,45}"
            r"(?:不扰动|不作扰动|保持固定|禁止扰动)",
            sentence,
            re.IGNORECASE,
        )
        for sentence in sensitivity_sentences
    ):
        fixed_sensitivity_conflicts.append("题面固定龙头速度")
    if "member_width_m" in guardrails and any(
        re.search(
            r"板宽[^。；;\n]{0,45}(?:[±＋]\s*\d+(?:\.\d+)?\s*%|扰动|变化)",
            sentence,
        )
        and not denies_all_fixed_source_perturbations(sentence)
        and not re.search(
            r"(?:不动|不扰动|不可扰动|不得扰动|不作扰动|禁止扰动|保持固定)"
            r"[^。；;\n]{0,24}板宽|"
            r"板宽[^。；;\n]{0,35}"
            r"(?:保持固定|固定|不扰动|不可扰动|不得扰动|不作扰动)",
            sentence,
        )
        for sentence in sensitivity_sentences
    ):
        fixed_sensitivity_conflicts.append("题面固定板宽")
    if guardrails.get("initial_spiral") and any(
        re.search(
            r"(?:问题\s*1(?:\s*/\s*2)?[^。；;\n]{0,25}螺距|"
            r"螺距[^。；;\n]{0,25}问题\s*1(?:\s*/\s*2)?)"
            r"[^。；;\n]{0,30}(?:[±＋]\s*\d+(?:\.\d+)?\s*%|扰动|变化)",
            sentence,
        )
        and not denies_all_fixed_source_perturbations(sentence)
        and not re.search(
            r"(?:问题\s*1(?:\s*/\s*2)?[^。；;\n]{0,25}螺距|"
            r"螺距[^。；;\n]{0,25}问题\s*1(?:\s*/\s*2)?)"
            r"[^。；;\n]{0,30}(?:不扰动|不作扰动|保持固定|固定源常量)",
            sentence,
        )
        for sentence in sensitivity_sentences
    ):
        fixed_sensitivity_conflicts.append("题面固定问题一/二螺距")
    source_has_named_initial_point = re.search(
        r"(?:初始时|初态|初始位置)[^。；\n]{0,70}"
        r"(?:[A-ZＡ-Ｚ]\s*点|点\s*[A-ZＡ-Ｚ])",
        source_text,
        re.IGNORECASE,
    )
    if source_has_named_initial_point and any(
        re.search(
            r"(?:问题\s*1[^。；;\n]{0,24})?"
            r"(?:螺线起点(?:相位)?|初始相位|初角|初始极角|"
            r"θ\s*_?\s*0|theta\s*_?\s*0)"
            rf"[^。；;\n]{{0,55}}(?:{perturbation_marker}|∈|范围|扫描)",
            sentence,
            re.IGNORECASE,
        )
        and not denies_all_fixed_source_perturbations(sentence)
        for sentence in sensitivity_sentences
    ):
        fixed_sensitivity_conflicts.append("题面固定命名初始点/初始相位")
    source_q3_match = re.search(
        r"问题\s*3(?P<body>.*?)(?=问题\s*4|$)", source_text, re.DOTALL
    )
    source_q3_has_initial = source_q3_match and re.search(
        r"(?:初始|初态|起点|第\s*\d+\s*圈)", source_q3_match.group("body")
    )
    if not source_q3_has_initial and any(
        re.search(
            r"问题\s*3[^。；;\n]{0,25}"
            r"(?:初始(?:位置|半径)|初态(?:位置|半径)?|起点(?:位置|半径)?)"
            r"[^。；;\n]{0,25}(?:[±＋]\s*\d+(?:\.\d+)?\s*%|扰动|变化)",
            sentence,
        )
        and not denies_all_fixed_source_perturbations(sentence)
        for sentence in sensitivity_sentences
    ):
        fixed_sensitivity_conflicts.append("问题三未授权初态")
    source_has_turning_diameter = re.search(
        r"调头空间[^。；\n]{0,40}直径(?:为)?\s*\d+(?:\.\d+)?\s*m",
        source_text,
    )
    if source_has_turning_diameter and any(
        re.search(
            r"(?:(?:调头空间)?直径|调头(?:空间)?圆?半径)"
            r"[^。；;\n]{0,30}"
            r"(?:[±＋]\s*\d+(?:\.\d+)?\s*%|扰动|变化)",
            sentence,
        )
        and not denies_all_fixed_source_perturbations(sentence)
        and not re.search(
            r"调头(?:空间)?圆?半径[^。；;\n]{0,18}固定"
            r"[^。；;\n]{0,8}(?:仅|只)[^。；;\n]{0,24}"
            r"(?:残差|数值求解)[^。；;\n]{0,12}容差",
            sentence,
        )
        for sentence in sensitivity_sentences
    ):
        fixed_sensitivity_conflicts.append("题面固定调头空间直径")
    source_has_speed_limit = re.search(
        r"速度[^。；\n]{0,25}(?:不超过|≤|<=)\s*\d+(?:\.\d+)?\s*m\s*/\s*s",
        source_text,
        re.IGNORECASE,
    )
    if source_has_speed_limit and any(
        re.search(
            r"(?:速度限制|限速)[^。；;\n]{0,45}"
            r"(?:[±＋]\s*\d+(?:\.\d+)?\s*%|扰动|变化)",
            sentence,
        )
        and not denies_all_fixed_source_perturbations(sentence)
        for sentence in sensitivity_sentences
    ):
        fixed_sensitivity_conflicts.append("题面固定速度上限")
    if fixed_sensitivity_conflicts:
        raise QualityGateError(
            "敏感性分析不得扰动" + "、".join(fixed_sensitivity_conflicts) + "；"
            "只能检查来源不确定量、允许决策域或数值容差"
        )

    source_ques1 = re.search(
        r"问题\s*1(?P<body>.*?)(?=问题\s*2|$)", source_text, re.DOTALL
    )
    ques1_text = value.get("ques1", "") if isinstance(value, dict) else ""
    if source_ques1 and ques1_text:
        initial_to_end = re.search(
            r"从\s*初始时刻\s*到\s*(?P<end>\d+)\s*s?.{0,80}?每秒",
            source_ques1.group("body"),
            re.DOTALL | re.IGNORECASE,
        )
        if initial_to_end:
            end_second = int(initial_to_end.group("end"))
            expected_groups = end_second + 1
            starts_at_one = re.search(
                rf"每秒\s*1\s*(?:[～~—-]|到|至)\s*{end_second}\s*s?",
                ques1_text,
                re.IGNORECASE,
            )
            for declared_match in re.finditer(
                r"共\s*(\d+)\s*(?:组(?!\s*(?:位置|速度))|帧|个时刻|个时间点)",
                ques1_text,
            ):
                declaration_context = ques1_text[
                    max(0, declared_match.start() - 80) : declared_match.end() + 40
                ]
                if re.search(r"论文|正文|论文表|正文表", declaration_context):
                    continue
                declared_groups = int(declared_match.group(1))
                if declared_groups != expected_groups:
                    raise QualityGateError(
                        f"逐秒输出遗漏端点：从初始时刻 0 s 到 {end_second} s"
                        f"（含两端）应有 {expected_groups} 组，而非 {declared_groups} 组"
                    )
            if starts_at_one:
                raise QualityGateError(
                    f"逐秒输出遗漏 0 s：题面要求从初始时刻到 {end_second} s，"
                    f"时间索引必须包含 0，共 {expected_groups} 个时刻"
                )
        uses_increasing_spiral_primitive = re.search(
            r"s\s*\(\s*t\s*\)\s*=\s*(?:t\s*\+\s*s\s*0|s\s*0\s*\+\s*t)"
            r"[^。；\n]{0,80}(?:s\s*0[^。；\n]{0,35}(?:θ|theta)[^。；\n]{0,20}"
            r"弧长(?:积分|原函数)|弧长(?:积分|原函数)[^。；\n]{0,35}s\s*0)",
            ques1_text,
            re.IGNORECASE,
        )
        if uses_increasing_spiral_primitive and re.search(
            r"(?:顺时针)?盘入", source_ques1.group("body")
        ):
            raise QualityGateError(
                "盘入方向与螺线弧长原函数符号矛盾：若 s0=S(θ0) 且 "
                "S(θ) 随正半径支路 θ 增大，则恒速向内盘入应满足 "
                "S(θ(t))=s0-vt，而不是 s0+vt"
            )
        wrong_inward_arc_increment = re.search(
            r"S\s*\(\s*θ?\s*\(\s*t\s*\)\s*\)\s*-\s*"
            r"S\s*\(\s*θ?\s*\(\s*t\s*-\s*(?:1|Δ\s*t)\s*\)\s*\)\s*"
            r"=\s*\+?\s*(?:1(?:\.0+)?\s*m|v\s*(?:Δ\s*t)?)",
            ques1_text,
            re.IGNORECASE,
        )
        wrong_inward_arc_residual = re.search(
            r"(?:[|｜]|abs\s*\()\s*S\s*\(\s*θ?\s*\(\s*t\s*\)\s*\)\s*-\s*"
            r"S\s*\(\s*θ?\s*\(\s*t\s*-\s*(?:1|Δ\s*t)\s*\)\s*\)\s*-\s*"
            r"(?:1(?:\.0+)?\s*m|v\s*(?:Δ\s*t)?)\s*(?:[|｜]|\))",
            ques1_text,
            re.IGNORECASE,
        )
        wrong_forward_inward_arc_increment = re.search(
            r"S\s*\(\s*θ?\s*\(\s*t\s*\+\s*(?:1|Δ\s*t)\s*\)\s*\)\s*[-−]\s*"
            r"S\s*\(\s*θ?\s*\(\s*t\s*\)\s*\)\s*"
            r"=\s*\+?\s*(?:1(?:\.0+)?\s*m|v\s*(?:[·*]?\s*Δ\s*t)?)",
            ques1_text,
            re.IGNORECASE,
        )
        if (
            wrong_inward_arc_increment
            or wrong_inward_arc_residual
            or wrong_forward_inward_arc_increment
        ) and re.search(
            r"(?:顺时针)?盘入", source_ques1.group("body")
        ):
            raise QualityGateError(
                "盘入时弧长闭合检验的差分方向写反：若 "
                "S(θ(t))=S(θ_0)-vt，则应检查 "
                "S(θ(t-1))-S(θ(t))=vΔt，或对差值取绝对值"
            )
        one_sided_arc_residual = re.search(
            r"(?<![|｜])(?:[|｜])\s*S\s*\([^。；\n]{0,35}\)\s*[-−]\s*"
            r"S\s*\([^。；\n]{0,35}\)\s*(?:[|｜])\s*[-−]\s*"
            r"(?:1(?:\.0+)?(?:\s*[·*]?\s*Δ\s*t)?|"
            r"v\s*(?:[·*]?\s*Δ\s*t)?)\s*(?:<=|≤|<)",
            ques1_text,
            re.IGNORECASE,
        )
        if one_sided_arc_residual:
            raise QualityGateError(
                "弧长闭合残差只写成 |ΔS|-vΔt≤tol，是单边不等式，"
                "负的大残差也会误通过；必须检查 ||ΔS|-vΔt|≤tol"
            )
        solves_periodic_node_root = re.search(
            r"(?:牛顿|二分|求根)[^。；\n]{0,45}θ\s*_?\{?i\s*\+\s*1\}?|"
            r"θ\s*_?\{?i\s*\+\s*1\}?[^。；\n]{0,45}(?:牛顿|二分|求根)",
            ques1_text,
            re.IGNORECASE,
        )
        controls_root_branch = re.search(
            r"(?:最近根|第一(?:个)?符号变化|局部(?:增量)?括界|根分支|"
            r"上一时刻[^。；\n]{0,20}(?:连续)?延拓|连续延拓|"
            r"相邻时刻[^。；\n]{0,25}(?:位移|速度积分))",
            ques1_text,
            re.IGNORECASE,
        )
        if solves_periodic_node_root and not controls_root_branch:
            raise QualityGateError(
                "周期曲线的节点距离方程存在多根：仅写牛顿/二分求 θ_{i+1} "
                "不能保证物理最近根；须声明局部增量括界、上一时刻连续延拓，"
                "并以位移--速度积分一致性阻断跳支"
            )
        inverses_follower_parameter = re.search(
            r"(?:节点|把手)\s*k[^。；\n]{0,120}(?:反解|求解)"
            r"[^。；\n]{0,20}(?:θ|theta)|"
            r"(?:反解|求解)[^。；\n]{0,20}(?:θ|theta)\s*[_ₖk]"
            r"[^。；\n]{0,80}(?:节点|把手)",
            ques1_text,
            re.IGNORECASE,
        )
        has_local_follower_root_rule = re.search(
            r"最近根|第一(?:个)?符号变化|局部(?:增量)?括界",
            ques1_text,
            re.IGNORECASE,
        )
        has_temporal_follower_continuation = re.search(
            r"上一时刻[^。；\n]{0,20}(?:连续)?延拓|连续延拓|"
            r"相邻时刻[^。；\n]{0,25}(?:位移|速度积分)",
            ques1_text,
            re.IGNORECASE,
        )
        if inverses_follower_parameter and not (
            has_local_follower_root_rule and has_temporal_follower_continuation
        ):
            raise QualityGateError(
                "周期螺线上的跟随节点反解存在多根；只写外侧序关系仍可能跳到"
                "远处圈。须同时指定允许方向的局部最近根括界和上一时刻连续延拓，"
                "并以位移--速度积分一致性检查阻断跳支"
            )
        if re.search(
            r"v[_ ]?i\s*=\s*v[_ ]?(?:i\s*[-−]\s*1|\{i-1\})"
            r"[^。；\n]{0,25}(?:ds|弧长)[^。；\n]{0,35}/"
            r"[^。；\n]{0,20}(?:ds|弧长)",
            ques1_text,
            re.IGNORECASE,
        ):
            raise QualityGateError(
                "速度传播公式遗漏刚性弦长约束的方向投影：不能只取两节点局部"
                "弧长度量之比；须对二维距离约束求导并保留弦向量与两端切向量点积"
            )

    source_ques2 = re.search(
        r"问题\s*2(?P<body>.*?)(?=问题\s*3|$)", source_text, re.DOTALL
    )
    ques2_text = value.get("ques2", "") if isinstance(value, dict) else ""
    for half_step_claim in re.finditer(
        r"(?:时间步(?:长)?|Δ\s*t)[^。；\n]{0,10}(?:减半|缩小一半)"
        r"[^。；\n]{0,12}[（(]\s*(?P<before>\d+(?:\.\d+)?)\s*"
        r"(?:→|->|至|到)\s*(?P<after>\d+(?:\.\d+)?)\s*s?\s*[）)]",
        ques2_text,
        re.IGNORECASE,
    ):
        before = float(half_step_claim.group("before"))
        after = float(half_step_claim.group("after"))
        if not math.isclose(after, before / 2, rel_tol=1e-9, abs_tol=1e-12):
            raise QualityGateError(
                "时间步长收敛说明的文字与数值不一致：声称减半时，"
                f"{before:g} 应变为 {before / 2:g}，而非 {after:g}"
            )
    if (
        source_ques2
        and ques2_text
        and not re.search(r"300\s*s", source_ques2.group("body"), re.IGNORECASE)
        and not re.search(
            r"(?:不限于|不止于|不局限于)\s*300\s*s", ques2_text, re.IGNORECASE
        )
        and re.search(
            r"若\s*300\s*s?[^。；\n]{0,30}无碰撞[^。；\n]{0,30}"
            r"(?:(?:终止|停止)[^。；\n]{0,15}300\s*s?|"
            r"(?:报告|取|记)[^。；\n]{0,15}300\s*s?[^。；\n]{0,12}"
            r"(?:暂定)?(?:终点|终止时刻))|"
            r"定义域[^。；\n]{0,20}(?:\[\s*0\s*,\s*300|300\s*s)|"
            r"(?:必须|应)[^。；\n]{0,25}(?:Q\s*1|问题\s*[一1])"
            r"[^。；\n]{0,15}300\s*s?[^。；\n]{0,15}(?:内)?(?:被)?确定",
            ques2_text,
            re.IGNORECASE,
        )
    ):
        raise QualityGateError(
            "问题二的首次碰撞搜索不能沿用问题一的 300 s 输出截点；"
            "题面未给该上限，若 300 s 内未碰撞必须继续搜索并包围真实首次事件"
        )
    source_requires_initial_collision_event = source_ques2 and re.search(
        r"碰撞[^。；\n]{0,100}(?:终止时刻|不能再继续盘入)|"
        r"(?:终止时刻|不能再继续盘入)[^。；\n]{0,100}碰撞",
        source_ques2.group("body"),
        re.IGNORECASE,
    )
    covers_initial_collision_interval = re.search(
        r"(?:从|自)\s*(?:题面)?(?:初始时刻|t\s*=\s*0|0\s*s)"
        r"[^。；\n]{0,45}(?:扫描|检测|回放|推进|SAT)|"
        r"(?:扫描|检测|回放|推进|SAT)[^。；\n]{0,45}"
        r"(?:从|自)\s*(?:题面)?(?:初始时刻|t\s*=\s*0|0\s*s)|"
        r"t\s*(?:=|从)?\s*0\s*(?:开始)?[^。；\n]{0,20}"
        r"(?:扫描|检测|回放|推进|SAT)|"
        r"(?:时间|时刻)[^。；\n]{0,12}(?:扫描|检测|回放|推进)?[^。；\n]{0,8}"
        r"(?:从|自)\s*0\s*(?:s)?\s*开始|"
        r"0\s*(?:≤|<=)\s*t\s*(?:<|≤|<=)\s*t\s*[_ ]?c",
        ques2_text,
        re.IGNORECASE,
    )
    if (
        source_requires_initial_collision_event
        and ques2_text
        and not covers_initial_collision_interval
    ):
        raise QualityGateError(
            "问题二要求首次碰撞/终止事件，但方案未明确从题面初始时刻 t=0 "
            "覆盖检测；必须扫描 0≤t<t_c 的完整前史并证明此前无碰撞，"
            "不能从任意中间时刻开始"
        )
    if re.search(
        r"result2\.xlsx[^。；\n]{0,80}"
        r"(?:0\s*(?:\.\.|…|至|到|[-−])\s*T\s*[_ ]?stop|逐秒)|"
        r"(?:0\s*(?:\.\.|…|至|到|[-−])\s*T\s*[_ ]?stop|逐秒)"
        r"[^。；\n]{0,80}result2\.xlsx",
        ques2_text,
        re.IGNORECASE,
    ):
        raise QualityGateError(
            "问题二的 result2.xlsx 只交付终止时刻的全部把手位置和速度；"
            "不能擅自改成 0..T_stop 的逐秒时间序列表或改变官方模板结构"
        )
    if re.search(
        r"(?:相邻(?:板凳|构件|实体)?(?:连接)?对|排除相邻(?:板凳|构件|实体)?对)"
        r"[^。；\n]{0,80}[（(]\s*223\s*[,，]\s*224\s*[）)]",
        ques2_text,
        re.IGNORECASE,
    ):
        raise QualityGateError(
            "2024A 只有 223 个板体实体，相邻板体索引对最多到 (222,223)；"
            "(223,224) 是末板的两个节点编号，不能作为第 224 个板体参与碰撞对排除"
        )
    finite_collision_search_cap = re.search(
        r"(?:延长|扩大)?(?:搜索|扫描)(?:时间|区间|上限)?\s*(?:至|到|=)"
        r"\s*\d+(?:\.\d+)?\s*s",
        ques2_text,
        re.IGNORECASE,
    )
    continues_beyond_numeric_cap = re.search(
        r"(?:题面|原题)[^。；\n]{0,12}未限时[^。；\n]{0,25}"
        r"(?:继续|逐次|反复)[^。；\n]{0,20}(?:搜索|扫描)"
        r"[^。；\n]{0,25}(?:包围|直到|直至)[^。；\n]{0,15}(?:事件|碰撞)|"
        r"(?:若|如)[^。；\n]{0,25}(?:仍)?(?:无|未)碰撞[^。；\n]{0,25}"
        r"(?:继续|再|逐次|反复|则)[^。；\n]{0,20}(?:延长|扩大|倍增|搜索|扫描)"
        r"[^。；\n]{0,25}(?:直到|直至|包围|碰撞)|"
        r"(?:逐次|反复|倍增)[^。；\n]{0,25}(?:搜索|扫描|时间上限|时间区间)"
        r"[^。；\n]{0,25}(?:直到|直至|包围)[^。；\n]{0,15}(?:事件|碰撞)",
        ques2_text,
        re.IGNORECASE,
    )
    if (
        source_ques2
        and ques2_text
        and finite_collision_search_cap
        and not re.search(r"\d+(?:\.\d+)?\s*s", source_ques2.group("body"))
        and not continues_beyond_numeric_cap
    ):
        raise QualityGateError(
            "问题二的首次碰撞搜索不能止于自行设定的有限时间上限；"
            "固定数值只能作为扩展检查点，若仍未碰撞必须继续扩大时间域，"
            "直至包围真实首次事件或给出无事件的解析证明"
        )

    ques3_text = "\n".join(collect_question_text(value, 3))
    q3_evidence_text = "\n".join(
        text
        for text in (
            ques3_text,
            value.get("sensitivity_analysis", "") if isinstance(value, dict) else "",
        )
        if text
    )
    source_ques3 = re.search(
        r"问题\s*3(?P<body>.*?)(?=问题\s*4|$)", source_text, re.DOTALL
    )
    uses_fixed_spatial_cutoff = re.search(
        r"(?:r\s*[_ ]?start|r\s*[_ ]?upper|R\s*[_ ]?T|截断半径)"
        r"\s*(?:=|取|为|（|\()[^。；\n]{0,18}\d+(?:\.\d+)?\s*m",
        ques3_text,
        re.IGNORECASE,
    )
    certifies_spatial_cutoff = re.search(
        r"(?:扩大|增加|增至|倍增|扩展)[^。；\n]{0,22}"
        r"(?:r\s*[_ ]?(?:start|upper)|R\s*[_ ]?T|截断半径|空间截断)"
        r"[^。；\n]{0,40}(?:临界值|临界\s*p|p\s*[_ ]?min|p\s*\*|最小螺距|结果)"
        r"[^。；\n]{0,20}(?:稳定|不变|变化|偏差\s*[<＜])|"
        r"(?:临界值|临界\s*p|p\s*[_ ]?min|p\s*\*|最小螺距|结果)[^。；\n]{0,25}"
        r"(?:对|随)[^。；\n]{0,20}(?:r\s*[_ ]?(?:start|upper)|截断半径)"
        r"[^。；\n]{0,20}(?:稳定|不敏感)|"
        r"(?:r\s*[_ ]?(?:start|upper)|R\s*[_ ]?T|截断半径)[^。；\n]{0,30}"
        r"(?:扩大|增加|增至|倍增|扩展)[^。；\n]{0,40}"
        r"(?:临界值|临界\s*p|p\s*[_ ]?min|p\s*\*|最小螺距|结果)"
        r"[^。；\n]{0,20}(?:稳定|不变|变化|偏差\s*[<＜])|"
        r"r\s*[_ ]?(?:start|upper)\s*=\s*\d+(?:\.\d+)?\s*"
        r"(?:/|、|,)\s*\d+(?:\.\d+)?\s*m[^。；\n]{0,30}"
        r"(?:复算|对比)[^。；\n]{0,35}(?:截断)?(?:稳定|不变|变化)|"
        r"(?:外侧)?截断半径[^。；\n]{0,20}(?:增加|扩大)"
        r"[^。；\n]{0,20}(?:p\s*[_ ]?min|最小螺距)"
        r"[^。；\n]{0,20}(?:影响|变化|偏差)[^。；\n]{0,15}[<＜]|"
        r"r\s*[_ ]?(?:start|upper)[^。；\n]{0,25}(?:和|与|及)"
        r"[^。；\n]{0,12}2\s*[×xX*]\s*r\s*[_ ]?(?:start|upper)"
        r"[^。；\n]{0,55}(?:分别)?(?:(?:重新)?求(?:解|得|出)?|复算|重算)"
        r"[^。；\n]{0,15}(?:p\s*[_ ]?min|最小螺距)"
        r"[^。；\n]{0,25}(?:差值|偏差|变化)[^。；\n]{0,15}[<＜]|"
        r"r\s*[_ ]?(?:start|upper)\s*(?:∈|=|取)?\s*"
        r"[（({[]?\s*\d+(?:\.\d+)?\s*(?:m\s*)?"
        r"(?:[/、,，]\s*\d+(?:\.\d+)?\s*(?:m\s*)?){1,4}"
        r"[）)}\]]?[^。；\n]{0,45}(?:p\s*[_ ]?min|p\s*\*|最小螺距)"
        r"[^。；\n]{0,20}(?:差值|偏差|变化|差)[^。；\n]{0,12}[<＜]|"
        r"(?:两个|两[档挡]|至少两个)[^。；\n]{0,25}(?:逐步)?(?:扩大|递增)"
        r"[^。；\n]{0,25}r\s*[_ ]?(?:start|upper)"
        r"[^。；\n]{0,90}(?:分别)?(?:重新)?求(?:解|得|出)?"
        r"[^。；\n]{0,18}(?:p\s*[_ ]?min|最小螺距)"
        r"[^。；\n]{0,45}(?:差值|偏差|变化|差)[^。；\n]{0,15}[<＜]",
        q3_evidence_text,
        re.IGNORECASE,
    )
    certifies_spatial_cutoff = certifies_spatial_cutoff or re.search(
        r"r\s*[_ ]?(?:start|upper)[^。；\n]{0,25}增大\s*(?:一|1)\s*倍"
        r"[^。；\n]{0,35}(?:计算|复算|求解)[^。；\n]{0,18}"
        r"(?:p\s*[_ ]?min|最小螺距)[^。；\n]{0,25}"
        r"(?:结果)?(?:差值|偏差|变化|差)\s*[≤<＜]",
        q3_evidence_text,
        re.IGNORECASE,
    )
    certifies_spatial_cutoff = certifies_spatial_cutoff or re.search(
        r"(?:两个|两[档挡]|至少两个)[^。；\n]{0,60}(?:逐步)?(?:扩大|递增)"
        r"[^。；\n]{0,60}r\s*[_ ]?(?:start|upper)"
        r"[^。；\n]{0,180}(?:分别)?(?:重新)?求(?:解|得|出)?"
        r"[^。；\n]{0,40}(?:p\s*[_ ]?min|最小螺距)"
        r"[^。；\n]{0,180}(?:差值|偏差|变化|差)"
        r"[^。；\n]{0,80}[<＜]",
        q3_evidence_text,
        re.IGNORECASE,
    )
    certifies_spatial_cutoff = certifies_spatial_cutoff or re.search(
        r"r\s*[_ ]?(?:start|upper)[^。；\n]{0,18}取?两[档挡]"
        r"[^。；\n]{0,30}\d+(?:\.\d+)?\s*[/、,，]\s*"
        r"\d+(?:\.\d+)?\s*m[^。；\n]{0,35}"
        r"(?:p\s*[_ ]?min|最小螺距)[^。；\n]{0,18}"
        r"(?:差值|偏差|变化|差)[^。；\n]{0,12}[<＜]",
        q3_evidence_text,
        re.IGNORECASE,
    )
    if uses_fixed_spatial_cutoff and not certifies_spatial_cutoff:
        raise QualityGateError(
            "问题三把单个外侧半径称为充分截断，却未验证空间截断稳定性；"
            "须扩大 r_start/r_upper 后复算并证明临界螺距稳定，"
            "不能用固定 60 m 等经验值代替无限外侧安全域证据"
        )
    source_authorizes_q1 = source_ques3 and re.search(
        r"(?:问题\s*1|第一问)", source_ques3.group("body")
    )
    initial_spiral = guardrails.get("initial_spiral")
    if ques3_text and initial_spiral:
        q1_turn_count = int(initial_spiral["turn_count"])
        uses_q1_turn_in_q3 = re.search(
            rf"第\s*{q1_turn_count}\s*圈|"
            rf"(?:r[_ ]?(?:init|start)|r[_ ]?0|初始半径)[^。；\n]{{0,18}}"
            rf"{q1_turn_count}\s*(?:[×xX*·]\s*)?p|"
            rf"(?:起点|起始圈数)[^。；\n]{{0,30}}(?:问题\s*1|继承)[^。；\n]{{0,20}}{q1_turn_count}|"
            rf"(?:问题\s*1|继承)[^。；\n]{{0,30}}(?:起点|起始圈数)[^。；\n]{{0,20}}{q1_turn_count}",
            ques3_text,
        )
        denies_q1_turn_in_q3 = re.search(
            rf"(?:不|未|不得|不能)[^。；\n]{{0,8}}"
            rf"(?:继承|采用|使用|假设|设为|取为)"
            rf"[^。；\n]{{0,20}}第?\s*{q1_turn_count}\s*圈|"
            rf"(?:不以|未以|不得以|不能以)[^。；\n]{{0,20}}"
            rf"第?\s*{q1_turn_count}\s*圈[^。；\n]{{0,25}}"
            r"(?:初态|初始|起点)",
            ques3_text,
        )
        if uses_q1_turn_in_q3 and not source_authorizes_q1 and not denies_q1_turn_in_q3:
            raise QualityGateError(
                f"问题三擅自采用问题一的第{q1_turn_count}圈起点；原问题三未引用问题一，"
                "不得以“题面未改”为由继承具体圈数，应覆盖充分外侧安全域并验证上界稳定"
            )
        q1_pitch = float(initial_spiral["pitch_m"])
        q1_pitch_pattern = re.escape(f"{q1_pitch:g}")
        uses_q1_pitch_as_q3_bound = re.search(
            rf"(?:p(?:[_ ]?min)?|最小螺距)[^。；\n]{{0,30}}"
            rf"(?:∈|>|≥|<|≤|上界|下界)[^。；\n]{{0,18}}{q1_pitch_pattern}|"
            rf"(?:候选域|搜索域|二分(?:法)?|区间)[^。；\n]{{0,45}}"
            rf"(?:\(\s*0\s*[,，]\s*{q1_pitch_pattern}\s*[\]\)]|"
            rf"{q1_pitch_pattern}[^。；\n]{{0,12}}(?:上界|高界))",
            ques3_text,
            re.IGNORECASE,
        )
        independently_proves_bound = re.search(
            rf"(?:SAT|有向矩形|实体碰撞|全链)[^。；\n]{{0,80}}"
            rf"{q1_pitch_pattern}[^。；\n]{{0,30}}(?:不可行|严格下界)|"
            rf"{q1_pitch_pattern}[^。；\n]{{0,30}}(?:不可行|严格下界)"
            r"[^。；\n]{0,80}(?:SAT|有向矩形|实体碰撞|全链)",
            ques3_text,
            re.IGNORECASE,
        )
        if (
            uses_q1_pitch_as_q3_bound
            and not source_authorizes_q1
            and not independently_proves_bound
        ):
            raise QualityGateError(
                "问题三把问题一的螺距数值用作优化域边界；原问题三未引用问题一，"
                "证据合同也必须保持跨问参数隔离。须覆盖完整正参数域，并独立证明"
                "计算上界充分或实体不可行下界"
            )
        q1_radius_pattern = re.escape(f"{float(initial_spiral['radius_m']):g}")
        uses_q1_radius_as_q3_initial = re.search(
            rf"(?:初始(?:位置|半径|状态)?|初态|起点|从)"
            rf"[^。；\n]{{0,18}}{q1_radius_pattern}\s*m",
            ques3_text,
            re.IGNORECASE,
        )
        source_authorizes_q3_initial = source_ques3 and re.search(
            r"(?:初始(?:位置|半径|圈数|状态)?|起点|从[^。；]{0,20}(?:圈|半径))",
            source_ques3.group("body"),
        )
        if (
            uses_q1_radius_as_q3_initial
            and not source_authorizes_q1
            and not source_authorizes_q3_initial
        ):
            raise QualityGateError(
                "问题三把问题一的初始半径用作本问初态；原问题三未给固定起点，"
                "应覆盖充分外侧计算域并通过扩大上界验证临界结果稳定"
            )
        borrows_q1_feasibility = re.search(
            rf"(?:问题\s*1|第一问)[^。；\n]{{0,45}}{q1_pitch_pattern}\s*m?"
            rf"[^。；\n]{{0,45}}(?:应)?可行|"
            rf"{q1_pitch_pattern}\s*m?[^。；\n]{{0,45}}(?:问题\s*1|第一问)"
            rf"[^。；\n]{{0,45}}(?:应)?可行|"
            rf"{q1_pitch_pattern}\s*m?[^。；\n]{{0,55}}"
            r"(?:问题\s*1|第一问)[^。；\n]{0,35}(?:下界证据|可行证据)|"
            rf"{q1_pitch_pattern}\s*m?[^。；\n]{{0,28}}"
            r"(?:应可行(?:[^。；\n]{0,12}(?:下界证据|可行证据))?|"
            r"(?:作为|提供)[^。；\n]{0,12}(?:下界证据|可行证据))",
            ques3_text,
            re.IGNORECASE,
        )
        independently_replays_q3_to_boundary = re.search(
            rf"{q1_pitch_pattern}\s*m?[^。；\n]{{0,80}}"
            r"(?:本问|问题\s*3)[^。；\n]{0,30}(?:完整|连续)"
            r"[^。；\n]{0,30}(?:回放|仿真|SAT)[^。；\n]{0,30}(?:4\.5|边界)",
            ques3_text,
            re.IGNORECASE,
        )
        explicitly_rejects_q1_parameter_borrow = re.search(
            rf"(?:不继承|不得继承|不沿用|不采用)[^。；\n]{{0,35}}"
            rf"(?:问题\s*1|第一问)[^。；\n]{{0,30}}{q1_pitch_pattern}\s*m?",
            ques3_text,
            re.IGNORECASE,
        )
        if (
            borrows_q1_feasibility
            and not source_authorizes_q1
            and not independently_replays_q3_to_boundary
            and not explicitly_rejects_q1_parameter_borrow
        ):
            raise QualityGateError(
                "问题一只验证其规定时间段，不能据此宣称同一螺距在问题三到达4.5 m边界前可行；"
                "问题三的可行/不可行括界必须独立回放至边界并复核全链实体碰撞"
            )

    if ques3_text and re.search(
        r"p[_ ]?max[^。；\n]{0,30}1\.7\s*m?[^。；\n]{0,35}"
        r"(?:问题\s*4|第四问|参考值)|"
        r"(?:问题\s*4|第四问)[^。；\n]{0,35}1\.7\s*m?"
        r"[^。；\n]{0,30}p[_ ]?max",
        ques3_text,
        re.IGNORECASE,
    ):
        raise QualityGateError(
            "问题三不能把问题四的1.7 m路径参数借作本问 p_max 来源；"
            "数值初始括界须明确为无来源偏好的计算种子，并通过独立扩展获得可行括界"
        )

    source_ques4 = re.search(
        r"问题\s*4(?P<body>.*?)(?=问题\s*5|$)", source_text, re.DOTALL
    )
    if (
        source_ques4
        and re.search(
            r"盘入螺线[^。；\n]{0,20}螺距(?:为)?\s*\d", source_ques4.group("body")
        )
        and re.search(
            r"问题\s*4[^。；\n]{0,28}"
            r"(?<!不)(?<!未)(?<!不会)(?<!不得)(?<!不能)(?:继承|沿用)[^。；\n]{0,20}"
            r"问题\s*3[^。；\n]{0,20}(?:临界螺距|最小螺距|p\s*\*)"
            r"(?:[^。；\n]{0,8}构型)?|"
            r"问题\s*4[^。；\n]{0,28}"
            r"(?<!不)(?<!未)(?<!不会)(?<!不得)(?<!不能)(?:继承|沿用)[^。；\n]{0,20}"
            r"问题\s*3(?:搜索到|所得|的)?[^。；\n]{0,8}(?:临界)?构型",
            text,
            re.IGNORECASE,
        )
    ):
        raise QualityGateError(
            "问题四只继承问题三的调头空间，不继承问题三临界螺距构型；"
            "题面已为问题四另行固定盘入螺距，必须按本问路径重新建立初态"
        )

    chinese_counts = {
        "一": 1,
        "二": 2,
        "三": 3,
        "四": 4,
        "五": 5,
        "六": 6,
        "七": 7,
        "八": 8,
        "九": 9,
        "十": 10,
    }
    for question_number in range(1, 10):
        source_question = re.search(
            rf"问题\s*{question_number}(?P<body>.*?)(?=问题\s*{question_number + 1}|$)",
            source_text,
            re.DOTALL,
        )
        if not source_question:
            continue
        listed_times = re.search(
            r"(?:论文|正文)[^。；\n]{0,30}(?:给出|列出)"
            r"(?P<body>[^。；\n]{0,220}?)(?:时刻?|处)",
            source_question.group("body"),
            re.IGNORECASE,
        )
        if not listed_times:
            continue
        expected_time_count = len(
            re.findall(
                r"(?<![\d.])(?:[-−]\s*)?\d+(?:\.\d+)?\s*s",
                listed_times.group("body"),
                re.IGNORECASE,
            )
        )
        if expected_time_count < 2:
            continue
        scoped_text = "\n".join(collect_question_text(value, question_number))
        declared_time_claims = list(
            re.finditer(
                r"(?:论文|正文|指定)[^。；\n]{0,140}?"
                r"(?P<count>[一二三四五六七八九十])\s*个?时刻(?:表|数据|结果)?",
                scoped_text,
                re.IGNORECASE,
            )
        )
        declared_time_claims.extend(
            re.finditer(
                r"(?:论文|正文|指定)[^。；\n]{0,140}?"
                r"(?P<count>\d+)\s*个时刻(?:表|数据|结果)?",
                scoped_text,
                re.IGNORECASE,
            )
        )
        for claim in declared_time_claims:
            raw_count = claim.group("count")
            declared_time_count = (
                int(raw_count) if raw_count.isdigit() else chinese_counts[raw_count]
            )
            if declared_time_count != expected_time_count:
                raise QualityGateError(
                    f"问题{question_number}的论文抽取清单列举了 "
                    f"{expected_time_count} 个时刻，而非 {declared_time_count} 个；"
                    "附件逐秒网格与论文点名时刻必须分开计数"
                )

    selected_nodes_match = re.search(
        r"龙头前把手.{0,160}?第\s*1[、,，]\s*51[、,，]\s*101"
        r"[、,，]\s*151[、,，]\s*201[^。]{0,80}?龙尾后把手",
        source_text,
        re.DOTALL,
    )
    for question_number in range(1, 10):
        source_question = re.search(
            rf"问题\s*{question_number}(?P<body>.*?)(?=问题\s*{question_number + 1}|$)",
            source_text,
            re.DOTALL,
        )
        if not source_question:
            continue
        source_body = source_question.group("body")
        result_file = re.search(r"result\d+\.xlsx", source_body, re.IGNORECASE)
        expects_full_team = re.search(
            r"(?:整个舞龙队|舞龙队)的位置和速度",
            source_body,
        )
        if not result_file or not expects_full_team:
            continue
        scoped_text = "\n".join(collect_question_text(value, question_number))
        file_pattern = re.escape(result_file.group(0))
        reduces_attachment_to_selected_objects = re.search(
            rf"{file_pattern}[^。；;\n]{{0,100}}"
            r"(?:每行\s*)?7\s*(?:列|个(?:对象|把手|节点))|"
            r"7\s*个(?:对象|把手|节点)"
            rf"[^。；;\n]{{0,80}}{file_pattern}",
            scoped_text,
            re.IGNORECASE,
        )
        if reduces_attachment_to_selected_objects:
            raise QualityGateError(
                f"{result_file.group(0)}必须填写题面要求的整个舞龙队"
                "所有把手；论文中的7个代表对象只是正文抽取清单，"
                "不能缩减官方结果附件"
            )
        attachment_contract_text = scoped_text
        if isinstance(value, dict) and isinstance(value.get("eda"), str):
            attachment_contract_text += "\n" + value["eda"]
        flattens_official_template = re.search(
            rf"{file_pattern}[^。；;\n]{{0,120}}(?:"
            r"224\s*[×xX*]\s*2[^。；;\n]{0,25}[+＋]\s*224|"
            r"301\s*[×xX*]\s*(?:672|\(\s*224\s*[×xX*]\s*2\s*[+＋]\s*224\s*\))|"
            r"301\s*行[^。；;\n]{0,80}672\s*列|"
            r"301\s*行[^。；;\n]{0,20}[×xX*]\s*224\s*(?:个)?节点|"
            r"224\s*(?:个)?节点[^。\n]{0,80}(?:共\s*)?301\s*行)",
            attachment_contract_text,
            re.IGNORECASE,
        )
        if flattens_official_template:
            raise QualityGateError(
                f"{result_file.group(0)} 是官方结果模板，不能把位置与速度扁平化改成"
                "自造的逐时刻单表或改写行列方向；必须保留原工作表、行列、表头和样式，"
                "只回填空白数值格"
            )
        invents_velocity_components = re.search(
            rf"{file_pattern}[^。；;\n]{{0,120}}"
            r"(?:速度|velocity)\s*[（(][^）)\n]{0,12}"
            r"v\s*[_ ]?x[^）)\n]{0,12}v\s*[_ ]?y[^）)\n]{0,12}[）)]",
            scoped_text,
            re.IGNORECASE,
        )
        source_authorizes_velocity_components = re.search(
            r"(?:速度|velocity)[^。；;\n]{0,25}"
            r"v\s*[_ ]?x[^。；;\n]{0,18}v\s*[_ ]?y",
            source_body,
            re.IGNORECASE,
        )
        if invents_velocity_components and not source_authorizes_velocity_components:
            raise QualityGateError(
                f"{result_file.group(0)} 的官方模板列结构必须由附件读取；"
                "题面未授权自行把速度扩成 vx/vy 分量，须保留原工作表、"
                "表头和标量/分量口径，只回填已有空白数值格"
            )
    if selected_nodes_match:
        expected_selected_nodes = 7
        selected_count_patterns = (
            r"(?:共|给出|列出)\s*(\d+)\s*个?(?:关键|指定)(?:节点|把手|对象)",
            r"(?:论文|正文|表(?:格)?\s*\d*).{0,40}?"
            r"(?:共|给出|列出)\s*(\d+)\s*个?(?:节点|把手|对象)",
        )
        for question_number in (1, 2, 4):
            scoped_text = "\n".join(collect_question_text(value, question_number))
            for pattern in selected_count_patterns:
                for claim in re.finditer(pattern, scoped_text):
                    declared = int(claim.group(1))
                    if declared != expected_selected_nodes:
                        raise QualityGateError(
                            f"论文指定对象数量错误：题面每个所列时刻要求龙头、5个龙身把手"
                            f"和龙尾后把手，共 {expected_selected_nodes} 个对象，而非 {declared} 个"
                        )
        for field in text_fields:
            for mapping in re.finditer(
                r"节点\s*(\d+)\s*[（(][^）)\n]{0,24}?"
                r"第\s*(\d+)\s*节龙身前(?:把手)?",
                field,
            ):
                declared_node = int(mapping.group(1))
                body_number = int(mapping.group(2))
                expected_node = body_number + 1
                if declared_node != expected_node:
                    raise QualityGateError(
                        f"论文指定把手索引错位：1基节点拓扑中第 {body_number} 节"
                        f"龙身前把手是节点 {expected_node}，不是节点 {declared_node}"
                    )

    if re.search(
        r"龙头前把手.{0,100}盘入到调头空间(?:的)?边界", source_text, re.DOTALL
    ):
        all_chain_inside_turning_space = (
            r"(?<!非)(?:整条龙|全体舞龙队|所有(?:\s*\d+\s*个)?(?:节点|把手)|"
            r"全部(?:\s*\d+\s*个)?(?:节点|把手))"
            r".{0,100}(?:全部)?(?:进入|位于|限制在).{0,30}"
            r"(?:调头空间|调头圆|圆形区域|圆内)",
            r"(?:最大节点极径|所有节点极径).{0,30}(?:<=|≤)\s*4\.5",
        )
        asserts_all_chain_inside = False
        for pattern in all_chain_inside_turning_space:
            for match in re.finditer(pattern, text, re.IGNORECASE):
                prefix = text[max(0, match.start() - 14) : match.start()]
                if re.search(
                    r"(?:不要求|无需|不必|不得要求|不能要求|未要求|并非)\s*$",
                    prefix,
                ):
                    continue
                asserts_all_chain_inside = True
                break
            if asserts_all_chain_inside:
                break
        if asserts_all_chain_inside:
            raise QualityGateError(
                "问题三擅自扩大题面约束主语：原文只要求龙头前把手到达"
                "调头空间边界，并未要求全链节点同时位于调头圆内"
            )
        invents_q1_inheritance = re.search(
            r"(?:继承|沿用|按).{0,12}问题\s*1.{0,25}(?:初始|第\s*16\s*圈)|"
            r"(?:初始|第\s*16\s*圈).{0,25}(?:继承|沿用|来自).{0,12}问题\s*1|"
            r"问题\s*1.{0,12}(?:初态|初始|第\s*16\s*圈)",
            ques3_text,
        )
        denies_q1_inheritance = re.search(
            r"(?:不|未|不得|不能).{0,6}(?:继承|沿用).{0,12}问题\s*1|"
            r"问题\s*1.{0,12}(?:初态|初始|第\s*16\s*圈).{0,8}(?:不|未)继承",
            ques3_text,
        )
        if (
            invents_q1_inheritance
            and not denies_q1_inheritance
            and not source_authorizes_q1
        ):
            raise QualityGateError(
                "问题三擅自继承问题一的第16圈/初始位置；原问题三未引用问题一，"
                "不得用跨问初态构造螺距下界，应从本问目标及前问已经定义的运动可行性推导"
            )
        source_authorizes_q3_initial = source_ques3 and re.search(
            r"(?:初始(?:位置|半径|圈数|状态)?|起点|从[^。；]{0,20}(?:圈|半径))",
            source_ques3.group("body"),
        )
        source_authorizes_figure4 = source_ques3 and re.search(
            r"(?:图\s*4|figure\s*4)", source_ques3.group("body"), re.IGNORECASE
        )
        claims_figure4_defines_q3_initial = re.search(
            r"(?:初态|初始(?:位置|相位|半径)?)[^。；]{0,45}"
            r"(?:读取|依赖|由|根据)?[^。；]{0,10}(?:图\s*4|figure\s*4)|"
            r"(?:图\s*4|figure\s*4)[^。；]{0,45}"
            r"(?:确定|读取|依赖)[^。；]{0,20}(?:初态|初始)",
            ques3_text,
            re.IGNORECASE,
        )
        if claims_figure4_defines_q3_initial and not source_authorizes_figure4:
            raise QualityGateError(
                "问题三原文只引用本问图示，不能借问题一的图4确定问题三初态/相位；"
                "图4仅约束问题一第16圈A点，问题三须覆盖充分外侧安全域并验证上界稳定性"
            )
        minimum_head_radius_to_fit_chain = re.search(
            r"(?:容纳|布置|排列|整链初始配置)[^。；]{0,45}"
            r"(?:最小(?:龙头)?半径|r[_ ]?min)|"
            r"(?:最小(?:龙头)?半径|r[_ ]?min)[^。；]{0,45}"
            r"(?:容纳|布置|排列|整链|全链|224)",
            ques3_text,
            re.IGNORECASE,
        )
        if minimum_head_radius_to_fit_chain:
            raise QualityGateError(
                "问题三的螺线向外无限延伸，不存在“容纳整链所需的最小龙头半径”；"
                "应从充分外侧截断到目标边界连续回放全部构型，并以扩大截断半径验证稳定性"
            )
        uses_pitch_bound_as_spatial_position = re.search(
            r"(?:充分外侧|外侧起点|启动位置)[^。；\n]{0,28}"
            r"p\s*[_ ]?\s*max\s*处|"
            r"(?:从|在)\s*p\s*[_ ]?\s*max\s*处[^。；\n]{0,25}"
            r"(?:向中心|盘入|作为起点)",
            ques3_text,
            re.IGNORECASE,
        )
        if uses_pitch_bound_as_spatial_position:
            raise QualityGateError(
                "问题三混淆了参数括界与空间截断：p_max 是螺距搜索上界，"
                "不能写成螺线上的“p_max处”；须另设有长度含义的 r_start/r_upper，"
                "并用扩大空间截断后的临界值稳定性核验"
            )
        uses_q3_bisection = re.search(r"二分", ques3_text)
        treats_bisection_as_monotonicity_check = re.search(
            r"二分(?:法|搜索)?\s*(?:或|/|或者)\s*"
            r"(?:全域|完整参数域|全区间)?(?:扫描|搜索)"
            r"[^。；\n]{0,18}(?:确认|验证|核验)"
            r"[^。；\n]{0,18}(?:可行性|可行谓词)[^。；\n]{0,8}单调|"
            r"(?:采用|使用|用|做)?\s*二分(?:法|搜索)?"
            r"[^。；\n]{0,12}(?:确认|验证|核验)"
            r"[^。；\n]{0,18}(?:可行性|可行谓词)[^。；\n]{0,8}单调",
            ques3_text,
            re.IGNORECASE,
        )
        if treats_bisection_as_monotonicity_check:
            raise QualityGateError(
                "问题三把二分搜索本身当作可行性单调性证书；"
                "二分只能在解析证明，或完整参数域扫描确认仅一次"
                "不可行到可行跃迁之后使用，不能与全域扫描并列二选一"
            )
        assumes_q3_monotonicity = re.search(
            r"(?:目标函数|可行性|可行域|可行判据)[^。；\n]{0,24}单调|"
            r"可行[^。；\n]{0,10}(?:减小|降低)\s*p"
            r"[^。；\n]{0,18}不可行[^。；\n]{0,10}(?:增大|提高)\s*p|"
            r"(?:随\s*)?p\s*(?:增大|提高)[^。；\n]{0,12}可行",
            ques3_text,
        )
        certifies_q3_monotonicity = re.search(
            r"(?:可行性|可行域|可行判据)[^。；\n]{0,24}单调性?"
            r"[^。；\n]{0,24}(?:证明|验证|核验|检查)|"
            r"(?:证明|验证|核验|检查|确认)[^。；\n]{0,24}"
            r"(?:可行性|可行域|可行判据)[^。；\n]{0,18}单调|"
            r"(?:完整|全(?:参数)?域|全区间)[^。；\n]{0,25}"
            r"(?:扫描|采样)[^。；\n]{0,35}"
            r"(?:(?:单次|单一)跃迁|无反转|仅一次[^。；\n]{0,8}(?:跃迁|翻转))|"
            r"(?:确认|验证|核验)[^。；\n]{0,18}可行性"
            r"[^。；\n]{0,18}(?:仅有一次|单次)[^。；\n]{0,10}"
            r"(?:跃迁|翻转)|"
            r"(?:解析|理论)[^。；\n]{0,18}证明[^。；\n]{0,18}单调",
            ques3_text,
            re.IGNORECASE,
        )
        if (
            uses_q3_bisection
            and assumes_q3_monotonicity
            and not certifies_q3_monotonicity
        ):
            raise QualityGateError(
                "问题三使用二分搜索前未证明或核验可行性关于螺距的单调性；"
                "目标值 p 本身单调不等于可行谓词单调。须给出解析证明，或先在完整参数域"
                "扫描并确认仅有一次不可行到可行跃迁；若出现反转则改用全域搜索"
            )
        denies_fixed_q3_initial = re.search(
            r"(?:不|未|不得|无需|没有)[^。；]{0,12}"
            r"(?:设定|指定|采用|选取|预设)[^。；]{0,12}"
            r"(?:任何)?(?:固定)?初始(?:位置|半径|圈数|起点)?|"
            r"原文[^。；]{0,12}未(?:给|指定)[^。；]{0,12}(?:起点|初始)",
            ques3_text,
            re.IGNORECASE,
        )
        invents_q3_initial = None
        if not denies_fixed_q3_initial:
            invents_q3_initial = re.search(
                r"(?:某外圈(?:初始)?位置|(?:自行|人为|任意|合理)[^。；]{0,12}初始|"
                r"(?:选|取|指定)[^。；]{0,10}初(?:态|始)(?:位置|相位|半径)?|"
                r"设(?:定)?[^。；]{0,10}初始(?:龙头)?(?:位置)?|"
                r"r0[^。；]{0,35}(?:图\s*5|最外圈)[^。；]{0,25}(?:唯一|确定)|"
                r"图\s*5[^。；]{0,35}(?:唯一|确定)[^。；]{0,20}r0|最外圈)",
                ques3_text,
                re.IGNORECASE,
            )
        proves_initial_independence = re.search(
            r"(?:与|对)[^。；]{0,12}初始(?:位置|半径|外圈)?[^。；]{0,18}"
            r"(?:无关|不敏感|不影响)|"
            r"(?:充分大|逐步扩大)[^。；]{0,18}(?:外圈|初始半径|搜索上界)"
            r"[^。；]{0,35}(?:无碰撞|安全域|结果不变|临界值稳定)|"
            r"(?:扩大|增加)[^。；]{0,12}(?:初始半径|外圈|搜索上界)"
            r"[^。；]{0,20}(?:结果不变|临界值稳定)|"
            r"对\s*r[_ ]?start[^。；]{0,18}(?:不敏感|无关|不影响)"
            r"[^。；]{0,30}(?:扩至|扩大|增加)[^。；]{0,18}"
            r"(?:p[_ ]?min|临界值|结果)[^。；]{0,18}"
            r"(?:变化\s*[<＜]|稳定|不变)|"
            r"(?:扩大|增加)[^。；]{0,12}(?:计算)?"
            r"(?:截断半径|r[_ ]?upper)[^。；]{0,40}"
            r"(?:变化\s*[<＜]|临界值稳定|结果不变)",
            ques3_text,
            re.IGNORECASE,
        )
        if not proves_initial_independence:
            defines_sufficient_outer_domain = re.search(
                r"充分外侧(?:起点|安全域)|外侧安全域|"
                r"r[_ ]?start[^。；]{0,45}(?:覆盖整条龙|充分外侧)",
                ques3_text,
                re.IGNORECASE,
            )
            verifies_outer_domain_stability = re.search(
                r"(?:扩大|加倍|增加|增至)[^。；]{0,20}"
                r"(?:p[_ ]?max|r[_ ]?start|r[_ ]?upper|搜索上界|起始半径|截断半径)"
                r"[^。；]{0,35}(?:稳定|不变)|"
                r"(?:p[_ ]?max|r[_ ]?start|r[_ ]?upper|搜索上界|起始半径|截断半径)"
                r"[^。；]{0,20}(?:扩大|加倍|增加|增至)[^。；]{0,35}(?:稳定|不变)",
                ques3_text,
                re.IGNORECASE,
            )
            proves_initial_independence = bool(
                defines_sufficient_outer_domain and verifies_outer_domain_stability
            )
        uses_outer_spatial_cutoff = re.search(
            r"r[_ ]?(?:start|upper)|(?:充分)?外侧截断|截断半径|从外侧[^。；\n]{0,12}截断",
            ques3_text,
            re.IGNORECASE,
        )
        has_outer_cutoff_stability = bool(
            proves_initial_independence or certifies_spatial_cutoff
        )
        if uses_outer_spatial_cutoff and not has_outer_cutoff_stability:
            raise QualityGateError(
                "问题三把 r_start/外侧位置作为无界螺线的计算截断，却没有扩大该空间截断"
                "并复算临界螺距；须报告 p_min 对至少两个外侧截断的稳定量"
            )
        if (
            ques3_text
            and invents_q3_initial
            and not source_authorizes_q3_initial
            and not has_outer_cutoff_stability
        ):
            raise QualityGateError(
                "问题三原文未给固定起点，不能自行设定某外圈或声称题图唯一确定最外圈；"
                "应直接覆盖从边界向外的连续构型，或证明充分外侧安全域及临界结果对搜索上界稳定"
            )
        width_based_lower_bound = re.search(
            r"(?:板宽|线间距)[^。；]{0,55}(?:0\.60|2\s*[×xX*]\s*0\.30)|"
            r"(?:0\.60|2\s*[×xX*]\s*0\.30)[^。；]{0,55}(?:板宽|线间距)",
            ques3_text,
            re.IGNORECASE,
        ) and re.search(
            r"(?:p(?:[_ ]?min)?\s*)?(?:下界|搜索从|起始|区间)[^。；]{0,40}0\.60|"
            r"p\s*[>≥]\s*(?:2\s*[×xX*]\s*0\.30|0\.60)",
            ques3_text,
            re.IGNORECASE,
        )
        retains_below_width_bound = re.search(
            r"(?:仍|同时|完整)[^。；]{0,15}(?:搜索|覆盖|保留)[^。；]{0,18}"
            r"(?:0\s*<\s*p\s*<\s*0\.60|0\.60\s*(?:以下|以内)|低于\s*0\.60)|"
            r"(?:证明|严格验证)[^。；]{0,30}p\s*<\s*0\.60[^。；]{0,20}不可行",
            ques3_text,
            re.IGNORECASE,
        )
        if width_based_lower_bound and not retains_below_width_bound:
            raise QualityGateError(
                "问题三不能仅凭两倍板宽把最小螺距搜索域裁到 p>=0.60 m；"
                "板宽关系只能作候选粗筛/初值，须保留更小螺距候选，"
                "或用有限长度、朝向和全链实体碰撞证明被裁区间必不可行"
            )
        positive_domain_cut = re.search(
            r"(?:二分|搜索|扫描|区间)[^。；]{0,20}p[^。；]{0,18}"
            r"[\[(]\s*(?P<bracket_lower>0?\.\d+|\d+(?:\.\d+)?)\s*[,，]|"
            r"(?:二分|搜索|扫描|区间)[^。；]{0,16}"
            r"[\[(]\s*(?P<bare_bracket_lower>0?\.\d+|\d+(?:\.\d+)?)"
            r"\s*[,，][^\])。；]{0,18}[\])](?:\s*m)?|"
            r"p\s*从\s*(?P<step_lower>0?\.\d+|\d+(?:\.\d+)?)\s*m?"
            r"[^。；]{0,16}(?:步长|递增|开始|起)",
            ques3_text,
            re.IGNORECASE,
        )
        lower = (
            positive_domain_cut.group("bracket_lower")
            or positive_domain_cut.group("bare_bracket_lower")
            or positive_domain_cut.group("step_lower")
            if positive_domain_cut
            else None
        )
        if lower and float(lower) > 0:
            retains_full_positive_domain = re.search(
                r"(?:0\s*<\s*p|p\s*[∈∊]\s*\(\s*0\s*[,，])|"
                rf"(?:保留|覆盖|搜索)[^。；]{{0,20}}p\s*<\s*{re.escape(lower)}|"
                rf"p\s*<\s*{re.escape(lower)}[^。；]{{0,35}}"
                r"(?:SAT|有向矩形|实体碰撞|全链)[^。；]{0,30}(?:不可行|严格下界)",
                ques3_text,
                re.IGNORECASE,
            )
            if not retains_full_positive_domain:
                raise QualityGateError(
                    f"问题三把最小螺距搜索域任意裁为 p>={lower} m；"
                    "题面未给正下界，须覆盖 0<p<=p_max，或用全链有限实体可行性"
                    "严格证明被裁区间全部不可行"
                )
        fixed_uniform_grid_over_open_domain = re.search(
            r"(?:完整参数域|全参数域|全域)[^。；\n]{0,45}"
            r"(?:0\s*<\s*p\s*(?:<=|≤)|p\s*[∈∊]\s*\(\s*0\s*[,，])"
            r"[^。；\n]{0,45}均匀网格|"
            r"均匀网格[^。；\n]{0,45}(?:完整参数域|全参数域|全域)"
            r"[^。；\n]{0,35}(?:0\s*<\s*p\s*(?:<=|≤)|"
            r"p\s*[∈∊]\s*\(\s*0\s*[,，])",
            ques3_text,
            re.IGNORECASE,
        )
        has_fixed_positive_pitch_step = re.search(
            r"(?:Δ\s*p|步长)[^。；\n]{0,10}(?:=|取|为)\s*"
            r"\d+(?:\.\d+)?\s*m",
            ques3_text,
            re.IGNORECASE,
        )
        proves_omitted_near_zero_interval = re.search(
            r"(?:0\s*<\s*p\s*<\s*Δ\s*p|小于网格步长的区间)"
            r"[^。；\n]{0,55}(?:解析证明|严格证明|全链SAT)"
            r"[^。；\n]{0,35}不可行|"
            r"(?:解析证明|严格证明|全链SAT)[^。；\n]{0,55}"
            r"(?:0\s*<\s*p\s*<\s*Δ\s*p|小于网格步长的区间)"
            r"[^。；\n]{0,25}不可行",
            ques3_text,
            re.IGNORECASE,
        )
        if (
            fixed_uniform_grid_over_open_domain
            and has_fixed_positive_pitch_step
            and not proves_omitted_near_zero_interval
        ):
            raise QualityGateError(
                "问题三声称用固定正步长均匀网格覆盖开放域 0<p<=p_max；"
                "有限网格必然遗漏靠近0的区间。须使用趋近0的可复核序列并给出"
                "尾段证明，或严格证明首个正网格点以下全部不可行"
            )
        sample_only_near_zero_certificate = re.search(
            r"(?:p\s*(?:→|->)\s*0\s*\+|0\s*(?:与|到|至)\s*首个(?:正)?网格点)"
            r"[^。；\n]{0,80}(?:加密)?(?:采样|网格)"
            r"[^。；\n]{0,35}(?:验证|确认|证明)"
            r"[^。；\n]{0,20}(?:不可行|无隐藏可行解)|"
            r"(?:无隐藏可行解|不可行)[^。；\n]{0,35}"
            r"0\s*(?:与|到|至)\s*首个(?:正)?网格点"
            r"[^。；\n]{0,35}(?:采样|网格)",
            q3_evidence_text,
            re.IGNORECASE,
        )
        proves_near_zero_interval = re.search(
            r"(?:区间算术|解析证明|严格证明|可验证下界|完备枚举)"
            r"[^。；\n]{0,70}(?:0\s*<\s*p|零附近|首个(?:正)?网格点以下)"
            r"[^。；\n]{0,35}不可行|"
            r"(?:0\s*<\s*p|零附近|首个(?:正)?网格点以下)"
            r"[^。；\n]{0,70}(?:区间算术|解析证明|严格证明|可验证下界|完备枚举)",
            q3_evidence_text,
            re.IGNORECASE,
        )
        if sample_only_near_zero_certificate and not proves_near_zero_interval:
            raise QualityGateError(
                "问题三用有限加密采样声称排除 p→0+ 开区间中的隐藏可行解；"
                "有限网格不能证明连续参数区间全部不可行，须给出解析/区间证书或"
                "趋近0的序列及其严格尾部界"
            )
        fixed_upper_bracket = re.search(
            r"p[_ ]?max\s*=\s*\d+(?:\.\d+)?\s*m?", ques3_text, re.IGNORECASE
        )
        certifies_upper_bracket = re.search(
            r"p[_ ]?max[^。；\n]{0,35}(?:验证|确认|证明)[^。；\n]{0,18}可行|"
            r"(?:验证|确认|证明)[^。；\n]{0,18}p[_ ]?max[^。；\n]{0,18}可行|"
            r"(?:若|如)[^。；\n]{0,18}不可行[^。；\n]{0,35}"
            r"(?:扩大|增加|倍增)[^。；\n]{0,12}p[_ ]?max|"
            r"(?:扩大|增加|倍增)\s*p[_ ]?max[^。；\n]{0,30}"
            r"(?:验证|检查|直至|直到)[^。；\n]{0,15}(?:稳定|可行)",
            ques3_text,
            re.IGNORECASE,
        )
        if fixed_upper_bracket and not certifies_upper_bracket:
            raise QualityGateError(
                "问题三的数值搜索上界必须先验证为可行括界；"
                "若初取 p_max 不可行，应继续扩展或倍增，直至获得"
                "不可行/可行两侧证据后再二分，不能把任意数值固定成完整定义域上界"
            )
        if re.search(
            r"(?:p\s*\*|p\s*[_ ]?min|最小螺距)[^。；\n]{0,20}"
            r"(?:应|预计|预期|必然)[^。；\n]{0,8}(?:远大于|明显大于)"
            r"[^。；\n]{0,12}\d+(?:\.\d+)?\s*m?",
            ques3_text,
            re.IGNORECASE,
        ):
            raise QualityGateError(
                "问题三在数值求解前预设了最小螺距的量级；"
                "必须保持完整搜索域并由实体碰撞与临界两侧证据确定 p_min，"
                "不能先声称其应远大于某个经验数值"
            )
        source_has_collision_feasibility = re.search(
            r"问题\s*2.{0,300}(?:板凳之间不发生碰撞|不能再继续盘入)",
            source_text,
            re.DOTALL,
        )
        if ques3_text and source_has_collision_feasibility:
            ignores_collision = re.search(
                r"(?:本问|此问|问题\s*3)?[^。；;]{0,25}"
                r"(?:不考虑|无需考虑|未限制|不要求)[^。；;，,但而]{0,18}碰撞|"
                r"(?:仅此|唯一)[^。；;]{0,12}(?:约束|条件)[^。；;]{0,50}16\s*p|"
                r"(?:p[_ ]?min|最小螺距)[^。；;]{0,30}"
                r"(?:4\.5\s*/\s*16|R[_ ]?turn\s*/\s*16)",
                ques3_text,
                re.IGNORECASE,
            )
            checks_body_feasibility = re.search(
                r"(?:有向矩形|分离轴|SAT|板凳实体|板体轮廓|实体碰撞|实体干涉)"
                r"[^。；]{0,120}(?:边界|4\.5)|"
                r"(?:边界|4\.5)[^。；]{0,120}"
                r"(?:有向矩形|分离轴|SAT|板凳实体|板体轮廓|实体碰撞|实体干涉)|"
                r"(?:调用|沿用|复用)[^。；]{0,35}问题\s*[12](?:\s*[/、和]\s*2)?"
                r"[^。；]{0,35}碰撞(?:检测|判定)|"
                r"(?:到达|边界|4\.5)[^。；]{0,100}(?:全链|全部板体)[^。；]{0,20}"
                r"(?:无碰撞|不发生碰撞)",
                ques3_text,
                re.IGNORECASE,
            )
            checks_body_feasibility = checks_body_feasibility or (
                re.search(
                    r"(?:有向矩形|分离轴|SAT|板凳实体|板体轮廓|实体碰撞|实体干涉)",
                    ques3_text,
                    re.IGNORECASE,
                )
                and re.search(
                    r"(?:碰撞前[^。；\n]{0,40}(?:到达|抵达|直至)[^。；\n]{0,25}"
                    r"(?:边界|4\.5)|(?:到达|抵达|直至)[^。；\n]{0,25}"
                    r"(?:边界|4\.5)[^。；\n]{0,40}(?:无碰撞|碰撞前))",
                    ques3_text,
                    re.IGNORECASE,
                )
            )
            if ignores_collision or not checks_body_feasibility:
                raise QualityGateError(
                    "问题三的最小螺距必须保证龙头到达调头边界前板凳实体仍可继续盘入；"
                    "不能只用初始圈半径给出零行程下界，须复用有限长宽板体碰撞/干涉判定"
                )
        if re.search(
            r"(?:碰撞|不可行)[^。；\n]{0,25}r\s*[<＜]\s*4\.5\s*m?"
            r"[^。；\n]{0,20}(?:之前|到达边界前)|"
            r"(?:到达边界前|之前)[^。；\n]{0,25}(?:碰撞|不可行)"
            r"[^。；\n]{0,25}r\s*[<＜]\s*4\.5\s*m?",
            ques3_text,
            re.IGNORECASE,
        ):
            raise QualityGateError(
                "问题三盘入到 r=4.5m 边界前的半径方向写反："
                "沿外侧向内盘入时，边界前状态满足 r>4.5m；"
                "r<4.5m 已越过题面目标边界，不能作为到达前碰撞证据"
            )

    ques5_text = value.get("ques5", "") if isinstance(value, dict) else ""
    source_ques5 = re.search(
        r"问题\s*5(?P<body>.*?)(?=问题\s*6|$)", source_text, re.DOTALL
    )
    if (
        source_ques5
        and re.search(r"问题\s*4[^。；\n]{0,20}路径", source_ques5.group("body"))
        and re.search(
            r"(?:基线|基准)[^。；\n]{0,12}(?:或|或者|皆可|任选)"
            r"[^。；\n]{0,12}(?:优化|调整)|"
            r"(?:优化|调整)[^。；\n]{0,12}(?:或|或者|皆可|任选)"
            r"[^。；\n]{0,12}(?:基线|基准)",
            ques5_text,
        )
    ):
        raise QualityGateError(
            "问题五必须继承问题四最终由数值比较确定的路径；若优化候选确实更短则"
            "使用该候选，否则保留基线，不能把基线/优化路径留作任意选择"
        )
    q5_multistart_global_claim = re.search(
        r"多(?:起点|初值|次重启)[^。；\n]{0,25}(?:爬坡|局部(?:搜索|优化))?"
        r"[^。；\n]{0,20}(?:确认|证明|保证)[^。；\n]{0,12}全局(?:最优|最大|极值)?|"
        r"(?:确认|证明|保证)[^。；\n]{0,12}全局(?:最优|最大|极值)?"
        r"[^。；\n]{0,35}多(?:起点|初值|次重启)",
        ques5_text,
    )
    q5_has_global_certificate = re.search(
        r"(?:解析(?:上界|证明)|可验证上界|区间算术|完备枚举|凸性证明|"
        r"网格间误差界)",
        ques5_text,
    )
    if q5_multistart_global_claim and not q5_has_global_certificate:
        raise QualityGateError(
            "问题五的多起点爬坡/局部搜索只能检查候选稳定性，不能单独确认全局极值；"
            "须遍历全部节点与路径并给出无限尾部定量上界等全局证据，或降低结论强度"
        )
    one_link_tail_bound = re.search(
        r"相邻节点[^。；\n]{0,80}(?:速度)?倍率[^。；\n]{0,35}"
        r"(?:<=|≤)\s*1\s*\+\s*C\s*/\s*r",
        ques5_text,
        re.IGNORECASE,
    )
    reuses_one_link_bound_for_whole_chain = re.search(
        r"(?:截断外|全链|整体)[^。；\n]{0,35}倍率[^。；\n]{0,35}"
        r"(?:<=|≤)\s*1\s*\+\s*C\s*/\s*r\s*[_ ]?cut",
        ques5_text,
        re.IGNORECASE,
    )
    accumulates_link_bounds = re.search(
        r"C[^。；\n]{0,35}(?:全链|223|逐段)(?:累积|求和|乘积)|"
        r"(?:全链|223|逐段)(?:累积|求和|乘积)[^。；\n]{0,35}C",
        ques5_text,
        re.IGNORECASE,
    )
    if (
        one_link_tail_bound
        and reuses_one_link_bound_for_whole_chain
        and not accumulates_link_bounds
    ):
        raise QualityGateError(
            "问题五把单个相邻节点的倍率界 1+C/r 直接当成全链倍率界；"
            "223段传播必须逐段累积为乘积/对数和并给出全链常数，"
            "不能复用同一个单段上界"
        )
    if re.search(
        r"(?:M|β|beta)\s*[_ ]?max\s*(?:<=|≤)\s*1[^。；\n]{0,65}"
        r"(?:无界|不限|不受限|无上限|任意大|∞|infinity)|"
        r"(?:无界|不限|不受限|无上限|任意大|∞|infinity)[^。；\n]{0,65}"
        r"(?:M|β|beta)\s*[_ ]?max\s*(?:<=|≤)\s*1|"
        r"(?:若|当)[^。；\n]{0,25}max[^。；\n]{0,20}m[^。；\n]{0,12}>\s*1"
        r"[^。；\n]{0,45}(?:否则|反之)[^。；\n]{0,35}"
        r"v\s*[_ ]?head[^。；\n]{0,15}(?:无界|不限|不受限|无上限|任意大)",
        ques5_text,
        re.IGNORECASE,
    ):
        raise QualityGateError(
            "问题五把速度倍率上界不超过1误写成龙头速度无界；"
            "龙头自身倍率恒为1，所以全节点M_max至少为1，且v_max=2/M_max≤2m/s"
        )
    if re.search(
        r"(?:M|β|beta)\s*[_ ]?max\s*(?:<=|≤)\s*1[^。；\n]{0,110}"
        r"(?:非活动|不活动|不是活动|无需达到)[^。；\n]{0,12}(?:约束)?|"
        r"(?:非活动|不活动|不是活动)[^。；\n]{0,80}"
        r"(?:M|β|beta)\s*[_ ]?max\s*(?:<=|≤)\s*1",
        ques5_text,
        re.IGNORECASE,
    ):
        raise QualityGateError(
            "问题五把 M_max≤1 误写成速度上限可能非活动；"
            "倍率集合包含龙头自身 m_1=1，所以实际只能有 M_max≥1；"
            "当 M_max=1 时龙头在 v_head=2m/s 恰好达到上限，约束仍为活动"
        )
    if re.search(
        r"(?:最小(?:节点)?(?:速度|速率)|(?:节点)?(?:速度|速率)最小值)"
        r"[^。；\n]{0,55}(?:达到|等于|触及)[^。；\n]{0,15}"
        r"\d+(?:\.\d+)?\s*m\s*/\s*s[^。；\n]{0,15}(?:上限|活动约束)",
        ques5_text,
        re.IGNORECASE,
    ):
        raise QualityGateError(
            "问题五把速度上限的活动约束误写成最小节点速度达到上限；"
            "在最大化龙头速度且要求所有节点不超过限速时，应由全节点最大速度"
            "达到上限，等价地由最大速度倍率 M_max 决定"
        )
    for max_range in re.finditer(
        r"(?:max|最大)[^。；\n]{0,90}i\s*=\s*1\s*(?:\.\.|~|～|至|到|-)\s*(\d+)",
        ques5_text,
        re.IGNORECASE,
    ):
        if int(max_range.group(1)) != expected_node_count:
            raise QualityGateError(
                "问题五的全节点速度极值索引未覆盖链尾："
                f"1 基节点必须遍历 i=1..{expected_node_count}，而非到 "
                f"{max_range.group(1)}"
            )
    source_has_unbounded_spiral_tails = (
        "盘入螺线" in source_text
        and "盘出螺线" in source_text
        and re.search(r"问题\s*5", source_text)
    )
    invents_finite_total_spiral_path = re.search(
        r"(?:路径)?总弧长[^。；\n]{0,45}(?:有限|开始|结束|起止)|"
        r"(?:有限|开始|结束|起止)[^。；\n]{0,45}(?:路径)?总弧长|"
        r"路径总弧长\s*/\s*[^。；\n]{0,20}(?:确定|得到)"
        r"[^。；\n]{0,20}(?:开始|结束|起止)时刻|"
        r"(?:仿真|扫描|回放)[^。；\n]{0,25}(?:一个)?完整周期"
        r"[^。；\n]{0,90}(?:盘入螺线|盘出螺线)",
        ques5_text,
    )
    provides_infinite_tail_bound = re.search(
        r"(?:无限延伸|无穷远|渐近|解析尾部界|尾部上界|外侧上界)|"
        r"(?:截断外|外侧新增壳层)[^。；\n]{0,55}(?:≤|<=)"
        r"[^。；\n]{0,55}(?:已扫描|有限区间|当前)"
        r"[^。；\n]{0,20}(?:η|eta|max|最大)|"
        r"(?:扩大|增加)[^。；\n]{0,18}(?:截断|时间区间|扫描区间)"
        r"[^。；\n]{0,30}(?:极值不再增加|最大值不变|上界稳定|不超过)|"
        r"(?:外侧)?截断(?:半径|区间)?[^。；\n]{0,35}(?:扩大|增加)"
        r"[^。；\n]{0,90}(?:极值|最大(?:值|倍率)?)[^。；\n]{0,25}"
        r"(?:稳定|变化\s*[<≤]|不再增加|不再增大)",
        ques5_text,
    )
    uses_limit_as_tail_bound = re.search(
        r"r\s*(?:→|->)\s*(?:∞|无穷)[^。；\n]{0,35}"
        r"曲率\s*(?:→|->)\s*0[^。；\n]{0,35}"
        r"(?:速度)?(?:系数|倍率)\s*(?:→|->)\s*1"
        r"[^。；\n]{0,35}(?:证明|所以|故)[^。；\n]{0,25}"
        r"(?:不产生更大|不会更大|为全局)",
        ques5_text,
        re.IGNORECASE,
    )
    has_quantitative_tail_dominance = re.search(
        r"(?:尾部|外侧)[^。；\n]{0,35}(?:上界|不超过|≤|<=)"
        r"[^。；\n]{0,35}(?:有限区间|当前|已得|G\s*[_ ]?max|最大值)|"
        r"(?:扩大|增加|倍增)[^。；\n]{0,20}(?:截断|扫描区间)"
        r"[^。；\n]{0,30}(?:极值稳定|最大值不变|不再增加)|"
        r"(?:外侧)?截断(?:半径|区间)?[^。；\n]{0,35}(?:扩大|增加)"
        r"[^。；\n]{0,90}(?:极值|最大(?:值|倍率)?)[^。；\n]{0,25}"
        r"(?:稳定|变化\s*[<≤]|不再增加|不再增大)",
        ques5_text,
        re.IGNORECASE,
    )
    if uses_limit_as_tail_bound and not has_quantitative_tail_dominance:
        raise QualityGateError(
            "问题五只给出速度倍率趋于 1 的极限，不能据此证明有限尾部处"
            "没有更大值；须给出截断半径以外的定量上界并与当前最大值比较，"
            "或扩大截断验证极值不再增加"
        )
    if re.search(
        r"(?:实际|预计|必然|应当|应得)[^。；\n]{0,18}"
        r"v\s*[_ ]?head\s*[_ ]?max\s*(?:>|大于)\s*1",
        ques5_text,
        re.IGNORECASE,
    ):
        raise QualityGateError(
            "问题五在计算前预设最大龙头速度大于 1 m/s；"
            "是否超过基线必须由全节点、全路径速度倍率和尾部证据决定"
        )
    finite_window_scan = re.search(
        r"(?:扫描|遍历|覆盖|检查)[^。；\n]{0,40}"
        r"(?:[-−]\s*)?\d+(?:\.\d+)?\s*(?:~|～|至|到|\.\.)\s*"
        r"(?:[-−]\s*)?\d+(?:\.\d+)?\s*s",
        ques5_text,
        re.IGNORECASE,
    )
    promotes_window_to_global_extremum = re.search(
        r"(?:取|得到|确定|作为)[^。；\n]{0,35}"
        r"(?:全路径|全时段|全时域|全局)[^。；\n]{0,20}最大|"
        r"(?:全路径|全时段|全时域|全局)[^。；\n]{0,25}最大(?:速度)?比|"
        r"(?:据此|从而)[^。；\n]{0,25}(?:计算|确定)[^。；\n]{0,15}"
        r"龙头(?:的)?最大(?:行进)?速度",
        ques5_text,
        re.IGNORECASE,
    )
    claims_global_path_extremum = re.search(
        r"(?:全路径|完整路径|全程)[^。；\n]{0,30}(?:扫描|回放)"
        r"[^。；\n]{0,35}(?:全局|最大|极值|[αa][ _-]?max|max[_ {])|"
        r"(?:全局|全路径)[^。；\n]{0,30}(?:最大|极值)|"
        r"C\s*=\s*max(?:imum)?\s*[_({]?[^{。；\n]{0,45}"
        r"(?:v\s*0\s*\*?\s*=\s*2\s*/\s*C|龙头[^。；\n]{0,12}最大速度)|"
        r"(?:β|beta)\s*[_ ]?\s*max\s*=\s*max[^。；\n]{0,80}"
        r"(?:沿路径|全路径|全部时刻|[_ {]t[,}])|"
        r"(?:ξ|xi)\s*[_ ]?\s*max\s*=\s*max[^。；\n]{0,80}"
        r"(?:沿路径|全路径|全部时刻|[_ {]t[,}])|"
        r"V\s*[_ ]?\s*max\s*=\s*2\s*/\s*max[^。；\n]{0,80}"
        r"(?:[_ {]k\s*[,，]\s*t|沿路径|全路径)",
        ques5_text,
        re.IGNORECASE,
    )
    claims_global_path_extremum = claims_global_path_extremum or (
        re.search(r"C\s*=\s*max", ques5_text, re.IGNORECASE)
        and re.search(r"v\s*0\s*\*?\s*=\s*2\s*/\s*C", ques5_text, re.IGNORECASE)
    )
    claims_global_path_extremum = claims_global_path_extremum or (
        re.search(
            r"(?:K|G)\s*[_ ]?max\s*=\s*max[^。；\n]{0,100}",
            ques5_text,
            re.IGNORECASE,
        )
        and re.search(
            r"v\s*\*|v\s*[_ ]?(?:head|0)?\s*max|最大允许龙头速度",
            ques5_text,
            re.IGNORECASE,
        )
    )
    if (
        source_has_unbounded_spiral_tails
        and claims_global_path_extremum
        and not provides_infinite_tail_bound
    ):
        raise QualityGateError(
            "问题五声称获得全路径极值，但盘入/盘出螺线具有无限尾部；"
            "必须给出解析渐近尾部界，或扩大截断后极值稳定且外侧不再增大的证据"
        )
    if (
        source_has_unbounded_spiral_tails
        and finite_window_scan
        and promotes_window_to_global_extremum
        and not provides_infinite_tail_bound
    ):
        raise QualityGateError(
            "问题五不能把有限时间窗扫描值直接提升为无限路径的全局极值；"
            "必须给出解析渐近尾部界，或扩大截断后极值稳定且外侧不再增大的证据"
        )
    if (
        source_has_unbounded_spiral_tails
        and invents_finite_total_spiral_path
        and not provides_infinite_tail_bound
    ):
        raise QualityGateError(
            "问题五不能把盘入/盘出螺线当作具有有限总弧长和起止时刻的路径；"
            "螺线向外无限延伸。应数值覆盖调头及高曲率有限区间，并给出"
            "外侧渐近尾部界或扩大截断后最大速度倍率不再增加的证据"
        )


def _find_unresolved_blockers(value: Any, path: str = "") -> list[str]:
    """Return non-empty ``blocked*`` fields from an evidence contract."""
    blockers: list[str] = []
    if isinstance(value, dict):
        for key, item in value.items():
            item_path = f"{path}.{key}" if path else str(key)
            if str(key).lower().startswith("blocked"):
                if isinstance(item, str):
                    normalized = item.strip().lower()
                    unresolved = normalized not in {
                        "",
                        "none",
                        "no",
                        "false",
                        "n/a",
                        "not_applicable",
                        "无",
                        "无阻塞",
                        "无阻塞项",
                    }
                else:
                    unresolved = item not in (None, False, [], {})
                if unresolved:
                    blockers.append(item_path)
            blockers.extend(_find_unresolved_blockers(item, item_path))
    elif isinstance(value, list):
        for index, item in enumerate(value):
            blockers.extend(_find_unresolved_blockers(item, f"{path}[{index}]"))
    return blockers


def _has_substantive_figure_facts(value: Any) -> bool:
    if isinstance(value, str):
        normalized = re.sub(r"\s+", "", value).lower()
        return len(normalized) >= 4 and normalized not in {
            "checked",
            "verified",
            "已检查",
            "已核验",
            "已查看",
            "见图",
        }
    if isinstance(value, dict):
        return any(_has_substantive_figure_facts(item) for item in value.values())
    if isinstance(value, list):
        return any(_has_substantive_figure_facts(item) for item in value)
    return False


def _figure_facts_are_resolved(value: Any) -> bool:
    text = (
        value
        if isinstance(value, str)
        else json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    )
    unresolved_patterns = (
        r"未(?:给出|定义|标明|说明|显示|核验|辨认|读取)",
        r"无法(?:确定|辨认|读取|核验)",
        r"(?:仅|只是)(?:为)?示意",
        r"(?:需要|需|只能|拟)(?:自行)?(?:约定|假设|猜测)",
        r"按(?:惯例|常见|通常).{0,12}(?:取|设|约定)",
    )
    return not any(re.search(pattern, text) for pattern in unresolved_patterns)


def _figure_facts_cover_requirement(value: Any, requirement: dict[str, Any]) -> bool:
    text = (
        value
        if isinstance(value, str)
        else json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    )
    context = str(requirement.get("context", ""))
    named_points = re.findall(r"([A-Za-zＡ-Ｚａ-ｚ])\s*点", context)
    orientation = (
        r"(?:正|负)?\s*[xyXY]\s*轴|左(?:侧|端|方)?|右(?:侧|端|方)?|"
        r"上(?:侧|端|方)?|下(?:侧|端|方)?|内端|外端|顺时针|逆时针|"
        r"坐标\s*[（(]?\s*-?\d"
    )
    for point in named_points:
        point_pattern = rf"{re.escape(point)}\s*点"
        if not (
            re.search(
                rf"{point_pattern}.{{0,12}}(?:位于|处于|在).{{0,30}}(?:{orientation})",
                text,
            )
            or re.search(
                rf"(?:{orientation}).{{0,20}}(?:为|是|标注(?:为)?).{{0,8}}{point_pattern}",
                text,
            )
        ):
            return False
    if not named_points and re.search(r"初始|起点|坐标|方向|方位|位于", context):
        return bool(re.search(orientation, text))
    return True


def _pdf_figure_caption_pages(path: Path, figure: int) -> list[int]:
    """Return one-based PDF pages containing the figure's actual caption."""
    try:
        import fitz

        with fitz.open(path) as document:
            pages = []
            for index, pdf_page in enumerate(document):
                text = pdf_page.get_text("text")
                if re.search(
                    rf"^\s*(?:图\s*{figure}|figure\s*{figure})(?!\d)\s+\S.{{0,80}}$",
                    text,
                    re.IGNORECASE | re.MULTILINE,
                ):
                    pages.append(index + 1)
            return pages
    except (OSError, RuntimeError, ValueError):
        return []


def _pdf_page_mentions_figure(path: Path, page: int, figure: int) -> bool:
    return page in _pdf_figure_caption_pages(path, figure)


def _pdf_named_point_axis_relations(
    path: Path, page: int, named_points: list[str]
) -> dict[str, str]:
    """Infer a labelled point's axis relation from PDF text coordinates.

    This only returns a relation when the origin, x-axis label and named point are
    unambiguously horizontally aligned. Ambiguous or image-only diagrams fall back
    to the ordinary evidence check instead of guessing.
    """
    try:
        import fitz

        with fitz.open(path) as document:
            if page < 1 or page > len(document):
                return {}
            labels: dict[str, list[tuple[float, float, float]]] = {}
            for block in document[page - 1].get_text("dict").get("blocks", []):
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        token = str(span.get("text", "")).strip()
                        if token not in {"O", "x", "X", *named_points}:
                            continue
                        x0, y0, x1, y1 = span["bbox"]
                        labels.setdefault(token.upper(), []).append(
                            ((x0 + x1) / 2, (y0 + y1) / 2, float(span["size"]))
                        )
            origins = labels.get("O", [])
            x_labels = labels.get("X", [])
            relations: dict[str, str] = {}
            for point in named_points:
                point_labels = labels.get(point.upper(), [])
                candidates = []
                for origin in origins:
                    for x_label in x_labels:
                        tolerance = max(origin[2], x_label[2]) * 0.45
                        if abs(origin[1] - x_label[1]) > tolerance:
                            continue
                        for point_label in point_labels:
                            point_tolerance = max(tolerance, point_label[2] * 0.45)
                            if abs(point_label[1] - origin[1]) <= point_tolerance:
                                candidates.append((origin, x_label, point_label))
                if len(candidates) != 1:
                    continue
                origin, x_label, point_label = candidates[0]
                if x_label[0] > origin[0] and point_label[0] > origin[0]:
                    relations[point] = "positive_x_axis"
                elif x_label[0] > origin[0] and point_label[0] < origin[0]:
                    relations[point] = "negative_x_axis"
            return relations
    except (KeyError, OSError, RuntimeError, TypeError, ValueError):
        return {}


def _facts_match_axis_relation(facts: Any, point: str, relation: str) -> bool:
    text = (
        facts
        if isinstance(facts, str)
        else json.dumps(facts, ensure_ascii=False, separators=(",", ":"))
    )
    compact = re.sub(r"\s+", "", text)
    point_pattern = rf"{re.escape(point)}点"
    if relation == "positive_x_axis":
        axis_pattern = r"(?:正x轴|x(?:轴)?(?:的)?正(?:向|方向|半轴))"
    elif relation == "negative_x_axis":
        axis_pattern = r"(?:负x轴|x(?:轴)?(?:的)?负(?:向|方向|半轴))"
    else:
        return True
    return bool(
        re.search(
            rf"{point_pattern}.{{0,30}}(?:位于|处于|在).{{0,30}}{axis_pattern}",
            compact,
            re.IGNORECASE,
        )
        or re.search(
            rf"{axis_pattern}.{{0,30}}(?:为|是|标注(?:为)?).{{0,12}}{point_pattern}",
            compact,
            re.IGNORECASE,
        )
    )


def _validate_source_figure_evidence(
    payload: dict[str, Any], source_text: str, work_dir: str | Path
) -> None:
    requirements = derive_source_figure_requirements(source_text)
    if not requirements:
        return
    data_scope = payload.get("data_scope")
    records = data_scope.get("source_figures") if isinstance(data_scope, dict) else None
    if not isinstance(records, list):
        actual_type = type(records).__name__ if records is not None else "missing"
        raise QualityGateError(
            "题面关键图示尚未核验：data_scope.source_figures 必须是 JSON 数组，"
            "每项记录 file、page、figure、facts；"
            f"当前类型为 {actual_type}。正确结构示例："
            '[{"file":"official.pdf","page":2,"figure":4,'
            '"facts":["命名点相对坐标轴/内外端/方向的明确关系"]}]；'
            "不得写成以 figure_4 为键的对象"
        )

    root = Path(work_dir).resolve()
    requirement_by_figure = {int(item["figure"]): item for item in requirements}
    covered: set[int] = set()
    diagnostics: list[str] = []
    for record in records:
        if not isinstance(record, dict):
            continue
        figure_value = str(record.get("figure", ""))
        figure_match = re.search(r"\d+", figure_value)
        file_value = record.get("file")
        page = record.get("page")
        facts = record.get("facts")
        if not (
            figure_match
            and isinstance(file_value, str)
            and file_value.strip()
            and isinstance(page, int)
            and not isinstance(page, bool)
            and page > 0
            and _has_substantive_figure_facts(facts)
        ):
            diagnostics.append(
                "来源记录字段无效：每项必须含非空 file、正整数 page、含数字的 "
                "figure 和具体 facts"
            )
            continue
        source_path = (root / file_value).resolve()
        if root not in source_path.parents or not source_path.is_file():
            diagnostics.append(f"图 {figure_match.group()}: 来源文件不存在或路径越界")
            continue
        if source_path.suffix.lower() not in {".pdf", ".png", ".jpg", ".jpeg"}:
            diagnostics.append(f"图 {figure_match.group()}: 来源文件必须是 PDF/PNG/JPG")
            continue
        figure_number = int(figure_match.group())
        if not _figure_facts_are_resolved(facts):
            diagnostics.append(
                f"图 {figure_number}: facts 仍含‘未给出/需约定/按惯例’等未解决表述"
            )
            continue
        requirement = requirement_by_figure.get(figure_number)
        layout_covers_requirement = False
        if source_path.suffix.lower() == ".pdf":
            caption_pages = _pdf_figure_caption_pages(source_path, figure_number)
            if page not in caption_pages:
                if caption_pages:
                    page_text = "、".join(f"第 {item} 页" for item in caption_pages)
                    diagnostics.append(
                        f"图 {figure_number}: {file_value} 第 {page} 页只有正文引用或无图注；"
                        f"实际图注页为{page_text}"
                    )
                else:
                    diagnostics.append(
                        f"图 {figure_number}: {file_value} 未提取到行首‘图 {figure_number} + 图名’图注"
                    )
                continue
            named_points = (
                re.findall(
                    r"([A-Za-zＡ-Ｚａ-ｚ])\s*点", str(requirement.get("context", ""))
                )
                if requirement
                else []
            )
            axis_relations = _pdf_named_point_axis_relations(
                source_path, page, named_points
            )
            mismatches = [
                (point, relation)
                for point, relation in axis_relations.items()
                if not _facts_match_axis_relation(facts, point, relation)
            ]
            if mismatches:
                descriptions = {
                    "positive_x_axis": "正 x 轴",
                    "negative_x_axis": "负 x 轴",
                }
                expected = "、".join(
                    f"{point} 点位于{descriptions[relation]}"
                    for point, relation in mismatches
                )
                diagnostics.append(
                    f"图 {figure_number}: PDF 文字标签坐标显示 {expected}；"
                    "facts 必须记录该来源关系，不得改写成象限或猜测方位"
                )
                continue
            layout_covers_requirement = bool(named_points) and all(
                point in axis_relations for point in named_points
            )
        if (
            requirement
            and not layout_covers_requirement
            and not _figure_facts_cover_requirement(facts, requirement)
        ):
            diagnostics.append(
                f"图 {figure_number}: facts 未明确写出命名点相对坐标轴、内外端或方向的关系；"
                f"题面上下文：{requirement['context']}"
            )
            continue
        covered.add(figure_number)

    required = {int(item["figure"]) for item in requirements}
    missing = sorted(required - covered)
    if missing:
        detail = (
            f"；诊断：{'；'.join(dict.fromkeys(diagnostics))}" if diagnostics else ""
        )
        raise QualityGateError(
            f"题面关键图示缺少可核验来源记录（文件/页码/图号/事实）: {missing}{detail}"
        )


def load_evidence_contract(
    work_dir: str | Path, subtask_title: str, *, source_text: str | None = None
) -> tuple[str, dict[str, Any]]:
    """Load and validate the machine-readable evidence chain for one subtask."""
    file_name = f"results_{subtask_title}.json"
    path = Path(work_dir) / file_name
    if not path.is_file():
        raise QualityGateError(f"缺少子任务证据合同: {file_name}")
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise QualityGateError(f"子任务证据合同不可读取: {file_name}: {exc}") from exc
    required = {
        "status",
        "subtask",
        "data_scope",
        "assumptions",
        "model",
        "numeric_results",
        "validation",
        "conclusion_bounds",
    }
    missing = sorted(required - set(payload))
    if missing:
        raise QualityGateError(f"子任务证据合同缺少字段: {missing}")
    if payload["status"] != "success" or payload["subtask"] != subtask_title:
        raise QualityGateError("子任务证据合同状态或 subtask 不匹配")
    unresolved_blockers = _find_unresolved_blockers(payload)
    if unresolved_blockers:
        raise QualityGateError(
            "子任务证据合同仍有未解决阻塞项，不得声明成功: "
            + ", ".join(unresolved_blockers)
        )
    if not _contains_finite_number(payload["numeric_results"]):
        raise QualityGateError("子任务证据合同没有可复算的有限数值结果")
    validation = payload["validation"]
    if not isinstance(validation, dict) or validation.get("passed") is not True:
        raise QualityGateError("子任务证据合同未声明验证通过")
    if not validation.get("checks"):
        raise QualityGateError("子任务证据合同缺少验证检查明细")
    if subtask_title == "eda" and isinstance(payload["conclusion_bounds"], dict):
        for key, value in payload["conclusion_bounds"].items():
            normalized_key = str(key).lower()
            is_q3_pitch_lower_bound = (
                any(token in normalized_key for token in ("q3", "ques3", "problem3", "问题3"))
                and any(token in normalized_key for token in ("pitch", "p_min", "螺距"))
                and any(token in normalized_key for token in ("lower", "下界", "起点"))
            )
            if (
                is_q3_pitch_lower_bound
                and isinstance(value, (int, float))
                and not isinstance(value, bool)
                and math.isfinite(float(value))
                and float(value) > 0
            ):
                raise QualityGateError(
                    "EDA 证据合同不得为题面未给正下界的最小螺距搜索擅设正下界；"
                    "应保留 0<p<=p_max，或在问题三求解阶段给出全链实体不可行性证明"
                )
        guardrails = derive_source_guardrails(source_text or "")
        initial_spiral = guardrails.get("initial_spiral") if guardrails else None
        source_ques3 = re.search(
            r"问题\s*3(?P<body>.*?)(?=问题\s*4|$)", source_text or "", re.DOTALL
        )
        source_authorizes_q1 = source_ques3 and re.search(
            r"(?:问题\s*1|第一问)", source_ques3.group("body")
        )

        def reject_q1_pitch_as_structured_q3_bound(
            item: Any, path: tuple[str, ...] = ()
        ) -> None:
            if not isinstance(item, dict):
                return
            normalized_path = "_".join(path).lower().replace(" ", "")
            is_q3_pitch_path = any(
                token in normalized_path
                for token in ("q3", "ques3", "question3", "problem3", "问题3")
            ) and any(
                token in normalized_path for token in ("pitch", "p_min", "螺距")
            )
            if is_q3_pitch_path and initial_spiral and not source_authorizes_q1:
                q1_pitch = float(initial_spiral["pitch_m"])
                for key, value in item.items():
                    normalized_key = str(key).lower()
                    if not any(
                        token in normalized_key for token in ("upper", "high", "上界")
                    ):
                        continue
                    if (
                        isinstance(value, (int, float))
                        and not isinstance(value, bool)
                        and math.isfinite(float(value))
                        and math.isclose(float(value), q1_pitch, abs_tol=1e-12)
                    ):
                        raise QualityGateError(
                            "问题三的结构化搜索边界把问题一的螺距数值写成上界；"
                            "原问题三未引用问题一，EDA 必须保留独立的 0<p<=p_max，"
                            "并把 p_max 的可行性证明留给问题三求解阶段"
                        )
            for key, value in item.items():
                reject_q1_pitch_as_structured_q3_bound(value, (*path, str(key)))

        reject_q1_pitch_as_structured_q3_bound(payload["conclusion_bounds"])
    if subtask_title == "eda":
        validation_text = json.dumps(payload["validation"], ensure_ascii=False)
        unsupported_scalar_feasibility = re.search(
            r"板宽[^。；\n]{0,35}(?:小于|远小于|<)[^。；\n]{0,35}"
            r"(?:物理|几何)?可行|"
            r"调头空间(?:半径)?[^。；\n]{0,40}(?:大于|>)[^。；\n]{0,40}"
            r"(?:最小可承载结构尺寸|(?:物理|几何)?可行)",
            validation_text,
        )
        if unsupported_scalar_feasibility:
            raise QualityGateError(
                "EDA 不能仅凭板宽小于节点距或调头半径大于某个未推导尺寸就声明"
                "物理/几何可行；这里只能核验题面常量与量纲，完整可行性须由对应"
                "子问题的有限尺寸实体回放和约束残差证明"
            )
    if source_text:
        _require_guardrail_numbers(payload["numeric_results"], source_text)
        _reject_guardrail_contradictions(
            payload, source_text, subtask_title=subtask_title
        )
        if (
            subtask_title == "ques1"
            and "等距螺线" in source_text
            and "顺时针盘入" in source_text
            and re.search(r"各把手中心均位于螺线上", source_text)
        ):
            numeric_results = payload["numeric_results"]
            required_metrics = {
                "max_distance_residual_m",
                "initial_head_radius_m",
                "initial_tail_radius_m",
            }
            missing_metrics = sorted(required_metrics - set(numeric_results))
            if missing_metrics:
                raise QualityGateError(
                    f"螺线刚性链证据缺少可机械复核指标: {missing_metrics}"
                )
            residual = numeric_results["max_distance_residual_m"]
            head_radius = numeric_results["initial_head_radius_m"]
            tail_radius = numeric_results["initial_tail_radius_m"]
            if not all(
                isinstance(value, (int, float))
                and not isinstance(value, bool)
                and math.isfinite(float(value))
                for value in (residual, head_radius, tail_radius)
            ):
                raise QualityGateError("螺线刚性链验证指标必须是有限数值")
            if float(residual) > 1e-6:
                raise QualityGateError(
                    f"刚性链相邻节点距离残差超过 1e-6 m: {float(residual):.6g} m"
                )
            if float(tail_radius) <= float(head_radius):
                raise QualityGateError(
                    "盘入初态串联次序矛盾：龙身/龙尾位于龙头后方外侧，"
                    "龙尾初始极径必须大于龙头初始极径"
                )
        if subtask_title == "eda":
            _validate_source_figure_evidence(payload, source_text, work_dir)
    return file_name, payload


def snapshot_spreadsheet_template(path: str | Path) -> dict[str, Any]:
    """Capture the immutable structure and current numeric-cell count of a template."""
    workbook = load_workbook(path, data_only=False)
    sheets: dict[str, Any] = {}
    numeric_count = 0
    for sheet in workbook.worksheets:
        constants: dict[str, Any] = {}
        formulas: dict[str, str] = {}
        styles: dict[str, int] = {}
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if cell.has_style:
                    styles[cell.coordinate] = cell.style_id
                if cell.data_type == "f":
                    formulas[cell.coordinate] = str(value)
                elif value is not None:
                    constants[cell.coordinate] = value
                    if isinstance(value, (int, float)) and not isinstance(value, bool):
                        numeric_count += int(math.isfinite(float(value)))
        sheets[sheet.title] = {
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "merged_ranges": sorted(str(item) for item in sheet.merged_cells.ranges),
            "constants": constants,
            "formulas": formulas,
            "styles": styles,
        }
    workbook.close()
    return {
        "sheet_names": workbook.sheetnames,
        "sheets": sheets,
        "numeric_count": numeric_count,
    }


def _validate_clockwise_initial_axis_direction(
    path: Path, source_text: str, work_dir: Path
) -> None:
    """Check the first reported movement against a source-defined axis direction."""
    if not (
        "顺时针" in source_text
        and re.search(r"初始.{0,80}([A-Za-zＡ-Ｚａ-ｚ])\s*点", source_text, re.DOTALL)
    ):
        return
    point = re.search(
        r"初始.{0,80}([A-Za-zＡ-Ｚａ-ｚ])\s*点", source_text, re.DOTALL
    ).group(1)
    relation = None
    for requirement in derive_source_figure_requirements(source_text):
        if point not in str(requirement.get("context", "")):
            continue
        figure = int(requirement["figure"])
        for source_path in sorted(work_dir.glob("*.pdf")):
            for page in _pdf_figure_caption_pages(source_path, figure):
                inferred = _pdf_named_point_axis_relations(
                    source_path, page, [point]
                ).get(point)
                if inferred:
                    relation = inferred
                    break
            if relation:
                break
        if relation:
            break
    if relation not in {"positive_x_axis", "negative_x_axis"}:
        return

    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook["位置"] if "位置" in workbook.sheetnames else workbook.active
        zero_column = next(
            (
                cell.column
                for cell in sheet[1]
                if re.fullmatch(r"\s*0\s*s\s*", str(cell.value), re.IGNORECASE)
            ),
            None,
        )
        next_column = next(
            (
                cell.column
                for cell in sheet[1]
                if re.fullmatch(r"\s*1\s*s\s*", str(cell.value), re.IGNORECASE)
            ),
            None,
        )
        head_y_row = next(
            (
                row[0].row
                for row in sheet.iter_rows(min_col=1, max_col=1)
                if re.search(r"龙头\s*y\s*[（(]?\s*m", str(row[0].value), re.IGNORECASE)
            ),
            None,
        )
        if None in {zero_column, next_column, head_y_row}:
            raise QualityGateError(
                f"结果模板缺少顺时针方位复核所需的 0 s/1 s 龙头 y 坐标: {path.name}"
            )
        y_zero = sheet.cell(head_y_row, zero_column).value
        y_next = sheet.cell(head_y_row, next_column).value
        if not all(
            isinstance(value, (int, float))
            and not isinstance(value, bool)
            and math.isfinite(float(value))
            for value in (y_zero, y_next)
        ):
            raise QualityGateError(
                f"结果模板未给出可复核的 0 s/1 s 龙头 y 坐标: {path.name}"
            )
        movement = float(y_next) - float(y_zero)
        wrong_direction = (relation == "positive_x_axis" and movement >= -1e-9) or (
            relation == "negative_x_axis" and movement <= 1e-9
        )
        if wrong_direction:
            descriptions = {
                "positive_x_axis": "正 x 轴",
                "negative_x_axis": "负 x 轴",
            }
            raise QualityGateError(
                f"顺时针坐标方向与官方题图矛盾：初始 {point} 点位于"
                f"{descriptions[relation]}，但 1 s 龙头 y 坐标沿反方向变化"
            )
    finally:
        workbook.close()


def _validate_time_series_motion_continuity(path: Path) -> None:
    """Reject branch jumps whose displacement is impossible at reported speeds."""
    workbook = load_workbook(path, data_only=True, read_only=False)
    try:
        if len(workbook.worksheets) < 2:
            return
        position, speed = workbook.worksheets[:2]
        node_count = speed.max_row - 1
        if node_count <= 0 or position.max_row != 2 * node_count + 1:
            return

        time_columns: list[tuple[int, float]] = []
        for cell in position[1]:
            match = re.fullmatch(
                r"\s*(-?\d+(?:\.\d+)?)\s*s\s*", str(cell.value), re.IGNORECASE
            )
            if match:
                time_columns.append((cell.column, float(match.group(1))))
        if len(time_columns) < 2:
            return
        speed_headers = {
            cell.column: str(cell.value).strip() for cell in speed[1] if cell.value
        }
        for column, reported_time in time_columns:
            if speed_headers.get(column) != position.cell(1, column).value:
                raise QualityGateError(
                    f"位置与速度工作表时间列不一致: {path.name} 第 {column} 列"
                )

        for (left_col, left_time), (right_col, right_time) in zip(
            time_columns, time_columns[1:]
        ):
            dt = right_time - left_time
            if dt <= 0:
                raise QualityGateError(f"结果模板时间列不是严格递增: {path.name}")
            for node in range(1, node_count + 1):
                row_x, row_y, row_speed = 2 * node, 2 * node + 1, node + 1
                values = (
                    position.cell(row_x, left_col).value,
                    position.cell(row_y, left_col).value,
                    position.cell(row_x, right_col).value,
                    position.cell(row_y, right_col).value,
                    speed.cell(row_speed, left_col).value,
                    speed.cell(row_speed, right_col).value,
                )
                if not all(
                    isinstance(value, (int, float))
                    and not isinstance(value, bool)
                    and math.isfinite(float(value))
                    for value in values
                ):
                    continue
                x0, y0, x1, y1, v0, v1 = map(float, values)
                displacement = math.hypot(x1 - x0, y1 - y0)
                speed_bound = max(abs(v0), abs(v1)) * dt * 1.1 + 0.02
                if displacement > speed_bound:
                    label = position.cell(row_x, 1).value or f"节点{node}"
                    raise QualityGateError(
                        "逐时位置存在与报告速度不相容的分支跳变："
                        f"{label} 在 {left_time:g}-{right_time:g}s 位移 "
                        f"{displacement:.6g}m，速度上界仅支持约 {speed_bound:.6g}m"
                    )
    finally:
        workbook.close()


def validate_spreadsheet_template_output(
    path: str | Path,
    before: dict[str, Any],
    *,
    source_text: str | None = None,
    subtask_title: str | None = None,
) -> None:
    """Require numerical fill-in while preserving the official template structure."""
    after = snapshot_spreadsheet_template(path)
    if after["sheet_names"] != before["sheet_names"]:
        raise QualityGateError(f"结果模板工作表变化: {Path(path).name}")
    for name in before["sheet_names"]:
        expected = before["sheets"][name]
        actual = after["sheets"][name]
        for field in ("max_row", "max_column", "merged_ranges", "formulas", "styles"):
            if actual[field] != expected[field]:
                raise QualityGateError(f"结果模板 {name} 的 {field} 被修改")
        for coordinate, value in expected["constants"].items():
            if actual["constants"].get(coordinate) != value:
                raise QualityGateError(f"结果模板固定单元格被修改: {name}!{coordinate}")
    if after["numeric_count"] <= before["numeric_count"]:
        raise QualityGateError(f"结果模板未新增数值结果: {Path(path).name}")
    resolved_path = Path(path).resolve()
    if source_text and re.search(
        re.escape(resolved_path.name), source_text, re.IGNORECASE
    ):
        workbook = load_workbook(resolved_path, data_only=True, read_only=True)
        try:
            missing_cells: list[str] = []
            for sheet in workbook.worksheets:
                header_columns = [
                    cell.column for cell in sheet[1][1:] if cell.value is not None
                ]
                if not header_columns:
                    continue
                for row in range(2, sheet.max_row + 1):
                    if sheet.cell(row, 1).value is None:
                        continue
                    for column in header_columns:
                        value = sheet.cell(row, column).value
                        if not (
                            isinstance(value, (int, float))
                            and not isinstance(value, bool)
                            and math.isfinite(float(value))
                        ):
                            missing_cells.append(f"{get_column_letter(column)}{row}")
                            if len(missing_cells) >= 8:
                                break
                    if len(missing_cells) >= 8:
                        break
                if len(missing_cells) >= 8:
                    break
            if missing_cells:
                raise QualityGateError(
                    f"结果模板未完整回填: {resolved_path.name}; "
                    f"缺失/非数值写入格示例 {missing_cells}"
                )
        finally:
            workbook.close()
    if source_text and subtask_title == "ques1":
        _validate_clockwise_initial_axis_direction(
            resolved_path, source_text, resolved_path.parent
        )
        if (
            "等距螺线" in source_text
            and "顺时针盘入" in source_text
            and re.search(r"各把手中心均位于螺线上", source_text)
        ):
            workbook = load_workbook(resolved_path, data_only=True, read_only=True)
            try:
                sheet = (
                    workbook["位置"]
                    if "位置" in workbook.sheetnames
                    else workbook.active
                )
                zero_column = next(
                    (
                        cell.column
                        for cell in sheet[1]
                        if re.fullmatch(r"\s*0\s*s\s*", str(cell.value), re.IGNORECASE)
                    ),
                    None,
                )
                radii: list[float] = []
                if zero_column is not None:
                    row = 2
                    while row < sheet.max_row:
                        x_label = str(sheet.cell(row, 1).value or "")
                        y_label = str(sheet.cell(row + 1, 1).value or "")
                        if re.search(
                            r"x\s*[()（）]?\s*m", x_label, re.IGNORECASE
                        ) and re.search(r"y\s*[()（）]?\s*m", y_label, re.IGNORECASE):
                            x_value = sheet.cell(row, zero_column).value
                            y_value = sheet.cell(row + 1, zero_column).value
                            if all(
                                isinstance(value, (int, float))
                                and not isinstance(value, bool)
                                and math.isfinite(float(value))
                                for value in (x_value, y_value)
                            ):
                                radii.append(math.hypot(float(x_value), float(y_value)))
                            row += 2
                        else:
                            row += 1
                if len(radii) >= 2:
                    guardrails = derive_source_guardrails(source_text)
                    initial_spiral = (
                        guardrails.get("initial_spiral") if guardrails else None
                    )
                    if initial_spiral and not math.isclose(
                        radii[0], float(initial_spiral["radius_m"]), abs_tol=2e-4
                    ):
                        raise QualityGateError(
                            "结果模板的龙头初始极径与题面圈数/螺距矛盾"
                        )
                    reversed_pair = next(
                        (
                            index
                            for index in range(len(radii) - 1)
                            if radii[index + 1] <= radii[index] + 1e-6
                        ),
                        None,
                    )
                    if reversed_pair is not None:
                        raise QualityGateError(
                            "盘入初态的节点空间次序反转：龙身/龙尾应在龙头"
                            "后方外侧，0 s 极径必须沿节点顺序逐点增大"
                        )
            finally:
                workbook.close()
    if subtask_title == "ques1":
        _validate_time_series_motion_continuity(Path(path).resolve())


def validate_modeler_result(
    result: ModelerToCoder,
    question_keys: set[str],
    *,
    source_text: str | None = None,
) -> None:
    """确保每个题目都有非空建模方案。"""
    required = {"eda", "sensitivity_analysis", *question_keys}
    missing = sorted(required - set(result.questions_solution))
    if missing:
        raise QualityGateError(f"建模方案缺少字段: {missing}")
    invalid = [
        key
        for key in result.questions_solution
        if not re.fullmatch(r"^(eda|sensitivity_analysis|ques\d+)$", key)
    ]
    if invalid:
        raise QualityGateError(f"建模方案包含未知字段: {invalid}")
    oversized = [
        key
        for key, value in result.questions_solution.items()
        if len(value) > (1600 if key == "eda" else 1200)
    ]
    if oversized:
        raise QualityGateError(f"建模方案过长且可能截断: {oversized}")
    if "source_parameter_audit" not in result.questions_solution["eda"]:
        raise QualityGateError("建模方案缺少 source_parameter_audit 来源参数审计")
    if source_text:
        figure_requirements = derive_source_figure_requirements(source_text)
        if figure_requirements and "source_figure_required" not in json.dumps(
            result.questions_solution, ensure_ascii=False
        ):
            raise QualityGateError(
                "建模方案缺少 source_figure_required：题面关键初始/边界几何依赖"
                "图示，必须交给代码阶段核验官方图号与事实，不能自行设定相位"
            )
        source_result_files = {
            item.lower()
            for item in re.findall(r"result\d+\.xlsx", source_text, re.IGNORECASE)
        }
        planned_result_files = _planned_result_files(
            json.dumps(result.questions_solution, ensure_ascii=False)
        )
        unexpected_result_files = sorted(planned_result_files - source_result_files)
        if source_result_files and unexpected_result_files:
            raise QualityGateError(
                "建模方案擅自增加题面未要求的结果文件: "
                + ", ".join(unexpected_result_files)
                + "；只生成题面指定附件"
            )
        for result_file in sorted(source_result_files):
            question_number_match = re.fullmatch(r"result(\d+)\.xlsx", result_file)
            if not question_number_match:
                continue
            question_key = f"ques{int(question_number_match.group(1))}"
            if question_key not in question_keys:
                continue
            if result_file not in _planned_result_files(
                result.questions_solution[question_key]
            ):
                raise QualityGateError(
                    f"{question_key} 遗漏题面指定附件 {result_file}；"
                    "必须在对应子问题中明确按官方模板回填全部要求的节点与时刻"
                )
        guardrails = derive_source_guardrails(source_text)
        eda_text = result.questions_solution["eda"]
        if guardrails:
            _reject_guardrail_contradictions(result.questions_solution, source_text)
            entity_count = int(guardrails["entity_count"])
            plan_text = json.dumps(result.questions_solution, ensure_ascii=False)
            assigns_rigid_distance_to_path_arc = re.search(
                r"(?:弧长差|沿螺线弧长|路径弧长)"
                r".{0,24}(?:=|等于|满足).{0,24}"
                r"(?:d_?\{?[iIkK]\}?|2\.86|1\.65|孔距|把手间距)",
                plan_text,
                re.IGNORECASE,
            )
            if assigns_rigid_distance_to_path_arc:
                raise QualityGateError(
                    "建模方案把刚性构件的欧氏弦长误作路径弧长差；"
                    "必须用两节点坐标的二维距离方程反解曲线参数"
                )
            contradiction_patterns = (
                rf"(?:共|总计|只有)\s*{entity_count}\s*(?:个)?"
                r"(?:不同)?(?:空间)?(?:把手)?(?:连接)?节点",
                r"(?:相邻)?(?:把手|连接)?节点(?:间距|(?:间|之间)(?:的)?"
                r"(?:弦长|距离|间距)).{0,30}(?:不固定|并非固定|可变|取决于|由.{0,8}夹角)",
                r"相邻把手(?:间|之间)(?:空间)?(?:距离|间距)"
                r".{0,30}(?:不固定|并非固定|可变|取决于|由.{0,8}夹角)",
                r"(?:conn(?:ection)?|连接).{0,100}(?:一半之和|半长之和|"
                r"\([^()]{0,30}\+[^()]{0,30}\)\s*/\s*2)",
                r"(?:沿[^。；]{0,20})?弧长.{0,30}"
                r"(?:恒等于|等于|=|近似(?:等于|用)?)"
                r".{0,30}(?:弦长|直线距离|节点距离|节点距)",
            )
            if any(re.search(pattern, eda_text) for pattern in contradiction_patterns):
                raise QualityGateError(
                    "建模方案与题面可直接推导的串联节点数或固定内部弦长矛盾"
                )
        _require_guardrail_numbers(
            [
                float(item)
                for item in re.findall(
                    r"-?\d+(?:\.\d+)?", result.questions_solution["eda"]
                )
            ],
            source_text,
            allow_centimetres=True,
        )


def validate_coder_result(
    result: CoderToWriter,
    work_dir: str | Path,
    *,
    subtask_title: str | None = None,
    source_text: str | None = None,
) -> None:
    """阻断失败、空结果和不存在的图表证据进入 Writer。"""
    if result.status != "success":
        raise QualityGateError(result.error or "代码阶段失败")
    if not (result.code_output or result.code_response):
        raise QualityGateError("代码阶段没有输出结果")
    if subtask_title is not None:
        load_evidence_contract(work_dir, subtask_title, source_text=source_text)
    root = Path(work_dir).resolve()
    for artifact in [*result.created_images, *result.artifacts]:
        path = (root / artifact).resolve()
        if root not in path.parents and path != root:
            raise QualityGateError(f"产物路径越界: {artifact}")
        if not path.exists():
            raise QualityGateError(f"代码声明的产物不存在: {artifact}")


def validate_competition_paper_text(
    text: str, *, section_name: str | None = None
) -> None:
    """拦截竞赛论文中的生成元话语和明显模板污染。"""
    meta_patterns = (
        r"(?:文档|论文|报告|PDF)?\s*(?:生成|验收)时间\s*[:：]",
        r"(?:Agent|模型|系统)\s*版本\s*[:：]",
        r"(?:由|使用).{0,12}(?:AI|Agent|大模型).{0,8}(?:自动)?(?:生成|撰写)",
        r"自动生成(?:的)?(?:论文|报告|文档)",
        r"本文的核心优势不是算法名称",
        r"(?:官方|标准)结果模板的写法",
        r"结果(?:仅)?另行保存",
    )
    for pattern in meta_patterns:
        if re.search(pattern, text, re.IGNORECASE):
            label = section_name or "正文"
            raise QualityGateError(f"{label} 包含生成过程、自我评价或模板元话语")

    for line in text.splitlines():
        is_heading = re.search(
            r"(?:\\(?:title|section|subsection|subsubsection)\b|^\s*#{1,6}\s+)",
            line,
        )
        has_color = re.search(
            r"(?:\\(?:color|textcolor|colorbox)\b|\bcolor\s*=|<font[^>]*\bcolor)",
            line,
            re.IGNORECASE,
        )
        if is_heading and has_color:
            raise QualityGateError("论文标题或章节标题包含非模板要求的显式颜色")

    reference_headings = re.findall(
        r"(?m)^\s*(?:#{1,6}\s*|\\(?:section\*?|chapter\*?)\s*\{)参考文献\}?\s*$",
        text,
    )
    if len(reference_headings) > 1:
        raise QualityGateError("论文包含重复的参考文献标题")


def validate_writer_result(
    result: WriterResponse, *, section_name: str | None = None
) -> None:
    """确保 Writer 至少返回可保存的非空正文。"""
    text = result.response_content
    if not isinstance(text, str) or not text.strip():
        raise QualityGateError("论文阶段返回空正文")
    if re.search(
        r"(?:TODO|PLACEHOLDER|待补充|待续写|搜索文献失败|执行过程中遇到错误)",
        text,
        re.IGNORECASE,
    ):
        raise QualityGateError("论文正文包含占位符或 Agent 错误信息")
    validate_competition_paper_text(text, section_name=section_name)
