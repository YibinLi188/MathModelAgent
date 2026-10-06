"""质量契约回归测试。"""

import json
from pathlib import Path

import fitz
import pytest
from openpyxl import Workbook, load_workbook

from app.core.agents.coder_agent import (
    classify_input_file_role,
    has_unresolved_solver_failure,
    is_source_figure_inspection_code,
)
from app.core.quality_gates import (
    QualityGateError,
    derive_source_guardrails,
    derive_source_figure_requirements,
    format_source_guardrails,
    load_evidence_contract,
    snapshot_spreadsheet_template,
    validate_coder_result,
    validate_competition_paper_text,
    validate_modeler_result,
    validate_spreadsheet_template_output,
    validate_writer_result,
)
from app.models.user_output import UserOutput
from app.core.prompts.coder import CODER_PROMPT
from app.core.prompts.modeler import MODELER_PROMPT
from app.core.prompts.writer import get_writer_prompt
from app.core.flows import Flows
from app.schemas.A2A import CoderToWriter, ModelerToCoder, WriterResponse
from app.schemas.enums import FormatOutPut


def test_modeler_requires_all_solution_sections():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: ok",
            "ques1": "ok",
            "sensitivity_analysis": "ok",
        }
    )
    validate_modeler_result(result, {"ques1"})
    with pytest.raises(QualityGateError):
        validate_modeler_result(result, {"ques1", "ques2"})


def test_failed_coder_cannot_cross_gate(tmp_path: Path):
    result = CoderToWriter(status="failed", error="syntax error")
    with pytest.raises(QualityGateError):
        validate_coder_result(result, tmp_path)


def test_coder_artifact_must_exist(tmp_path: Path):
    result = CoderToWriter(
        status="success", code_output="ok", artifacts=["results.json"]
    )
    with pytest.raises(QualityGateError):
        validate_coder_result(result, tmp_path)


def test_evidence_contract_requires_numeric_results_and_validation(tmp_path: Path):
    contract = {
        "status": "success",
        "subtask": "ques1",
        "data_scope": "official statement",
        "assumptions": [],
        "model": "rigid chain",
        "numeric_results": {"critical_time_s": 412.47},
        "validation": {"passed": True, "checks": ["distance residual"]},
        "conclusion_bounds": "fixed geometry",
    }
    (tmp_path / "results_ques1.json").write_text(json.dumps(contract), encoding="utf-8")
    _, loaded = load_evidence_contract(tmp_path, "ques1")
    assert loaded["numeric_results"]["critical_time_s"] == 412.47

    contract["numeric_results"] = {"note": "no calculation"}
    (tmp_path / "results_ques1.json").write_text(json.dumps(contract), encoding="utf-8")
    with pytest.raises(QualityGateError):
        load_evidence_contract(tmp_path, "ques1")


def test_success_contract_rejects_unresolved_blocked_fields(tmp_path: Path):
    contract = {
        "status": "success",
        "subtask": "eda",
        "data_scope": {
            "source": "official statement",
            "blocked_source_facts": "图4中A点极角仍未核验",
        },
        "assumptions": [],
        "model": "rigid chain",
        "numeric_results": {"node_count": 224},
        "validation": {"passed": True, "checks": ["topology"]},
        "conclusion_bounds": "pending figure fact",
    }
    path = tmp_path / "results_eda.json"
    path.write_text(json.dumps(contract), encoding="utf-8")
    with pytest.raises(QualityGateError, match="未解决阻塞项"):
        load_evidence_contract(tmp_path, "eda")

    contract["data_scope"]["blocked_source_facts"] = []
    path.write_text(json.dumps(contract), encoding="utf-8")
    load_evidence_contract(tmp_path, "eda")


FIGURE_DEPENDENT_SOURCE = (
    "初始时，运动端点位于轨迹第 16 圈 A 点处（见图 4）。构件结构示意见图 2。"
)


def test_source_figure_requirements_only_capture_geometric_facts():
    requirements = derive_source_figure_requirements(FIGURE_DEPENDENT_SOURCE)
    assert [item["figure"] for item in requirements] == [4]


def test_source_figure_requirements_include_boundary_geometry():
    source = "调头空间为直径9m的圆形区域（见图5），目标是到达该区域边界。"
    requirements = derive_source_figure_requirements(source)
    assert [item["figure"] for item in requirements] == [5]


def test_eda_requires_verifiable_source_figure_evidence(tmp_path: Path):
    document = fitz.open()
    document.new_page().insert_text((72, 72), "problem statement")
    document.new_page().insert_text((72, 72), "supporting table")
    figure_page = document.new_page()
    figure_page.insert_text((72, 72), "Figure 4 trajectory")
    figure_page.insert_text((100, 200), "O")
    figure_page.insert_text((200, 200), "A")
    figure_page.insert_text((240, 200), "x")
    figure_page.insert_text((100, 110), "y")
    document.save(tmp_path / "official.pdf")
    document.close()
    contract = {
        "status": "success",
        "subtask": "eda",
        "data_scope": "official statement",
        "assumptions": [],
        "model": "curve kinematics",
        "numeric_results": {"initial_turn": 16},
        "validation": {"passed": True, "checks": ["source audit"]},
        "conclusion_bounds": "source facts only",
    }
    path = tmp_path / "results_eda.json"
    path.write_text(json.dumps(contract), encoding="utf-8")
    with pytest.raises(QualityGateError, match="source_figures"):
        load_evidence_contract(tmp_path, "eda", source_text=FIGURE_DEPENDENT_SOURCE)

    contract["data_scope"] = {
        "source_figures": {
            "figure_4": {
                "file": "official.pdf",
                "page": 3,
                "figure": 4,
                "facts": ["A 点位于正 x 轴外端"],
            }
        }
    }
    path.write_text(json.dumps(contract), encoding="utf-8")
    with pytest.raises(QualityGateError, match="必须是 JSON 数组.*当前类型为 dict"):
        load_evidence_contract(tmp_path, "eda", source_text=FIGURE_DEPENDENT_SOURCE)

    contract["data_scope"] = {
        "source_figures": [
            {"file": "official.pdf", "page": 3, "figure": "图4", "facts": "已核验"}
        ]
    }
    path.write_text(json.dumps(contract), encoding="utf-8")
    with pytest.raises(QualityGateError, match="缺少可核验来源记录"):
        load_evidence_contract(tmp_path, "eda", source_text=FIGURE_DEPENDENT_SOURCE)

    contract["data_scope"]["source_figures"][0]["facts"] = (
        "图中未给出 A 点方向，需要自行约定极角"
    )
    path.write_text(json.dumps(contract), encoding="utf-8")
    with pytest.raises(QualityGateError, match="缺少可核验来源记录"):
        load_evidence_contract(tmp_path, "eda", source_text=FIGURE_DEPENDENT_SOURCE)

    contract["data_scope"]["source_figures"][0]["page"] = 1
    contract["data_scope"]["source_figures"][0]["facts"] = (
        "A 点位于正 x 轴外端，轨迹沿顺时针方向向内"
    )
    path.write_text(json.dumps(contract), encoding="utf-8")
    with pytest.raises(QualityGateError, match="实际图注页为第 3 页"):
        load_evidence_contract(tmp_path, "eda", source_text=FIGURE_DEPENDENT_SOURCE)

    contract["data_scope"]["source_figures"][0]["page"] = 3
    contract["data_scope"]["source_figures"][0]["facts"] = (
        "A 点标注于轨迹上；图中含 x 轴和 y 轴"
    )
    path.write_text(json.dumps(contract), encoding="utf-8")
    with pytest.raises(QualityGateError, match="PDF 文字标签坐标显示 A 点位于正 x 轴"):
        load_evidence_contract(tmp_path, "eda", source_text=FIGURE_DEPENDENT_SOURCE)

    contract["data_scope"]["source_figures"][0]["facts"] = (
        "A 点位于第一象限，沿顺时针方向向内"
    )
    path.write_text(json.dumps(contract), encoding="utf-8")
    with pytest.raises(QualityGateError, match="PDF 文字标签坐标显示 A 点位于正 x 轴"):
        load_evidence_contract(tmp_path, "eda", source_text=FIGURE_DEPENDENT_SOURCE)

    contract["data_scope"]["source_figures"][0]["facts"] = (
        "A 点位于正 x 轴外端，轨迹沿顺时针方向向内"
    )
    path.write_text(json.dumps(contract), encoding="utf-8")
    load_evidence_contract(tmp_path, "eda", source_text=FIGURE_DEPENDENT_SOURCE)

    contract["data_scope"]["source_figures"][0]["facts"] = (
        "A点位于螺线第16圈上，处于x正半轴，龙头前把手初始位于A点"
    )
    path.write_text(json.dumps(contract), encoding="utf-8")
    load_evidence_contract(tmp_path, "eda", source_text=FIGURE_DEPENDENT_SOURCE)


SERIAL_MEMBER_SOURCE = (
    "某板凳龙由 223 节板凳组成。龙头的板长为 341 cm，"
    "龙身和龙尾的板长均为 220 cm。每节板凳上均有两个孔，"
    "孔的中心距离最近的板头 27.5 cm。相邻两条板凳通过把手连接。"
)


def test_source_guardrail_derives_member_internal_distances():
    guardrail = derive_source_guardrails(SERIAL_MEMBER_SOURCE)
    assert guardrail is not None
    assert guardrail["entity_count"] == 223
    assert guardrail["node_count"] == 224
    distances = [item["node_distance_m"] for item in guardrail["member_constraints"]]
    assert distances == pytest.approx([2.86, 1.65])
    assert "不是参考答案" in format_source_guardrails(SERIAL_MEMBER_SOURCE)


def test_source_guardrail_derives_complete_handle_identity_contract():
    source = SERIAL_MEMBER_SOURCE + ("第1节为龙头，后面221节为龙身，最后1节为龙尾。")
    guardrail = derive_source_guardrails(source)
    assert guardrail is not None
    assert guardrail["handle_identity_contract"] == {
        "total_output_handles": 224,
        "head_front_node": 1,
        "body_front_nodes": [2, 222],
        "body_indices": [1, 221],
        "tail_front_node": 223,
        "tail_rear_node": 224,
        "tail_rear_included_in_total": True,
    }
    formatted = format_source_guardrails(source)
    assert '"tail_rear_node":224' in formatted
    assert '"tail_rear_included_in_total":true' in formatted


def test_source_guardrail_exposes_s_curve_external_tangency_contract():
    source = SERIAL_MEMBER_SOURCE + (
        "问题4调头路径是S形曲线，由两段圆弧相切连接，前一段圆弧半径是后一段的2倍。"
    )
    guardrail = derive_source_guardrails(source)
    assert guardrail is not None
    assert guardrail["s_curve_tangency_contract"] == {
        "required": ["反向转向的外切分支", "圆心距 R1+R2"],
        "forbidden": ["内切分支", "R1-R2", "R1±R2"],
    }


def test_source_guardrail_exposes_adjustment_before_shortening_baseline():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 调头路径为双圆弧。能否调整圆弧并保持相切，使调头曲线变短？"
        "问题 5 求最大速度。"
    )
    guardrail = derive_source_guardrails(source)
    assert guardrail is not None
    contract = guardrail["shortening_comparison_contract"]
    assert "调整前双圆弧构型" in contract["required"][0]
    assert "任取一个可行半径" in contract["forbidden"][0]


def test_source_guardrail_exposes_unbounded_path_extremum_contract():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 路径包含盘入螺线和盘出螺线。"
        "问题 5 请确定龙头的最大行进速度。"
    )
    guardrail = derive_source_guardrails(source)
    assert guardrail is not None
    contract = guardrail["unbounded_path_extremum_contract"]
    assert "逐次扩大外侧截断" in contract["required"][1]
    assert "速度倍率趋一" in contract["forbidden"][0]


def test_modeler_accepts_explicit_rejection_of_internal_s_curve_tangency():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 调头路径是S形曲线，由两段圆弧相切连接，"
        "前一段圆弧半径是后一段的2倍。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223节构件、224个节点，"
                "两类内部节点距为2.86m和1.65m。"
            ),
            "ques4": (
                "两弧选择外切分支，圆心距R1+R2；不得用内切R1-R2。"
                "位置与切向C1连续，连接点曲率发生跳变。"
            ),
            "sensitivity_analysis": "缩小相切方程求解容差。",
        }
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_run46_inconsistent_node_decomposition():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224个节点"
                "（1个龙头前把手+221个龙身前把手+1个龙尾后把手）。"
            ),
            "ques1": "用二维距离方程求224个节点。",
            "sensitivity_analysis": "检查网格。",
        }
    )
    with pytest.raises(QualityGateError, match="分项合计 223"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_run49_wrong_total_rigid_segment_count():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223节板凳、224个节点；"
                "223节板凳构成222个连接段，第1段2.86m，其余222段1.65m。"
            ),
            "ques1": "用二维距离方程递推全部节点。",
            "sensitivity_analysis": "检查网格。",
        }
    )
    with pytest.raises(QualityGateError, match="刚性段总数|223"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_allows_local_rigid_segment_subcounts():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223节板凳、224个节点。"
                "板内刚性段总数为223段：节点1-2为2.86m；"
                "节点2-3至节点222-223均为1.65m（共221段）；"
                "节点223-224为1.65m（1段）。"
            ),
            "ques1": "用二维距离方程递推全部节点。",
            "sensitivity_analysis": "检查网格。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_allows_total_count_followed_by_node_explanation_and_sum():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223节板凳。板内刚性段总数："
                "224节点间共223段；刚性段总数=1+221+1=223段。"
                "首段2.86m，其余为1.65m。"
            ),
            "ques1": "用二维距离方程递推全部节点。",
            "sensitivity_analysis": "检查网格。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_wrong_aggregate_count_with_breakdown():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223节板凳、224个节点；"
                "共222段：1段2.86m，其余均为1.65m。"
            ),
            "ques1": "用二维距离方程递推全部节点。",
            "sensitivity_analysis": "检查网格。",
        }
    )
    with pytest.raises(QualityGateError, match="刚性段总数|223"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_source_guardrail_surfaces_cross_question_and_curve_contracts():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 请确定不能再继续盘入且板凳之间不发生碰撞的时刻。"
        "问题 3 请使龙头前把手盘入到调头空间的边界。"
        "问题 4 调头路径由两段圆弧相切连接，前一段半径是后一段的2倍。"
    )
    guardrail = derive_source_guardrails(source)
    assert guardrail is not None
    assert "cross_question_feasibility" in guardrail
    assert "piecewise_curve_continuity" in guardrail


def test_source_guardrail_accepts_pdf_extraction_line_break_after_each():
    source = SERIAL_MEMBER_SOURCE.replace("每节板凳", "每\n节板凳")
    guardrail = derive_source_guardrails(source)
    assert guardrail is not None
    assert guardrail["node_count"] == 224


def test_modeler_rejects_cross_member_average_for_serial_members():
    wrong = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个把手，间距2.255和1.65 m",
            "ques1": "candidate",
            "sensitivity_analysis": "candidate",
        }
    )
    with pytest.raises(QualityGateError, match="节点数/内部节点距离"):
        validate_modeler_result(wrong, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)

    correct = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques1": "candidate",
            "sensitivity_analysis": "candidate",
        }
    )
    validate_modeler_result(correct, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)

    correct_centimetres = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长286 cm和165 cm",
            "ques1": "candidate",
            "sensitivity_analysis": "candidate",
        }
    )
    validate_modeler_result(
        correct_centimetres, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE
    )


def test_modeler_allows_half_length_sum_as_collision_candidate_filter():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques2": (
                "相邻板体在共享把手处允许接触且不视为碰撞。"
                "仅检查非相邻板体；候选对粗筛使用两板矩形中心距离"
                "小于两板半长之和加板宽，随后用含真实长度和朝向的"
                "SAT精确判定矩形相交。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "所有板凳的板宽为30cm。问题2检查碰撞。"
    validate_modeler_result(result, {"ques2"}, source_text=source)


@pytest.mark.parametrize(
    "unsafe_filter",
    [
        "候选对粗筛使用两板矩形中心距离小于两板半宽之和的0.9倍，随后SAT精判。",
        "候选对粗筛间距阈值如<0.4m，随后SAT精确判定矩形相交。",
    ],
)
def test_modeler_rejects_collision_prefilter_that_can_miss_long_bodies(
    unsafe_filter: str,
):
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques2": (
                "相邻连接板体共享把手处允许接触，仅检查非相邻板体；" + unsafe_filter
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "所有板凳的板宽均为30cm。问题2检查碰撞。"
    with pytest.raises(QualityGateError, match="粗筛不保证无漏检"):
        validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_allows_half_diagonal_collision_prefilter():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques2": (
                "相邻连接板体共享把手处允许接触，仅检查非相邻板体；"
                "候选对粗筛使用两板半对角线包围圆半径之和，"
                "随后以有向矩形SAT精确判定相交。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "所有板凳的板宽均为30cm。问题2检查碰撞。"
    validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_rejects_hole_distance_inside_collision_half_diagonal():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques2": (
                "相邻连接板体共享把手处允许接触，仅检查非相邻板体；"
                "粗筛包围圆半径取龙头sqrt(2.86^2+0.30^2)/2、"
                "龙身sqrt(1.65^2+0.30^2)/2，随后用有向矩形SAT精判。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "所有板凳的板宽均为30cm。问题2检查碰撞。"
    with pytest.raises(QualityGateError, match="包围圆也必须使用题面板体原长"):
        validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_rejects_parenthesized_hole_distance_in_collision_half_diagonal():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques2": (
                "相邻连接板体共享把手处允许接触，仅检查非相邻板体；"
                "粗筛半对角=sqrt((2.86)^2+0.30^2)/2，随后用有向矩形SAT精判。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "所有板凳的板宽均为30cm。问题2检查碰撞。"
    with pytest.raises(QualityGateError, match="包围圆也必须使用题面板体原长"):
        validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_rejects_inconsistent_simple_arithmetic_claim():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224个节点；"
                "串链总长=2.86+222×1.65+0.275=369.71 m。"
            ),
            "ques1": "使用二维弦长约束递推。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="算术等式不成立"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_accepts_product_in_first_arithmetic_term():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224个节点；"
                "孔距闭合为2*0.275+2.86=3.41m，2×0.275+1.65=2.20m。"
            ),
            "ques1": "使用二维弦长方程。",
            "sensitivity_analysis": "检查数值容差。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_wrong_product_in_first_arithmetic_term():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点；孔距闭合为2*0.275+2.86=3.50m。",
            "ques1": "使用二维弦长方程。",
            "sensitivity_analysis": "检查数值容差。",
        }
    )
    with pytest.raises(QualityGateError, match="算术等式不成立"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_accepts_subtraction_and_product_arithmetic():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224个节点；"
                "孔距闭合为3.41-2*0.275=2.86m，2.20-2×0.275=1.65m。"
            ),
            "ques1": "使用二维弦长方程。",
            "sensitivity_analysis": "检查数值容差。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_inconsistent_subtraction_chain():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224个节点；"
                "龙身孔距=2.20-0.275=2.20-0.275-0.275×2=1.65m。"
            ),
            "ques1": "使用二维弦长方程。",
            "sensitivity_analysis": "检查数值容差。",
        }
    )
    with pytest.raises(QualityGateError, match="算术等式不成立"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_body_node_range_that_includes_tail_front_handle():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m；"
                "节点1为龙头前把手，节点k（2≤k≤223）=第k−1节龙身前把手，"
                "节点224为龙尾后把手。"
            ),
            "ques1": "使用二维弦长约束递推。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="把龙尾前把手误标为龙身"):
        validate_modeler_result(
            result,
            {"ques1"},
            source_text=SERIAL_MEMBER_SOURCE + "后面221节为龙身，最后1节为龙尾。",
        )


def test_modeler_rejects_body_index_range_that_includes_nonexistent_body():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m；"
                "节点编号1=龙头前把手，k=1..222为第k节龙身前把手，"
                "223=龙尾前把手，224=龙尾后把手。"
            ),
            "ques1": "使用二维弦长约束递推。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="只有 221 节龙身"):
        validate_modeler_result(
            result,
            {"ques1"},
            source_text=SERIAL_MEMBER_SOURCE + "后面221节为龙身，最后1节为龙尾。",
        )


def test_modeler_rejects_run56_duplicate_tail_front_node_identity():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m；"
                "节点222=第221节龙身前把手=龙尾前把手，"
                "节点223=龙尾前把手，节点224=龙尾后把手。"
            ),
            "ques1": "使用二维弦长约束递推。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="龙尾前把手.*节点 223.*节点 222"):
        validate_modeler_result(
            result,
            {"ques1"},
            source_text=SERIAL_MEMBER_SOURCE + "后面221节为龙身，最后1节为龙尾。",
        )


def test_modeler_rejects_tail_front_equated_to_last_body_front_without_node_number():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m；"
                "龙尾前把手即第221节龙身的前把手。"
            ),
            "ques1": "使用二维弦长约束递推。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="不能等同于最后一节龙身前把手"):
        validate_modeler_result(
            result,
            {"ques1"},
            source_text=SERIAL_MEMBER_SOURCE + "后面221节为龙身，最后1节为龙尾。",
        )


def test_modeler_rejects_explicit_body_node_mapping_shifted_into_tail():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m；"
                "节点223=第221节龙身前把手即龙尾前把手；"
                "节点224=龙尾后把手。"
            ),
            "ques1": "使用二维弦长约束递推。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="第 221 节龙身前把手应为节点 222"):
        validate_modeler_result(
            result,
            {"ques1"},
            source_text=SERIAL_MEMBER_SOURCE + "后面221节为龙身，最后1节为龙尾。",
        )


def test_modeler_rejects_numeric_body_range_shifted_into_tail():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m；"
                "节点2~223依次为第1~221节龙身的前把手；"
                "龙尾前后把手身份另列。"
            ),
            "ques1": "使用二维弦长约束递推。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="显式龙身节点区间"):
        validate_modeler_result(
            result,
            {"ques1"},
            source_text=SERIAL_MEMBER_SOURCE + "后面221节为龙身，最后1节为龙尾。",
        )


def test_modeler_rejects_wrong_explicit_tail_handle_nodes():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m；"
                "节点224为龙尾前把手，节点225为龙尾后把手。"
            ),
            "ques1": "使用二维弦长约束递推。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="龙尾前把手.*节点 223.*节点 224"):
        validate_modeler_result(
            result,
            {"ques1"},
            source_text=SERIAL_MEMBER_SOURCE + "后面221节为龙身，最后1节为龙尾。",
        )


@pytest.mark.parametrize(
    "output_contract",
    [
        "输出301×225位置速度。",
        "共224节点+龙尾后把手。",
        "龙尾后把手是第225个待输出对象。",
    ],
)
def test_modeler_rejects_extra_225th_output_handle(output_contract):
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques1": f"使用二维弦长约束递推；{output_contract}",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="矛盾的把手/节点总数"):
        validate_modeler_result(
            result,
            {"ques1"},
            source_text=SERIAL_MEMBER_SOURCE + "后面221节为龙身，最后1节为龙尾。",
        )


def test_modeler_rejects_unbounded_body_node_identity_mapping():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m；"
                "节点k为第k-1节前把手，龙头前把手为节点1，龙尾后把手为节点224。"
            ),
            "ques1": "使用二维弦长约束递推。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="缺少龙身区间|不能把无界 k"):
        validate_modeler_result(
            result,
            {"ques1"},
            source_text=SERIAL_MEMBER_SOURCE + "后面221节为龙身，最后1节为龙尾。",
        )


@pytest.mark.parametrize(
    "mapping",
    [
        "节点k=第k-1节龙身前把手，仅限2≤k≤222",
        "节点k=第k-1节龙身前把手，k=2..222",
        "节点k=第k-1节龙身前把手，k∈[2,222]",
    ],
)
def test_modeler_accepts_postposed_body_node_identity_range(mapping):
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m；"
                f"节点1=龙头前把手；{mapping}；"
                "节点223=龙尾前把手；节点224=龙尾后把手。"
            ),
            "ques1": "使用二维弦长约束递推。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    validate_modeler_result(
        result,
        {"ques1"},
        source_text=SERIAL_MEMBER_SOURCE + "后面221节为龙身，最后1节为龙尾。",
    )


def test_modeler_accepts_t_from_zero_collision_history_word_order():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques2": "t从0推进并逐步做SAT检测，完整扫描确认T_stop前无碰撞。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    validate_modeler_result(
        result,
        {"ques2"},
        source_text=(
            SERIAL_MEMBER_SOURCE
            + "问题2 确定终止时刻，使得板凳之间不发生碰撞，即不能再继续盘入。"
        ),
    )


def test_modeler_rejects_result2_full_time_history_contract():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques2": (
                "从t=0推进并做SAT检测，完整扫描确认T_stop前无碰撞；"
                "输出result2.xlsx（T_stop时刻全部节点+0..T_stop逐秒）。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="只交付终止时刻"):
        validate_modeler_result(
            result,
            {"ques2"},
            source_text=SERIAL_MEMBER_SOURCE + "问题2 将此时结果存放到result2.xlsx。",
        )


def test_modeler_rejects_224th_board_in_adjacent_collision_pairs():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques2": (
                "从t=0推进并做SAT检测，完整扫描确认T_stop前无碰撞；"
                "排除相邻板凳对(1,2),(2,3),...,(223,224)。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="最多到 \\(222,223\\)"):
        validate_modeler_result(
            result,
            {"ques2"},
            source_text=SERIAL_MEMBER_SOURCE + "问题2 确定首次碰撞终止时刻。",
        )


def test_modeler_accepts_braced_spatial_cutoff_sequence_with_stability_metric():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques3": "从外侧截断连续回放至目标边界，搜索完整0<p≤p_max。",
            "sensitivity_analysis": (
                "取r_start∈{30,60,120}m分别复算，验证p_min差<1e-4m。"
            ),
        }
    )
    validate_modeler_result(
        result,
        {"ques3"},
        source_text=SERIAL_MEMBER_SOURCE + "问题3 确定最小螺距。",
    )


def test_modeler_accepts_symbolically_indexed_expanding_spatial_cutoffs():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "从外侧截断连续回放至目标边界，搜索完整0<p≤p_max。"
                "采用两个逐步扩大的外侧截断r_start(1)与"
                "r_start(2)=r_start(1)+Δr（Δr≥50m），"
                "分别求解p_min^(1)与p_min^(2)，"
                "报告差值<1e-4m。"
            ),
            "sensitivity_analysis": "检查截断稳定性。",
        }
    )
    validate_modeler_result(
        result,
        {"ques3"},
        source_text=SERIAL_MEMBER_SOURCE + "问题3 确定最小螺距。",
    )


def test_modeler_rejects_inverse_order_wrong_tail_front_identity():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m；"
                "第k节龙身前把手为节点k+1（1≤k≤221），"
                "龙尾前把手为节点222，龙尾后把手为节点224。"
            ),
            "ques1": "使用二维弦长约束递推。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="龙尾前把手.*节点 223.*节点 222"):
        validate_modeler_result(
            result,
            {"ques1"},
            source_text=SERIAL_MEMBER_SOURCE + "后面221节为龙身，最后1节为龙尾。",
        )


def test_modeler_allows_independently_computed_full_chain_speeds():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques1": "用二维欧氏弦长和切向投影递推速度。",
            "ques5": (
                "逐秒回放全链节点速度（各节点速度由刚性距离约束求导独立计算，"
                "相邻节点速度比等于弦向量在两端切向上的投影比，不假定所有节点等速）。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    validate_modeler_result(
        result, {"ques1", "ques5"}, source_text=SERIAL_MEMBER_SOURCE
    )


def test_modeler_does_not_read_scan_step_as_collision_body_length():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques2": (
                "仅检查非相邻板凳对（|i-j|>1），有向矩形长=3.41m或2.20m，"
                "宽=0.30m，并用SAT精确判定相交；板凳越挤，先大步长0.05m扫描。"
            ),
            "sensitivity_analysis": "检查扫描步长。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "所有板凳的板宽均为30cm。问题2检查碰撞。"
    validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_rejects_run56_presupposed_sat_center_heuristic_order():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques2": (
                "排除相邻板体后，用有向矩形SAT检测非相邻板凳的实体碰撞；"
                "与中心点距离阈值对比，确认SAT结果不早于中心点法。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "所有板凳的板宽均为30cm。问题2检查碰撞。"
    with pytest.raises(QualityGateError, match="不存在固定排序"):
        validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_accepts_sat_center_comparison_without_presupposed_order():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques2": (
                "排除相邻板体后，用有向矩形SAT检测非相邻板凳的实体碰撞；"
                "另算中心点距离启发式并报告实际时刻差异，不预设两者先后。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "所有板凳的板宽均为30cm。问题2检查碰撞。"
    validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_rejects_run57_presupposed_collision_location():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques2": (
                "排除相邻板体后，用有向矩形SAT扫描全部非相邻板凳；"
                "碰撞必然出现在内圈高曲率处。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "所有板凳的板宽均为30cm。问题2检查碰撞。"
    with pytest.raises(QualityGateError, match="不能.*预设"):
        validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_accepts_high_curvature_as_collision_check_priority():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques2": (
                "排除相邻板体后，用有向矩形SAT扫描全部非相邻板凳；"
                "高曲率区仅作为重点检查区域，不预设碰撞对或时刻。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "所有板凳的板宽均为30cm。问题2检查碰撞。"
    validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_accepts_nonadjacent_pair_inequality_as_neighbour_exclusion():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques2": (
                "每时刻检查非相邻板凳对（|i-j|>1）的有向矩形是否相交；"
                "矩形长=3.41m或2.20m、宽=0.30m，最终用SAT精确判定。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "所有板凳的板宽均为30cm。问题2检查碰撞。"
    validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_rejects_local_arc_metric_as_exact_rigid_chord():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques1": "由距离约束反解 √(r²+(dr/dθ)²)Δθ=节点间距。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="固定弦长/弧长关系"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_allows_head_travel_arc_length_before_chord_residual_clause():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques1": (
                "300s时龙头总弧长=300m积分吻合，"
                "相邻节点距严格由二维欧氏弦长方程求解，残差<1e-6m。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_allows_arc_length_parameterization_before_euclidean_chord_clause():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m。"
                "轨迹模拟用弧长参数化，刚性链用相邻节点欧氏距离=1.65m"
                "或2.86m反解曲线参数。"
            ),
            "ques1": "路径弧长仅参数化龙头运动，节点始终满足二维欧氏弦长方程。",
            "sensitivity_analysis": "检查残差。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_positive_arc_primitive_increment_for_inward_motion():
    source = SERIAL_MEMBER_SOURCE + ("问题 1 沿等距螺线顺时针盘入，龙头速度为1m/s。")
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m。",
            "ques1": (
                "令S(θ)随θ递增，并使S(θ(t))=S(θ_0)-vt。验证S(θ(t))-S(θ(t-1))=1m。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="差分方向写反"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_positive_arc_increment_inside_absolute_residual():
    source = SERIAL_MEMBER_SOURCE + ("问题 1 沿等距螺线顺时针盘入，龙头速度为1m/s。")
    wrong = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m。",
            "ques1": ("使S(θ(t))=S(θ_0)-vt，并验证|S(θ(t))-S(θ(t-Δt))-1m|<1e-6。"),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="差分方向写反"):
        validate_modeler_result(wrong, {"ques1"}, source_text=source)

    correct = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m。",
            "ques1": ("使S(θ(t))=S(θ_0)-vt，并验证|S(θ(t))-S(θ(t-Δt))|=1m。"),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    validate_modeler_result(correct, {"ques1"}, source_text=source)


def test_modeler_allows_fixed_q1_q2_pitch_explicitly_not_perturbed():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 1 沿螺距为55cm的等距螺线盘入。问题 2 沿用问题1的螺线。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m。",
            "ques1": "按题面固定螺距求解。",
            "ques2": "沿用问题1轨迹并求首次碰撞。",
            "sensitivity_analysis": (
                "问题1/2的螺距55cm为固定源常量，不扰动；只检查数值求解容差。"
            ),
        }
    )
    validate_modeler_result(result, {"ques1", "ques2"}, source_text=source)


def test_modeler_allows_expanding_r_start_to_verify_critical_p_stability():
    source = SERIAL_MEMBER_SOURCE + ("问题 3 请确定龙头盘入至4.5m边界的最小螺距。")
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m。",
            "ques3": (
                "从r_start=60m的外侧截断回放至4.5m边界，再扩大r_start验证临界p稳定。"
            ),
            "sensitivity_analysis": "扩大空间截断并检查结果稳定性。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_allows_q3_cutoff_stability_metric_in_sensitivity_section():
    source = SERIAL_MEMBER_SOURCE + ("问题 3 请确定龙头盘入至4.5m边界的最小螺距。")
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m。",
            "ques3": (
                "从r_start=50m的外侧截断回放至4.5m边界，增加截断半径复核。"
            ),
            "sensitivity_analysis": (
                "问题3将r_start扩大至1.5-2倍，确认最小螺距变化<1%。"
            ),
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_allows_radius_first_q3_cutoff_impact_bound_in_sensitivity():
    source = SERIAL_MEMBER_SOURCE + "问题 3 请确定龙头盘入至4.5m边界的最小螺距。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65m。",
            "ques3": (
                "从充分外侧r_start=20m连续回放到边界，用SAT检查全链碰撞，"
                "并以扩大后的截断复算。"
            ),
            "sensitivity_analysis": "问题3外侧截断半径增加50%对p_min影响（需<1e-4m）。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_allows_expanded_cutoff_where_extremum_stops_increasing():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 路径包含盘入螺线、调头曲线和盘出螺线。"
        "问题 5 求使所有把手速度不超过2m/s的龙头最大速度。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m。",
            "ques4": "按题面路径求解。",
            "ques5": (
                "扫描所有节点和路径段得到全路径最大倍率，并扩大截断验证极值不再增加。"
            ),
            "sensitivity_analysis": "检查空间截断与步长。",
        }
    )
    validate_modeler_result(result, {"ques4", "ques5"}, source_text=source)


def test_modeler_rejects_trailing_theta_subtraction_formula():
    source = SERIAL_MEMBER_SOURCE + "沿等距螺线顺时针盘入。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques1": ("采用r=pθ/(2π)，龙头盘入θ递减；后续节点递推取θ_{i+1}=θ_i-Δθ。"),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="跟随次序矛盾"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_direct_reversed_trailing_theta_inequality():
    source = SERIAL_MEMBER_SOURCE + "沿等距螺线顺时针盘入。"
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223个构件、224节点，"
                "内部弦长2.86m和1.65m；采用r=pθ/(2π)。"
            ),
            "ques1": "对节点i已知θ_i，求θ_{i+1}<θ_i满足刚性弦长约束。",
            "sensitivity_analysis": "检查时间步长和弦长残差。",
        }
    )
    with pytest.raises(QualityGateError, match="盘入方向|跟随次序"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_allows_decreasing_theta_between_time_steps():
    source = SERIAL_MEMBER_SOURCE + "沿等距螺线顺时针盘入。"
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223个构件、224节点，"
                "内部弦长2.86m和1.65m；采用r=pθ/(2π)。"
            ),
            "ques1": (
                "龙头满足S(θ(t))=S(θ_0)-v·t，用牛顿法对弧长方程求"
                "θ_{i+1}，确保θ_{i+1}<θ_i；链上节点按固定弦长逐点求解，"
                "且后续节点满足θ_k>θ_{k-1}。"
            ),
            "sensitivity_analysis": "检查时间步长和弦长残差。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_allows_explicit_rejection_of_decreasing_chain_theta():
    source = SERIAL_MEMBER_SOURCE + "沿等距螺线顺时针盘入。"
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223个构件、224节点，内部弦长2.86m和1.65m；"
                "采用r=pθ/(2π)，后续节点在外侧，不得写成θ_{i+1}=θ_i-Δθ。"
            ),
            "ques1": "按固定弦长逐点求解，并令后续节点满足θ_k>θ_{k-1}。",
            "sensitivity_analysis": "检查弦长残差。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=source)


@pytest.mark.parametrize(
    "contradiction",
    [
        "共223个空间节点",
        "逐节求解后得到全部226个把手位置",
        "相邻节点间距并非固定，取决于两节夹角",
        "相邻把手间空间距离取决于两板夹角",
    ],
)
def test_modeler_rejects_source_topology_contradictions(contradiction: str):
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223个构件、224节点，"
                f"内部弦长2.86和1.65 m；{contradiction}"
            ),
            "ques1": "candidate",
            "sensitivity_analysis": "candidate",
        }
    )
    with pytest.raises(QualityGateError, match="矛盾"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_wrong_per_time_output_object_count():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86m和1.65m。",
            "ques1": "输出契约：每个时刻223个对象，按固定弦长递推。",
            "sensitivity_analysis": "检查节点数。",
        }
    )
    with pytest.raises(QualityGateError, match="把手/节点总数|224"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_all_rigid_nodes_have_head_speed():
    source = (
        SERIAL_MEMBER_SOURCE
        + "龙头前把手的行进速度保持 1 m/s，请给出整个舞龙队各把手的速度。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques1": "因所有把手刚性连接，任意把手速率相同为 1 m/s。",
            "sensitivity_analysis": "检查时间步长",
        }
    )
    with pytest.raises(QualityGateError, match="等速"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_member_chord_used_as_arc_coordinate_gap():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques1": (
                "建立弧长坐标；龙头后把手沿螺线后退弧长 Δs=2.86 m，"
                "后续节点弧长间隔取为孔距1.65 m。"
            ),
            "sensitivity_analysis": "检查时间步长",
        }
    )
    with pytest.raises(QualityGateError, match="弧长关系"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_oversized_sections():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: " + "x" * 1601,
            "ques1": "candidate",
            "sensitivity_analysis": "candidate",
        }
    )
    with pytest.raises(QualityGateError, match="过长"):
        validate_modeler_result(result, {"ques1"})


def test_eda_evidence_rechecks_source_topology_numbers(tmp_path: Path):
    contract = {
        "status": "success",
        "subtask": "eda",
        "data_scope": "official statement",
        "assumptions": [],
        "model": "serial rigid members",
        "numeric_results": {"node_count": 224, "distances_m": [2.255, 1.65]},
        "validation": {"passed": True, "checks": ["topology"]},
        "conclusion_bounds": "source facts only",
    }
    path = tmp_path / "results_eda.json"
    path.write_text(json.dumps(contract), encoding="utf-8")
    with pytest.raises(QualityGateError, match="2.86"):
        load_evidence_contract(tmp_path, "eda", source_text=SERIAL_MEMBER_SOURCE)

    contract["numeric_results"]["distances_m"] = [2.86, 1.65]
    path.write_text(json.dumps(contract), encoding="utf-8")
    load_evidence_contract(tmp_path, "eda", source_text=SERIAL_MEMBER_SOURCE)


def test_eda_rejects_structured_wrong_node_count_even_when_224_appears_elsewhere(
    tmp_path: Path,
):
    contract = {
        "status": "success",
        "subtask": "eda",
        "data_scope": "official statement",
        "assumptions": ["题面应推导 224 个节点"],
        "model": "serial rigid members",
        "numeric_results": {
            "node_count": 223,
            "distances_m": [2.86, 1.65],
            "unrelated_check": 224,
        },
        "validation": {"passed": True, "checks": ["topology"]},
        "conclusion_bounds": "source facts only",
    }
    path = tmp_path / "results_eda.json"
    path.write_text(json.dumps(contract), encoding="utf-8")
    with pytest.raises(QualityGateError, match="题面拓扑矛盾"):
        load_evidence_contract(tmp_path, "eda", source_text=SERIAL_MEMBER_SOURCE)


def test_eda_rejects_cross_member_half_sum_even_with_correct_numbers(tmp_path: Path):
    contract = {
        "status": "success",
        "subtask": "eda",
        "data_scope": "official statement",
        "assumptions": [],
        "model": "serial rigid members",
        "numeric_results": {
            "node_count": 224,
            "distances_m": [2.86, 1.65],
            "conn_head_body": "(2.86+1.65)/2 = 2.255 m",
        },
        "validation": {"passed": True, "checks": ["topology"]},
        "conclusion_bounds": "source facts only",
    }
    path = tmp_path / "results_eda.json"
    path.write_text(json.dumps(contract), encoding="utf-8")
    with pytest.raises(QualityGateError, match="半长平均"):
        load_evidence_contract(tmp_path, "eda", source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_inward_spiral_with_reversed_parameter_order():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224个节点，内部距离2.86m和1.65m；"
                "采用r=(0.55/(2π))θ。"
            ),
            "ques1": ("龙头盘入时极角随时间增大；后续节点取较小θ根，即θ_i<θ_{i-1}。"),
            "sensitivity_analysis": "检查步长和残差。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "沿螺距55cm的等距螺线顺时针盘入。"
    with pytest.raises(QualityGateError, match="盘入方向|跟随次序"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_geodesic_distance_and_angle_increase_synonyms():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224个节点，内部距离2.86m和1.65m；"
                "采用r=(0.55/(2π))θ。"
            ),
            "ques1": (
                "第k+1节点由测地距离约束s(θ_{k+1})-s(θ_k)=d_k求解。"
                "验证龙头角度随盘入单调递增。"
            ),
            "sensitivity_analysis": "检查步长和残差。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "沿螺距55cm的等距螺线顺时针盘入。"
    with pytest.raises(QualityGateError, match="弦长|弧长关系"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_alternating_head_and_body_segment_lengths():
    result = ModelerToCoder(
        questions_solution={
            "eda": ("source_parameter_audit: 224个节点，内部距离2.86m和1.65m。"),
            "ques1": "使用2.86m和1.65m交替的刚性杆串联。",
            "sensitivity_analysis": "检查步长和残差。",
        }
    )
    with pytest.raises(QualityGateError, match="弦长|弧长关系"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_wrong_remaining_segment_count():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m；"
                "龙头段2.86m，后续221段1.65m。"
            ),
            "ques1": "使用欧氏距离方程递推。",
            "sensitivity_analysis": "检查残差。",
        }
    )
    with pytest.raises(QualityGateError, match="后续段数量|应有 222 段"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_wrong_repeated_segment_multiplier():
    result = ModelerToCoder(
        questions_solution={
            "eda": ("source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。"),
            "ques1": "总链长按2.86+221×1.65=367.51m核对。",
            "sensitivity_analysis": "检查残差。",
        }
    )
    with pytest.raises(QualityGateError, match="后续段乘数|应乘 222"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_arc_length_claimed_identical_to_chord():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques1": "任意相邻把手中心之间的螺线弧长恒等于两点弦长。",
            "sensitivity_analysis": "检查残差。",
        }
    )
    with pytest.raises(QualityGateError, match="弦长|弧长关系"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_turn_count_divided_by_two_pi_again():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 舞龙队沿螺距为 55 cm 的等距螺线盘入。"
        "初始时，龙头位于螺线第 16 圈。"
    )
    wrong = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques1": "第16圈起点半径约R=16×0.55/(2π)=1.4m。",
            "sensitivity_analysis": "检查残差。",
        }
    )
    with pytest.raises(QualityGateError, match="重复除以 2π|半径不一致"):
        validate_modeler_result(wrong, {"ques1"}, source_text=source)

    correct = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques1": ("第16圈起点半径R=16×0.55=8.8m；总链长为2.86+222×1.65=369.16m。"),
            "sensitivity_analysis": "检查残差。",
        }
    )
    validate_modeler_result(correct, {"ques1"}, source_text=source)


def test_modeler_does_not_confuse_turning_radius_with_initial_radius():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 舞龙队沿螺距为 55 cm 的等距螺线盘入。"
        "初始时，龙头位于螺线第 16 圈。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques1": "初始半径R0=16p必须大于调头空间半径4.5m。",
            "sensitivity_analysis": "检查残差。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_does_not_confuse_later_width_with_initial_radius():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 舞龙队沿螺距为 55 cm 的等距螺线盘入。"
        "初始时，龙头位于螺线第 16 圈。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques1": ("终止时刻受第16圈半径8.8m起点、螺距0.55m和板宽0.30m共同决定。"),
            "sensitivity_analysis": "检查残差。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_does_not_confuse_boundary_after_symbolic_initial_radius():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 舞龙队沿螺距为 55 cm 的等距螺线盘入。"
        "初始时，龙头位于螺线第 16 圈。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques1": ("调头空间半径R=4.5m；龙头从第16圈（半径16p）盘入至r=4.5m边界。"),
            "sensitivity_analysis": "检查残差。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_does_not_bind_later_turning_radius_in_sensitivity_list():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 舞龙队沿螺距为 55 cm 的等距螺线盘入。"
        "初始时，龙头位于螺线第 16 圈。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques1": "第16圈起点半径为8.8m。",
            "sensitivity_analysis": (
                "禁止扰动：问题1/4固定螺距、龙头初始圈数（第16圈对应r=8.8m）、"
                "龙头速度1m/s、调头圆半径4.5m、R1:R2=2:1。"
            ),
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_does_not_bind_q3_initial_to_q1_turn_radius():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 舞龙队沿螺距为 55 cm 的等距螺线盘入。"
        "初始时，龙头位于螺线第 16 圈。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到半径4.5m的调头空间边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques1": "第16圈起点半径为8.8m。",
            "ques3": (
                "从边界向外覆盖连续构型，逐步扩大搜索上界直至临界值稳定；"
                "以有向矩形SAT检查全链到达半径4.5m边界前的碰撞。"
            ),
            "sensitivity_analysis": "检查搜索上界。",
        }
    )
    validate_modeler_result(result, {"ques1", "ques3"}, source_text=source)


def test_modeler_rejects_q3_r_start_borrowed_from_q1_turn_count():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 舞龙队沿螺距为55cm的等距螺线盘入，"
        "初始时龙头位于螺线第16圈。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到半径4.5m的调头空间边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques1": "第16圈起点半径为8.8m。",
            "ques3": (
                "给定p时从充分外侧r_start（如16p）回放全链到4.5m边界，"
                "再扩大外侧截断并复算p_min稳定性。"
            ),
            "sensitivity_analysis": "检查空间截断。",
        }
    )
    with pytest.raises(QualityGateError, match="问题三擅自采用问题一的第16圈起点"):
        validate_modeler_result(result, {"ques1", "ques3"}, source_text=source)


def test_modeler_rejects_arc_gap_and_all_chain_equal_speed_variants():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques1": ("相邻节点间弧长=节点距；弧长约束求导后，全链速度大小均为1m/s。"),
            "sensitivity_analysis": "检查残差。",
        }
    )
    with pytest.raises(QualityGateError, match="弦长|弧长关系|等速"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)

    correct = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques1": "由欧氏距离约束求导，各节点速度方向沿切向，大小不相等。",
            "sensitivity_analysis": "检查残差。",
        }
    )
    validate_modeler_result(correct, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_allows_explicit_rejection_of_all_chain_equal_speed():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques1": (
                "由固定弦长约束求导，速度按弦向量与两端切向投影比传播，"
                "不可直接令全链速度相等。"
            ),
            "sensitivity_analysis": "检查速度递推与位置差分。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_q3_requires_body_feasibility_before_turning_boundary():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 请确定盘入终止时刻，使得板凳之间不发生碰撞，"
        "即舞龙队不能再继续盘入的时间。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    wrong = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": "仅此约束为16p≥4.5，故p_min=4.5/16。",
            "sensitivity_analysis": "检查边界。",
        }
    )
    with pytest.raises(QualityGateError, match="板凳实体|零行程下界"):
        validate_modeler_result(wrong, {"ques3"}, source_text=source)

    correct = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "对候选螺距递推全链位形，在龙头到达4.5m边界前，"
                "用有向矩形SAT复核板凳实体碰撞，再二分最小可行螺距。"
            ),
            "sensitivity_analysis": "检查边界。",
        }
    )
    validate_modeler_result(correct, {"ques3"}, source_text=source)


def test_modeler_accepts_q3_sat_and_boundary_event_in_adjacent_clauses():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到半径4.5m的调头空间边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "对候选p从充分外侧沿螺线回放整链SAT碰撞；"
                "若碰撞前龙头前把手到达r=4.5m边界，则该p可行。"
            ),
            "sensitivity_analysis": "扩大外侧截断并复算临界螺距。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_unreferenced_q1_initial_state_in_q3():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 初始时龙头位于第16圈。"
        "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "继承问题1的初始第16圈，令16p≥4.5；同时用有向矩形SAT检查到达边界前碰撞。"
            ),
            "sensitivity_analysis": "检查边界。",
        }
    )
    with pytest.raises(QualityGateError, match="擅自继承问题一"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_accepts_explicit_denial_of_q1_turn_count_in_q3():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 初始时龙头位于第16圈。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "题面未给问题3起点，不得假设为第16圈；从充分外侧安全域回放，"
                "用有向矩形SAT检查全链；扩大截断半径后临界值稳定，"
                "p_min变化<1e-4m。"
            ),
            "sensitivity_analysis": "检查扩大外侧截断后的临界值变化。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_eda_contract_rejects_q1_pitch_as_q3_lower_bound(tmp_path: Path):
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题 1 舞龙队沿螺距为 55 cm 的等距螺线盘入，初始位于第16圈。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    contract = {
        "status": "success",
        "subtask": "eda",
        "data_scope": "official statement",
        "assumptions": [],
        "model": "serial rigid members",
        "numeric_results": {
            "node_count": 224,
            "distances_m": [2.86, 1.65],
        },
        "validation": {"passed": True, "checks": ["topology"]},
        "conclusion_bounds": {
            "问题1": "螺距0.55m，初始第16圈。",
            "问题3": "最小螺距 p_min∈(0.55, 若干m]。",
        },
    }
    path = tmp_path / "results_eda.json"
    path.write_text(json.dumps(contract), encoding="utf-8")
    with pytest.raises(QualityGateError, match="问题一的螺距数值"):
        load_evidence_contract(tmp_path, "eda", source_text=source)


def test_eda_contract_allows_q3_domain_independent_of_q1_pitch(tmp_path: Path):
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题 1 舞龙队沿螺距为 55 cm 的等距螺线盘入，初始位于第16圈。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间边界。"
    )
    contract = {
        "status": "success",
        "subtask": "eda",
        "data_scope": "official statement",
        "assumptions": [],
        "model": "serial rigid members",
        "numeric_results": {
            "node_count": 224,
            "distances_m": [2.86, 1.65],
        },
        "validation": {"passed": True, "checks": ["topology"]},
        "conclusion_bounds": {
            "问题1": "螺距0.55m，初始第16圈。",
            "问题3": "最小螺距在完整候选域 p∈(0,p_max] 上由全链SAT独立判定。",
        },
    }
    path = tmp_path / "results_eda.json"
    path.write_text(json.dumps(contract), encoding="utf-8")
    load_evidence_contract(tmp_path, "eda", source_text=source)


def test_eda_contract_rejects_run85_q1_pitch_as_structured_q3_upper_bound(
    tmp_path: Path,
):
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题 1 舞龙队沿螺距为 55 cm 的等距螺线盘入，初始位于第16圈。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    contract = {
        "status": "success",
        "subtask": "eda",
        "data_scope": "official statement",
        "assumptions": [],
        "model": "serial rigid members",
        "numeric_results": {"node_count": 224, "distances_m": [2.86, 1.65]},
        "validation": {"passed": True, "checks": ["topology"]},
        "conclusion_bounds": {
            "question3_pitch_search_range_m": {
                "lower": 0.0,
                "upper": 0.55,
                "note": "搜索更小螺距",
            }
        },
    }
    (tmp_path / "results_eda.json").write_text(
        json.dumps(contract), encoding="utf-8"
    )
    with pytest.raises(QualityGateError, match="结构化搜索边界|问题一"):
        load_evidence_contract(tmp_path, "eda", source_text=source)


def test_eda_contract_rejects_run85_scalar_only_physical_feasibility_claim(
    tmp_path: Path,
):
    contract = {
        "status": "success",
        "subtask": "eda",
        "data_scope": "official statement",
        "assumptions": [],
        "model": "serial rigid members",
        "numeric_results": {"node_count": 224, "distances_m": [2.86, 1.65]},
        "validation": {
            "passed": True,
            "checks": ["板宽30cm远小于节点距，物理可行"],
        },
        "conclusion_bounds": "source facts only",
    }
    (tmp_path / "results_eda.json").write_text(
        json.dumps(contract), encoding="utf-8"
    )
    with pytest.raises(QualityGateError, match="不能仅凭板宽"):
        load_evidence_contract(tmp_path, "eda", source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_invented_outer_initial_state_in_q3():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "设初始龙头位于某外圈，r0由图5唯一确定为最外圈；"
                "用有向矩形SAT检查到达边界前的碰撞。"
            ),
            "sensitivity_analysis": "检查边界。",
        }
    )
    with pytest.raises(QualityGateError, match="未给固定起点"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_allows_q3_outer_domain_with_stability_proof():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "从边界向外连续搜索，用有向矩形SAT检查全链碰撞；"
                "逐步扩大搜索上界到充分大外圈，验证临界值稳定且外侧为安全域。"
            ),
            "sensitivity_analysis": "检查搜索上界。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_allows_explicit_denial_of_fixed_q3_initial_with_r_upper_stability():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "原文未给起点，因此不设定任何固定初始位置。"
                "从4.5m边界向外覆盖连续构型，并对每个候选螺距全程用有向矩形SAT检查碰撞；"
                "扩大计算截断半径r_upper为50/100/200m复算，临界值变化<1%。"
            ),
            "sensitivity_analysis": "检查截断半径与网格收敛。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_minimum_head_radius_claimed_to_fit_chain_in_q3():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "不设固定初始位置；对每个p计算容纳整链所需的最小龙头半径r_min(p)，"
                "从r_min盘入到4.5m边界并用SAT检查碰撞；"
                "扩大截断半径r_upper后临界值变化<1%。"
            ),
            "sensitivity_analysis": "检查截断半径。",
        }
    )
    with pytest.raises(QualityGateError, match="不存在“容纳整链所需的最小龙头半径”"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_q1_figure_as_q3_initial_source():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 初始位置见图4。"
        "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 调头空间见图5，请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224个节点，内部距离2.86m和1.65m；"
                "source_figure_required: 代码阶段分别核验图4与图5。"
            ),
            "ques3": (
                "Q3不继承Q1圈数，但初态由代码读取图4/图5确定；"
                "用有向矩形SAT检查全链到达边界前碰撞。"
            ),
            "sensitivity_analysis": "检查搜索上界。",
        }
    )
    with pytest.raises(QualityGateError, match="图4确定问题三"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_unreferenced_q1_turn_count_in_q3():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 螺距为55cm，初始时位于第16圈。"
        "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "起点第16圈为问题1继承值（题面未改）；"
                "完整搜索0<p<=p_max，用SAT检查到达边界前全链碰撞。"
            ),
            "sensitivity_analysis": "检查搜索上界。",
        }
    )
    with pytest.raises(QualityGateError, match="第16圈起点"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_run45_q1_pitch_as_q3_upper_bound():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 螺距为55cm，初始时位于第16圈。"
        "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "用二分法在(0,0.55]内搜索，高界0.55m；"
                "全程复用问题2有限长宽板体SAT碰撞判据。"
            ),
            "sensitivity_analysis": "检查数值网格。",
        }
    )
    with pytest.raises(QualityGateError, match="优化域边界"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_run45_q1_radius_as_q3_initial_state():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 螺距为55cm，初始时位于第16圈。"
        "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "龙头从8.8m递减到调头边界；候选域为0<p<=p_max，"
                "全程复用问题2有限长宽板体SAT碰撞判据。"
            ),
            "sensitivity_analysis": "扩大计算上界并检查稳定性。",
        }
    )
    with pytest.raises(QualityGateError, match="初始半径"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_run46_q1_turn_rewritten_as_q3_radius_formula():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 螺距为55cm，初始时位于第16圈。"
        "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 问题三不继承问题一螺距。",
            "ques3": (
                "候选域0<p<=p_max，但令初始半径r_init=16p>4.5m；"
                "全程复用问题2有限长宽板体SAT碰撞判据。"
            ),
            "sensitivity_analysis": "扩大计算上界并检查稳定性。",
        }
    )
    with pytest.raises(QualityGateError, match="第16圈起点"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_wrong_selected_node_count():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 请在论文中给出0s时龙头前把手、龙头后面第"
        "1、51、101、151、201节龙身前把手和龙尾后把手的位置速度。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques1": "论文每个时刻给出共13个关键节点位置和速度。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="共 7 个对象"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_off_by_one_selected_body_handle_index():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 请在论文中给出龙头前把手、龙头后面第"
        "1、51、101、151、201节龙身前把手和龙尾后把手的位置速度。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 节点1是龙头前把手，"
                "节点2是龙头后把手/第1节龙身前把手，共224节点。"
            ),
            "ques1": "论文列节点1、节点3（第1节龙身前把手）等共7个指定把手。",
            "sensitivity_analysis": "检查索引。",
        }
    )
    with pytest.raises(QualityGateError, match="节点 2，不是节点 3"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_allows_full_attachment_nodes_and_seven_paper_objects():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 请在论文中给出0s时龙头前把手、龙头后面第"
        "1、51、101、151、201节龙身前把手和龙尾后把手的位置速度。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques1": (
                "每秒224个节点的位置和速度写入result1.xlsx；"
                "论文给出指定时刻的7个指定把手位置和速度。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=source)


@pytest.mark.parametrize(
    ("question_number", "file_name"),
    [(1, "result1.xlsx"), (2, "result2.xlsx"), (4, "result4.xlsx")],
)
def test_modeler_rejects_reducing_full_team_attachment_to_seven_objects(
    question_number: int, file_name: str
):
    source = (
        SERIAL_MEMBER_SOURCE
        + f"问题 {question_number} 请给出舞龙队的位置和速度，"
        f"将结果保存到{file_name}；论文另列7个代表对象。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223个构件、224节点，"
                "内部弦长2.86m和1.65m。"
            ),
            f"ques{question_number}": f"{file_name}给出7个对象的位置和速度。",
            "sensitivity_analysis": "检查数值容差。",
        }
    )
    with pytest.raises(QualityGateError, match="所有把手|缩减"):
        validate_modeler_result(
            result, {f"ques{question_number}"}, source_text=source
        )


def test_modeler_rejects_run86_missing_required_result4_attachment():
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题 4 从-100s到100s每秒给出整个舞龙队的位置和速度，"
        "将结果存放到文件result4.xlsx中。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques4": (
                "求双圆弧基线与优化长度，报告是否可缩短，并给出R1、R2、"
                "切点和新弧长。"
            ),
            "sensitivity_analysis": "检查相切残差。",
        }
    )
    with pytest.raises(QualityGateError, match="遗漏题面指定附件 result4.xlsx"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_run87_reversed_flattened_result1_description():
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题 1 从0s到300s每秒给出整个舞龙队的位置和速度，"
        "将结果保存到result1.xlsx。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques1": (
                "result1.xlsx列含时间、224节点x,y,vx,vy，共301行，保留6位小数。"
            ),
            "sensitivity_analysis": "检查数值容差。",
        }
    )
    with pytest.raises(QualityGateError, match="官方结果模板|改写行列方向"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_run87_semicolon_flattened_result1_description():
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题 1 从0s到300s每秒给出整个舞龙队的位置和速度，"
        "将结果保存到result1.xlsx。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques1": (
                "result1.xlsx列含时间、224节点x,y,vx,vy；共301行；保留6位小数。"
            ),
            "sensitivity_analysis": "检查数值容差。",
        }
    )
    with pytest.raises(QualityGateError, match="官方结果模板|改写行列方向"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_allows_full_team_attachment_and_seven_paper_objects_same_question():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 请给出舞龙队的位置和速度，"
        "将结果保存到result1.xlsx；论文另列7个代表对象。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223个构件、224节点，"
                "内部弦长2.86m和1.65m。"
            ),
            "ques1": (
                "result1.xlsx填写全部224个节点的位置和速度；"
                "论文另给出7个代表对象。"
            ),
            "sensitivity_analysis": "检查数值容差。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_missing_initial_second_in_inclusive_output_window():
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题 1 从初始时刻到 300 s 为止，请给出每秒整个舞龙队的位置和速度。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques1": "每秒1～300s共300组，每组224节点，写入result1.xlsx。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="301 组|遗漏 0 s"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_accepts_inclusive_initial_second_output_window():
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题 1 从初始时刻到 300 s 为止，请给出每秒整个舞龙队的位置和速度。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques1": "从0s到300s（含端点）共301组，每组224节点，写入result1.xlsx。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_allows_six_paper_times_beside_full_second_grid():
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题 1 从初始时刻到300s为止，每秒输出位置和速度；"
        "论文另给出0、60、120、180、240、300s的结果。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65m。",
            "ques1": (
                "附件从0s到300s含端点共301个时刻；"
                "论文表给出0、60、120、180、240、300s共6个时刻。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_wrong_count_for_listed_question_times():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 4 从-100s到100s每秒输出；同时在论文中给出"
        "-100s、-50s、0s、50s、100s时的指定把手位置和速度。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques4": "附件共201时刻；论文给出指定-100、-50、0、50、100六个时刻表。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="列举了 5 个时刻|而非 6"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_allows_listed_question_times_without_declared_count():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 4 从-100s到100s每秒输出；同时在论文中给出"
        "-100s、-50s、0s、50s、100s时的指定把手位置和速度。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques4": "附件共201个时刻；论文指定-100、-50、0、50、100时刻表。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_q2_collision_search_capped_by_q1_output_window():
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题 1 从初始时刻到 300 s 为止，请给出每秒的位置和速度。"
        "问题 2 沿问题1设定的螺线盘入，请确定不能再继续盘入的终止时刻。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques1": "从0s到300s共301组。",
            "ques2": "若300s内无碰撞，终止时刻为300s；定义域[0,300s]。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="不能沿用问题一的 300 s"):
        validate_modeler_result(result, {"ques1", "ques2"}, source_text=source)


def test_modeler_rejects_q2_temporary_endpoint_at_q1_output_window():
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题 1 从初始时刻到 300 s 为止，请给出每秒的位置和速度。"
        "问题 2 沿问题1设定的螺线盘入，请确定不能再继续盘入的终止时刻。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques1": "从0s到300s共301组。",
            "ques2": "从t=0扫描；若300s内无碰撞则报告300s为暂定终点并说明。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="300 s 输出截点|继续搜索"):
        validate_modeler_result(result, {"ques1", "ques2"}, source_text=source)


def test_modeler_rejects_run42_q2_collision_search_capped_by_q1_window():
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题 1 从初始时刻到 300 s 为止，请给出每秒的位置和速度。"
        "问题 2 沿问题1设定的螺线盘入，请确定不能再继续盘入的终止时刻。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques1": "从0s到300s共301组。",
            "ques2": (
                "t_coll必须>0且在Q1全程300s内被确定（若300s内无碰撞则报告"
                "'300s内未碰撞'并给出最近间距）。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="不能沿用问题一的 300 s"):
        validate_modeler_result(result, {"ques1", "ques2"}, source_text=source)


def test_modeler_rejects_first_collision_search_without_initial_time_coverage():
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题 1 从初始时刻到300s给出位置。"
        "问题 2 板凳发生碰撞时不能再继续盘入，请确定终止时刻。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65m。",
            "ques1": "从0s到300s共301组。",
            "ques2": "逐帧执行SAT，检测到相交后加密并报告首次碰撞时刻。",
            "sensitivity_analysis": "细化时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="初始时刻 t=0|完整前史"):
        validate_modeler_result(result, {"ques1", "ques2"}, source_text=source)


def test_modeler_accepts_first_collision_search_from_initial_time():
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题 1 从初始时刻到300s给出位置。"
        "问题 2 板凳发生碰撞时不能再继续盘入，请确定终止时刻。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65m。",
            "ques1": "从0s到300s共301组。",
            "ques2": (
                "从题面初始时刻t=0开始逐帧执行SAT，首次相交后加密；"
                "若300s内无碰撞则继续向后扫描并包围真实首次事件。"
            ),
            "sensitivity_analysis": "细化时间步长。",
        }
    )
    validate_modeler_result(result, {"ques1", "ques2"}, source_text=source)


def test_modeler_accepts_numeric_collision_checkpoint_followed_by_until_collision():
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题 1 从初始时刻到300s给出位置。"
        "问题 2 板凳发生碰撞时不能再继续盘入，请确定终止时刻。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65m。",
            "ques1": "从0s到300s共301组。",
            "ques2": (
                "从题面初始时刻t=0开始扫描完整前史；"
                "若搜索至300s未碰撞，则延长扫描直至碰撞。"
            ),
            "sensitivity_analysis": "细化时间步长。",
        }
    )
    validate_modeler_result(result, {"ques1", "ques2"}, source_text=source)


def test_modeler_accepts_q2_explicitly_not_limited_to_q1_window():
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题 1 从初始时刻到300s给出位置。"
        "问题 2 板凳发生碰撞时不能再继续盘入，请确定终止时刻。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65m。",
            "ques1": "从0s到300s共301组。",
            "ques2": (
                "从题面初始时刻t=0开始逐帧执行SAT；定义域不限于300s，"
                "须继续推进直至实际碰撞或证明无碰撞。"
            ),
            "sensitivity_analysis": "细化时间步长。",
        }
    )
    validate_modeler_result(result, {"ques1", "ques2"}, source_text=source)


def test_modeler_rejects_first_collision_search_capped_at_500_seconds():
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题 1 从初始时刻到300s给出位置。"
        "问题 2 板凳发生碰撞时不能再继续盘入，请确定终止时刻。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65m。",
            "ques1": "从0s到300s共301组。",
            "ques2": (
                "从题面初始时刻t=0开始执行SAT；若300s前未碰撞，"
                "延长搜索至500s确认。"
            ),
            "sensitivity_analysis": "细化时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="有限时间上限|继续扩大时间域"):
        validate_modeler_result(result, {"ques1", "ques2"}, source_text=source)


def test_modeler_accepts_numeric_collision_checkpoint_with_open_continuation():
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题 1 从初始时刻到300s给出位置。"
        "问题 2 板凳发生碰撞时不能再继续盘入，请确定终止时刻。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65m。",
            "ques1": "从0s到300s共301组。",
            "ques2": (
                "从题面初始时刻t=0开始执行SAT；若300s前未碰撞，"
                "延长搜索至500s作为检查点；题面未限时，须继续搜索并包围真实碰撞事件。"
            ),
            "sensitivity_analysis": "细化时间步长。",
        }
    )
    validate_modeler_result(result, {"ques1", "ques2"}, source_text=source)


def test_modeler_accepts_q3_reuse_of_prior_finite_body_collision_kernel():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 板凳之间不发生碰撞，即舞龙队不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques3": (
                "搜索域0<p<=p_max，从充分外侧正行程开始；"
                "调用问题1/2递推与碰撞检测，只有全链无碰撞到达边界才可行，"
                "并扩大搜索上界验证稳定。"
            ),
            "sensitivity_analysis": "检查搜索上界。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


@pytest.mark.parametrize(
    "wrong_q4, expected",
    [
        ("R1=2ρ，R2=ρ；与等半径基线对比弧长。C1连续且曲率跳变。", "等半径基线"),
        (
            "R1=2ρ，R2=ρ；L_s=ρθ_span1+0.5ρθ_span2。C1连续且曲率跳变。",
            "长度公式",
        ),
        ("R1=2ρ，R2=ρ；曲率最小处为1/ρ。C1连续且曲率跳变。", "曲率大小写反"),
    ],
)
def test_modeler_rejects_inconsistent_unequal_arc_baseline_or_formula(
    wrong_q4: str, expected: str
):
    source = SERIAL_MEMBER_SOURCE + "问题 4 两段圆弧相切，前一段半径是后一段的2倍。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques4": wrong_q4,
            "sensitivity_analysis": "检查半径。",
        }
    )
    with pytest.raises(QualityGateError, match=expected):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_infeasible_single_arc_as_q4_baseline():
    source = SERIAL_MEMBER_SOURCE + "问题 4 两段圆弧相切，前一段半径是后一段的2倍。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques4": (
                "R1=2R2；位置和切向C1连续，连接点曲率跳变；"
                "与基线单圆弧（取两段均值半径）比较缩短百分比。"
            ),
            "sensitivity_analysis": "检查半径。",
        }
    )
    with pytest.raises(QualityGateError, match="2:1 双圆弧"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_run57_entire_supporting_circle_inside_constraint():
    source = SERIAL_MEMBER_SOURCE + "问题 4 两段圆弧相切，前一段半径是后一段的2倍。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques4": (
                "R1=2R2；位置和切向C1连续，连接点曲率由1/R1跳变到1/R2；"
                "约束圆心到调头圆心距离+半径≤4.5m。"
            ),
            "sensitivity_analysis": "检查弧段采样密度。",
        }
    )
    with pytest.raises(QualityGateError, match="支撑圆"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_accepts_actual_arc_points_inside_turning_circle():
    source = SERIAL_MEMBER_SOURCE + "问题 4 两段圆弧相切，前一段半径是后一段的2倍。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques4": (
                "R1=2R2；位置和切向C1连续，连接点曲率由1/R1跳变到1/R2；"
                "仅对实际弧段求最大极径，并加密采样确认所有弧上点极径≤4.5m。"
            ),
            "sensitivity_analysis": "检查弧段采样密度。",
        }
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_run42_q4_variable_radius_ratio():
    source = SERIAL_MEMBER_SOURCE + "问题 4 两段圆弧相切，前一段半径是后一段的2倍。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: candidate",
            "ques4": (
                "前段半径R1=2R，后段半径R2=R（初始基线，优化中可调比值）；"
                "决策变量为R1/R2比值及角度；两弧连接点C1连续但曲率跳变。"
            ),
            "sensitivity_analysis": "检查半径。",
        }
    )
    with pytest.raises(QualityGateError, match="固定为 2:1"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_accepts_run43_explicit_curvature_discontinuity_denial():
    source = SERIAL_MEMBER_SOURCE + "问题 4 两段圆弧相切，前一段半径是后一段的2倍。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques4": (
                "连续性声明：两段圆弧相切仅保证位置与切向C1连续，"
                "连接点两侧曲率分别为1/R1和1/R2，存在跳变，"
                "不得声明C2连续或曲率连续。R1=2R，R2=R。"
            ),
            "sensitivity_analysis": "扰动圆弧尺度R并检查切点残差。",
        }
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_allows_sampled_position_group_count_beside_full_time_count():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 从初始时刻到300s，每秒给出全部把手的位置和速度。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques1": (
                "从0s到300s含端点共301个时刻，逐秒输出全部224个节点；"
                "论文表格另给出7个指定对象在6个指定时刻的结果，共42组位置和42组速度。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_accepts_run56_fixed_ratio_with_common_scale_perturbation():
    source = SERIAL_MEMBER_SOURCE + "问题 4 两段圆弧相切，前一段半径是后一段的2倍。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques4": ("R1=2R2；两段圆弧位置与切向C1连续，连接点曲率发生跳变。"),
            "sensitivity_analysis": (
                "半径比例固定为R1:R2=2:1，不得扰动该比值；"
                "R2在±10%范围内变化，同时R1自动按2R2变化。"
            ),
        }
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_run43_perturbation_of_fixed_q4_radius_ratio():
    source = SERIAL_MEMBER_SOURCE + "问题 4 两段圆弧相切，前一段半径是后一段的2倍。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques4": ("R1=2R，R2=R；两段圆弧位置与切向C1连续，连接点曲率发生跳变。"),
            "sensitivity_analysis": (
                "圆弧半径比（严格2:1，检查±2%鲁棒性）影响调头弧长与速度分布。"
            ),
        }
    )
    with pytest.raises(QualityGateError, match="固定约束"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_run57_multistart_global_claim_in_sensitivity():
    source = SERIAL_MEMBER_SOURCE + "问题 4 两段圆弧相切，前一段半径是后一段的2倍。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques4": "R1=2R2；位置和切向C1连续，连接点曲率由1/R1跳变到1/R2。",
            "sensitivity_analysis": (
                "问题4圆弧优化多起点，确认最优弧长和参数收敛，证明全局最优。"
            ),
        }
    )
    with pytest.raises(QualityGateError, match="不能单独证明全局最优"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


@pytest.mark.parametrize(
    ("sensitivity", "message"),
    [
        ("板宽0.3m扰动±10%，比较首次碰撞时刻。", "题面固定板宽"),
        ("速度限制2m/s扰动±10%，比较最大龙头速度。", "题面固定速度上限"),
    ],
)
def test_modeler_rejects_run45_fixed_source_constraint_sensitivity(
    sensitivity: str, message: str
):
    source = (
        SERIAL_MEMBER_SOURCE
        + "所有板凳的板宽均为30cm。问题 5 要求各把手速度均不超过2m/s。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 板宽0.3m、限速2m/s均为题面固定约束。",
            "ques5": "遍历全部节点与时间后计算最大许可速度。",
            "sensitivity_analysis": sensitivity,
        }
    )
    with pytest.raises(QualityGateError, match=message):
        validate_modeler_result(result, {"ques5"}, source_text=source)


def test_modeler_rejects_run67_fixed_width_perturbation_inside_question_plan():
    source = SERIAL_MEMBER_SOURCE + "所有板凳的板宽均为30cm。问题 2 确定首次碰撞。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 板宽0.3m为题面固定约束。",
            "ques2": "用有向矩形SAT定位首碰；板宽扰动±1mm，复算终止时刻。",
            "sensitivity_analysis": "仅缩小时间步长和SAT数值容差。",
        }
    )
    with pytest.raises(QualityGateError, match="题面固定构件尺寸"):
        validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_accepts_q3_circle_scope_denial_before_collision_requirement():
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题 2 请确定板凳不能再继续盘入的时刻。"
        "问题 3 请确定使龙头到达调头空间边界的最小螺距。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223节构件、224个节点，"
                "内部节点距2.86m和1.65m。"
            ),
            "ques2": "用有向矩形SAT定位首次碰撞。",
            "ques3": (
                "从充分外侧连续回放至4.5m边界，全程用SAT检查全链无碰撞。"
                "不要求全链在圆内; 到达前必须整链无碰撞，且存在正行程。"
            ),
            "sensitivity_analysis": "扩大空间截断并缩小时间步长。",
        }
    )
    validate_modeler_result(result, {"ques2", "ques3"}, source_text=source)


@pytest.mark.parametrize(
    ("sensitivity", "message"),
    [
        ("问题1/2螺距±10%，比较终止时刻。", "问题一/二螺距"),
        ("问题3初始半径±2%，比较最小螺距。", "问题三未授权初态"),
        ("调头空间直径9m±2%，比较弧长。", "调头空间直径"),
    ],
)
def test_modeler_rejects_run46_other_fixed_source_sensitivity(
    sensitivity: str, message: str
):
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 螺距为55cm，初始时位于第16圈。"
        "问题 3 调头空间为直径9m圆形区域，请确定最小螺距。"
        "问题 4 在该调头空间内设计路径。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 固定来源参数已冻结。",
            "ques1": "按题面螺距求解。",
            "ques3": "覆盖0<p<=p_max并扩大上界检查稳定性。",
            "ques4": "在固定调头空间中优化路径。",
            "sensitivity_analysis": sensitivity,
        }
    )
    with pytest.raises(QualityGateError, match=message):
        validate_modeler_result(result, {"ques1", "ques3", "ques4"}, source_text=source)


def test_modeler_allows_run49_fixed_width_denial_before_other_checks():
    source = SERIAL_MEMBER_SOURCE + "所有板凳的板宽均为30cm。"
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223个构件、224节点，"
                "内部弦长2.86m和1.65m；板宽0.3m为题面固定约束。"
            ),
            "ques1": "按题面参数建立模型。",
            "sensitivity_analysis": (
                "仅检查来源不确定量与数值容差，不动题面固定板宽/板长/孔距；"
                "时间步长从0.01s缩小到0.005s。"
            ),
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_accepts_postposed_fixed_geometry_cannot_be_perturbed():
    source = SERIAL_MEMBER_SOURCE + "所有板凳的板宽均为30cm。"
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223节构件、224个节点，"
                "内部节点距2.86m和1.65m。"
            ),
            "ques1": "按题面固定几何建立刚性链模型。",
            "sensitivity_analysis": (
                "题面固定几何量（板长、板宽、孔径、端部偏移）不可扰动；"
                "仅缩小时间步长和求解容差。"
            ),
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_run49_reversed_fixed_pitch_sensitivity_wording():
    source = SERIAL_MEMBER_SOURCE + "问题 1 螺距为55cm，初始时位于第16圈。"
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223个构件、224节点，"
                "内部弦长2.86m和1.65m；问题1螺距0.55m为题面固定参数。"
            ),
            "ques1": "按题面固定螺距求解。",
            "sensitivity_analysis": "螺距p（问题1/2，±10%范围0.495-0.605m）。",
        }
    )
    with pytest.raises(QualityGateError, match="问题一/二螺距"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_run51_impossible_spiral_turn_count_for_timed_motion():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 螺距为55cm，龙头前把手速度始终保持1m/s，"
        "初始时位于第16圈，请给出到300s为止的位置和速度。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86m和1.65m。",
            "ques1": (
                "采用r=pθ/(2π)，从θ(0)=32π向内积分；"
                "300s行进300m弧长，对应约46圈，半径仍大于0。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="行进圈数.*弧长"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_flattening_official_position_speed_template():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 请给出0s到300s舞龙队的位置和速度，"
        "将结果保存到result1.xlsx；论文另列7个代表对象。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223个构件、224节点，"
                "内部弦长2.86m和1.65m。"
            ),
            "ques1": (
                "result1.xlsx含301行×(224×2坐标+224速度)列，"
                "同时填写0s到300s所有节点。"
            ),
            "sensitivity_analysis": "检查数值容差。",
        }
    )
    with pytest.raises(QualityGateError, match="官方结果模板.*逐时刻单表"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_reversing_official_template_rows_and_time_columns():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 请给出0s到300s舞龙队的位置和速度，"
        "将结果保存到result1.xlsx。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点。",
            "ques1": "result1.xlsx每秒0~300s共301行×224节点位置速度。",
            "sensitivity_analysis": "检查数值容差。",
        }
    )
    with pytest.raises(QualityGateError, match="官方结果模板.*行列方向"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


@pytest.mark.parametrize(
    ("sensitivity", "message"),
    [
        (
            "允许扰动测量容差：板长±1cm、孔位27.5cm±0.2cm、板宽30cm±0.1cm。",
            "构件尺寸/孔位",
        ),
        (
            "龙头速度不确定性：问题1速度1m/s±2%，用于鲁棒性分析。",
            "龙头速度",
        ),
    ],
)
def test_modeler_rejects_run51_fixed_geometry_or_speed_sensitivity(
    sensitivity: str, message: str
):
    source = (
        SERIAL_MEMBER_SOURCE + "所有板凳的板宽均为30cm。问题 1 螺距为55cm，"
        "龙头前把手速度始终保持1m/s，初始时位于第16圈。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 题面尺寸和速度均冻结。",
            "ques1": "按题面固定参数求解。",
            "sensitivity_analysis": sensitivity,
        }
    )
    with pytest.raises(QualityGateError, match=message):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_accepts_scientific_notation_tolerance_for_fixed_speed_problem():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 龙头前把手速度始终保持1m/s。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223个构件、224节点，"
                "内部弦长2.86m和1.65m；题面速度已冻结。"
            ),
            "ques1": "按题面固定速度求解。",
            "sensitivity_analysis": (
                "问题1速度递推的求解容差取1e-10至1e-8，"
                "要求位置残差小于1e-6m。"
            ),
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_accepts_postposed_denial_of_fixed_speed_perturbation():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 龙头前把手速度始终保持1m/s。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223个构件、224节点，"
                "内部弦长2.86m和1.65m；题面速度已冻结。"
            ),
            "ques1": "按题面固定速度求解。",
            "sensitivity_analysis": (
                "问题1-2中龙头速度1m/s为题面固定，不扰动；"
                "仅细化时间步长。"
            ),
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_allows_fixed_speed_as_an_observed_output_metric():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 龙头前把手速度始终保持1m/s。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223个构件、224节点，"
                "内部弦长2.86m和1.65m；题面速度已冻结。"
            ),
            "ques1": "按题面固定速度求解。",
            "sensitivity_analysis": (
                "时间步长扰动0.005/0.02s，评估问题1位置速度、"
                "问题2碰撞时刻和问题4-5极值的变化；"
                "禁止扰动龙头速度1m/s。"
            ),
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_accepts_explicit_denial_of_fixed_geometry_perturbation():
    source = SERIAL_MEMBER_SOURCE + "所有板凳的板宽均为30cm。"
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223个构件、224节点，"
                "内部弦长2.86m和1.65m；题面尺寸均冻结。"
            ),
            "ques1": "按题面参数建立模型。",
            "sensitivity_analysis": (
                "保持板长、孔位和板宽固定，不作扰动；仅检查时间步长与求解容差。"
            ),
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_accepts_grouped_postposed_denial_for_fixed_parameters():
    source = (
        SERIAL_MEMBER_SOURCE + "所有板凳板宽30cm。问题1龙头速度始终保持1m/s。"
        "问题3调头空间直径为9m。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223个构件、224节点，"
                "内部弦长2.86m和1.65m。"
            ),
            "ques1": "按题面固定参数求解。",
            "ques3": "覆盖完整正参数域并检查全链可行性。",
            "sensitivity_analysis": (
                "所有题面固定参数（板长3.41/2.20m、孔距、板宽、"
                "龙头速度1m/s、调头空间直径9m）均作为确定性常量，"
                "不做任何扰动或不确定性假设。"
            ),
        }
    )
    validate_modeler_result(result, {"ques1", "ques3"}, source_text=source)


def test_modeler_rejects_positive_perturbation_despite_grouped_denial():
    source = SERIAL_MEMBER_SOURCE + "所有板凳板宽30cm。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 题面尺寸均冻结。",
            "ques1": "按题面固定参数求解。",
            "sensitivity_analysis": (
                "所有题面固定参数不做扰动，但板宽±1cm用于鲁棒性分析。"
            ),
        }
    )
    with pytest.raises(QualityGateError, match="构件尺寸"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_run51_opening_fixed_q4_radius_ratio():
    source = SERIAL_MEMBER_SOURCE + "问题 4 两段圆弧相切，前一段半径是后一段的2倍。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques4": (
                "基线固定R1:R2=2:1，两段圆弧位置与切向C1连续且曲率跳变；"
                "再放开比例寻找更短路径，并把新比例作为可行答案。"
            ),
            "sensitivity_analysis": "保持2:1不扰动，检查相切残差。",
        }
    )
    with pytest.raises(QualityGateError, match="固定为 2:1"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_run51_bare_wrong_node_count_in_conclusion():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223个构件、224节点，"
                "内部弦长2.86m和1.65m；结论为2124节点拓扑链沿螺线运动。"
            ),
            "ques1": "按刚性链模型求解。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="节点总数"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_run46_node_distance_symbol_as_body_length():
    source = SERIAL_MEMBER_SOURCE + "所有板凳板宽30cm。问题 2 请确定板凳首次碰撞时刻。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: L_1=2.86m，其余L_i=1.65m为节点距。",
            "ques2": (
                "每节板凳建模为长L_i、宽0.3m的有向矩形；"
                "排除相邻连接对，对非相邻矩形用SAT精判。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="复用了刚性节点距符号"):
        validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_rejects_run46_adjacent_boards_as_collision_definition():
    source = SERIAL_MEMBER_SOURCE + "所有板凳板宽30cm。问题 2 请确定板凳首次碰撞时刻。"
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 碰撞定义为相邻板凳实体轮廓互相干涉，"
                "板体采用真实外形。"
            ),
            "ques2": "对有向矩形使用SAT。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="共享把手装配关系"):
        validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_rejects_unrequested_result_attachment():
    source = "问题 1 请计算结果并写入result1.xlsx。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 核对题面输出文件。",
            "ques1": "写入result1.xlsx，并按统一处理另存result3.xlsx。",
            "sensitivity_analysis": "检查数值精度。",
        }
    )
    with pytest.raises(QualityGateError, match="result3.xlsx"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


@pytest.mark.parametrize(
    "denial",
    [
        "问题5不生成result5.xlsx（题面未要求）。",
        "问题5无需另行输出 result5.xlsx。",
        "题面未要求 result5.xlsx。",
    ],
)
def test_modeler_accepts_negated_unrequested_result_attachment(denial: str):
    source = "问题 1 请计算结果并写入result1.xlsx。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 核对题面输出文件。",
            "ques1": "写入result1.xlsx。" + denial,
            "sensitivity_analysis": "检查数值精度。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_run67_arbitrary_feasible_q4_baseline():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 两段圆弧保持相切，前一段半径是后一段的2倍。"
        "能否调整圆弧使调头曲线变短？"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223节构件、224个节点，"
                "两类内部节点距为2.86m和1.65m；题面约束已冻结。"
            ),
            "ques4": (
                "两弧位置与切向C1连续，连接点曲率发生跳变。"
                "建立基线：取一满足相切的R2基准，计算基线长度s0，"
                "再与优化长度s*比较。"
            ),
            "sensitivity_analysis": "缩小求解容差。",
        }
    )
    with pytest.raises(QualityGateError, match="任取一个可行半径"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_accepts_run44_long_curvature_jump_statement():
    source = SERIAL_MEMBER_SOURCE + "问题 4 两段圆弧相切，前一段半径是后一段的2倍。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques4": (
                "两弧在P2处位置与切向C1连续；P2两侧曲率分别为1/(2ρ)和1/ρ，"
                "两者不等并跳变，不得称为C2连续。R1=2ρ，R2=ρ。"
            ),
            "sensitivity_analysis": "保持2:1不扰动，检查公共尺度ρ与切向容差。",
        }
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


@pytest.mark.parametrize(
    "wrong_geometry",
    [
        "约束为P1P2=2ρ，P2P3=ρ（弦长）。",
        "P1、P2、P3共线，圆心O1、O2与P2共线。",
    ],
)
def test_modeler_rejects_run44_arc_endpoints_confused_with_radii(
    wrong_geometry: str,
):
    source = SERIAL_MEMBER_SOURCE + "问题 4 两段圆弧相切，前一段半径是后一段的2倍。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65 m",
            "ques4": (
                "前弧起点P1与盘入螺线相切，后弧终点P3与盘出螺线相切，"
                "两弧在P2相切；位置与切向C1连续且曲率跳变。"
                + wrong_geometry
                + "R1=2ρ，R2=ρ。"
            ),
            "sensitivity_analysis": "保持2:1不扰动，检查公共尺度ρ。",
        }
    )
    with pytest.raises(QualityGateError, match="弧端点、连接点和圆心混淆"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_multistart_as_global_optimality_proof():
    source = SERIAL_MEMBER_SOURCE + "问题 4 两段圆弧相切，前一段半径是后一段的2倍。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: candidate",
            "ques4": (
                "R1=2R2；位置和切向C1连续，连接点曲率跳变；"
                "采用20个随机初值，多起点确认全局最小。"
            ),
            "sensitivity_analysis": "检查初值。",
        }
    )
    with pytest.raises(QualityGateError, match="不能单独证明全局最优"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_accepts_run55_denial_of_whole_chain_circle_constraint():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 盘入螺线螺距为1.7m，两段圆弧相切，前一段半径是后一段的2倍。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques4": (
                "R1=2R2；位置和切向C1连续，连接点曲率跳变；"
                "调头轨迹位于圆内，不需全链同时位于圆内。"
            ),
            "sensitivity_analysis": "保持2:1不扰动，检查相切残差。",
        }
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_run55_naked_numeric_r2_baseline():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 盘入螺线螺距为1.7m，两段圆弧相切，前一段半径是后一段的2倍。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques4": (
                "R1=2R2；位置和切向C1连续，连接点曲率跳变；"
                "与基准配置R2=1.5m比较路径长度。"
            ),
            "sensitivity_analysis": "保持2:1不扰动，检查相切残差。",
        }
    )
    with pytest.raises(QualityGateError, match="裸写一个 R2 数值"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_run55_finite_grid_as_global_certificate():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 盘入螺线螺距为1.7m，两段圆弧相切，前一段半径是后一段的2倍。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques4": (
                "R1=2R2；位置和切向C1连续，连接点曲率跳变；"
                "在可行区间以0.01m网格扫描后局部精修，小步长穷举确认全局最小。"
            ),
            "sensitivity_analysis": "保持2:1不扰动，检查相切残差。",
        }
    )
    with pytest.raises(QualityGateError, match="有限网格.*不能单独证明"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_presolved_speed_bottleneck_claim():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 两段圆弧相切，前一段半径是后一段的2倍。"
        "问题 5 请确定龙头最大速度，使各把手速度均不超过2m/s。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: candidate",
            "ques4": "R1=2R2；位置和切向C1连续，连接点曲率跳变。",
            "ques5": (
                "遍历全部224节点与完整时间窗计算k_max；"
                "节点224通常速度更高，图中标注节点224尾部峰值；实际k_max>1。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="预设了瓶颈节点"):
        validate_modeler_result(result, {"ques4", "ques5"}, source_text=source)


def test_modeler_rejects_reversed_arc_curvature_ranking_in_followup_question():
    source = SERIAL_MEMBER_SOURCE + "问题 4 两段圆弧相切，前一段半径是后一段的2倍。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: candidate",
            "ques4": "R1=2ρ，R2=ρ；C1连续且曲率跳变。",
            "ques5": "检查S形内圈曲率最小处（曲率1/ρ处）是否主导速度上界。",
            "sensitivity_analysis": "检查半径。",
        }
    )
    with pytest.raises(QualityGateError, match="曲率大小写反"):
        validate_modeler_result(result, {"ques4", "ques5"}, source_text=source)


@pytest.mark.parametrize("reversed_jump", ["1/R2→1/R1", "1/R2→1/(2R2)", "1/ρ→1/(2ρ)"])
def test_modeler_rejects_reversed_arc_curvature_jump_order(reversed_jump: str):
    source = SERIAL_MEMBER_SOURCE + "问题 4 两段圆弧相切，前一段半径是后一段的2倍。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: candidate",
            "ques4": (
                "前弧R1=2R2，后弧R2；位置与切向C1连续，"
                f"验证连接点曲率跳变{reversed_jump}。"
            ),
            "sensitivity_analysis": "检查半径。",
        }
    )
    with pytest.raises(QualityGateError, match="曲率跳变顺序写反"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_hole_distance_as_collision_length_without_numbers():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques2": (
                "有向矩形使用题面板长和板宽，以SAT检查非相邻板体碰撞；"
                "相邻连接对允许装配接触。边界：孔距决定矩形长度。"
            ),
            "sensitivity_analysis": "检查板宽。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "板宽为30cm，请检查板凳碰撞。"
    with pytest.raises(QualityGateError, match="孔中心距只定义运动学"):
        validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_rejects_two_handle_endpoints_as_rectangle_center():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques2": (
                "矩形长3.41m/2.2m、宽0.3m，中心为两把手终点；"
                "排除相邻连接对后用SAT精确判定非相邻板体碰撞。"
            ),
            "sensitivity_analysis": "检查板宽。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "板宽为30cm，请检查板凳碰撞。"
    with pytest.raises(QualityGateError, match="两孔/把手坐标的中点"):
        validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_rejects_speed_propagation_using_only_local_arc_metrics():
    source = SERIAL_MEMBER_SOURCE + "问题 1 从初始时刻到300s为止，每秒给出位置速度。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques1": (
                "从0s到300s含端点共301组；"
                "速度传播取v_i=v_{i-1}×(ds_{i-1}/dθ_{i-1})/(ds_i/dθ_i)。"
            ),
            "sensitivity_analysis": "检查差分。",
        }
    )
    with pytest.raises(QualityGateError, match="方向投影"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_reversed_chord_tangent_speed_projection_ratio():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques1": (
                "由刚性弦长约束求导，速度递推"
                "v_{i+1}=v_i*(弦向量·切向_{i+1})/(弦向量·切向_i)。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="分子分母写反"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_accepts_correct_chord_tangent_speed_projection_ratio():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques1": (
                "由刚性弦长约束求导，速度递推"
                "v_{i+1}=v_i*(弦向量·切向_i)/(弦向量·切向_{i+1})。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_uncomputed_monotone_chain_speed_claim():
    source = SERIAL_MEMBER_SOURCE + "问题 1 请给出整个舞龙队的位置和速度。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques1": (
                "对二维弦长约束求导传播各节点速度；"
                "验证策略：速度沿链从龙头向龙尾单调衰减。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="不能在计算前预设速度沿链单调"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_run49_cosine_only_speed_propagation():
    source = SERIAL_MEMBER_SOURCE + "问题 1 请给出整个舞龙队的位置和速度。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques1": (
                "由刚性约束对时间求导，速度传播简化为"
                "v_{i+1}=v_i*cos(Δα)，其中Δα为两节点切线角之差。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="弦向量与两端切向的投影比"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_run49_monotone_nonincreasing_speed_claim():
    source = SERIAL_MEMBER_SOURCE + "问题 1 请给出整个舞龙队的位置和速度。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques1": "先传播各节点速度，再验证速度沿链单调不增。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="不能在计算前预设速度沿链单调"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_invented_source_r2_baseline():
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题 4 调头路径由两段圆弧相切连接，前一段半径是后一段2倍；"
        "能否调整圆弧使调头曲线变短？"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques4": (
                "保持R1=2R2和三处相切，优化连接点；"
                "双圆弧连接点位置与切向C1连续，曲率1/R1到1/R2发生跳变；"
                "与题面给出的R2基准比较弧长。"
            ),
            "sensitivity_analysis": "检查求解容差。",
        }
    )
    with pytest.raises(QualityGateError, match="题面没有给出数值 R2 基准"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_run50_entire_chain_inside_turning_circle():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 4 舞龙队在半径4.5m调头空间内完成调头，"
        "路径由两段圆弧相切连接，前一段半径是后一段2倍。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques4": (
                "双圆弧位置与切向C1连续，连接点曲率从1/R1跳变到1/R2；"
                "调头圆内路径须确保整条链的板体在运动全程不越出边界。"
            ),
            "sensitivity_analysis": "检查相切残差。",
        }
    )
    with pytest.raises(QualityGateError, match="调头轨迹|整条链"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_run50_arc_center_and_two_distinct_points_collinear():
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题 4 调头路径由两段圆弧相切连接，前一段半径是后一段2倍。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques4": (
                "两弧位置与切向C1连续，连接点曲率从1/R1跳变到1/R2；"
                "前弧圆心O1、入口切点P1和公共切点P三点共线，"
                "后弧圆心O2、出口切点P2和公共切点P三点共线。"
            ),
            "sensitivity_analysis": "检查相切残差。",
        }
    )
    with pytest.raises(QualityGateError, match="同一圆弧|两个不同弧上点"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_run50_finite_total_length_for_infinite_spiral_path():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 4 路径由盘入螺线、两段圆弧和中心对称盘出螺线组成。"
        "问题 5 沿问题4设定路径确定龙头最大速度。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques5": (
                "按路径总弧长除以1m/s确定有限的开始和结束时刻，"
                "在该全程扫描全部节点速度倍率。"
            ),
            "sensitivity_analysis": "检查时间网格。",
        }
    )
    with pytest.raises(QualityGateError, match="螺线向外无限延伸|尾部界"):
        validate_modeler_result(result, {"ques5"}, source_text=source)


def test_modeler_accepts_finite_transition_scan_with_infinite_tail_bound():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 4 路径由盘入螺线、两段圆弧和中心对称盘出螺线组成。"
        "问题 5 沿问题4设定路径确定龙头最大速度。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques5": (
                "螺线两端向外无限延伸；数值扫描覆盖调头段及其两侧高曲率区，"
                "并用切向投影比趋近1的解析尾部界证明扩大截断区间后"
                "外侧速度倍率不超过当前最大值。"
            ),
            "sensitivity_analysis": "扩大时间截断检查上界稳定。",
        }
    )
    validate_modeler_result(result, {"ques5"}, source_text=source)


def test_modeler_rejects_finite_window_as_global_extremum_on_infinite_path():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 4 路径由盘入螺线、两段圆弧和中心对称盘出螺线组成。"
        "问题 5 沿问题4设定路径确定龙头最大速度。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques5": (
                "对每个时刻计算全部节点的速度比，扫描-100~100 s，"
                "取全路径全局最大速度比并据此计算龙头最大速度。"
            ),
            "sensitivity_analysis": "把时间步从1 s细化至0.5 s。",
        }
    )
    with pytest.raises(QualityGateError, match="有限时间窗|无限尾部"):
        validate_modeler_result(result, {"ques5"}, source_text=source)


def test_modeler_accepts_finite_window_with_explicit_infinite_tail_evidence():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 4 路径由盘入螺线、两段圆弧和中心对称盘出螺线组成。"
        "问题 5 沿问题4设定路径确定龙头最大速度。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques5": (
                "先扫描-100~100 s覆盖调头有限区间；螺线两端向外无限延伸，"
                "再用切向投影比趋近1的解析尾部界证明外侧速度倍率"
                "不超过有限区间已得最大值。"
            ),
            "sensitivity_analysis": "扩大截断区间并检查上界稳定。",
        }
    )
    validate_modeler_result(result, {"ques5"}, source_text=source)


def test_modeler_rejects_bare_global_scan_on_unbounded_spiral_path():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 4 路径由盘入螺线、两段圆弧和中心对称盘出螺线组成。"
        "问题 5 沿问题4设定路径确定龙头最大速度。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques5": "回放Q4全程并全路径扫描全部节点，取全局最大速度放大系数。",
            "sensitivity_analysis": "把时间步减半检查收敛。",
        }
    )
    with pytest.raises(QualityGateError, match="全路径极值|无限尾部"):
        validate_modeler_result(result, {"ques5"}, source_text=source)


def test_modeler_allows_denial_of_all_chain_turning_circle_requirement():
    source = SERIAL_MEMBER_SOURCE + "问题 3 请使龙头前把手盘入到调头空间的边界。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "搜索0<p<=p_max，从充分外侧起点到边界全程用SAT检查板体碰撞；"
                "目标仅为龙头前把手到边界，非整条龙全部进入调头圆。"
            ),
            "sensitivity_analysis": "扩大上界检查稳定。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_allows_constructed_outer_domain_with_stability_expansion():
    source = SERIAL_MEMBER_SOURCE + "问题 3 请使龙头前把手盘入到调头空间的边界。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "搜索0<p<=p_max；取初始启动半径r_start=70m作为充分外侧安全域，"
                "从该正行程到边界全程用SAT检查板体碰撞；"
                "扩大p_max至两倍并相应增加r_start，确认临界值稳定。"
            ),
            "sensitivity_analysis": "检查r_start扰动。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_run56_pitch_bound_used_as_spatial_outer_position():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "搜索0<p<=p_max；p螺线从充分外侧（如p_max处）向中心盘入，"
                "直到龙头到达边界，并全程用有向矩形SAT检查实体碰撞。"
            ),
            "sensitivity_analysis": "将p_max扩大两倍检查结果稳定。",
        }
    )
    with pytest.raises(QualityGateError, match="参数括界与空间截断"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_run56_q3_bisection_without_feasibility_monotonicity():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "在0<p<=p_max内用二分搜索；目标函数（最小p）单调，"
                "可行则减小p，不可行则增大p。从充分外侧r_start到边界"
                "全程用有向矩形SAT检查实体碰撞，并扩大r_start验证稳定。"
            ),
            "sensitivity_analysis": "细化碰撞检测步长。",
        }
    )
    with pytest.raises(QualityGateError, match="可行谓词单调"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_accepts_q3_bisection_after_full_domain_monotonicity_check():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "先在完整参数域对数采样，用有向矩形SAT从充分外侧r_start回放到边界；"
                "确认可行性仅有一次不可行到可行跃迁后，才在该括区间二分。"
                "扩大r_start并复核临界值稳定；若扫描发现反转则改用全域搜索。"
            ),
            "sensitivity_analysis": "细化参数扫描和碰撞检测步长。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_accepts_q3_bisection_after_confirming_monotonicity():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "完整搜索0<p<=p_max，从充分外侧r_start到边界"
                "全程用有向矩形SAT检查实体碰撞；"
                "用全域扫描确认可行性单调后再二分，"
                "并扩大r_start后重算p_min，复核临界螺距稳定性。"
            ),
            "sensitivity_analysis": "扩大r_start并细化螺距网格。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_run88_bisection_or_scan_as_monotonicity_check():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "在0<p<=p_max内用有向矩形SAT从充分外侧r_start回放到边界；"
                "验证p_max可行后，再做二分或全域扫描确认可行性单调，"
                "并扩大r_start复算临界值稳定性。"
            ),
            "sensitivity_analysis": "细化参数扫描和碰撞检测步长。",
        }
    )
    with pytest.raises(QualityGateError, match="二分搜索本身"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_accepts_run50_r_t_outer_cutoff_stability_evidence():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "完整搜索0<p<=p_max；从边界向外构造连续构型，"
                "取R_T=4.5+Kp为充分外侧截断，全程用有向矩形SAT检查；"
                "将R_T从4.5+50p扩大到4.5+200p，p_min变化<0.5%。"
            ),
            "sensitivity_analysis": "检查R_T扩展与p_min稳定性。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_run50_uncertified_fixed_q3_upper_bracket():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "搜索域0<p<=p_max=1.0m（初定上界），直接在该区间二分；"
                "从充分外侧到边界全程用有向矩形SAT检查碰撞。"
            ),
            "sensitivity_analysis": "检查网格。",
        }
    )
    with pytest.raises(QualityGateError, match="搜索上界.*可行|扩展"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_accepts_q3_upper_bracket_with_feasibility_expansion():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "完整定义域0<p；先取p_max=1.0m并用全链SAT验证上界可行，"
                "若不可行则将p_max倍增，直至得到不可行/可行括区间后再二分；"
                "从充分外侧到边界全程用有向矩形SAT回放，"
                "并扩大截断半径确认临界值稳定。"
            ),
            "sensitivity_analysis": "检查网格和截断半径。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_accepts_run42_q3_r_start_stability_evidence():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "完整搜索0<p<=p_max；从初始半径R_start=10m（充分外侧）开始，"
                "在龙头到达半径4.5m边界前全程用有向矩形SAT检查224节点板体碰撞；"
                "验证结果对R_start不敏感（扩至15m重算p_min变化<0.001m）。"
            ),
            "sensitivity_analysis": "检查R_start扩展与p_min稳定性。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_arbitrary_positive_q3_search_lower_bound():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "二分p，初始区间如[0.3,1.0]m；用有向矩形SAT检查全链到达边界前碰撞。"
            ),
            "sensitivity_analysis": "检查网格。",
        }
    )
    with pytest.raises(QualityGateError, match="任意裁为"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_bare_positive_q3_scan_bracket_without_lower_proof():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "对候选p从充分外侧回放全链SAT至边界；"
                "先栅格扫描[0.01,2]m确认可行性跃迁。"
            ),
            "sensitivity_analysis": "扩大r_start并检查p_min稳定。",
        }
    )
    with pytest.raises(QualityGateError, match="任意裁为"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_positive_q3_step_scan_start_without_lower_proof():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "p从0.05m步长0.05m递增找首个可行解；"
                "从充分外侧到边界用有向矩形SAT检查全链碰撞。"
            ),
            "sensitivity_analysis": "扩大r_start并检查p_min稳定。",
        }
    )
    with pytest.raises(QualityGateError, match="任意裁为"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_allows_positive_q3_domain_with_computational_bracket():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "完整候选域0<p<=p_max；可在已覆盖全部更小候选后用[0.3,1.0]作数值括区间；"
                "用有向矩形SAT检查全链到达边界前碰撞，并扩大p_max验证稳定。"
            ),
            "sensitivity_analysis": "检查搜索上界。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_width_only_q3_pitch_lower_bound():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "板宽为0.30m，故取p下界为2×0.30=0.60m，搜索从0.60m开始；"
                "每个候选仍用有向矩形SAT检查全链到达边界前的碰撞。"
            ),
            "sensitivity_analysis": "检查网格。",
        }
    )
    with pytest.raises(QualityGateError, match="两倍板宽"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_allows_width_coarse_filter_when_lower_domain_is_retained():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "以两倍板宽0.60m作搜索初值，但仍完整覆盖0<p<0.60的候选；"
                "用有向矩形SAT检查全链到达边界前的碰撞。"
            ),
            "sensitivity_analysis": "检查网格。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_q1_initial_state_parenthetical_in_q3():
    source = (
        SERIAL_MEMBER_SOURCE + "问题 1 初始时龙头位于第16圈。"
        "问题 2 板凳之间碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques3": (
                "龙头初始位于8.8m（问题1初态），用有向矩形SAT检查到达边界前碰撞。"
            ),
            "sensitivity_analysis": "检查边界。",
        }
    )
    with pytest.raises(QualityGateError, match="擅自继承问题一"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_radial_difference_as_polar_chord():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques1": "由二维距离方程 |r(θ_{i+1})-r(θ_i)|=L_i 反解后续极角。",
            "sensitivity_analysis": "检查残差。",
        }
    )
    with pytest.raises(QualityGateError, match="固定弦长"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_rigid_chord_assigned_to_path_arc_length():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques1": (
                "对后续节点求解θ_i>θ_{i-1}，使沿螺线弧长差=d_i，再得到节点坐标。"
            ),
            "sensitivity_analysis": "检查距离残差。",
        }
    )

    with pytest.raises(QualityGateError, match="欧氏弦长误作路径弧长差"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_handle_distance_as_collision_body_length():
    source = (
        SERIAL_MEMBER_SOURCE
        + "所有板凳的板宽为30cm。问题 2 请确定板凳之间不发生碰撞的终止时刻。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques2": (
                "板凳建模为有向矩形，长2.86m或1.65m、宽0.3m；用SAT判定实体碰撞。"
            ),
            "sensitivity_analysis": "检查网格。",
        }
    )
    with pytest.raises(QualityGateError, match="孔中心距误作板凳外形长度"):
        validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_rejects_one_based_endpoint_number_that_stops_early():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m；"
                "龙头前把手为节点1，龙尾后把手为节点223。"
            ),
            "ques1": "用二维欧氏距离递推。",
            "sensitivity_analysis": "检查节点数。",
        }
    )
    with pytest.raises(QualityGateError, match="应编号 224"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_does_not_bind_tail_rear_to_first_node_in_output_list():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m；"
                "龙尾后把手为节点224。"
            ),
            "ques1": (
                "论文给出龙头、第1/51/101/151/201节龙身、"
                "龙尾后把手（节点1/2/52/102/152/202/224）位速。"
            ),
            "sensitivity_analysis": "检查节点数。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_equals_endpoint_number_that_stops_early():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m；"
                "龙头前把手=节点1，龙尾后把手=节点223。"
            ),
            "ques1": "用二维欧氏距离递推。",
            "sensitivity_analysis": "检查节点数。",
        }
    )
    with pytest.raises(QualityGateError, match="应编号 224"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_repeated_segment_index_range_that_stops_early():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 共224个节点；"
                "第i节点到第i+1节点中，i=1用2.86m，i=2..222用1.65m。"
            ),
            "ques1": "用二维欧氏距离递推。",
            "sensitivity_analysis": "检查节点数。",
        }
    )
    with pytest.raises(QualityGateError, match="统一后续段索引|223"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_accepts_backward_endpoint_index_convention():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 共224个节点；对i=2..224建立"
                "|P_i-P_{i-1}|=d_i，其中d_2=2.86m，i≥3时d_i=1.65m。"
            ),
            "ques1": "用二维欧氏距离递推。",
            "sensitivity_analysis": "检查节点数。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_shifted_backward_endpoint_distance_index():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 共224个节点；对i=2..224建立"
                "|P_i-P_{i-1}|=d_i，其中d_1=2.86m，i≥2时d_i=1.65m。"
            ),
            "ques1": "用二维欧氏距离递推。",
            "sensitivity_analysis": "检查节点数。",
        }
    )
    with pytest.raises(QualityGateError, match="下标与端点公式错位|d_2"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_run85_explicit_shifted_backward_distance_range():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques1": (
                "对后续节点i=2,...,224，由刚性约束"
                "|P_i-P_{i-1}|=d_i求解，其中d_1=2.86m，"
                "d_2...d_224=1.65m。"
            ),
            "sensitivity_analysis": "检查节点数。",
        }
    )
    with pytest.raises(QualityGateError, match="整体错移|d_2"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_requires_source_figure_handoff_for_named_initial_point():
    source = SERIAL_MEMBER_SOURCE + FIGURE_DEPENDENT_SOURCE
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224个节点，内部距离2.86m和1.65m；"
                "题面未给A点角度，设初始极角为0。"
            ),
            "ques1": "用二维欧氏距离递推。",
            "sensitivity_analysis": "改变初始极角。",
        }
    )
    with pytest.raises(QualityGateError, match="source_figure_required"):
        validate_modeler_result(result, {"ques1"}, source_text=source)

    result.questions_solution["eda"] += " source_figure_required: 图4待官方PDF核验。"
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_sensitivity_on_named_initial_point_phase():
    source = SERIAL_MEMBER_SOURCE + FIGURE_DEPENDENT_SOURCE
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224个节点，内部距离2.86m和1.65m；"
                "source_figure_required: 已核验图4中A点位于正x轴。"
            ),
            "ques1": "以A点为固定初态，用二维欧氏距离递推。",
            "sensitivity_analysis": (
                "题面未给精确角，对问题1螺线起点相位θ0作±0.1 rad扰动。"
            ),
        }
    )
    with pytest.raises(QualityGateError, match="固定命名初始点|初始相位"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_conditional_sensitivity_on_named_initial_point_phase():
    source = SERIAL_MEMBER_SOURCE + FIGURE_DEPENDENT_SOURCE
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224个节点，内部距离2.86m和1.65m；"
                "source_figure_required: 图4中A点相位待官方PDF核验。"
            ),
            "ques1": "从A点固定初态出发，用欧氏弦长递推。",
            "sensitivity_analysis": (
                "初始相位φ（Q1/Q2，若图4未唯一确定）：φ∈[0,2π)，"
                "检查龙头位置与碰撞时刻变化。"
            ),
        }
    )
    with pytest.raises(QualityGateError, match="固定命名初始点|初始相位"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_curvature_continuity_for_unequal_tangent_arcs():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques4": "两弧相切，路径连续且曲率连续（一阶连续即相切）。",
            "sensitivity_analysis": "检查相切残差。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE
        + "调头路径由两段圆弧相切连接，前一段圆弧半径是后一段的2倍。"
    )
    with pytest.raises(QualityGateError, match="不能声明曲率连续"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_requires_curvature_jump_for_unequal_tangent_arcs():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques4": "两圆弧相切并检查切向残差。",
            "sensitivity_analysis": "检查相切残差。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE
        + "调头路径由两段圆弧相切连接，前一段圆弧半径是后一段的2倍。"
    )
    with pytest.raises(QualityGateError, match="必须同时明确"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_allows_c1_tangent_arcs_with_curvature_jump():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques4": "两弧位置和切向C1连续，但半径不同导致曲率跳变。",
            "sensitivity_analysis": "检查相切残差。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE
        + "调头路径由两段圆弧相切连接，前一段圆弧半径是后一段的2倍。"
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_allows_curvature_continuity_audit_heading_with_jump_result():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques4": (
                "模型选择：两段不等半径相切圆弧与曲率连续性审计。"
                "两弧位置和切向C1连续，连接点两侧曲率1/R1和1/R2不同并发生跳变。"
            ),
            "sensitivity_analysis": "检查相切残差。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE
        + "调头路径由两段圆弧相切连接，前一段圆弧半径是后一段的2倍。"
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_run59_preanswered_q4_shortening_direction():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques4": (
                "R1=2R2；两弧位置和切向C1连续，连接点曲率发生跳变。"
                "题目输出契约：回答‘能否缩短’为肯定并给出缩短比例。"
            ),
            "sensitivity_analysis": "保持半径比固定，检查相切残差。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE + "问题 4 两段圆弧相切，前一段圆弧半径是后一段2倍。"
        "能否调整圆弧并保\n持相切，使调头曲线变短？"
    )
    with pytest.raises(QualityGateError, match="不能在优化计算前.*预设为肯定"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_allows_q4_shortening_decided_after_baseline_comparison():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques4": (
                "R1=2R2；两弧位置和切向C1连续，连接点曲率发生跳变。"
                "由实际切点求圆心角Δφ1、Δφ2，按L=R1·Δφ1+R2·Δφ2计算长度。"
                "先复算题面基线长度，再求约束内候选解；由长度差与误差界决定能否缩短。"
            ),
            "sensitivity_analysis": "保持半径比固定，检查相切残差。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE + "问题 4 两段圆弧相切，前一段圆弧半径是后一段2倍。"
        "能否调整圆弧并保持相切，使调头曲线变短？"
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_q4_shortening_without_numeric_baseline_comparison():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques4": (
                "R1=2R2；两弧位置和切向C1连续，连接点曲率发生跳变。"
                "基准为几何方程求得的可行R2集合，求最小弧长L_opt。"
            ),
            "sensitivity_analysis": "保持半径比固定，检查相切残差。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE + "问题 4 两段圆弧相切，前一段圆弧半径是后一段2倍。"
        "能否调整圆弧并保持相切，使调头曲线变短？"
    )
    with pytest.raises(QualityGateError, match="调整前基线长度"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_ambiguous_inner_outer_tangent_branch_for_s_curve():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques4": (
                "R1=2R2；两弧位置和切向C1连续，连接点曲率发生跳变。"
                "圆心距=R1±R2，视内外切选择分支。"
                "复算调整前基线长度L_base，并与L_opt比较得到缩短比例。"
            ),
            "sensitivity_analysis": "保持半径比固定，检查相切残差。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE + "问题 4 S形曲线由两段圆弧相切连接，"
        "前一段圆弧半径是后一段2倍。能否调整圆弧使曲线变短？"
    )
    with pytest.raises(QualityGateError, match=r"外切分支|R1\+R2"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_accepts_position_plus_tangent_c1_shorthand():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques4": (
                "R1=2R2；两弧满足位置+切向C1，连接点曲率跳变1/R1→1/R2。"
                "选择外切分支，圆心距R1+R2。"
            ),
            "sensitivity_analysis": "保持半径比固定，检查相切残差。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE + "问题 4 S形曲线由两段圆弧相切连接，"
        "前一段圆弧半径是后一段2倍。"
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_accepts_external_center_distance_three_times_base_radius():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques4": (
                "前段半径2R、后段半径R，两弧选择外切连接，圆心距=3R。"
                "两弧位置和切向C1连续，连接点曲率由1/(2R)跳变到1/R。"
            ),
            "sensitivity_analysis": "保持半径比固定，检查相切残差。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE + "问题 4 S形曲线由两段圆弧相切连接，"
        "前一段圆弧半径是后一段2倍。"
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_internal_tangency_label_with_external_center_distance():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques4": (
                "R1=2R2；两弧位置与切向C1连续，连接点曲率跳变。"
                "采用内切分支相切，圆心距R1+R2。"
            ),
            "sensitivity_analysis": "保持半径比固定，检查相切残差。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE + "问题 4 S形曲线由两段圆弧相切连接，"
        "前一段圆弧半径是后一段2倍。"
    )
    with pytest.raises(QualityGateError, match="外切分支.*不得.*内切"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_wrong_collision_entity_count():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223节板凳、224节点，内部距离2.86m和1.65m。",
            "ques2": "用SAT检查222个有向矩形板体，排除相邻装配接触。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="碰撞实体总数.*223"):
        validate_modeler_result(result, {"ques2"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_fractional_collision_pair_count_formula():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223节板凳、224节点，内部距离2.86m和1.65m。",
            "ques2": (
                "用SAT检查223个有向矩形板体；排除相邻装配接触后，"
                "遍历全部221×223/2对非相邻实体。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="半整数|不是整数"):
        validate_modeler_result(result, {"ques2"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_accepts_q3_full_domain_single_transition_wording():
    source = SERIAL_MEMBER_SOURCE + "问题 3 求使龙头到达4.5m边界的最小螺距。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223节板凳、224节点，内部距离2.86m和1.65m。",
            "ques3": (
                "用SAT回放全链至边界，二分搜索最小螺距；"
                "先在全域扫描确认可行性仅有单一跃迁，若反转则改用全域搜索。"
            ),
            "sensitivity_analysis": "扩大r_start并检查p_min稳定。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_does_not_treat_fixed_ratio_as_numeric_r2_baseline():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques4": (
                "R1=2R2；两弧位置和切向C1连续，连接点曲率发生跳变。"
                "基线保持固定比例R1:R2=2:1，由几何方程复算调整前长度L_base；"
                "由实际切点求圆心角Δφ1、Δφ2，按L=R1·Δφ1+R2·Δφ2计算长度；"
                "采用外切分支和圆心距R1+R2，将L_opt与L_base比较并报告缩短比例。"
            ),
            "sensitivity_analysis": "保持半径比固定，检查相切残差。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE + "问题 4 S形曲线由两段圆弧相切连接，"
        "前一段圆弧半径是后一段2倍。能否调整圆弧使曲线变短？"
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_assuming_turning_arcs_are_semicircles():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques4": (
                "R1=2R2；两弧位置和切向C1连续，连接点曲率发生跳变。"
                "采用外切分支和圆心距R1+R2；复算基线长度L_base，"
                "若两弧为半圆，目标弧长L=πR1+πR2=3πR2，"
                "并与L_opt比较缩短比例。"
            ),
            "sensitivity_analysis": "保持半径比固定，检查相切残差。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE + "问题 4 S形曲线由两段圆弧相切连接，"
        "前一段圆弧半径是后一段2倍。能否调整圆弧使曲线变短？"
    )
    with pytest.raises(QualityGateError, match="半圆|3πR2"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_accepts_turning_arc_lengths_from_solved_central_angles():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques4": (
                "R1=2R2；两弧位置和切向C1连续，连接点曲率发生跳变。"
                "采用外切分支和圆心距R1+R2；由实际相切点和连接点"
                "求圆心角Δφ1、Δφ2，以L=R1Δφ1+R2Δφ2"
                "复算调整前基线长度L_base，"
                "再与L_opt比较缩短比例。"
            ),
            "sensitivity_analysis": "保持半径比固定，检查相切残差。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE + "问题 4 S形曲线由两段圆弧相切连接，"
        "前一段圆弧半径是后一段2倍。能否调整圆弧使曲线变短？"
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_accepts_q4_central_angle_length_after_fixed_ratio_substitution():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques4": (
                "R1=2R2；两弧位置和切向C1连续，连接点曲率发生跳变。"
                "采用外切分支和圆心距R1+R2；由实际相切点求圆心角Δφ1、Δφ2，"
                "按L=2R2·Δφ1+R2·Δφ2复算基线长度L_base，"
                "再与L_opt比较缩短比例。"
            ),
            "sensitivity_analysis": "保持半径比固定，检查相切残差。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE + "问题 4 S形曲线由两段圆弧相切连接，"
        "前一段圆弧半径是后一段2倍。能否调整圆弧使曲线变短？"
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_accepts_q4_central_angle_length_with_named_small_radius():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques4": (
                "设后一段圆弧半径r，前一段圆弧半径R1=2r；"
                "两弧位置和切向C1连续，连接点曲率发生跳变，"
                "采用外切分支且圆心距R1+R2。"
                "由实际切点求圆心角Δφ1、Δφ2，目标L=2rΔφ1+rΔφ2；"
                "数值复算调整前基线长度L_base，并与L_opt比较缩短比例和误差界。"
            ),
            "sensitivity_analysis": "保持半径比固定，检查相切残差。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE + "问题 4 S形曲线由两段圆弧相切连接，"
        "前一段圆弧半径是后一段2倍。能否调整圆弧使曲线变短？"
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_accepts_run87_radius_function_central_angle_length():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques4": (
                "固定R1:R2=2:1，记R2=R；两弧位置和切向C1连续，"
                "连接点曲率跳变，采用外切分支和圆心距R1+R2。"
                "由实际切点求圆心角，解算L(R)=2R·Δφ1(R)+R·Δφ2(R)，"
                "并与按相同口径复算的调整前基线L_base比较缩短量。"
            ),
            "sensitivity_analysis": "保持半径比固定，检查相切残差。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE + "问题 4 S形曲线由两段圆弧相切连接，"
        "前一段圆弧半径是后一段2倍。能否调整圆弧使曲线变短？"
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_q5_ambiguous_problem4_path_inheritance():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques5": "沿用问题4路径，基线或优化后构型皆可，需明确后扫描速度倍率。",
            "sensitivity_analysis": "扩大外侧截断检查倍率稳定。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 调整双圆弧并比较是否缩短。"
        "问题 5 舞龙队沿问题 4 设定的路径行进，确定龙头最大速度。"
    )
    with pytest.raises(QualityGateError, match="继承问题四最终|任意选择"):
        validate_modeler_result(result, {"ques5"}, source_text=source)


def test_modeler_rejects_q5_multistart_as_global_certificate():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques5": (
                "使用问题4最终确定的路径，扫描全部节点和时间；"
                "多起点爬坡确认全局。"
            ),
            "sensitivity_analysis": "扩大外侧截断检查倍率稳定。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "问题 5 求龙头最大速度。"
    with pytest.raises(QualityGateError, match="多起点.*不能单独确认全局"):
        validate_modeler_result(result, {"ques5"}, source_text=source)


def test_modeler_rejects_minimum_node_speed_as_q5_active_constraint():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 共224个节点，内部距离2.86m和1.65m。",
            "ques5": (
                "计算M_max=max_{k,t}(v_k/v_head)，并取v_head,max=2/M_max；"
                "重跑确认最小节点速率达到2m/s上限即活动约束。"
            ),
            "sensitivity_analysis": "细化时间步长。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "问题 5 求所有把手速度不超过2m/s时的最大龙头速度。"
    with pytest.raises(QualityGateError, match="最小节点速度达到上限"):
        validate_modeler_result(result, {"ques5"}, source_text=source)


def test_modeler_rejects_explicit_backward_distance_indices_shifted_by_one():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223节板凳、224节点，孔距2.86m和1.65m。",
            "ques1": (
                "对后续节点i=2..224，由|节点i−节点i−1|=d_i递推；"
                "取d₁=2.86m，d₂..d₂₂₃=1.65m。"
            ),
            "sensitivity_analysis": "减小时间步长检查残差。",
        }
    )
    with pytest.raises(QualityGateError, match="显式下标|d_2=2.86"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_explicit_backward_length_indices_shifted_by_one():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223节板凳、224节点，孔距2.86m和1.65m。",
            "ques1": (
                "对后续节点由|P_i-P_{i-1}|=L_i递推，i=2..224；"
                "取L₁=2.86m，L₂..L₂₂₃=1.65m。"
            ),
            "sensitivity_analysis": "减小时间步长检查残差。",
        }
    )
    with pytest.raises(QualityGateError, match="整体错移|首段应为 d_2"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_explicit_forward_distance_indices_shifted_by_one():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223节板凳、224节点，孔距2.86m和1.65m。",
            "ques1": (
                "节点2~224由刚性约束|P_{i+1}-P_i|=d_i递推；"
                "取d₂=2.86m，d₃..d₂₂₄=1.65m。"
            ),
            "sensitivity_analysis": "减小时间步长检查残差。",
        }
    )
    with pytest.raises(QualityGateError, match="前向端点公式整体错移|d_1=2.86"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_eda_contract_rejects_unproved_positive_q3_pitch_lower_bound(tmp_path: Path):
    contract = {
        "status": "success",
        "subtask": "eda",
        "data_scope": "official statement",
        "assumptions": [],
        "model": "rigid chain",
        "numeric_results": {
            "entity_count": 223,
            "node_count": 224,
            "head_node_distance_m": 2.86,
            "body_node_distance_m": 1.65,
        },
        "validation": {"passed": True, "checks": ["topology"]},
        "conclusion_bounds": {"Q3_min_pitch_searcher_lower_bound_m": 0.05},
    }
    (tmp_path / "results_eda.json").write_text(json.dumps(contract), encoding="utf-8")
    with pytest.raises(QualityGateError, match="不得.*擅设正下界"):
        load_evidence_contract(tmp_path, "eda", source_text=SERIAL_MEMBER_SOURCE)


def test_eda_contract_does_not_require_q4_numeric_baseline_comparison(tmp_path: Path):
    contract = {
        "status": "success",
        "subtask": "eda",
        "data_scope": "official statement",
        "assumptions": [],
        "model": "source audit only",
        "numeric_results": {
            "entity_count": 223,
            "node_count": 224,
            "head_node_distance_m": 2.86,
            "body_node_distance_m": 1.65,
        },
        "validation": {"passed": True, "checks": ["source facts"]},
        "conclusion_bounds": "source facts only",
    }
    (tmp_path / "results_eda.json").write_text(json.dumps(contract), encoding="utf-8")
    source = SERIAL_MEMBER_SOURCE + (
        "问题4的双圆弧前一段半径是后一段2倍，能否调整圆弧使曲线变短？"
    )
    load_evidence_contract(tmp_path, "eda", source_text=source)


def test_modeler_allows_following_nodes_to_take_increasing_theta_root():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223节板凳、224节点，孔距2.86m和1.65m。",
            "ques1": (
                "采用r=cθ且c>0，龙头向中心盘入时θ随时间递减；"
                "由后向约束|P_i-P_{i-1}|=d_i，d_2=2.86m，d_3..d_224=1.65m。"
                "后方节点位于龙头外侧，沿θ增大方向取根，并由上一时刻连续延拓。"
            ),
            "sensitivity_analysis": "减小时间步长检查残差。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "等距螺线盘入。"
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_allows_symbolic_doubled_spatial_cutoff_recalculation():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223节板凳、224节点，孔距2.86m和1.65m。",
            "ques3": (
                "从充分外侧截断r_start=60m连续回放到目标边界。"
                "取r_start和2×r_start两个截断，分别重新求解p_min，"
                "并报告差值<1e-6m。搜索域保持0<p<=p_max。"
            ),
            "sensitivity_analysis": "比较两个空间截断下的临界量。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "问题3求最小螺距。"
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_bare_wrong_total_rigid_segment_count():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223节板凳、224节点，孔距2.86m和1.65m。",
            "ques1": (
                "由|P_i-P_{i-1}|=d_i递推，d_2=2.86m，d_3..d_224=1.65m；"
                "验证总段数=1+221=222段。"
            ),
            "sensitivity_analysis": "检查距离残差。",
        }
    )
    with pytest.raises(QualityGateError, match="总数|223"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_malformed_archimedean_arc_primitive():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223节板凳、224节点，孔距2.86m和1.65m。",
            "ques1": (
                "螺线r=pθ/(2π)，使用弧长原函数"
                "S(θ)=0.5×sqrt((p/2π)²θ²+r²)+(p/4π)ln(θ+sqrt(θ²+r²))。"
            ),
            "sensitivity_analysis": "减小时间步长检查残差。",
        }
    )
    with pytest.raises(QualityGateError, match="弧长原函数漏乘|不能从"):
        validate_modeler_result(result, {"ques1"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_rejects_wrong_double_arc_length_coefficients():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223节板凳、224节点，孔距2.86m和1.65m。",
            "ques4": (
                "设R1=2ρ、R2=ρ，选择外切分支且圆心距R1+R2；"
                "两弧位置与切向C1连续、曲率发生跳变。由实际切点确定圆心角，"
                "却写总长L(ρ)=3ρΔφ1+1.5ρΔφ2。"
                "由题面相切约束复算调整前基线长度L_base，"
                "与L_opt数值比较缩短比例和误差界。"
            ),
            "sensitivity_analysis": "保持半径比固定并检查相切残差。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 S形曲线由两段圆弧相切连接，前一段半径是后一段2倍。"
        "能否调整圆弧使调头曲线变短？"
    )
    with pytest.raises(QualityGateError, match="长度系数|2ρΔφ1"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


@pytest.mark.parametrize(
    "speed_claim",
    [
        "若M_max≤1，则v_max无界且可任意大。",
        "当max m_i>1时v_head,max=2/max m_i，否则v_head不限。",
    ],
)
def test_modeler_rejects_unbounded_speed_when_max_multiplier_at_most_one(
    speed_claim,
):
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223节板凳、224节点，孔距2.86m和1.65m。",
            "ques5": (
                "扫描全部224节点，定义M_max=max m_i(t)，v_max=2/M_max。"
                + speed_claim
            ),
            "sensitivity_analysis": "扩大扫描区间检查倍率。",
        }
    )
    with pytest.raises(QualityGateError, match="龙头自身倍率恒为1|至少为1"):
        validate_modeler_result(result, {"ques5"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_allows_adjacent_assembly_contact_said_as_not_counted_collision():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223节板凳、224节点，内部距离2.86m和1.65m。",
            "ques2": (
                "以SAT检查223个有向矩形板体；碰撞对象为所有非相邻板凳对，"
                "相邻板凳在共享把手处的装配接触不算碰撞。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    validate_modeler_result(result, {"ques2"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_does_not_treat_nonadjacent_boards_as_adjacent_collision():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个构件、224节点，内部弦长2.86和1.65m。",
            "ques2": (
                "用SAT判定任意两不相邻板凳矩形是否相交；"
                "相邻连接对的共享把手装配接触排除。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    validate_modeler_result(result, {"ques2"}, source_text=SERIAL_MEMBER_SOURCE)


def test_modeler_allows_open_ratio_only_as_infeasible_comparison():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223节板凳、224节点，内部距离2.86m和1.65m。",
            "ques4": (
                "R1=2R2；两弧位置和切向C1连续，连接点曲率发生跳变；"
                "采用外切分支和圆心距R1+R2。复算基线长度L_base并与L_opt比较；"
                "由实际切点求圆心角Δφ1、Δφ2，按L=R1·Δφ1+R2·Δφ2计算长度；"
                "等半径、单圆弧、开放比例的方案仅作不可行对照。"
            ),
            "sensitivity_analysis": "保持2:1固定，不放开该比例。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE + "问题 4 S形曲线由两段圆弧相切连接，"
        "前一段圆弧半径是后一段2倍。能否调整圆弧使曲线变短？"
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_allows_opening_geometry_while_fixed_ratio_is_retained():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques4": (
                "R1=2R2；两弧位置和切向C1连续，连接点曲率由1/R1跳变到1/R2。"
                "选择外切分支，圆心距=R1+R2。由题面几何复算基线长度L_base；"
                "由实际切点求圆心角Δφ1、Δφ2，按L=R1·Δφ1+R2·Δφ2计算长度；"
                "放开几何构型，在保持R1:R2=2:1的约束内调整切点，"
                "将L_opt与L_base比较并报告缩短比例。"
            ),
            "sensitivity_analysis": "保持半径比固定，检查相切残差。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 S形曲线由两段圆弧相切连接，"
        "前一段圆弧半径是后一段2倍。能否调整圆弧使曲线变短？"
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_s_curve_without_explicit_external_tangent_branch():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques4": (
                "R1=2R2；两弧位置和切向C1连续，连接点曲率由1/R1跳变到1/R2。"
                "由题面几何复算基线长度L_base，并与L_opt比较缩短比例。"
            ),
            "sensitivity_analysis": "保持半径比固定，检查相切残差。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 S形曲线由两段圆弧相切连接，"
        "前一段圆弧半径是后一段2倍。能否调整圆弧使曲线变短？"
    )
    with pytest.raises(QualityGateError, match=r"外切分支|R1\+R2"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_q5_maximum_that_omits_last_node():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 路径含盘入螺线和盘出螺线。问题 5 求各把手不超过2m/s时的最大龙头速度。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223节板凳、224节点，内部距离2.86m和1.65m。",
            "ques5": (
                "取K_max=max_{t,i=1..223}g_i(t)，令v*=2/K_max。"
                "扩大截断后验证极值不再增加。"
            ),
            "sensitivity_analysis": "细化时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match=r"i=1\.\.224|链尾"):
        validate_modeler_result(result, {"ques5"}, source_text=source)


def test_modeler_rejects_finite_kmax_without_unbounded_tail_evidence():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 路径含盘入螺线和盘出螺线。问题 5 求各把手不超过2m/s时的最大龙头速度。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223节板凳、224节点，内部距离2.86m和1.65m。",
            "ques5": "取K_max=max_{t∈[-100,100],i=1..224}g_i(t)，令最大允许龙头速度v*=2/K_max。",
            "sensitivity_analysis": "细化时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="无限尾部|解析渐近尾部界"):
        validate_modeler_result(result, {"ques5"}, source_text=source)


def test_modeler_rejects_inward_trailing_nodes_described_as_inside():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224个节点，内部距离2.86m和1.65m；"
                "采用r=a+bθ且b>0。"
            ),
            "ques1": (
                "龙头向中心盘入；龙头后把手位于半径方向内侧，"
                "后续每个节点继续沿螺线内侧递推。"
            ),
            "sensitivity_analysis": "检查步长和残差。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "沿螺距55cm的等距螺线顺时针盘入。"
    with pytest.raises(QualityGateError, match="盘入方向|跟随次序"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_run55_chain_tail_described_as_inside():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224个节点，内部距离2.86m和1.65m；"
                "采用r=pθ/(2π)，链尾可延伸至第16圈内侧多层。"
            ),
            "ques1": "龙头向中心盘入，后续节点由欧氏弦长方程递推。",
            "sensitivity_analysis": "检查步长和残差。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "沿螺距55cm的等距螺线顺时针盘入。"
    with pytest.raises(QualityGateError, match="盘入方向|跟随次序"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_positive_theta_rate_for_inward_positive_spiral():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224个节点，内部距离2.86m和1.65m；"
                "采用r=(0.55/(2π))θ，顺时针盘入时θ递减。"
            ),
            "ques1": "龙头速率v=1m/s，取dθ/dt=v/sqrt(r^2+c^2)。",
            "sensitivity_analysis": "检查步长和残差。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "沿螺距55cm的等距螺线顺时针盘入。"
    with pytest.raises(QualityGateError, match="角速度公式符号"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_run49_negative_theta_start_that_moves_outward():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224个节点，内部距离2.86m和1.65m；"
                "采用r=pθ/(2π)，第16圈半径为8.8m。"
            ),
            "ques1": "龙头从θ_0=-32π开始盘入，随后θ继续递减。",
            "sensitivity_analysis": "检查步长和残差。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "沿螺距55cm的等距螺线顺时针盘入。"
    with pytest.raises(QualityGateError, match="负角初值|向外"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_direction_gate_does_not_join_different_sentences():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224个节点，内部距离2.86m和1.65m；"
                "采用r=(0.55/(2π))θ。"
            ),
            "ques1": (
                "龙头前把手沿螺线顺时针盘入，极角随时间递减。"
                "其余223个节点由欧氏距离方程递推，取θ增大的外侧解。"
            ),
            "sensitivity_analysis": "检查步长和残差。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "沿螺距55cm的等距螺线顺时针盘入。"
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_direction_gate_does_not_join_head_and_trailing_node_clauses():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224个节点，内部距离2.86m和1.65m；"
                "采用r=(0.55/(2π))θ。"
            ),
            "ques1": (
                "顺时针盘入时龙头极角θ递减，后续节点位于龙头外侧，"
                "极角θ逐点增大；盘出螺线与盘入关于中心对称（θ递增盘出）。"
            ),
            "sensitivity_analysis": "检查步长和残差。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "沿螺距55cm的等距螺线顺时针盘入。"
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_allows_c1_not_c2_shorthand_for_tangent_unequal_arcs():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques4": (
                "两段圆弧外切，半径比R1:R2=2:1，连接点满足C1非C2，"
                "因此切向连续但曲率1/R1与1/R2不同并发生跳变。"
            ),
            "sensitivity_analysis": "检查相切残差。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "调头曲线由半径比2:1的两段圆弧相切连接。"
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_allows_unequal_arc_c1_contract_in_global_audit():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224个节点，内部距离2.86m和1.65m；"
                "双圆弧连接点为C1连续，曲率1/R1与1/R2不同并发生跳变。"
            ),
            "ques4": "两段圆弧外切，半径比R1:R2=2:1，并检查相切残差。",
            "sensitivity_analysis": "检查相切残差。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "调头曲线由半径比2:1的两段圆弧相切连接。"
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_allows_c1_not_c2_without_repeating_curvature_formula():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224个节点，内部距离2.86m和1.65m；"
                "双圆弧采用外切分支（圆心距R1+R2），连接点仅位置和切向C1连续非C2。"
            ),
            "ques4": "两段圆弧半径比R₁:R₂=2:1，并检查相切残差。",
            "sensitivity_analysis": "检查相切残差。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "调头曲线由半径比2:1的两段圆弧相切连接。"
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_allows_position_tangent_continuous_not_c2_shorthand():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224个节点，内部距离2.86m和1.65m；"
                "双圆弧采用外切分支（圆心距R1+R2），连接点仅位置+切向连续非C2。"
            ),
            "ques4": "两段圆弧半径比R1:R2=2:1，并检查相切残差。",
            "sensitivity_analysis": "检查相切残差。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "调头曲线由半径比2:1的两段圆弧相切连接。"
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_allows_external_tangency_contract_in_global_audit():
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224个节点，内部距离2.86m和1.65m；"
                "双圆弧采用反向外切分支（圆心距R1+R2），"
                "连接点C1连续且曲率1/R1与1/R2跳变。"
            ),
            "ques4": "两段圆弧半径比R₁:R₂=2:1，并检查相切残差。",
            "sensitivity_analysis": "检查相切残差。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题4调头路径为S形曲线，由半径比2:1的两段圆弧相切连接。"
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_allows_unicode_subscripts_in_external_tangency_contract():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques4": (
                "R₁=2R₂；选择外切分支，圆心距O₁O₂=R₁+R₂；"
                "两弧位置与切向C1连续，连接点曲率由1/R₁跳变到1/R₂。"
            ),
            "sensitivity_analysis": "保持半径比固定，检查相切残差。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE + "问题4调头路径为S形曲线，由两段圆弧相切连接，"
        "前一段圆弧半径是后一段的2倍。"
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


@pytest.mark.parametrize(
    "direction_text",
    [
        "龙头随时间θ递减；沿链向外满足θ_{龙头}<θ_{i}<θ_{i+1}<...，即θ逐点增大。",
        "从龙头向后逐点求解224个节点，取全部θ_i单调递增的外侧根。",
        "从龙头前把手出发，沿θ增大方向（向外）逐点解出后续节点。",
    ],
)
def test_modeler_direction_gate_allows_increasing_chain_phrasings(
    direction_text: str,
):
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224个节点，内部距离2.86m和1.65m；"
                "采用r=(0.55/(2π))θ。"
            ),
            "ques1": direction_text,
            "sensitivity_analysis": "检查步长和残差。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "沿螺距55cm的等距螺线顺时针盘入。"
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_allows_circle_arc_angle_recovered_from_its_chord():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques4": (
                "双圆弧总弧长=2ρΔθ1+ρΔθ2；每段圆心角由该圆弧端点弦长"
                "与对应半径的圆几何关系独立反解，不涉及板凳孔距。"
            ),
            "sensitivity_analysis": "检查相切残差。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "问题 4 调头路径由两段圆弧相切连接。"
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_allows_unicode_subscripts_in_actual_central_angle_length():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques4": (
                "两弧外切且圆心距R1+R2，连接点C1连续、曲率跳变；"
                "按实际切点确定圆心角，调整前基线长度"
                "L_base=R₁Δφ₁+R₂Δφ₂；"
                "优化后计算L_opt，并比较L_opt与L_base判断能否缩短。"
            ),
            "sensitivity_analysis": "检查相切残差。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE + "问题4的S形曲线由半径比2:1的两段圆弧相切连接，"
        "能否调整圆弧使调头曲线变短？"
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_point_distance_as_final_bench_collision_test():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques2": (
                "碰撞判据为任意两个非相邻把手节点间欧氏距离小于0.30m，"
                "首次达到阈值即为终止时刻。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "板宽为30cm，请确定板凳之间不发生碰撞的终止时刻。"
    with pytest.raises(QualityGateError, match="节点距离只能用于粗筛"):
        validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_allows_safe_distance_prefilter_before_oriented_rectangles():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques2": (
                "先以两板半对角线包围圆半径之和作中心距离粗筛并生成候选对，"
                "排除共享把手的相邻连接对后，再用含真实长度、宽度和朝向的"
                "有向矩形分离轴精确判定非相邻板体碰撞。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "板宽为30cm，请确定板凳之间不发生碰撞的终止时刻。"
    validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_rejects_vague_node_circle_or_grid_collision_prefilter():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques2": (
                "排除共享把手的相邻连接对；粗筛用节点包围圆/空间网格生成候选对，"
                "再用真实板长、板宽和朝向的有向矩形SAT精确判定。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "板宽为30cm。问题2检查板凳碰撞。"
    with pytest.raises(QualityGateError, match="候选生成无漏检"):
        validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_rejects_half_step_claim_with_tenfold_numeric_change():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques2": "时间步减半（0.1→0.01s）后复算首次碰撞时刻。",
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "问题2请确定首次碰撞终止时刻。"
    with pytest.raises(QualityGateError, match="文字与数值不一致"):
        validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_rejects_run55_half_length_sum_as_oriented_contact_distance():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques2": (
                "排除相邻连接对，以真实板长、板宽和朝向建立有向矩形，"
                "用SAT精判首次碰撞；绘制碰撞对中心距曲线，最小值接近两半长之和。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "所有板凳的板宽为30cm。问题2检查板凳碰撞。"
    with pytest.raises(QualityGateError, match="没有统一.*中心距"):
        validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_rejects_collision_search_without_connected_neighbour_rule():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques2": (
                "将板凳表示为真实长宽的有向矩形，用SAT检查任意两节矩形；"
                "矩形重叠即碰撞。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "板宽为30cm，请确定板凳之间不发生碰撞的终止时刻。"
    with pytest.raises(QualityGateError, match="连接邻居|共享把手"):
        validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_accepts_connected_neighbour_allowed_overlap_wording():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques2": (
                "板凳使用3.41/2.20m、宽0.30m的有向矩形并以SAT精判碰撞；"
                "相邻板体把手连接处允许重叠且不视为碰撞，仅检查非相邻板体。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "所有板凳的板宽为30cm。问题2检查碰撞。"
    validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_accepts_run86_adjacent_pairs_are_allowed_assembly_wording():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques2": (
                "223实体中相邻对(1,2)..(222,223)为允许装配，"
                "其余非相邻板体用真实长宽有向矩形SAT精判碰撞。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "所有板凳的板宽为30cm。问题2检查碰撞。"
    validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_accepts_run87_assembly_contact_is_allowed_wording():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques2": (
                "223个有向矩形板体用SAT精判；相邻板(i,i+1)因共用把手，"
                "装配接触为允许，不判碰撞；其余非相邻板对检查实体重叠。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "所有板凳的板宽为30cm。问题2检查碰撞。"
    validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_accepts_run87_time_advances_from_zero_wording():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques2": (
                "223个有向矩形板体用SAT精判，排除相邻装配接触；"
                "时间推进从0开始用小步更新全部节点，定位首次非相邻碰撞。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题2确定板凳之间不能再继续盘入的首次碰撞终止时刻。"
    )
    validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_accepts_run86_two_spatial_cutoff_levels_wording():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques3": (
                "从充分外侧r_start布置全链，取两挡r_start（两倍递增量）"
                "分别重新求解p_min并检查截断稳定性。"
            ),
            "sensitivity_analysis": "r_start取两挡50/80m，验证p_min差<1e-6m。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_accepts_run87_doubled_cutoff_result_difference_wording():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques3": (
                "从外侧r_start连续回放到目标边界；对不同r_start计算p_min判断收敛。"
            ),
            "sensitivity_analysis": (
                "将r_start增大一倍计算p_min，结果差≤0.001m。"
            ),
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_accepts_run87_no_collision_until_turn_boundary_wording():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 2 碰撞后不能继续盘入。问题 3 求使龙头到达4.5m边界的最小螺距。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques3": (
                "用有向矩形SAT判据检查整段。可行定义为直至r=4.5m无碰撞；"
                "从充分外侧连续回放并扩大截断后复算p_min，结果差<1e-6m。"
            ),
            "sensitivity_analysis": "细化碰撞检查步长。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_collision_body_length_that_rounds_source_value():
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224个节点，内部距离2.86m和1.65m。",
            "ques2": (
                "板凳为长340/220cm、宽30cm的有向矩形；排除相邻连接对后，"
                "用SAT精确判定非相邻板体碰撞。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    source = SERIAL_MEMBER_SOURCE + "板宽为30cm，请确定板凳之间不发生碰撞的终止时刻。"
    with pytest.raises(QualityGateError, match="外形长度与题面原长不一致"):
        validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_rejects_expanding_head_boundary_to_whole_chain():
    result = ModelerToCoder(
        questions_solution={
            "eda": ("source_parameter_audit: 224个节点，内部距离2.86m和1.65m。"),
            "ques3": (
                "最小化螺距；整条龙224个节点全部进入圆内，并要求最大节点极径≤4.5。"
            ),
            "sensitivity_analysis": "检查边界扰动。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 板凳之间不发生碰撞时可继续盘入。"
        "问题 3 调头空间为直径9m的圆形区域，请确定最小螺距，使得"
        "龙头前把手能够沿着相应的螺线盘入到调头空间的边界。"
    )
    with pytest.raises(QualityGateError, match="扩大题面约束主语"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_allows_head_boundary_with_whole_chain_collision_check():
    result = ModelerToCoder(
        questions_solution={
            "eda": ("source_parameter_audit: 224个节点，内部距离2.86m和1.65m。"),
            "ques3": (
                "以龙头前把手到达调头空间边界为可达条件，同时复核全链非相邻板体不碰撞。"
            ),
            "sensitivity_analysis": "检查边界扰动。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE + "调头空间为直径9m的圆形区域，请确定最小螺距，使得"
        "龙头前把手能够沿着相应的螺线盘入到调头空间的边界。"
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_allows_arrival_with_all_bodies_collision_free():
    result = ModelerToCoder(
        questions_solution={
            "eda": ("source_parameter_audit: 224个节点，内部距离2.86m和1.65m。"),
            "ques3": (
                "龙头前把手到达半径4.5m的圆边界；"
                "到达前全部板体无碰撞，并用问题2的SAT判据完整复核。"
            ),
            "sensitivity_analysis": "检查边界扰动。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE + "问题 2 板凳之间不发生碰撞时可继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入调头空间边界。"
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_allows_explicit_denial_of_whole_chain_circle_constraint():
    result = ModelerToCoder(
        questions_solution={
            "eda": ("source_parameter_audit: 224个节点，内部距离2.86m和1.65m。"),
            "ques3": (
                "目标主语仅为龙头前把手，不要求全部节点进入圆内，"
                "但全链板体需无碰撞直至目标事件。"
            ),
            "sensitivity_analysis": "检查边界扰动。",
        }
    )
    source = (
        SERIAL_MEMBER_SOURCE + "调头空间为直径9m的圆形区域，请确定最小螺距，使得"
        "龙头前把手能够沿着相应的螺线盘入到调头空间的边界。"
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_run55_q1_window_as_q3_feasibility_evidence():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 1 螺距为55cm，初始时位于第16圈，计算到300s。"
        "问题 2 请确定板凳之间不发生碰撞的终止时刻。"
        "问题 3 调头空间半径4.5m，请确定最小螺距。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques3": (
                "搜索0<p<=p_max，对候选螺距以真实板长宽有向矩形SAT复核到达4.5m边界；"
                "p=0.55m因问题1在300s内无碰撞，应可行并提供下界证据。"
            ),
            "sensitivity_analysis": "扩大外侧截断并检查稳定性。",
        }
    )
    with pytest.raises(QualityGateError, match="问题一只验证其规定时间段"):
        validate_modeler_result(result, {"ques3"}, source_text=source)

    result.questions_solution["ques3"] = (
        "搜索0<p<=p_max，对候选螺距以真实板长宽有向矩形SAT复核到达4.5m边界；"
        "p=0.55m应可行并提供下界证据。"
    )
    with pytest.raises(QualityGateError, match="问题一只验证其规定时间段"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_run55_q4_pitch_as_q3_bracket_source():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 3 调头空间半径4.5m，请确定最小螺距。问题 4 盘入螺线的螺距为1.7m。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques3": (
                "搜索0<p<=p_max；p_max初始取1.7m（问题4参考值），"
                "若不可行则倍增直至获得可行括界。"
            ),
            "sensitivity_analysis": "扩大外侧截断并检查稳定性。",
        }
    )
    with pytest.raises(QualityGateError, match="问题四的1.7 m"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_run55_q3_critical_pitch_configuration_in_q4():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 3 调头空间为直径9m的圆。"
        "问题 4 盘入螺线的螺距为1.7m，两段圆弧相切，前弧半径是后弧2倍。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224节点，内部弦长2.86m和1.65m；"
                "问题4初态继承问题3搜索到的临界螺距构型。"
            ),
            "ques4": "R1=2R2；位置和切向C1连续，连接点两侧曲率跳变。",
            "sensitivity_analysis": "保持2:1不扰动，检查相切残差。",
        }
    )
    with pytest.raises(QualityGateError, match="只继承问题三的调头空间"):
        validate_modeler_result(result, {"ques4"}, source_text=source)

    result.questions_solution["eda"] = (
        "source_parameter_audit: 224节点，内部弦长2.86m和1.65m；"
        "问题4仅继承问题3调头空间，并按本问1.7m螺距重新建立构型。"
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_allows_explicit_denial_of_q4_inheriting_q3_pitch():
    source = SERIAL_MEMBER_SOURCE + ("问题 3 求最小螺距。问题 4 盘入螺线的螺距为1.7m。")
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224节点，内部弦长2.86m和1.65m；"
                "问题4不继承问题3的最小螺距数值，仅继承调头空间区域几何。"
            ),
            "ques4": "R1=2R2；位置和切向C1连续，连接点两侧曲率跳变。",
            "sensitivity_analysis": "保持2:1不扰动，检查相切残差。",
        }
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_run55_global_speed_factor_without_infinite_tail_bound():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 路径含盘入螺线和盘出螺线。"
        "问题 5 确定龙头最大速度，使各把手速度不超过2m/s。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques5": "以v0=1仿真，取C=max_{k,t}|v_k|，则v0*=2/C。",
            "sensitivity_analysis": "细化时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="无限尾部"):
        validate_modeler_result(result, {"ques5"}, source_text=source)


def test_modeler_rejects_run56_complete_cycle_beta_max_without_tail_bound():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 路径含盘入螺线和盘出螺线。"
        "问题 5 确定龙头最大速度，使各把手速度不超过2m/s。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques5": (
                "取β_max=max_{i∈链条, 沿路径}(v_i/v_head)，令龙头以1m/s行进，"
                "仿真一个完整周期（包括盘入螺线段、S形调头段、盘出螺线段），"
                "再计算V_max=2/β_max。"
            ),
            "sensitivity_analysis": "细化弧长步长。",
        }
    )
    with pytest.raises(QualityGateError, match="无限尾部|螺线向外无限延伸"):
        validate_modeler_result(result, {"ques5"}, source_text=source)


def test_modeler_accepts_beta_max_with_explicit_infinite_tail_bound():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 路径含盘入螺线和盘出螺线。"
        "问题 5 确定龙头最大速度，使各把手速度不超过2m/s。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques5": (
                "取β_max=max_{i∈链条, 沿路径}(v_i/v_head)。螺线两端向外无限延伸；"
                "用切向投影比趋近1的解析尾部界证明外侧速度倍率不超过有限区间最大值。"
            ),
            "sensitivity_analysis": "扩大截断区间检查上界稳定。",
        }
    )
    validate_modeler_result(result, {"ques5"}, source_text=source)


def test_modeler_rejects_run57_xi_global_factor_without_infinite_tail_bound():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 路径含盘入螺线和盘出螺线。"
        "问题 5 确定龙头最大速度，使各把手速度不超过2m/s。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques5": (
                "取ξ_max=max_{k,t}ξ_k(t)，覆盖盘入、调头和盘出全时间区间，"
                "故V_max=2/max_{k,t}ξ_k(t)。"
            ),
            "sensitivity_analysis": "细化时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="无限尾部"):
        validate_modeler_result(result, {"ques5"}, source_text=source)


def test_modeler_accepts_xi_global_factor_with_infinite_tail_bound():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 路径含盘入螺线和盘出螺线。"
        "问题 5 确定龙头最大速度，使各把手速度不超过2m/s。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques5": (
                "取ξ_max=max_{k,t}ξ_k(t)，故V_max=2/max_{k,t}ξ_k(t)。"
                "螺线两端向外无限延伸，用切向投影比趋近1的解析尾部界证明"
                "外侧速度倍率不超过有限区间最大值。"
            ),
            "sensitivity_analysis": "扩大截断区间检查上界稳定。",
        }
    )
    validate_modeler_result(result, {"ques5"}, source_text=source)


def test_modeler_rejects_run58_inward_spiral_arc_length_plus_sign():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 1 沿等距螺线顺时针盘入，龙头速度为1m/s。问题 2 求首次碰撞。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques1": (
                "采用r=0.55θ/(2π)，令s(t)=t+s0，s0由θ=32π处弧长积分求得。"
                "节点距离方程用局部增量括取最近根，并以上一时刻连续延拓。"
            ),
            "sensitivity_analysis": "细化时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="弧长原函数符号|s0-vt"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_run85_forward_difference_with_positive_inward_arc_length():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 1 沿等距螺线顺时针盘入，龙头速度为1m/s。问题 2 求首次碰撞。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques1": (
                "采用r=0.55θ/(2π)，以S(θ(t))=s0-vt保证向内运动；"
                "核对S(θ(t+Δt))−S(θ(t))=v·Δt。"
                "节点距离方程使用局部增量括界并连续延拓。"
            ),
            "sensitivity_analysis": "细化时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="差分方向写反"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_run86_one_sided_arc_length_residual():
    source = SERIAL_MEMBER_SOURCE + "问题 1 沿等距螺线顺时针盘入。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques1": "验证弧长闭合差 |S(θ(t))-S(θ(t-Δt))|-1*Δt≤1e-6。",
            "sensitivity_analysis": "细化时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="单边不等式"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_accepts_two_sided_arc_length_residual():
    source = SERIAL_MEMBER_SOURCE + "问题 1 沿等距螺线顺时针盘入。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques1": "验证弧长闭合残差 ||S(θ(t))-S(θ(t-1))|-1|≤1e-6。",
            "sensitivity_analysis": "细化时间步长。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_run86_fixed_grid_claiming_to_cover_open_pitch_domain():
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题 3 请确定最小螺距，使龙头前把手能够沿相应螺线盘入到调头空间边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques3": (
                "在完整参数域p∈(0,p_max]采用均匀网格扫描，"
                "网格步长Δp=0.005m，并称从p→0+开始检查。"
            ),
            "sensitivity_analysis": "细化网格。",
        }
    )
    with pytest.raises(QualityGateError, match="固定正步长.*开放域"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_run86_reusing_one_link_tail_bound_for_full_chain():
    source = SERIAL_MEMBER_SOURCE + "问题 5 求所有把手限速下的龙头最大速度。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques5": (
                "扫描全路径所有节点倍率并取M_max，令v_head,max=2/M_max。"
                "相邻节点速度倍率≤1+C/r，故截断外整体倍率≤1+C/r_cut。"
            ),
            "sensitivity_analysis": "扩大截断半径。",
        }
    )
    with pytest.raises(QualityGateError, match="单个相邻节点.*全链"):
        validate_modeler_result(result, {"ques5"}, source_text=source)


def test_modeler_rejects_run58_newton_node_root_without_branch_control():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 1 沿等距螺线顺时针盘入，龙头速度为1m/s。问题 2 求首次碰撞。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques1": (
                "采用r=0.55θ/(2π)，以S(θ(t))=s0-t保证向内运动；"
                "每个时刻用牛顿法求θ_{i+1}并检查弦长残差。"
            ),
            "sensitivity_analysis": "细化时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="物理最近根|连续延拓"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_accepts_run58_correct_inward_sign_and_root_continuation():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 1 沿等距螺线顺时针盘入，龙头速度为1m/s。问题 2 求首次碰撞。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques1": (
                "采用r=0.55θ/(2π)，以S(θ(t))=s0-t保证向内运动；"
                "从当前参数用局部增量括取θ_{i+1}的最近根，"
                "并以上一时刻根连续延拓，检查位移与速度积分一致。"
            ),
            "sensitivity_analysis": "细化时间步长。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_run58_fixed_outer_cutoff_without_stability_check():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 1 沿螺距为55cm的等距螺线盘入，初始时位于第16圈。"
        "问题 2 板凳碰撞后不能再继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques3": (
                "完整搜索0<p<=p_max，先验证p_max可行，不可行就倍增；"
                "取r_start=60m为充分外侧，从该处到4.5m边界全程用SAT检查碰撞。"
            ),
            "sensitivity_analysis": "细化螺距网格和碰撞步长。",
        }
    )
    with pytest.raises(QualityGateError, match="空间截断稳定性|固定 60 m"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_unnumbered_outer_cutoff_without_stability_check():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques3": (
                "从长度量纲变量r_start所表示的充分外侧截断向内连续回放全链至边界，"
                "用SAT检查全程碰撞。"
            ),
            "sensitivity_analysis": "细化时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="扩大该空间截断"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_accepts_run59_radius_first_spatial_cutoff_stability_check():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques3": (
                "取r_start=50m为充分外侧并回放全链；"
                "将r_start从50m扩大到100m复算，并要求p_min在1e-4m内稳定后才接受截断。"
            ),
            "sensitivity_analysis": "检查空间截断半径。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_accepts_r_start_increased_to_numeric_bound_and_recomputed():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques3": (
                "取r_start=50m为充分外侧，从该处到4.5m边界全程用SAT检查；"
                "r_start从50m增至100m重算，确认p_min偏差<0.005m。"
            ),
            "sensitivity_analysis": "检查空间截断半径。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_accepts_explicit_spatial_cutoff_radius_series():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques3": (
                "设r_start=20m为外侧截断；"
                "对所得p_min用r_start=25/30m复算确认截断稳定，p_min变化<0.1%。"
            ),
            "sensitivity_analysis": "检查空间截断半径。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_accepts_spatial_cutoff_series_with_p_star_notation():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques3": (
                "取r_start=50m为充分外侧并回放全链；"
                "截断r_start取50/70/90m三档验证p*变化<1e-4。"
            ),
            "sensitivity_analysis": "检查空间截断半径。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_prejudged_q3_minimum_pitch_magnitude():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques3": "搜索0<p<=p_max并用SAT验证；p*应远大于0.1m量级。",
            "sensitivity_analysis": "检查搜索网格。",
        }
    )
    with pytest.raises(QualityGateError, match="预设了最小螺距的量级"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_accepts_expanded_spatial_cutoff_with_p_star_notation():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques3": (
                "给定p时从r_start=20m回放全链至边界；"
                "扩大r_start=25m复核p*稳定后才接受空间截断。"
            ),
            "sensitivity_analysis": "继续用r_start=30m检查截断稳定性。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_accepts_explicit_rejection_of_q1_pitch_borrow():
    source = SERIAL_MEMBER_SOURCE + (
        "问题1沿螺距为55cm的等距螺线盘入，初始时位于第16圈。"
        "问题3请确定使龙头到达4.5m边界的最小螺距。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques3": (
                "不继承问题1的p=0.55m或初始圈数，只继承物理可行性判据；"
                "给定p时从r_start=20m连续回放全链至4.5m边界，"
                "再扩大r_start并复核p*稳定。"
            ),
            "sensitivity_analysis": "检查空间截断。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_accepts_not_using_q1_turn_as_q3_initial_state():
    source = SERIAL_MEMBER_SOURCE + (
        "问题1沿螺距为55cm的等距螺线盘入，初始时位于第16圈。"
        "问题3请确定使龙头到达4.5m边界的最小螺距。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques3": (
                "不以第16圈或任何前问数值为初态；从r_start=20m回放全链至边界，"
                "再扩大r_start并复核p*稳定。"
            ),
            "sensitivity_analysis": "检查空间截断。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_run58_limit_only_as_infinite_tail_bound():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 路径含盘入螺线和盘出螺线。问题 5 求各把手不超过2m/s时的最大龙头速度。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques5": (
                "取G_max=max_{i,t}g_i(t)，并令v_head_max=2/G_max。"
                "r→∞时曲率→0、速度系数→1，所以证明无限尾部不产生更大系数。"
            ),
            "sensitivity_analysis": "细化时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="极限|定量上界"):
        validate_modeler_result(result, {"ques5"}, source_text=source)


def test_modeler_accepts_run67_explicit_outer_cutoff_radius_series():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 路径含盘入螺线和盘出螺线。"
        "问题 5 求各把手不超过2m/s时的最大龙头速度。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques5": (
                "取G_max=max_{i,t}g_i(t)，并令v_head_max=2/G_max。"
                "曲率趋零、倍率趋一只作渐近参考；将外侧截断半径由20m依次扩大为"
                "40m和80m，要求G_max变化<1e-4且新增外侧壳层最大倍率不再增加。"
            ),
            "sensitivity_analysis": "细化时间步长。",
        }
    )
    validate_modeler_result(result, {"ques5"}, source_text=source)


def test_modeler_accepts_quantitative_bound_for_outer_cutoff_shell():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 路径含盘入螺线和盘出螺线。"
        "问题 5 求各把手不超过2m/s时的最大龙头速度。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques5": (
                "全程扫描所有节点求η_max=max_{i,t}η_i(t)，并令"
                "v_head,max=2/η_max。对螺线r>R_cut推导解析上界，"
                "须证明截断外新增壳层η≤1+C/R_cut≤已扫描η_max。"
            ),
            "sensitivity_analysis": "扩大R_cut并细化时间步长。",
        }
    )
    validate_modeler_result(result, {"ques5"}, source_text=source)


def test_modeler_rejects_run58_preassumed_maximum_head_speed_direction():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 路径含盘入螺线和盘出螺线。问题 5 求各把手不超过2m/s时的最大龙头速度。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 224节点，内部弦长2.86m和1.65m。",
            "ques5": (
                "取G_max=max_{i,t}g_i(t)，并令v_head_max=2/G_max。"
                "用解析尾部界证明外侧倍率不超过有限区间最大值；"
                "实际应得v_head_max>1。"
            ),
            "sensitivity_analysis": "扩大截断并细化时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="预设最大龙头速度"):
        validate_modeler_result(result, {"ques5"}, source_text=source)


def test_question_contract_rejects_reversed_trailing_spiral_order(tmp_path: Path):
    contract = {
        "status": "success",
        "subtask": "ques1",
        "data_scope": "official statement",
        "assumptions": [],
        "model": "r=(0.55/(2π))θ；后续节点取较小θ根",
        "numeric_results": {
            "node_count": 224,
            "distances_m": [2.86, 1.65],
            "head_x_m": 8.8,
        },
        "validation": {"passed": True, "checks": ["distance residual"]},
        "conclusion_bounds": "fixed geometry",
    }
    (tmp_path / "results_ques1.json").write_text(json.dumps(contract), encoding="utf-8")
    source = SERIAL_MEMBER_SOURCE + "沿螺距55cm的等距螺线顺时针盘入。"
    with pytest.raises(QualityGateError, match="盘入方向|跟随次序"):
        load_evidence_contract(tmp_path, "ques1", source_text=source)


def test_solver_failure_warning_is_not_a_successful_execution():
    assert has_unresolved_solver_failure(
        "WARNING: 无法在θ∈[0,3]找到距离1.65m的解，使用回退值"
    )
    assert has_unresolved_solver_failure("ERROR solver failed: no valid root")
    assert not has_unresolved_solver_failure(
        "distance residual max=8.2e-10; all 224 roots bracketed"
    )
    assert has_unresolved_solver_failure("max distance error = 1.650000e+00 m")
    assert has_unresolved_solver_failure("最大弦长残差: 2.4e-3")
    assert not has_unresolved_solver_failure("max distance error = 8.2e-10 m")


def test_inward_spiral_contract_requires_residual_and_outward_tail(tmp_path: Path):
    source = (
        SERIAL_MEMBER_SOURCE + "问题1舞龙队沿螺距55cm的等距螺线顺时针盘入，"
        "各把手中心均位于螺线上。"
    )
    contract = {
        "status": "success",
        "subtask": "ques1",
        "data_scope": "official statement",
        "assumptions": [],
        "model": "r=(0.55/(2π))θ；后续节点在外侧",
        "numeric_results": {
            "node_count": 224,
            "distances_m": [2.86, 1.65],
        },
        "validation": {"passed": True, "checks": ["弦长残差"]},
        "conclusion_bounds": "0..300s",
    }
    path = tmp_path / "results_ques1.json"
    path.write_text(json.dumps(contract), encoding="utf-8")
    with pytest.raises(QualityGateError, match="缺少可机械复核指标"):
        load_evidence_contract(tmp_path, "ques1", source_text=source)

    contract["numeric_results"].update(
        {
            "max_distance_residual_m": 2e-4,
            "initial_head_radius_m": 8.8,
            "initial_tail_radius_m": 12.1,
        }
    )
    path.write_text(json.dumps(contract), encoding="utf-8")
    with pytest.raises(QualityGateError, match="残差超过"):
        load_evidence_contract(tmp_path, "ques1", source_text=source)

    contract["numeric_results"]["max_distance_residual_m"] = 8e-10
    contract["numeric_results"]["initial_tail_radius_m"] = 3.55
    path.write_text(json.dumps(contract), encoding="utf-8")
    with pytest.raises(QualityGateError, match="初始极径"):
        load_evidence_contract(tmp_path, "ques1", source_text=source)

    contract["numeric_results"]["initial_tail_radius_m"] = 12.1
    path.write_text(json.dumps(contract), encoding="utf-8")
    load_evidence_contract(tmp_path, "ques1", source_text=source)


def test_spreadsheet_template_requires_fill_without_structure_change(tmp_path: Path):
    path = tmp_path / "result1.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "位置"
    sheet["A1"] = "固定说明"
    sheet["B1"] = "0 s"
    sheet["A2"] = "龙头x (m)"
    sheet["B2"].number_format = "0.000000"
    workbook.save(path)
    before = snapshot_spreadsheet_template(path)

    with pytest.raises(QualityGateError, match="未新增数值结果"):
        validate_spreadsheet_template_output(path, before)

    workbook = load_workbook(path)
    workbook["位置"]["B2"] = 8.8
    workbook.save(path)
    validate_spreadsheet_template_output(path, before)

    workbook = load_workbook(path)
    workbook["位置"]["A1"] = "修改说明"
    workbook.save(path)
    with pytest.raises(QualityGateError, match="固定单元格"):
        validate_spreadsheet_template_output(path, before)


def test_spreadsheet_requires_complete_numeric_rectangle_and_outward_node_order(
    tmp_path: Path,
):
    path = tmp_path / "result1.xlsx"
    workbook = Workbook()
    position = workbook.active
    position.title = "位置"
    speed = workbook.create_sheet("速度")
    for sheet in (position, speed):
        sheet["B1"] = "0 s"
        sheet["C1"] = "1 s"
    for row, label in enumerate(
        ("龙头x (m)", "龙头y (m)", "龙尾（后）x (m)", "龙尾（后）y (m)"), 2
    ):
        position.cell(row, 1).value = label
    speed["A2"] = "龙头 (m/s)"
    speed["A3"] = "龙尾（后） (m/s)"
    workbook.save(path)
    before = snapshot_spreadsheet_template(path)
    source = (
        "问题1舞龙队沿等距螺线顺时针盘入，各把手中心均位于螺线上，"
        "将结果保存到result1.xlsx。"
    )

    workbook = load_workbook(path)
    workbook["位置"]["B2"] = 8.8
    workbook["位置"]["B3"] = 0.0
    workbook.save(path)
    with pytest.raises(QualityGateError, match="未完整回填"):
        validate_spreadsheet_template_output(
            path, before, source_text=source, subtask_title="ques1"
        )

    workbook = load_workbook(path)
    values = {
        "B2": 8.8,
        "B3": 0.0,
        "B4": 8.0,
        "B5": 0.0,
        "C2": 8.75,
        "C3": -0.99,
        "C4": 7.94,
        "C5": -0.99,
    }
    for coordinate, value in values.items():
        workbook["位置"][coordinate] = value
    for coordinate in ("B2", "C2", "B3", "C3"):
        workbook["速度"][coordinate] = 1.0
    workbook.save(path)
    with pytest.raises(QualityGateError, match="空间次序反转"):
        validate_spreadsheet_template_output(
            path, before, source_text=source, subtask_title="ques1"
        )

    workbook = load_workbook(path)
    workbook["位置"]["B4"] = 9.2
    workbook["位置"]["C4"] = 9.14
    workbook.save(path)
    validate_spreadsheet_template_output(
        path, before, source_text=source, subtask_title="ques1"
    )


def test_spreadsheet_rejects_time_series_branch_jump(tmp_path: Path):
    path = tmp_path / "result1.xlsx"
    workbook = Workbook()
    position = workbook.active
    position.title = "位置"
    speed = workbook.create_sheet("速度")
    for sheet in (position, speed):
        sheet["B1"] = "0 s"
        sheet["C1"] = "1 s"
    position["A2"] = "龙头x (m)"
    position["A3"] = "龙头y (m)"
    speed["A2"] = "龙头 (m/s)"
    workbook.save(path)
    before = snapshot_spreadsheet_template(path)

    workbook = load_workbook(path)
    workbook["位置"]["B2"] = 0.0
    workbook["位置"]["B3"] = 0.0
    workbook["位置"]["C2"] = 0.8
    workbook["位置"]["C3"] = 0.0
    workbook["速度"]["B2"] = 1.0
    workbook["速度"]["C2"] = 1.0
    workbook.save(path)
    validate_spreadsheet_template_output(path, before, subtask_title="ques1")

    workbook = load_workbook(path)
    workbook["位置"]["C2"] = 8.0
    workbook.save(path)
    with pytest.raises(QualityGateError, match="分支跳变"):
        validate_spreadsheet_template_output(path, before, subtask_title="ques1")


def test_spreadsheet_rejects_clockwise_path_mirrored_across_x_axis(tmp_path: Path):
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Figure 4 trajectory")
    page.insert_text((100, 200), "O")
    page.insert_text((200, 200), "A")
    page.insert_text((240, 200), "x")
    document.save(tmp_path / "official.pdf")
    document.close()

    path = tmp_path / "result1.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "位置"
    sheet["B1"] = "0 s"
    sheet["C1"] = "1 s"
    sheet["A2"] = "龙头y (m)"
    workbook.save(path)
    before = snapshot_spreadsheet_template(path)

    source = "顺时针盘入。初始时，龙头位于 A 点处（见图 4）。"
    workbook = load_workbook(path)
    workbook["位置"]["B2"] = 0.0
    workbook["位置"]["C2"] = 0.1
    workbook.save(path)
    with pytest.raises(QualityGateError, match="顺时针坐标方向"):
        validate_spreadsheet_template_output(
            path, before, source_text=source, subtask_title="ques1"
        )

    workbook = load_workbook(path)
    workbook["位置"]["C2"] = -0.1
    workbook.save(path)
    validate_spreadsheet_template_output(
        path, before, source_text=source, subtask_title="ques1"
    )


def test_writer_rejects_placeholders():
    with pytest.raises(QualityGateError):
        validate_writer_result(WriterResponse(response_content="待补充"))


@pytest.mark.parametrize(
    "text",
    [
        "论文生成时间：2026-08-13",
        "Agent 版本：v2",
        "本文的核心优势不是算法名称，而是证据链。",
    ],
)
def test_writer_rejects_generation_meta_text(text: str):
    with pytest.raises(QualityGateError):
        validate_writer_result(
            WriterResponse(response_content=text), section_name="firstPage"
        )


def test_paper_rejects_colored_heading():
    with pytest.raises(QualityGateError):
        validate_competition_paper_text(r"\section{\textcolor{blue}{问题分析}}")


def test_paper_rejects_duplicate_reference_heading():
    with pytest.raises(QualityGateError):
        validate_competition_paper_text("## 参考文献\n正文\n# 参考文献")


def test_normal_competition_abstract_passes():
    validate_writer_result(
        WriterResponse(
            response_content=(
                "# 板凳龙运动轨迹与调头路径模型\n\n"
                "摘要：建立等距约束模型，求得首次碰撞时刻并完成全区间复核。"
            )
        ),
        section_name="firstPage",
    )


def test_prompts_route_abc_and_require_domain_audit():
    assert "A/B/C 差异化先验" in MODELER_PROMPT


def test_modeler_prompt_requires_initial_event_history_and_template_preservation():
    assert "题面初态 `t=0`" in MODELER_PROMPT
    assert "`r_start=16p`" in MODELER_PROMPT
    assert "至少两个逐步扩大的截断" in MODELER_PROMPT
    assert "基线或优化后构型皆可" in MODELER_PROMPT
    assert "多起点/爬坡确认全局" in MODELER_PROMPT
    assert "`d_2`" in MODELER_PROMPT
    assert "S(θ)=a/2" in MODELER_PROMPT
    assert "M_max≥1" in MODELER_PROMPT
    assert "原工作表、行列、表头和样式" in MODELER_PROMPT
    assert "content_override" in MODELER_PROMPT
    assert "连续域与全题覆盖验收" in CODER_PROMPT
    assert "多个相互分离的候选峰" in CODER_PROMPT
    assert "source_parameter_audit" in MODELER_PROMPT
    assert "L-e_1-e_2" in MODELER_PROMPT
    assert "同一个空间节点" in MODELER_PROMPT
    assert "题目原文" in CODER_PROMPT
    assert "节点--构件关联表" in CODER_PROMPT
    assert "局部增量" in CODER_PROMPT
    assert "绝对参数上界" in CODER_PROMPT
    assert "弦约束微分" in CODER_PROMPT
    assert "有限差分只能作为独立交叉核验" in CODER_PROMPT


def test_modeler_requires_source_parameter_audit():
    result = ModelerToCoder(
        questions_solution={"eda": "ok", "ques1": "ok", "sensitivity_analysis": "ok"}
    )
    with pytest.raises(QualityGateError):
        validate_modeler_result(result, {"ques1"})


def test_coder_flow_includes_original_problem_as_fact_source():
    questions = {
        "background": "实体尺寸背景",
        "ques_count": 1,
        "ques1": "孔中心分别距两端 0.275 m",
    }
    modeler = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: checked",
            "ques1": "candidate",
            "sensitivity_analysis": "candidate",
        }
    )
    flows = Flows(questions).get_solution_flows(questions, modeler)
    assert "唯一事实来源" in flows["eda"]["coder_prompt"]
    assert "孔中心分别距两端 0.275 m" in flows["eda"]["coder_prompt"]
    assert "唯一事实来源" in flows["ques1"]["coder_prompt"]
    assert "孔中心分别距两端 0.275 m" in flows["ques1"]["source_text"]


@pytest.mark.parametrize(
    ("file_name", "expected_role"),
    [
        ("result1.xlsx", "output_template_candidate"),
        ("提交结果模板.xlsx", "output_template_candidate"),
        ("附件1.xlsx", "observational_data_candidate"),
        ("A题.pdf", "problem_statement_or_instructions"),
    ],
)
def test_coder_freezes_input_file_roles(file_name: str, expected_role: str):
    assert classify_input_file_role(file_name) == expected_role


def test_coder_recognizes_source_figure_inspection_code():
    assert is_source_figure_inspection_code("doc = fitz.open('official.pdf')")
    assert is_source_figure_inspection_code("text = page.get_text()")
    assert is_source_figure_inspection_code("Image.open('fig4_crop.png')")
    assert is_source_figure_inspection_code(
        "text = open('problem.txt').read(); print('图 4 图注', text)"
    )
    assert not is_source_figure_inspection_code("open('problem.txt').read()")
    assert not is_source_figure_inspection_code("load_workbook('result1.xlsx')")


def test_coder_prompt_stops_blank_template_eda():
    assert "输出模板中的空白数值格" in CODER_PROMPT
    assert "最多用两次工具调用" in CODER_PROMPT
    assert "source_figures` **JSON 数组**" in CODER_PROMPT


def test_eda_writing_template_is_not_fixed_to_descriptive_statistics():
    template = (
        Path(__file__).parents[2] / "app" / "config" / "md_template.toml"
    ).read_text(encoding="utf-8")
    assert "## 4.2 数据与参数说明" in template
    assert "## 4.2 描述性统计" not in template
    assert "不讨论空白模板" in template


def test_writer_prompt_avoids_template_bloat():
    prompt = get_writer_prompt(FormatOutPut.LaTeX)
    assert "全文合计：13-18张" not in prompt
    assert "每幅图表至少配3行" not in prompt
    assert "各级标题使用黑色" in prompt
    assert "生成/验收时间" in prompt


def test_repeated_reference_keeps_one_number(tmp_path: Path):
    output = UserOutput(str(tmp_path), ques_count=1)
    first = WriterResponse(response_content="结论 {[^1]: Source}")
    second = WriterResponse(response_content="再次引用 {[^1]: Source}")
    output.set_res("firstPage", first)
    output.set_res("ques1", second)
    # 完整顺序所需的空章节以空字符串补齐。
    for key in output.seq:
        output.res.setdefault(key, {"response_content": "", "footnotes": None})
    text = output.get_result_to_save()
    assert text.count("[^1]") == 3  # 正文两次 + 文末脚注定义一次
    assert text.count("Source") == 1


def test_modeler_rejects_run84_radius_search_after_fixing_turn_endpoints():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 盘入、盘出螺线中心对称，调头路径由两段圆弧相切连接，"
        "前弧半径是后弧的2倍。能否调整圆弧使调头曲线变短？"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223个板体、224节点，"
                "内部节点距2.86m与1.65m，题面几何已冻结。"
            ),
            "ques4": (
                "固定盘入边界点M与盘出边界点N；M处切向为T_M，N处切向为T_N。"
                "两弧在连接点位置和切向C1连续，R1=2R2，外切且圆心距为R1+R2；"
                "两侧曲率1/R1与1/R2不同并发生跳变。变量为R2、圆心和切点位置，"
                "用黄金分割扫描R2求最小长度L_opt。复算题面基线长度L_base，"
                "比较L_opt与L_base后回答能否缩短。"
            ),
            "sensitivity_analysis": "只检查方程残差和数值容差。",
        }
    )
    with pytest.raises(QualityGateError, match="自由度定义矛盾"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_accepts_direct_solution_for_fixed_turn_endpoints():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 盘入、盘出螺线中心对称，调头路径由两段圆弧相切连接，"
        "前弧半径是后弧的2倍。能否调整圆弧使调头曲线变短？"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223个板体、224节点，"
                "内部节点距2.86m与1.65m，题面几何已冻结。"
            ),
            "ques4": (
                "固定入口点M和出口点N；M处切向为T_M，N处切向为T_N。"
                "R1=2R2；两弧反向转向外切，"
                "圆心距为R1+R2。先列圆心、连接点和半径未知量与独立约束，"
                "直接联立求解并检验可行解是否唯一，不再扫描R2。两弧在连接点位置和"
                "切向C1连续，但曲率1/R1与1/R2不同并发生跳变。复算基线长度L_base，"
                "各弧按实际切点与连接点的圆心角计算L=R1·Δφ1+R2·Δφ2；"
                "令约束内最优值L_opt与L_base作数值比较，再由长度差回答能否缩短。"
            ),
            "sensitivity_analysis": "只检查方程残差和数值容差。",
        }
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_accepts_explicitly_movable_turn_tangent_points():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 盘入、盘出螺线中心对称，调头路径由两段圆弧相切连接，"
        "前弧半径是后弧的2倍。能否调整圆弧使调头曲线变短？"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223个板体、224节点，"
                "内部节点距2.86m与1.65m，题面几何已冻结。"
            ),
            "ques4": (
                "入口和出口切点不固定；以u、v参数化两切点，使其可沿各自螺线移动，"
                "并明确移动域。保持R1=2R2，两弧反向转向外切且圆心距为R1+R2；"
                "连接点位置和切向C1连续，曲率1/R1与1/R2不同并发生跳变。"
                "在完整相切和圆内约束下扫描u、v，按实际圆心角计算"
                "L=R1·Δφ1+R2·Δφ2并得到L_opt；"
                "复算题面调整前基线长度L_base并比较两者后回答能否缩短。"
            ),
            "sensitivity_analysis": "只检查切点网格和方程残差。",
        }
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_run84_fixed_dimensions_as_manufacturing_uncertainty():
    source = SERIAL_MEMBER_SOURCE + "问题 1 按题面固定板长与孔位建立运动学。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个板体、224节点，题面尺寸已冻结。",
            "ques1": "使用欧氏弦长约束并细化时间步长。",
            "sensitivity_analysis": (
                "节点距离/板长按孔心距±0.5%制造误差扰动；"
                "同时声明题面板长和孔位禁止扰动。"
            ),
        }
    )
    with pytest.raises(QualityGateError, match="题面固定构件尺寸"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_run87_all_nodes_checked_as_unit_arc_speed():
    source = SERIAL_MEMBER_SOURCE + "问题 1 龙头前把手速度保持1m/s。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个板体、224节点，孔距2.86m与1.65m。",
            "ques1": (
                "速度由刚性弦长约束的两端切向投影递推；逐秒检验"
                "||S(θ_i(t))-S(θ_i(t-1))|-1|<1e-6。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="任意节点.*龙头1m/s"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_accepts_head_only_unit_arc_speed_check():
    source = SERIAL_MEMBER_SOURCE + "问题 1 龙头前把手速度保持1m/s。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个板体、224节点，孔距2.86m与1.65m。",
            "ques1": (
                "速度由刚性弦长约束的两端切向投影递推；龙头节点1逐秒检验"
                "||S(θ_i(t))-S(θ_i(t-1))|-1|<1e-6，其中i=1。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_run87_bisection_after_bare_pitch_monotonic_claim():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 3 求最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个板体、224节点，孔距2.86m与1.65m。",
            "ques3": (
                "搜索0<p≤p_max，随p增大可行；直接用二分求p_min。"
                "从充分外侧回放并扩大r_start复算p_min，结果差<1e-6m。"
            ),
            "sensitivity_analysis": "检查时间步长。",
        }
    )
    with pytest.raises(QualityGateError, match="二分搜索前未证明"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_run87_inward_boundary_radius_direction_reversed():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 2 碰撞后不能继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个板体、224节点，孔距2.86m与1.65m。",
            "ques3": (
                "从充分外侧沿螺线向内连续回放全链，用有向矩形SAT判定碰撞；"
                "检查碰撞发生在r<4.5m之前。"
            ),
            "sensitivity_analysis": (
                "将r_start增大一倍计算p_min，结果差≤0.001m；细化碰撞时间步。"
            ),
        }
    )
    with pytest.raises(QualityGateError, match=r"边界前.*r>4\.5"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_accepts_run87_inward_boundary_radius_direction():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 2 碰撞后不能继续盘入。"
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个板体、224节点，孔距2.86m与1.65m。",
            "ques3": (
                "从充分外侧沿螺线向内连续回放全链，用有向矩形SAT判定碰撞；"
                "若在到达边界前发生碰撞，该状态应满足r>4.5m。"
            ),
            "sensitivity_analysis": (
                "将r_start增大一倍计算p_min，结果差≤0.001m；细化碰撞时间步。"
            ),
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_run87_ambiguous_pre_adjustment_turning_baseline():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 调头路径由两段圆弧相切连接成S形曲线，前一段半径是"
        "后一段的2倍。能否调整圆弧使调头曲线变短？"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个板体、224节点，孔距2.86m与1.65m。",
            "ques4": (
                "保持R1:R2=2:1；两弧位置和切向C1连续，连接点曲率跳变，"
                "采用外切分支且圆心距R1+R2。由实际切点求圆心角，"
                "按L(R)=2R·Δφ1(R)+R·Δφ2(R)计算长度。"
                "基线取完整几何方程的唯一或优选可行解，再与L_opt比较缩短量。"
            ),
            "sensitivity_analysis": "保持比例固定，检查相切残差和求解容差。",
        }
    )
    with pytest.raises(QualityGateError, match=r"调整前基线.*唯一或优选"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_accepts_source_determined_pre_adjustment_turning_baseline():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 调头路径由两段圆弧相切连接成S形曲线，前一段半径是"
        "后一段的2倍。能否调整圆弧使调头曲线变短？"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个板体、224节点，孔距2.86m与1.65m。",
            "ques4": (
                "调整前入口、出口及切点由题面几何唯一确定，不在候选族中择优；"
                "保持R1:R2=2:1，记R2=R；两弧位置和切向C1连续，连接点曲率跳变，"
                "采用外切分支且圆心距R1+R2。由实际切点求圆心角，"
                "按L(R)=2R·Δφ1(R)+R·Δφ2(R)复算调整前基线L_base并报告相切残差，"
                "再与约束内候选L_opt比较缩短量。若方程存在多解则停止并核对题面，"
                "不自行选择基线。"
            ),
            "sensitivity_analysis": "保持比例固定，检查相切残差和求解容差。",
        }
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_accepts_run88_half_diagonal_broad_phase_formula():
    source = SERIAL_MEMBER_SOURCE + "问题 2 确定板凳之间首次碰撞时刻。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个板体、224节点，孔距2.86m与1.65m。",
            "ques2": (
                "粗筛用两矩形中心距小于两板半对角线之和；"
                "半对角线=√(半长²+半宽²)，随后用有向矩形SAT精判。"
            ),
            "sensitivity_analysis": "细化事件时间步长。",
        }
    )
    validate_modeler_result(result, {"ques2"}, source_text=source)


def test_modeler_accepts_run88_global_fixed_ratio_denial():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 调头路径由两段圆弧相切连接，前段半径是后段的2倍。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223个板体、224节点，"
                "内部距离2.86m和1.65m，半径比已冻结。"
            ),
            "ques4": "两弧位置与切向C1连续，连接点曲率跳变。",
            "sensitivity_analysis": (
                "圆弧半径比R1:R2=2:1为来源常量，不做敏感性扰动；"
                "只将数值求解容差改变±1%。"
            ),
        }
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_run88_nonexistent_222nd_body_section():
    source = SERIAL_MEMBER_SOURCE + "后面221节为龙身，最后1节为龙尾。"
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223个板体、224节点，"
                "内部距离2.86m和1.65m；"
                "节点223=第222节龙身后把手=龙尾前把手。"
            ),
            "sensitivity_analysis": "检查节点映射。",
        }
    )
    with pytest.raises(QualityGateError, match="只有 221 节龙身"):
        validate_modeler_result(result, set(), source_text=source)


def test_modeler_rejects_run88_fixed_turning_radius_perturbation():
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题 3 调头空间是直径9m的圆形区域。问题 4 在该调头空间内设计路径。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224节点，内部距离2.86m和1.65m；"
                "调头空间直径9m，半径4.5m。"
            ),
            "ques4": "在固定调头空间内求解路径。",
            "sensitivity_analysis": "问题4将调头空间圆半径4.5m按±0.1%变化并比较长度。",
        }
    )
    with pytest.raises(QualityGateError, match="调头空间直径"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_accepts_turning_radius_fixed_with_residual_tolerance_change():
    source = (
        SERIAL_MEMBER_SOURCE
        + "问题 3 调头空间是直径9m的圆形区域。问题 4 在该调头空间内设计路径。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224节点，内部距离2.86m和1.65m；"
                "调头空间直径9m，半径4.5m。"
            ),
            "ques4": "在固定调头空间内求解路径。",
            "sensitivity_analysis": (
                "题面固定物理量（调头空间直径9m、半径比2:1）为来源常量，"
                "不做敏感性扰动。"
                "保持调头空间圆半径4.5m固定，只将圆域约束的数值残差容差"
                "从1e-8改到1e-10、入口和出口切点移动域扩缩±1%。"
            ),
        }
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_run88_finite_sampling_as_open_tail_certificate():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个板体、224节点，孔距2.86m与1.65m。",
            "ques3": "搜索域保持0<p≤p_max，对每个候选用全链SAT判定可行性。",
            "sensitivity_analysis": (
                "检验p→0+方向无隐藏可行解，在0与首个网格点间加密采样验证不可行。"
            ),
        }
    )
    with pytest.raises(QualityGateError, match="有限加密采样"):
        validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_accepts_interval_certificate_for_open_pitch_tail():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个板体、224节点，孔距2.86m与1.65m。",
            "ques3": "搜索域保持0<p≤p_max，对每个候选用全链SAT判定可行性。",
            "sensitivity_analysis": (
                "问题3的p_max从初始候选扩大1.5倍验证括界；"
                "对0<p<p1用区间算术与解析下界严格证明全链不可行；"
                "有限采样只作交叉核验。"
            ),
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_accepts_run89_symbolic_double_cutoff_recalculation():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 3 请确定最小螺距，使龙头前把手盘入到调头空间的边界。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个板体、224节点，孔距2.86m与1.65m。",
            "ques3": (
                "保持0<p≤p_max，从外侧截断r_start盘入到边界并用全链SAT判定；"
                "在r_start与2×r_start分别重算p_min并报告差值<1e-4m。"
            ),
            "sensitivity_analysis": "检查空间截断扩展收敛。",
        }
    )
    validate_modeler_result(result, {"ques3"}, source_text=source)


def test_modeler_rejects_run88_fixed_sampling_as_arc_containment_proof():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 调头路径由两段圆弧相切连接成S形，前段半径是后段的2倍。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223个板体、224节点，"
                "内部距离2.86m和1.65m，半径比已冻结。"
            ),
            "ques4": (
                "两弧位置与切向C1连续，连接点曲率跳变；采用外切分支，"
                "圆心距R1+R2。全曲线位于4.5m调头圆内，"
                "采样2000点检查最大极径≤4.5m。"
            ),
            "sensitivity_analysis": "细化采样点。",
        }
    )
    with pytest.raises(QualityGateError, match="有限采样点"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_accepts_analytic_arc_containment_with_sampling_cross_check():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 调头路径由两段圆弧相切连接成S形，前段半径是后段的2倍。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 223个板体、224节点，内部距离2.86m和1.65m；"
                "半径比已冻结。"
            ),
            "ques4": (
                "两弧位置与切向C1连续，连接点曲率跳变；采用外切分支，"
                "圆心距R1+R2。对每段圆弧极径解析求导，检查端点与驻点，"
                "证明全曲线位于4.5m调头圆内；采样2000点只作交叉核验。"
            ),
            "sensitivity_analysis": "检查解析极值与采样结果一致。",
        }
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_run88_mmax_one_as_inactive_speed_constraint():
    source = SERIAL_MEMBER_SOURCE + "问题 5 各把手速度均不超过2m/s。"
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224节点、内部距离2.86m和1.65m，"
                "均受速度上限约束。"
            ),
            "ques5": (
                "M_max=max_{k=1..224}|v_k/v_head|，v_head,max=2/M_max。"
                "若M_max≤1，则v_head≤2m/s可能非活动约束。"
            ),
            "sensitivity_analysis": "加密时间步。",
        }
    )
    with pytest.raises(QualityGateError, match="可能非活动"):
        validate_modeler_result(result, {"ques5"}, source_text=source)


def test_modeler_accepts_mmax_one_as_active_head_speed_constraint():
    source = SERIAL_MEMBER_SOURCE + "问题 5 各把手速度均不超过2m/s。"
    result = ModelerToCoder(
        questions_solution={
            "eda": (
                "source_parameter_audit: 224节点，内部距离2.86m和1.65m，"
                "均受速度上限约束。"
            ),
            "ques5": (
                "M_max=max_{k=1..224}|v_k/v_head|且因龙头自身m_1=1而M_max≥1。"
                "v_head,max=2/M_max；当M_max=1时，龙头速度恰好达到2m/s，约束仍活动。"
            ),
            "sensitivity_analysis": "加密时间步。",
        }
    )
    validate_modeler_result(result, {"ques5"}, source_text=source)


def test_modeler_rejects_run89_missing_turning_path_containment_certificate():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 舞龙队在调头空间内完成调头，路径由两段圆弧相切连接成S形，"
        "前段半径是后段的2倍。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个板体、224节点，孔距2.86m与1.65m。",
            "ques4": (
                "两弧位置与切向C1连续，连接点曲率跳变；采用外切分支，"
                "圆心距R1+R2。可视化中画出直径9m调头圆与双圆弧。"
            ),
            "sensitivity_analysis": "检查相切方程残差。",
        }
    )
    with pytest.raises(QualityGateError, match="连续区域证书"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_accepts_required_turning_path_with_analytic_containment():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 舞龙队在调头空间内完成调头，路径由两段圆弧相切连接成S形，"
        "前段半径是后段的2倍。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个板体、224节点，孔距2.86m与1.65m。",
            "ques4": (
                "两弧位置与切向C1连续，连接点曲率跳变；采用外切分支，"
                "圆心距R1+R2。对每段圆弧极径解析求导并检查端点与驻点，"
                "证明全曲线位于4.5m调头圆内。"
            ),
            "sensitivity_analysis": "检查解析极值和相切方程残差。",
        }
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_rejects_run89_invented_velocity_component_columns():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 1 给出每秒整个舞龙队的位置和速度，将结果保存到result1.xlsx。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个板体、224节点，孔距2.86m与1.65m。",
            "ques1": "result1.xlsx逐秒回填全部224节点位置(x,y)与速度(vx,vy)。",
            "sensitivity_analysis": "检查数值残差。",
        }
    )
    with pytest.raises(QualityGateError, match="自行把速度扩成 vx/vy"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_accepts_official_template_without_invented_velocity_columns():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 1 给出每秒整个舞龙队的位置和速度，将结果保存到result1.xlsx。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个板体、224节点，孔距2.86m与1.65m。",
            "ques1": (
                "result1.xlsx保留官方工作表、行列和表头，只回填全部224节点"
                "已有的位置与速度数值格。"
            ),
            "sensitivity_analysis": "检查数值残差。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_run89_follower_inverse_without_full_branch_control():
    source = SERIAL_MEMBER_SOURCE + "问题 1 沿等距螺线顺时针盘入。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个板体、224节点，孔距2.86m与1.65m。",
            "ques1": (
                "节点k=2..224由欧氏弦长方程在螺线上反解θₖ，"
                "取外侧支θₖ>θₖ₋₁。"
            ),
            "sensitivity_analysis": "检查刚性残差。",
        }
    )
    with pytest.raises(QualityGateError, match="局部最近根括界"):
        validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_accepts_follower_inverse_with_spatial_and_temporal_branch_control():
    source = SERIAL_MEMBER_SOURCE + "问题 1 沿等距螺线顺时针盘入。"
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个板体、224节点，孔距2.86m与1.65m。",
            "ques1": (
                "节点k=2..224由欧氏弦长方程在螺线上反解θₖ；"
                "在允许外侧方向取局部增量括界中的最近根，并从上一时刻连续延拓，"
                "用相邻时刻位移与速度积分一致性检查阻断跳支。"
            ),
            "sensitivity_analysis": "检查刚性残差。",
        }
    )
    validate_modeler_result(result, {"ques1"}, source_text=source)


def test_modeler_rejects_run89_missing_turn_start_time_origin():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 调头路径由两段圆弧相切连接成S形，前段半径是后段的2倍；"
        "以调头开始时间为零时刻，给出从-100s到100s的位置和速度。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个板体、224节点，孔距2.86m与1.65m。",
            "ques4": (
                "两弧位置与切向C1连续，连接点曲率跳变，采用外切分支；"
                "整条路径参数化，在-100~100s逐秒推进并输出位置速度。"
            ),
            "sensitivity_analysis": "检查时间步长收敛。",
        }
    )
    with pytest.raises(QualityGateError, match="遗漏题面时间原点"):
        validate_modeler_result(result, {"ques4"}, source_text=source)


def test_modeler_accepts_turn_start_mapped_to_zero_time():
    source = SERIAL_MEMBER_SOURCE + (
        "问题 4 调头路径由两段圆弧相切连接成S形，前段半径是后段的2倍；"
        "以调头开始时间为零时刻，给出从-100s到100s的位置和速度。"
    )
    result = ModelerToCoder(
        questions_solution={
            "eda": "source_parameter_audit: 223个板体、224节点，孔距2.86m与1.65m。",
            "ques4": (
                "两弧位置与切向C1连续，连接点曲率跳变，采用外切分支，"
                "圆心距为R1+R2。"
                "明确t=0对应调头开始；由该事件向前回放盘入段，"
                "向后推进调头与盘出段，在-100~100s逐秒输出。"
            ),
            "sensitivity_analysis": "检查时间步长收敛。",
        }
    )
    validate_modeler_result(result, {"ques4"}, source_text=source)
